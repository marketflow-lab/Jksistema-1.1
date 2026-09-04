import csv
import json
import sqlite3
import tempfile
import unittest
from datetime import date, datetime, timedelta
from pathlib import Path
from unittest.mock import patch

import pandas as pd

from backend.services import codex_actions, codex_assistant_storage, codex_reports_advanced, ia_tools_vendas
from backend.services.codex.assistant import analysis_management as assistant_management
from backend.services.codex.assistant import analysis_stock as assistant_stock
from backend.services.codex.assistant import reports_artifacts as assistant_artifacts
from backend.services.codex.assistant import reports_html as assistant_html
from backend.services.codex.assistant import runtime as assistant_runtime
from backend.services.sales_tools import inventory as sales_inventory


class CodexAdvancedReportTest(unittest.TestCase):
    def _tenant(self, root: str, client_id: str = "000001") -> Path:
        tenant = Path(root, client_id)
        tenant.mkdir(parents=True, exist_ok=True)
        return tenant

    def test_cost_map_fails_closed_for_duplicate_store_display_names(self):
        with tempfile.TemporaryDirectory() as root:
            tenant = self._tenant(root)
            with (tenant / "cadastro_custos_lojas.csv").open(
                "w", encoding="utf-8", newline=""
            ) as handle:
                writer = csv.DictWriter(
                    handle,
                    fieldnames=["store_id", "loja_sync", "sku", "custo"],
                )
                writer.writeheader()
                writer.writerows(
                    [
                        {"store_id": "store-a", "loja_sync": "Mesmo nome", "sku": "SKU-1", "custo": "10"},
                        {"store_id": "store-b", "loja_sync": "Mesmo nome", "sku": "SKU-1", "custo": "99"},
                    ]
                )

            by_store, _generic = codex_reports_advanced._load_cost_maps(str(tenant))

        self.assertNotIn((codex_reports_advanced._text_key("Mesmo nome"), "SKU-1"), by_store)

    def test_cost_map_prefers_exact_identity_over_later_legacy_shadow(self):
        with tempfile.TemporaryDirectory() as root:
            tenant = self._tenant(root)
            with (tenant / "cadastro_custos_lojas.csv").open(
                "w", encoding="utf-8", newline=""
            ) as handle:
                writer = csv.DictWriter(
                    handle,
                    fieldnames=["store_id", "loja_sync", "sku", "custo"],
                )
                writer.writeheader()
                writer.writerows(
                    [
                        {"store_id": "store-a", "loja_sync": "Loja A", "sku": "SKU-1", "custo": "10"},
                        {"store_id": "", "loja_sync": "Loja A", "sku": "SKU-1", "custo": "999"},
                    ]
                )

            by_store, _generic = codex_reports_advanced._load_cost_maps(str(tenant))

        self.assertEqual(
            by_store[(codex_reports_advanced._text_key("Loja A"), "SKU-1")]["cost"],
            10,
        )

    def _seed_operational_data(self, root: str) -> tuple[Path, dict]:
        tenant = self._tenant(root)
        (tenant / "lojas_config.json").write_text(
            json.dumps([{"store_id": "store-1", "nome": "Loja 1"}]),
            encoding="utf-8",
        )
        period = codex_reports_advanced._period_for_profile("daily_exceptions")
        current_start = date.fromisoformat(period["start"])
        current_end = date.fromisoformat(period["end"])
        db = sqlite3.connect(tenant / "vendas_historico.db")
        db.executescript(
            """
            CREATE TABLE vendas (
                id_unico TEXT PRIMARY KEY, data TEXT, loja_conta TEXT, numero TEXT,
                sku TEXT, produto TEXT, quantidade REAL, valor REAL, situacao TEXT,
                canal TEXT, comprador TEXT, devolucao INTEGER DEFAULT 0
            );
            CREATE TABLE notas_entrada_itens (
                id_unico TEXT PRIMARY KEY, data_emissao TEXT, loja_conta TEXT,
                numero_nota TEXT, sku TEXT, descricao TEXT, quantidade REAL,
                valor_total REAL, fornecedor TEXT, devolucao INTEGER DEFAULT 0
            );
            """
        )
        sequence = 0
        history_start = current_end - timedelta(days=179)
        day = history_start
        while day <= current_end:
            sequence += 1
            db.execute(
                "INSERT INTO vendas VALUES (?,?,?,?,?,?,?,?,?,?,?,0)",
                (f"hist-{sequence}", day.isoformat(), "Loja 1", f"H{sequence}", "SKU-A", "Roçadeira", 2, 200, "Concluída", "ML", "Cliente"),
            )
            day += timedelta(days=7)
        previous_day = current_start - timedelta(days=10)
        db.execute(
            "INSERT INTO vendas VALUES (?,?,?,?,?,?,?,?,?,?,?,0)",
            ("previous-b", previous_day.isoformat(), "Loja 1", "PB", "SKU-B", "Caçamba", 20, 2000, "Concluída", "ML", "Cliente"),
        )
        db.execute(
            "INSERT INTO vendas VALUES (?,?,?,?,?,?,?,?,?,?,?,0)",
            ("current-b", current_start.isoformat(), "Loja 1", "CB", "SKU-B", "Caçamba", 2, 200, "Concluída", "ML", "Cliente"),
        )
        db.execute(
            "INSERT INTO notas_entrada_itens VALUES (?,?,?,?,?,?,?,?,?,0)",
            ("receipt-a", (current_start + timedelta(days=2)).isoformat(), "Loja 1", "NF1", "SKU-A", "Roçadeira", 10, 500, "Fornecedor A"),
        )
        db.commit()
        db.close()

        stock_db = sqlite3.connect(tenant / "estoque_historico.db")
        stock_db.execute(
            """CREATE TABLE estoque_historico (
                data_ref TEXT, recorded_at TEXT, loja_sync TEXT, sku TEXT,
                nome_bling TEXT, saldo_loja REAL, saldo_full REAL
            )"""
        )
        opening = current_start - timedelta(days=1)
        for sku, local_stock, full_stock in (("SKU-A", 20, 999), ("SKU-B", 30, 0)):
            stock_db.execute(
                "INSERT INTO estoque_historico VALUES (?,?,?,?,?,?,?)",
                (opening.isoformat(), opening.isoformat(), "Loja 1", sku, sku, local_stock, full_stock),
            )
            stock_db.execute(
                "INSERT INTO estoque_historico VALUES (?,?,?,?,?,?,?)",
                (current_end.isoformat(), current_end.isoformat(), "Loja 1", sku, sku, 1 if sku == "SKU-A" else 50, full_stock),
            )
        stock_db.commit()
        stock_db.close()

        with (tenant / "cadastro_produtos.csv").open("w", encoding="utf-8-sig", newline="") as fh:
            writer = csv.DictWriter(fh, fieldnames=["sku", "nome", "custo", "preco", "imposto"], delimiter=";")
            writer.writeheader()
            writer.writerows(
                [
                    {"sku": "SKU-A", "nome": "Roçadeira", "custo": "50", "preco": "100", "imposto": "10"},
                    {"sku": "SKU-B", "nome": "Caçamba", "custo": "60", "preco": "100", "imposto": "10"},
                ]
            )
        (tenant / "listas_pedidos.json").write_text(
            json.dumps(
                [
                    {
                        "id": "open-all",
                        "nome_lista": "Reposição geral",
                        "loja": "__todas",
                        "status": "Pedido Aprovado",
                        "itens": [{"SKU": "SKU-A", "Quantidade": 5}],
                    }
                ],
                ensure_ascii=False,
            ),
            encoding="utf-8",
        )
        return tenant, period

    def test_missing_values_do_not_become_zero_and_stock_full_stays_separate(self):
        with tempfile.TemporaryDirectory() as root:
            _tenant, period = self._seed_operational_data(root)
            margin_rows = [
                {"_valor_num": 1200, "_lucro_num": 240, "_margem_completa": True},
            ]
            result = codex_reports_advanced.build_profile_context(
                info_base=root,
                client_id="000001",
                profile="daily_exceptions",
                margin_rows=margin_rows,
            )

        self.assertEqual(result["scope"]["period_start"], period["start"])
        self.assertTrue(result["financial_coverage"]["contribution_margin_valid"])
        self.assertFalse(result["financial_coverage"]["net_margin_after_ads_valid"])
        self.assertIsNone(result["financial_summary"]["advertising_brl"])
        sku_a = next(row for row in result["inventory_rows"] if row["sku"] == "SKU-A")
        self.assertEqual(sku_a["full_stock"], 999)
        self.assertEqual(sku_a["open_purchase"], 5)
        self.assertGreater(sku_a["suggested_purchase"], 0, "Full nao pode compensar ruptura local")
        self.assertIsNotNone(sku_a["sell_through_pct"])
        self.assertIsNotNone(sku_a["inventory_turnover"])
        self.assertIsNotNone(sku_a["gmroi"])
        self.assertLessEqual(len(result["top_actions"]), 5)
        self.assertTrue(any(action["action_type"] == "price_review" for action in result["top_actions"]))

    def test_homonymous_store_blocks_stock_read_in_advanced_report(self):
        with tempfile.TemporaryDirectory() as root:
            tenant, _period = self._seed_operational_data(root)
            (tenant / "lojas_config.json").write_text(
                json.dumps([
                    {"store_id": "store-a", "nome": "Loja 1"},
                    {"store_id": "store-b", "nome": "loja 1"},
                ]),
                encoding="utf-8",
            )
            with (
                patch.object(
                    codex_reports_advanced,
                    "_load_latest_stock",
                    side_effect=AssertionError("estoque ambiguo nao deve ser lido"),
                ),
                patch.object(
                    codex_reports_advanced,
                    "_load_opening_stock",
                    side_effect=AssertionError("abertura ambigua nao deve ser lida"),
                ),
            ):
                result = codex_reports_advanced.build_profile_context(
                    info_base=root,
                    client_id="000001",
                    profile="weekly_sales_stock",
                    store="Loja 1",
                )

        self.assertTrue(any("lojas homônimas" in item for item in result["data_quality"]["warnings"]))
        estoque_source = next(
            item
            for item in result["data_quality"]["source_health"]
            if item["source"] == "Histórico de estoque"
        )
        self.assertEqual(estoque_source["status"], "unavailable")

    def test_below_95_percent_keeps_consolidated_margin_unavailable(self):
        with tempfile.TemporaryDirectory() as root:
            tenant, _period = self._seed_operational_data(root)
            rows = list(csv.DictReader((tenant / "cadastro_produtos.csv").open("r", encoding="utf-8-sig"), delimiter=";"))
            with (tenant / "cadastro_produtos.csv").open("w", encoding="utf-8-sig", newline="") as fh:
                writer = csv.DictWriter(fh, fieldnames=rows[0].keys(), delimiter=";")
                writer.writeheader()
                writer.writerow(rows[1])
            result = codex_reports_advanced.build_profile_context(
                info_base=root,
                client_id="000001",
                profile="daily_exceptions",
                margin_rows=[],
            )

        self.assertLess(result["financial_coverage"]["cost_by_revenue_pct"], 95)
        self.assertFalse(result["financial_coverage"]["contribution_margin_valid"])
        self.assertIsNone(result["financial_coverage"]["consolidated_profit_brl"])
        self.assertTrue(any(action["action_type"] == "data_quality" for action in result["top_actions"]))

    def test_confirmed_zero_cost_is_distinct_from_missing_cost(self):
        with tempfile.TemporaryDirectory() as root:
            tenant, _period = self._seed_operational_data(root)
            with (tenant / "cadastro_produtos.csv").open("w", encoding="utf-8-sig", newline="") as fh:
                writer = csv.DictWriter(fh, fieldnames=["sku", "nome", "custo", "preco", "imposto"], delimiter=";")
                writer.writeheader()
                writer.writerows(
                    [
                        {"sku": "SKU-A", "nome": "Roçadeira", "custo": "0", "preco": "100", "imposto": "10"},
                        {"sku": "SKU-B", "nome": "Caçamba", "custo": "0", "preco": "100", "imposto": "10"},
                    ]
                )
            result = codex_reports_advanced.build_profile_context(
                info_base=root,
                client_id="000001",
                profile="daily_exceptions",
                margin_rows=[],
            )

        self.assertEqual(result["financial_coverage"]["cost_by_revenue_pct"], 100)
        self.assertTrue(all(row["unit_cost"] == 0 for row in result["sales_rows"]))
        self.assertIsNone(codex_reports_advanced._optional_float(""))
        self.assertEqual(codex_reports_advanced._optional_float("0"), 0)

    def test_unavailable_sources_never_publish_false_zeroes(self):
        with tempfile.TemporaryDirectory() as root:
            tenant = self._tenant(root)
            (tenant / "lojas_config.json").write_text(
                json.dumps([{"store_id": "store-1", "nome": "Loja 1"}]),
                encoding="utf-8",
            )
            stock_db = sqlite3.connect(tenant / "estoque_historico.db")
            stock_db.execute(
                "CREATE TABLE estoque_historico (data_ref TEXT, recorded_at TEXT, loja_sync TEXT, sku TEXT, nome_bling TEXT, saldo_loja REAL, saldo_full REAL)"
            )
            stock_db.execute(
                "INSERT INTO estoque_historico VALUES (?,?,?,?,?,?,?)",
                (date.today().isoformat(), date.today().isoformat(), "Loja 1", "A", "Produto A", 0, 0),
            )
            stock_db.commit()
            stock_db.close()
            result = codex_reports_advanced.build_profile_context(
                info_base=root,
                client_id="000001",
                profile="daily_exceptions",
                margin_rows=[],
            )

        self.assertIsNone(result["financial_summary"]["gross_revenue_brl"])
        self.assertIsNone(result["financial_summary"]["returns_brl"])
        self.assertIsNone(result["financial_summary"]["net_revenue_brl"])
        self.assertIsNone(result["financial_coverage"]["cost_by_revenue_pct"])
        inventory = result["inventory_rows"][0]
        self.assertEqual(inventory["local_stock"], 0)
        self.assertIsNone(inventory["average_daily"])
        self.assertIsNone(inventory["suggested_purchase"])
        sales_source = next(item for item in result["data_quality"]["source_health"] if item["source"] == "Histórico de vendas")
        self.assertEqual(sales_source["status"], "unavailable")
        self.assertTrue(any("nao foram interpretados" in warning.lower() for warning in result["data_quality"]["warnings"]))

    def test_abc_xyz_rounding_and_mojibake_are_deterministic(self):
        rows = [
            {"store": "Loja", "sku": "A", "revenue": 800},
            {"store": "Loja", "sku": "B", "revenue": 150},
            {"store": "Loja", "sku": "C", "revenue": 50},
        ]
        codex_reports_advanced._assign_abc(rows, 0.8, 0.95)
        self.assertEqual([row["abc"] for row in rows], ["A", "B", "C"])
        self.assertEqual(codex_reports_advanced._round_order_quantity(17, 20, 6), 24)
        self.assertEqual(codex_reports_advanced.fix_mojibake("RoÃ§adeira"), "Roçadeira")
        self.assertEqual(codex_reports_advanced.fix_mojibake("CaÃ§amba"), "Caçamba")
        self.assertEqual(codex_reports_advanced.fix_mojibake("CotaÃ§Ã£o"), "Cotação")

    def test_import_allocations_close_and_scenarios_change_decisions(self):
        settings = codex_reports_advanced.normalize_report_settings({})
        order = {
            "id": "import-1",
            "nome_lista": "Container julho",
            "loja": "Loja 1",
            "supplier": "Fornecedor A",
            "currency": "USD",
            "incoterm": "FOB",
            "exchange_rate": 5,
            "itens": [
                {"SKU": "A", "Quantidade": 10, "Valor unidade": 10, "Valor total": 100, "Frete Internacional": 10, "M3": 1, "NCM": "8201", "II": 10, "IPI": 5, "PIS": 2, "COFINS": 8, "ICMS": 18},
                {"SKU": "B", "Quantidade": 20, "Valor unidade": 10, "Valor total": 200, "Frete Internacional": 20, "M3": 2, "NCM": "8202", "II": 10, "IPI": 5, "PIS": 2, "COFINS": 8, "ICMS": 18},
            ],
        }
        inventory = [
            {"sku": "A", "sale_price": 250, "coverage_days": 240},
            {"sku": "B", "sale_price": 250, "coverage_days": 200},
        ]
        result = codex_reports_advanced._landed_cost_analysis("", order, settings, inventory, 500_000)

        self.assertAlmostEqual(sum(row["landed_total_brl"] for row in result["items"]), result["cash_required_brl"], places=2)
        baseline = next(item for item in result["scenarios"] if item["kind"] == "baseline")
        exchange = next(item for item in result["scenarios"] if item["kind"] == "exchange")
        delay = next(item for item in result["scenarios"] if item["kind"] == "delay")
        self.assertGreater(exchange["cash_required_brl"], baseline["cash_required_brl"])
        self.assertGreater(exchange["minimum_sale_price_brl"], baseline["minimum_sale_price_brl"])
        self.assertLess(exchange["estimated_margin_pct"], baseline["estimated_margin_pct"])
        self.assertLess(delay["next_order_deadline"], baseline["next_order_deadline"])
        self.assertEqual(result["tax_band_pct"], 9.5)

    def test_output_profiles_limit_executive_html_and_keep_xlsx_details(self):
        rows = [{"sku": f"SKU-{index:05d}", "store": "Loja", "revenue": index} for index in range(120)]
        context = {
            "report_type": "weekly_sales_stock",
            "scope": {"period_start": "2026-07-01", "period_end": "2026-07-07", "comparison_start": "2026-06-24", "comparison_end": "2026-06-30"},
            "data_quality": {"confidence": "alta", "score": 90, "source_health": [], "warnings": []},
            "financial_summary": {"gross_revenue_brl": 1000},
            "financial_coverage": {"contribution_margin_valid": True},
            "top_actions": [],
            "sales_rows": rows,
            "inventory_rows": [],
            "store_summaries": [],
            "purchase_pipeline_rows": [],
        }
        html = assistant_html._assistant_build_advanced_report_html("Relatório Roçadeira", context)
        self.assertIn("Relatório Roçadeira", html)
        self.assertIn("SKU-00019", html)
        self.assertNotIn("SKU-00119", html)
        with tempfile.TemporaryDirectory() as root:
            path = Path(root, "report.xlsx")
            assistant_artifacts._assistant_write_advanced_xlsx(str(path), context)
            import openpyxl

            workbook = openpyxl.load_workbook(path, read_only=True)
            self.assertEqual(
                workbook.sheetnames,
                ["Resumo", "Margens_MLB", "Historico_Margens", "Ações", "Lojas", "Vendas_SKU", "Estoque", "Compras_Transito", "Importacao", "Fontes", "Fontes_ML", "Avisos"],
            )
            self.assertEqual(workbook["Vendas_SKU"].max_row, 121)
            workbook.close()

    def test_settings_adjustments_and_internal_queue_are_tenant_scoped_and_audited(self):
        with tempfile.TemporaryDirectory() as root:
            saved = codex_reports_advanced.report_settings_save(
                root,
                "tenant-a",
                {"global": {"target_margin_pct": 25}, "stores": {"Loja 1": {"monthly_sales_target_brl": 90000}}},
                "admin",
            )
            adjustment = codex_reports_advanced.codex_assistant_storage.codex_assistant_financial_adjustment_save(
                root,
                "tenant-a",
                {"kind": "advertising", "store": "Loja 1", "period_start": "2026-07-01", "period_end": "2026-07-07", "amount": 123.45, "note": "Importado"},
                created_by="admin",
            )
            action = codex_reports_advanced.create_queue_action(
                info_base=root,
                client_id="tenant-a",
                username="admin",
                payload={"report_id": "report-1", "action_type": "replenishment", "status": "queued", "skus": ["A"]},
            )
            updated = codex_reports_advanced.update_queue_action(
                info_base=root,
                client_id="tenant-a",
                action_id=action["action_id"],
                username="buyer",
                updates={"status": "in_progress", "note": "Cotando"},
            )
            other_settings = codex_reports_advanced.report_settings_get(root, "tenant-b")
            other_actions = codex_reports_advanced.queue_actions_list(root, "tenant-b")

        self.assertEqual(saved["updated_by"], "admin")
        self.assertEqual(adjustment["created_by"], "admin")
        self.assertEqual(updated["status"], "in_progress")
        self.assertEqual(updated["audit"][-1]["actor"], "buyer")
        self.assertEqual(other_settings["global"]["target_margin_pct"], 20)
        self.assertEqual(other_actions, [])

    def test_approved_replenishment_creates_only_internal_records(self):
        action = {
            "action_id": "action-1",
            "action_type": "replenishment",
            "title": "Repor A",
            "store": "Loja 1",
            "skus": ["A"],
            "items": [{"sku": "A", "quantity": 12}],
            "queueable": True,
        }
        with tempfile.TemporaryDirectory() as root:
            codex_assistant_storage.codex_assistant_report_save(
                root,
                "tenant-a",
                {"report_id": "report-1", "title": "Relatorio", "top_actions": [action]},
            )
            proposal = {
                "client_id": "tenant-a",
                "created_by": "admin",
                "action": {"id": "reports.queue_replenishment"},
                "params": {"report_id": "report-1", "report_action": action},
            }
            with (
                patch.object(assistant_runtime, "info_base", return_value=root),
                patch.object(codex_actions, "PASTA_INFO", root, create=True),
            ):
                result = codex_actions._execute_internal_report_queue("run-1", proposal)

            lists = json.loads(Path(root, "tenant-a", "listas_pedidos.json").read_text(encoding="utf-8"))
            queues = codex_reports_advanced.queue_actions_list(root, "tenant-a")

        self.assertFalse(result["external_mutation"])
        self.assertEqual(lists[0]["status"], "Lista gerada")
        self.assertEqual(lists[0]["itens"][0]["Quantidade"], 12)
        self.assertEqual(queues[0]["report_id"], "report-1")

    def test_stale_stock_alert_exposes_scope_age_capital_and_top_items(self):
        result = {
            "loja": "JK Pecas",
            "data_referencia": "2026-07-12",
            "itens": [
                {
                    "sku": "A-1",
                    "produto": "Produto sem historico de venda",
                    "saldo_loja": 10,
                    "saldo_full": 50,
                    "status": "nunca_vendeu",
                    "dias_sem_vender": None,
                    "ultima_venda": "",
                    "custo_cadastrado": True,
                    "custo_unitario": 12.5,
                    "valor_custo_estoque_loja": 125,
                },
                {
                    "sku": "B-2",
                    "produto": "Produto muito antigo",
                    "saldo_loja": 5,
                    "saldo_full": 0,
                    "status": "sem_venda_recente",
                    "dias_sem_vender": 220,
                    "ultima_venda": "2025-12-04",
                    "custo_cadastrado": False,
                },
                {
                    "sku": "C-3",
                    "produto": "Produto em observacao",
                    "saldo_loja": 2,
                    "saldo_full": 3,
                    "status": "sem_venda_recente",
                    "dias_sem_vender": 75,
                    "ultima_venda": "2026-04-28",
                    "custo_cadastrado": True,
                    "custo_unitario": 20,
                    "valor_custo_estoque_loja": 40,
                },
            ],
        }
        suggestions = assistant_management._assistant_suggestions_from_results(
            [
                {
                    "function": "get_days_without_sale_top",
                    "result": {
                        "itens": [
                            {"sku": "MENOR", "produto": "Amostra menor", "saldo_loja": 1, "status": "nunca_vendeu", "dias_sem_vender": None}
                        ]
                    },
                },
                {"function": "get_days_without_sale_top", "result": result},
            ],
            "proactive",
        )
        alert = next(item for item in suggestions if item["kind"] == "stale_stock")

        self.assertEqual(sum(1 for item in suggestions if item["kind"] == "stale_stock"), 1)
        self.assertEqual(alert["severity"], "warning")
        self.assertIn("3 SKU(s) e 17 unidade(s)", alert["detail"])
        self.assertIn("1 nunca venderam", alert["detail"])
        self.assertIn("1 com 180+ dias", alert["detail"])
        self.assertIn("1 com 30–89 dias", alert["detail"])
        self.assertIn("R$ 165,00 em 2/3 SKU(s)", alert["detail"])
        self.assertIn("`A-1`", alert["detail"])
        self.assertIn("Saldo: 10 un. • nunca vendeu", alert["detail"])
        self.assertIn("JK Pecas • posicao 2026-07-12", alert["detail"])
        self.assertIn("todos os SKUs envolvidos", alert["report_prompt"])
        self.assertIn("Nao trate custo ausente como zero", alert["report_prompt"])
        self.assertTrue(alert["recommendation"])

    def test_stale_stock_report_rows_are_prioritized_and_explain_each_action(self):
        context = {
            "tool_results": [
                {
                    "function": "get_days_without_sale_top",
                    "result": {
                        "loja": "JK Pecas",
                        "itens": [
                            {"sku": "A", "produto": "Nunca vendeu", "saldo_loja": 10, "saldo_full": 4, "status": "nunca_vendeu", "dias_sem_vender": None, "custo_cadastrado": True, "custo_unitario": 10},
                            {"sku": "B", "produto": "Muito antigo", "saldo_loja": 5, "saldo_full": 0, "status": "sem_venda_recente", "dias_sem_vender": 220, "ultima_venda": "2025-12-04", "custo_cadastrado": False},
                            {"sku": "C", "produto": "Atencao", "saldo_loja": 2, "saldo_full": 1, "status": "sem_venda_recente", "dias_sem_vender": 75, "ultima_venda": "2026-04-28", "custo_cadastrado": True, "custo_unitario": 20},
                            {"sku": "D", "produto": "Venda recente", "saldo_loja": 8, "saldo_full": 0, "status": "sem_venda_recente", "dias_sem_vender": 5, "ultima_venda": "2026-07-07"},
                        ],
                    },
                }
            ]
        }
        rows = assistant_stock._assistant_collect_stale_stock_rows(context, limit=20)

        self.assertEqual([row["sku"] for row in rows], ["A", "B", "C"])
        self.assertEqual(rows[0]["prioridade"], "1 - Imediata")
        self.assertEqual(rows[0]["situacao"], "Nunca vendeu")
        self.assertEqual(rows[0]["saldo_full"], "4")
        self.assertEqual(rows[0]["capital_custo"], "R$ 100,00")
        self.assertIn("bloquear nova compra", rows[0]["acao_recomendada"].lower())
        self.assertEqual(rows[1]["capital_custo"], "indisponivel")
        self.assertIn("180 dias", rows[1]["motivo"])
        self.assertEqual(rows[2]["prioridade"], "3 - Atencao")

    def test_stale_stock_tool_uses_average_registered_store_cost_without_inventing_missing_cost(self):
        products = pd.DataFrame(
            [
                {"sku_norm": "A", "nome_tool": "Produto A", "saldo_total": "10", "saldo_loja": "10", "saldo_full": "0", "custo": "", "preco": ""},
                {"sku_norm": "B", "nome_tool": "Produto B", "saldo_total": "5", "saldo_loja": "5", "saldo_full": "0", "custo": "", "preco": ""},
            ]
        )
        store_costs = pd.DataFrame(
            [
                {"loja_sync": "Loja 1", "sku": "A", "custo": "10", "preco": "20"},
                {"loja_sync": "Loja 2", "sku": "A", "custo": "14", "preco": "22"},
                {"loja_sync": "Loja 1", "sku": "B", "custo": "", "preco": ""},
            ]
        )
        with tempfile.TemporaryDirectory() as root:
            db_path = Path(root, "vendas_historico.db")
            connection = sqlite3.connect(db_path)
            connection.execute("CREATE TABLE vendas (sku TEXT, data TEXT, devolucao INTEGER, situacao TEXT)")
            connection.execute("INSERT INTO vendas VALUES ('B', '2025-12-24', 0, 'Concluida')")
            connection.commit()
            connection.close()
            with (
                patch.object(sales_inventory, "_ia_carregar_produtos_tool_df", return_value=products),
                patch.object(sales_inventory, "_listar_bancos_vendas_tenant", return_value=[str(db_path)]),
                patch.object(sales_inventory, "_sql_filtro_loja_vendas", return_value=("", [])),
                patch.object(sales_inventory, "_ia_obter_data_referencia_vendas", return_value=date(2026, 7, 12)),
                patch.object(sales_inventory, "_cadastro_ler_custos_lojas", return_value=store_costs),
            ):
                payload = ia_tools_vendas._ia_tool_get_days_without_sale_top(
                    "tenant-a",
                    None,
                    500,
                    True,
                    False,
                    True,
                )

        result = payload["result"]
        item_a = next(item for item in result["itens"] if item["sku"] == "A")
        item_b = next(item for item in result["itens"] if item["sku"] == "B")
        self.assertTrue(item_a["custo_cadastrado"])
        self.assertEqual(item_a["custo_unitario"], 12)
        self.assertEqual(item_a["valor_custo_estoque_loja"], 120)
        self.assertEqual(item_a["custo_origem"], "media do cadastro por loja")
        self.assertFalse(item_b["custo_cadastrado"])
        self.assertIsNone(item_b["valor_custo_estoque_loja"])
        self.assertEqual(result["resumo_estoque_parado"]["custos_cobertos"], 1)
        self.assertEqual(result["resumo_estoque_parado"]["capital_custo_conhecido"], 120)

    def test_stale_stock_store_filter_fails_closed_for_duplicate_display_name(self):
        maps = {
            "products": {"A": "Produto A"},
            "costs": {},
            "prices": {},
            "known_costs": set(),
            "known_prices": set(),
            "cost_sources": {},
            "price_sources": {},
        }
        store_costs = pd.DataFrame(
            [
                {"store_id": "store-a", "loja_sync": "Mesmo nome", "sku": "A", "custo": "10"},
                {"store_id": "store-b", "loja_sync": "Mesmo nome", "sku": "A", "custo": "99"},
            ]
        )

        with patch.object(
            sales_inventory,
            "_cadastro_ler_custos_lojas",
            return_value=store_costs,
        ):
            sales_inventory._merge_store_costs("tenant-a", "Mesmo nome", maps)

        self.assertNotIn("A", maps["costs"])

    def test_stale_stock_store_filter_prefers_exact_row_over_legacy_shadow(self):
        maps = {
            "products": {"A": "Produto A"},
            "costs": {},
            "prices": {},
            "known_costs": set(),
            "known_prices": set(),
            "cost_sources": {},
            "price_sources": {},
        }
        store_costs = pd.DataFrame(
            [
                {"store_id": "store-a", "loja_sync": "Loja A", "sku": "A", "custo": "10"},
                {"store_id": "", "loja_sync": "Loja A", "sku": "A", "custo": "999"},
            ]
        )

        with patch.object(
            sales_inventory,
            "_cadastro_ler_custos_lojas",
            return_value=store_costs,
        ):
            sales_inventory._merge_store_costs("tenant-a", "Loja A", maps)

        self.assertEqual(maps["costs"]["A"], 10)

    def test_marketplace_margin_keeps_same_sku_separate_by_store_mlb_and_variation(self):
        base = {
            "financial_summary": {"gross_revenue_brl": 1000, "returns_brl": 0, "advertising_brl": 0},
            "financial_coverage": {"minimum_required_pct": 95},
            "store_summaries": [
                {"store": "Loja A", "gross_revenue_brl": 500, "advertising_brl": 0},
                {"store": "Loja B", "gross_revenue_brl": 500, "advertising_brl": 0},
            ],
            "sales_rows": [
                {"store": "Loja A", "sku": "SKU-1", "unit_cost": 40, "tax_pct": 10},
                {"store": "Loja B", "sku": "SKU-1", "unit_cost": 55, "tax_pct": 8},
            ],
            "inventory_rows": [],
            "data_quality": {"warnings": [], "source_health": []},
        }
        snapshot = {
            "status": "ok",
            "collected_at": "2026-07-20T12:00:00-03:00",
            "listing_rows": [
                {"store": "Loja A", "item_id": "MLB1", "variation_id": "V1", "sku": "SKU-1", "price": 100, "sale_fee_amount": 12, "shipping_seller_cost": 8},
                {"store": "Loja A", "item_id": "MLB2", "variation_id": "V2", "sku": "SKU-1", "price": 120, "sale_fee_amount": 14, "shipping_seller_cost": 9},
                {"store": "Loja B", "item_id": "MLB1", "variation_id": "V9", "sku": "SKU-1", "price": 130, "sale_fee_amount": 16, "shipping_seller_cost": 10},
            ],
            "ledger_rows": [],
        }

        result = codex_reports_advanced.apply_marketplace_commercial(base, snapshot)

        keys = {(row["store"], row["mlb"], row["variation_id"]) for row in result["listing_margin_rows"]}
        self.assertEqual(keys, {("Loja A", "MLB1", "V1"), ("Loja A", "MLB2", "V2"), ("Loja B", "MLB1", "V9")})
        self.assertEqual(result["listing_margin_rows"][0]["unit_cost_brl"], 40)
        self.assertEqual(result["listing_margin_rows"][2]["unit_cost_brl"], 55)

    def test_marketplace_free_shipping_without_confirmed_cost_remains_incomplete(self):
        base = {
            "financial_summary": {"gross_revenue_brl": 100, "returns_brl": 0},
            "financial_coverage": {"minimum_required_pct": 95},
            "store_summaries": [{"store": "Loja A", "gross_revenue_brl": 100, "advertising_brl": None}],
            "sales_rows": [{"store": "Loja A", "sku": "A", "unit_cost": 20, "tax_pct": 10}],
            "inventory_rows": [],
            "data_quality": {"warnings": [], "source_health": []},
        }
        result = codex_reports_advanced.apply_marketplace_commercial(
            base,
            {"status": "partial", "listing_rows": [{"store": "Loja A", "item_id": "MLB1", "sku": "A", "price": 100, "sale_fee_amount": 12, "free_shipping": True}]},
        )
        row = result["listing_margin_rows"][0]
        self.assertEqual(row["margin_status"], "unavailable")
        self.assertIn("frete", row["missing_components"])
        self.assertIsNone(row["unit_contribution_brl"])

    def test_marketplace_financial_threshold_is_exactly_95_percent(self):
        def build(gross_covered):
            base = {
                "financial_summary": {"gross_revenue_brl": 1000, "returns_brl": 0, "advertising_brl": 0},
                "financial_coverage": {"minimum_required_pct": 95},
                "store_summaries": [{"store": "Loja A", "gross_revenue_brl": 1000, "advertising_brl": 0}],
                "sales_rows": [],
                "inventory_rows": [],
                "data_quality": {"warnings": [], "source_health": []},
            }
            ledger = {
                "store": "Loja A", "order_id": str(gross_covered), "item_id": "MLB1", "sku": "A",
                "quantity": gross_covered / 100, "unit_price": 100, "gross_amount": gross_covered,
                "unit_cost": 20, "tax_pct": 10, "sale_fee_amount": 12,
                "seller_shipping_cost": 0, "pack_item_count": 1,
            }
            return codex_reports_advanced.apply_marketplace_commercial(base, {"status": "ok", "listing_rows": [], "ledger_rows": [ledger]})

        below = build(949.9)
        accepted = build(950)

        self.assertEqual(below["financial_coverage"]["complete_margin_by_revenue_pct"], 94.99)
        self.assertFalse(below["financial_coverage"]["contribution_margin_valid"])
        self.assertIsNone(below["financial_summary"]["contribution_profit_brl"])
        self.assertEqual(accepted["financial_coverage"]["complete_margin_by_revenue_pct"], 95.0)
        self.assertTrue(accepted["financial_coverage"]["contribution_margin_valid"])
        self.assertIsNotNone(accepted["financial_summary"]["contribution_profit_brl"])

    def test_multi_item_pack_shipping_is_never_silently_allocated(self):
        base = {
            "financial_summary": {"gross_revenue_brl": 200, "returns_brl": 0, "advertising_brl": 0},
            "financial_coverage": {"minimum_required_pct": 95},
            "store_summaries": [{"store": "Loja A", "gross_revenue_brl": 200, "advertising_brl": 0}],
            "sales_rows": [], "inventory_rows": [], "data_quality": {"warnings": [], "source_health": []},
        }
        ledger = {
            "store": "Loja A", "order_id": "1", "pack_id": "P1", "item_id": "MLB1", "sku": "A",
            "quantity": 1, "unit_price": 200, "gross_amount": 200, "unit_cost": 50, "tax_pct": 10,
            "sale_fee_amount": 20, "seller_shipping_cost": 15, "pack_item_count": 2,
        }
        result = codex_reports_advanced.apply_marketplace_commercial(base, {"status": "ok", "ledger_rows": [ledger]})
        row = result["historical_margin_ledger"][0]
        self.assertEqual(row["shipping_scope"], "pack_unallocated")
        self.assertEqual(row["margin_status"], "unavailable")
        self.assertIsNone(row["contribution_total_brl"])


if __name__ == "__main__":
    unittest.main()
