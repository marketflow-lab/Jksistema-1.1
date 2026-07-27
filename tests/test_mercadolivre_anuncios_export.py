from __future__ import annotations

import io
import logging

import pytest
from fastapi import HTTPException
from openpyxl import load_workbook

from backend.routers import mercado_livre as mercado_livre_router
from backend.services import mercadolivre_anuncios as anuncios_service


class _Response:
    def __init__(self, status_code: int, payload):
        self.status_code = status_code
        self._payload = payload
        self.text = str(payload)

    def json(self):
        return self._payload


def _raw_item(index: int, *, title: str | None = None, sku: str | None = None, variations=None):
    return {
        "id": f"MLB{index:09d}",
        "title": title or f"Produto {index}",
        "seller_sku": sku or f"SKU-{index}",
        "status": "active",
        "price": 100.5 + index,
        "available_quantity": index,
        "sold_quantity": index * 2,
        "listing_type_id": "gold_special",
        "condition": "new",
        "category_id": "MLB1234",
        "domain_id": "MLB-AUTO_PARTS",
        "date_created": "2026-07-01T10:30:00.000-03:00",
        "last_updated": "2026-07-27T11:45:00Z",
        "start_time": "2026-07-01T10:30:00Z",
        "stop_time": "2026-12-01T10:30:00Z",
        "variations": variations or [],
    }


def _normalized_item(raw: dict):
    variations = [
        {
            "id": variation.get("id"),
            "title": variation.get("title"),
            "sku": variation.get("sku"),
            "parent_sku": raw.get("seller_sku"),
            "inventory_id": variation.get("inventory_id"),
            "price": variation.get("price"),
            "available_quantity": variation.get("available_quantity", 0),
            "sold_quantity": variation.get("sold_quantity", 0),
            "picture_id": variation.get("picture_id"),
        }
        for variation in raw.get("variations") or []
    ]
    return {
        "id": raw["id"],
        "title": raw["title"],
        "thumbnail": f"https://img.example/{raw['id']}.jpg",
        "family_name": "Familia teste",
        "user_product_id": f"UP-{raw['id']}",
        "catalog_listing": True,
        "catalog_product_id": f"CAT-{raw['id']}",
        "listing_type_id": raw["listing_type_id"],
        "listing_type_name": "Classico",
        "item_condition": "Novo",
        "price": raw["price"],
        "standard_price": raw["price"] + 10,
        "original_price": raw["price"] + 10,
        "discount_pct": 5.5,
        "shipping_cost": 12.25,
        "shipping_buyer_cost": 0.0,
        "shipping_list_cost": 20.0,
        "shipping_base_cost": 18.0,
        "ad_cost": 15.75,
        "fixed_fee_amount": 6.5,
        "listing_fee_amount": 0.0,
        "sale_fee_pct": 12.0,
        "free_shipping": True,
        "logistic_type": "fulfillment",
        "shipping_mode": "me2",
        "available_quantity": raw["available_quantity"],
        "sold_quantity": raw["sold_quantity"],
        "status": raw["status"],
        "sku": raw["seller_sku"],
        "has_promotion": True,
        "promotion_id": "PROMO-1",
        "promotion_type": "DEAL",
        "channels": ["marketplace"],
        "permalink": f"https://produto.mercadolivre.com.br/{raw['id']}",
        "variations": variations,
    }


class _Context:
    logger = logging.getLogger("test-mercadolivre-export")

    def __init__(self, responses: list[_Response], items: dict[str, dict]):
        self.responses = list(responses)
        self.items = items
        self.calls = []
        self.detail_calls = []

    def obter_cfg_ml(self, client_id, loja):
        self.scope = (client_id, loja)
        return {"user_id": "SELLER-123", "site_id": "MLB"}

    def ml_api_request_com_retry(self, client_id, loja, cfg, method, url, **kwargs):
        self.calls.append({"client_id": client_id, "loja": loja, "url": url, **kwargs})
        assert self.responses, "chamada inesperada ao provedor"
        return self.responses.pop(0), cfg

    def ml_buscar_itens_batch(self, client_id, loja, cfg, item_ids):
        self.detail_calls.append((client_id, loja, list(item_ids)))
        return [self.items[item_id] for item_id in item_ids if item_id in self.items], cfg

    def ml_montar_detalhe_anuncio_listagem(self, client_id, loja, cfg, item, sku_filter):
        return _normalized_item(item)

    @staticmethod
    def ml_parse_error_detail(resp, fallback):
        return (resp.json() or {}).get("message") or fallback


def _patch_context(monkeypatch, context: _Context):
    monkeypatch.setattr(anuncios_service, "_ctx", lambda: context)


def _ids(first: int, last: int):
    return [f"MLB{index:09d}" for index in range(first, last + 1)]


def test_exporta_todos_os_anuncios_em_scan_com_abas_campos_e_tipos(monkeypatch):
    item_ids = _ids(1, 45)
    variation = {
        "id": 9001,
        "title": "@Cor: Azul",
        "sku": "\tVAR-1",
        "inventory_id": "INV-1",
        "price": 123.45,
        "available_quantity": 7,
        "sold_quantity": 9,
        "picture_id": "PIC-1",
    }
    items = {item_id: _raw_item(index) for index, item_id in enumerate(item_ids, start=1)}
    items[item_ids[0]] = _raw_item(1, title="=1+1", sku="+SKU-1", variations=[variation])
    context = _Context(
        [
            _Response(200, {"paging": {"total": 45}, "results": item_ids[:25], "scroll_id": "scroll-1"}),
            _Response(200, {"paging": {"total": 45}, "results": item_ids[25:]}),
        ],
        items,
    )
    montar_detalhe_original = context.ml_montar_detalhe_anuncio_listagem

    def montar_detalhe_com_texto_em_campo_numerico(client_id, loja, cfg, item, sku_filter):
        detalhe = montar_detalhe_original(client_id, loja, cfg, item, sku_filter)
        if item["id"] == item_ids[0]:
            detalhe["price"] = "=2+2"
            detalhe["available_quantity"] = "+1"
        return detalhe

    context.ml_montar_detalhe_anuncio_listagem = montar_detalhe_com_texto_em_campo_numerico
    _patch_context(monkeypatch, context)

    output, filename = anuncios_service.exportar_anuncios_ativos_mercado_livre("000002", "JK Pecas")
    workbook = load_workbook(output, data_only=False)

    assert context.scope == ("000002", "JK Pecas")
    assert context.detail_calls == [("000002", "JK Pecas", item_ids)]
    assert context.calls[0]["params"] == {"search_type": "scan", "limit": 100, "status": "active"}
    assert context.calls[1]["params"] == {"search_type": "scan", "limit": 100, "scroll_id": "scroll-1"}
    assert workbook.sheetnames == ["Anuncios", "Variacoes"]
    assert workbook["Anuncios"].max_row == 46
    assert workbook["Variacoes"].max_row == 2
    assert filename.startswith("anuncios-ativos-JK-Pecas-") and filename.endswith(".xlsx")

    anuncios = workbook["Anuncios"]
    assert anuncios.freeze_panes == "A2"
    assert anuncios.auto_filter.ref == "A1:AO46"
    header_index = {cell.value: cell.column for cell in anuncios[1]}
    assert "Vendidas acumuladas" in header_index
    assert "Link" in header_index
    assert anuncios.cell(2, header_index["Titulo"]).value == "'=1+1"
    assert anuncios.cell(2, header_index["SKU"]).value == "'+SKU-1"
    assert anuncios.cell(2, header_index["Preco atual"]).value == "'=2+2"
    assert anuncios.cell(2, header_index["Preco atual"]).data_type != "f"
    assert anuncios.cell(2, header_index["Estoque disponivel"]).value == "'+1"
    assert anuncios.cell(2, header_index["Estoque disponivel"]).data_type != "f"
    assert anuncios.cell(3, header_index["Preco atual"]).data_type == "n"
    assert anuncios.cell(3, header_index["Estoque disponivel"]).data_type == "n"
    assert anuncios.cell(2, header_index["Data de criacao"]).is_date

    variacoes = workbook["Variacoes"]
    variation_header = {cell.value: cell.column for cell in variacoes[1]}
    assert variacoes.cell(2, variation_header["Variacao"]).value == "'@Cor: Azul"
    assert variacoes.cell(2, variation_header["SKU da variacao"]).value == "'\tVAR-1"
    assert variacoes.cell(2, variation_header["Preco"]).data_type == "n"
    assert all(
        cell.data_type != "f"
        for worksheet in workbook.worksheets
        for row in worksheet.iter_rows(min_row=2)
        for cell in row
    )


def test_scan_deduplica_sem_aceitar_cobertura_menor_que_total(monkeypatch):
    first_page = _ids(1, 20)
    second_page = _ids(20, 30)
    unique_ids = _ids(1, 30)
    items = {item_id: _raw_item(index) for index, item_id in enumerate(unique_ids, start=1)}
    context = _Context(
        [
            _Response(200, {"paging": {"total": 30}, "results": first_page, "scroll_id": "scroll-1"}),
            _Response(200, {"paging": {"total": 30}, "results": second_page, "scroll_id": "scroll-2"}),
        ],
        items,
    )
    _patch_context(monkeypatch, context)

    output, _filename = anuncios_service.exportar_anuncios_ativos_mercado_livre("tenant-a", "Loja A")
    workbook = load_workbook(output)

    assert workbook["Anuncios"].max_row == 31
    assert context.detail_calls[0][2] == unique_ids


def test_zero_ativos_gera_workbook_valido_com_cabecalhos(monkeypatch):
    context = _Context([_Response(200, {"paging": {"total": 0}, "results": []})], {})
    _patch_context(monkeypatch, context)

    output, _filename = anuncios_service.exportar_anuncios_ativos_mercado_livre("tenant-a", "Loja Vazia")
    workbook = load_workbook(output)

    assert workbook.sheetnames == ["Anuncios", "Variacoes"]
    assert workbook["Anuncios"].max_row == 1
    assert workbook["Variacoes"].max_row == 1
    assert context.detail_calls == []


def test_fallback_offset_mantem_cobertura_e_enriquece_em_lotes_limitados(monkeypatch):
    item_ids = _ids(1, 145)
    items = {item_id: _raw_item(index) for index, item_id in enumerate(item_ids, start=1)}
    context = _Context(
        [
            _Response(400, {"message": "scan unavailable"}),
            _Response(200, {"paging": {"total": 145}, "results": item_ids[:100]}),
            _Response(200, {"paging": {"total": 145}, "results": item_ids[100:]}),
        ],
        items,
    )
    _patch_context(monkeypatch, context)

    output, _filename = anuncios_service.exportar_anuncios_ativos_mercado_livre("tenant-b", "Loja B")
    workbook = load_workbook(output)

    assert workbook["Anuncios"].max_row == 146
    assert [call["params"] for call in context.calls] == [
        {"search_type": "scan", "limit": 100, "status": "active"},
        {"offset": 0, "limit": 100, "status": "active"},
        {"offset": 100, "limit": 100, "status": "active"},
    ]
    assert all(call["client_id"] == "tenant-b" and call["loja"] == "Loja B" for call in context.calls)
    assert [len(call[2]) for call in context.detail_calls] == [100, 45]


def test_falha_intermediaria_nao_entrega_arquivo_parcial(monkeypatch):
    first_page = _ids(1, 20)
    context = _Context(
        [
            _Response(200, {"paging": {"total": 30}, "results": first_page, "scroll_id": "scroll-1"}),
            _Response(503, {"message": "provider unavailable"}),
        ],
        {item_id: _raw_item(index) for index, item_id in enumerate(first_page, start=1)},
    )
    _patch_context(monkeypatch, context)

    with pytest.raises(HTTPException) as exc_info:
        anuncios_service.exportar_anuncios_ativos_mercado_livre("tenant-a", "Loja A")

    assert exc_info.value.status_code == 502
    assert "provider unavailable" in str(exc_info.value.detail)
    assert context.detail_calls == []


def test_falha_quando_detalhes_nao_cobrem_todos_os_ids(monkeypatch):
    item_ids = _ids(1, 2)
    context = _Context(
        [_Response(200, {"paging": {"total": 2}, "results": item_ids})],
        {item_ids[0]: _raw_item(1)},
    )
    _patch_context(monkeypatch, context)

    with pytest.raises(HTTPException) as exc_info:
        anuncios_service.exportar_anuncios_ativos_mercado_livre("tenant-a", "Loja A")

    assert exc_info.value.status_code == 502
    assert "faltaram detalhes" in str(exc_info.value.detail)


def test_bloqueia_exportacoes_concorrentes_da_mesma_loja(monkeypatch):
    context = _Context([], {})
    _patch_context(monkeypatch, context)
    chave = ("tenant-a", "loja a")
    with anuncios_service._ML_EXPORT_ACTIVE_LOCK:
        anuncios_service._ML_EXPORT_ACTIVE_KEYS.add(chave)
    try:
        with pytest.raises(HTTPException) as exc_info:
            anuncios_service.exportar_anuncios_ativos_mercado_livre("tenant-a", "Loja A")
    finally:
        with anuncios_service._ML_EXPORT_ACTIVE_LOCK:
            anuncios_service._ML_EXPORT_ACTIVE_KEYS.discard(chave)

    assert exc_info.value.status_code == 409
    assert "em andamento" in str(exc_info.value.detail)
    assert not hasattr(context, "scope")


def test_rota_estatica_precede_item_dinamico_e_define_headers(monkeypatch):
    monkeypatch.setattr(
        mercado_livre_router,
        "exportar_anuncios_ativos_mercado_livre",
        lambda client_id, loja: (io.BytesIO(b"xlsx"), "anuncios-ativos-Loja-A-20260727-120000.xlsx"),
    )
    router = mercado_livre_router.create_mercado_livre_router(
        mercado_livre_router.MercadoLivreRouterConfig(get_tenant_id=lambda: "tenant-a")
    )
    paths = [route.path for route in router.routes]
    export_route = next(route for route in router.routes if route.path == "/api/mercadolivre/anuncios/exportar")

    response = export_route.endpoint(loja="Loja A", client_id="tenant-a")

    assert paths.index("/api/mercadolivre/anuncios/exportar") < paths.index("/api/mercadolivre/anuncios/{item_id}")
    assert response.media_type == anuncios_service.ML_EXPORT_XLSX_MEDIA_TYPE
    assert response.headers["content-disposition"] == 'attachment; filename="anuncios-ativos-Loja-A-20260727-120000.xlsx"'
    assert response.headers["cache-control"] == "no-store"
