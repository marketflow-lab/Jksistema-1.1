from __future__ import annotations

import csv
import io
import json
import logging
import sqlite3
from datetime import datetime
from types import SimpleNamespace

import pytest

from backend.modules.vendas import analytics, general_report, performance, reports
from backend.modules.vendas.compositions import CompositionCatalog
from backend.services import bling_vendas, cadastro_lojas_produtos


STORES = [{"store_id": "store-a", "nome": "Loja A"}, {"store_id": "store-b", "nome": "Loja B"}]
SCHEMA = """
CREATE TABLE vendas (
    id_unico TEXT, data TEXT, loja_conta TEXT, canal TEXT, numero TEXT,
    situacao TEXT, devolucao INTEGER, sku TEXT, produto TEXT,
    quantidade REAL, valor REAL, comprador TEXT, unidade_negocio TEXT,
    nota_fiscal_id TEXT, store_id TEXT
);
CREATE TABLE notas_entrada_itens (
    id_unico TEXT, data_emissao TEXT, loja_conta TEXT, numero_nota TEXT,
    origem_codigo TEXT, sku TEXT, descricao TEXT, quantidade REAL,
    valor_total REAL, unidade_negocio_virtual TEXT, unidade_negocio TEXT,
    natureza_operacao TEXT, finalidade_operacao TEXT, devolucao INTEGER, store_id TEXT
);
"""


def product(sku, components=None, store_id="store-a"):
    row = {"sku": sku, "store_id": store_id, "produto": "Produto " + sku}
    if components is not None:
        row.update(formato_bling="E", componentes_bling_json=json.dumps([
            {"produto": {"codigo": child}, "quantidade": factor}
            for child, factor in components
        ]))
    return row


def sale(uid, sku, quantity, value, *, day="2026-01-10", store="Loja A", store_id="store-a", unit="Loja", state="Faturado", invoice=""):
    return (uid, day, store, "Bling", uid, state, 0, sku, "Produto " + sku,
            quantity, value, "Comprador", unit, invoice, store_id)


def returned(uid, sku, quantity, value, *, day="2026-02-10", store="Loja A", store_id="store-a", unit="Loja"):
    return (uid, day, store, uid, uid, sku, "Produto " + sku, quantity, value,
            unit, unit, "Devolucao", "", 1, store_id)


@pytest.fixture
def sources(tmp_path, monkeypatch):
    performance.invalidate_vendas_cache()
    db_path = tmp_path / "sales.db"
    with sqlite3.connect(db_path) as db:
        db.executescript(SCHEMA)
    catalog = CompositionCatalog([
        product("479"), product("B"), product("kit2", [("479", 2)]),
        product("KIT3", [("479", 3)]), product("KIT5", [("479", 5)]),
        product("MIX", [("kit2", 3), ("B", 2)]),
        product("479", store_id="store-b"),
        product("kit2", [("479", 10)], store_id="store-b"),
    ], stores=STORES)
    for module in (analytics, reports, general_report):
        monkeypatch.setattr(module, "_listar_bancos_vendas_tenant", lambda *_: [str(db_path)])
        monkeypatch.setattr(module, "load_composition_catalog", lambda _: catalog)
    for module in (analytics, general_report, performance, cadastro_lojas_produtos):
        monkeypatch.setattr(module, "get_tenant_path", lambda _: str(tmp_path))
    monkeypatch.setattr(reports, "_deduplicar_vendas_consolidadas", bling_vendas._deduplicar_vendas_consolidadas)
    monkeypatch.setattr(reports, "_deve_excluir_venda_ebazar", lambda *_: False)
    monkeypatch.setattr(analytics, "_deve_excluir_venda_ebazar", lambda *_: False)
    monkeypatch.setattr(analytics, "logger", logging.getLogger("test-composition-reports"))
    monkeypatch.setattr(analytics, "_filtrar_vendas_lojas_ativas", lambda rows, *_: rows)
    monkeypatch.setattr(analytics, "_sql_filtro_loja_vendas", lambda loja: (" AND loja_conta = ?", [loja]))
    monkeypatch.setattr(analytics, "_sql_filtro_loja_notas_entrada", lambda _cursor, loja:
                        (" AND loja_conta = ?", [loja]) if loja and loja != "__todas" else ("", []))
    monkeypatch.setattr(analytics, "_carregar_mapeamento_lojas_virtuais_cliente", lambda _: {})
    monkeypatch.setattr(analytics, "_sql_filtro_unidade_com_mapa", lambda unit, _map: (" AND unidade_negocio = ?", [unit]))
    monkeypatch.setattr(analytics, "_sql_filtro_unidade_devolucao", lambda unit:
                        (" AND unidade_negocio = ?", [unit]) if unit and unit != "__todos" else ("", []))
    monkeypatch.setattr(analytics, "_vendas_series_estoque_historico", lambda **_: {"estoque_meta": {}})
    monkeypatch.setattr(general_report, "_read_current_local_stock", lambda *_: {
        "items": {"479": {"sku": "479", "produto": "Produto 479", "saldo_loja": 59}},
        "available": True, "source": "fixture", "snapshot_at": "2026-03-31",
        "event_id": "fixture", "total_skus_snapshot": 1, "errors": [],
    })

    def insert(sales=(), returns=()):
        with sqlite3.connect(db_path) as db:
            db.executemany("INSERT INTO vendas VALUES (?,?,?,?,?,?,?,?,?,?,?,?,?,?,?)", sales)
            db.executemany("INSERT INTO notas_entrada_itens VALUES (?,?,?,?,?,?,?,?,?,?,?,?,?,?,?)", returns)

    yield SimpleNamespace(path=db_path, root=tmp_path, catalog=catalog, insert=insert)
    performance.invalidate_vendas_cache()


def graph(**kwargs):
    return analytics.grafico_vendas.__wrapped__(
        client_id="fixture", data_inicio="2026-01-01", data_fim="2026-03-31",
        intervalo="mes", **kwargs,
    )


def general():
    return general_report.generate_general_sku_report.__wrapped__(
        client_id="fixture", loja="Loja A", periodo="personalizado",
        data_inicio="2026-01-01", data_fim="2026-03-31",
    )


def pareto():
    return reports.generate_pareto_report.__wrapped__(
        client_id="fixture", loja="Loja A", periodo="personalizado",
        data_inicio="2026-01-01", data_fim="2026-03-31",
    )


def example(sources):
    direct = [sale("d" + str(i), "479", 2, 20) for i in range(10)]
    kits = [sale("k2", "kit2", 9, 180, day="2026-02-12", unit="Full"),
            sale("k3", "KIT3", 2, 60, day="2026-03-12"),
            sale("k5", "KIT5", 3, 150, day="2026-03-12")]
    sources.insert(direct + kits, [returned("r-kit", "kit2", 2, 40), returned("r-direct", "479", 1, 10)])
    return direct, kits


def test_graph_59_units_commercial_target_and_returns_without_extra_money(sources, monkeypatch):
    direct, kits = example(sources)
    # Source mirrors must be consolidated before projecting both sales and returns.
    monkeypatch.setattr(analytics, "_listar_bancos_vendas_tenant", lambda *_: [str(sources.path), str(sources.path)])
    all_sales = graph(loja="Loja A")
    assert all_sales["quantidades_vendas"] == [20, 18, 21]
    assert sum(all_sales["valores_vendas"]) == 590
    assert sum(all_sales["quantidades_devolucoes"]) == 5
    assert sum(all_sales["valores_devolucoes"]) == 50
    component = graph(loja="Loja A", sku="479")
    assert sum(component["quantidades_vendas"]) == 59
    assert sum(component["valores_vendas"]) == 200
    assert sum(component["quantidades_devolucoes"]) == 5
    assert sum(component["valores_devolucoes"]) == 10
    kit = graph(loja="Loja A", sku="kit2")
    assert sum(kit["quantidades_vendas"]) == 9
    assert sum(kit["valores_vendas"]) == 180
    assert sum(kit["quantidades_devolucoes"]) == 2
    assert sum(kit["valores_devolucoes"]) == 40


def test_general_monthly_coverage_records_and_stock_use_component_units(sources):
    example(sources)
    result = general()
    rows = {row["sku"].upper(): row for row in result["skus"]}
    assert result["resumo"]["unidades_vendidas"] == 59
    assert result["resumo"]["faturamento"] == 590
    assert result["resumo"]["quantidade_devolvida"] == 5
    assert result["resumo"]["valor_devolvido"] == 50
    assert rows["479"]["vendas_mensais"] == {"2026-01": 20, "2026-02": 18, "2026-03": 21}
    assert rows["479"]["media_mensal"] == pytest.approx(59 / 3)
    assert rows["479"]["estoque_loja"] == 59
    assert rows["479"]["cobertura_meses"] == 3
    assert rows["479"]["registros_vendas"] == 10
    assert sum(row["registros_vendas"] for row in rows.values()) == 13
    assert rows["KIT2"]["unidades_vendidas"] == 0
    assert rows["KIT2"]["quantidade_comercial"] == 9
    assert rows["KIT2"]["faturamento"] == 180
    assert sum(row["ocorrencias_devolucao"] for row in rows.values()) == 2
    assert result["auditoria"]["devolucoes_lidas"] == 2
    assert result["auditoria"]["vendas"]["linhas_elegiveis"] == 13


def test_pareto_preserves_financial_ranking_and_commercial_counts(sources, monkeypatch):
    example(sources)
    monkeypatch.setattr(reports, "load_composition_catalog", lambda _: CompositionCatalog())
    baseline = pareto()
    monkeypatch.setattr(reports, "load_composition_catalog", lambda _: sources.catalog)
    result = pareto()
    assert result["resumo"] == baseline["resumo"]
    assert result["canais"] == baseline["canais"]
    assert result["mensal"] == baseline["mensal"]
    assert result["analise_executiva"] == baseline["analise_executiva"]
    assert [(r["sku"], r["rank"], r["faturamento"], r["registros"]) for r in result["todos_skus"]] == [
        (r["sku"], r["rank"], r["faturamento"], r["registros"]) for r in baseline["todos_skus"]]
    assert sum(r["quantidade"] for r in result["todos_skus"]) == 59
    assert result["todos_skus"][0]["quantidade"] == 59


def test_nested_mixed_components_keep_revenue_and_no_synthetic_orders(sources):
    sources.insert([sale("nested", "MIX", 2, 100), sale("negative", "NEG", 1, -2)],
                   [returned("nested-return", "MIX", 1, 50)])
    result = general()
    rows = {row["sku"]: row for row in result["skus"]}
    assert rows["479"]["unidades_vendidas"] == 12
    assert rows["B"]["unidades_vendidas"] == 4
    assert rows["MIX"]["unidades_vendidas"] == 0
    assert rows["MIX"]["quantidade_comercial"] == 2
    assert rows["479"]["faturamento"] == rows["B"]["faturamento"] == 0
    assert rows["479"]["registros_vendas"] == rows["B"]["registros_vendas"] == 0
    assert rows["479"]["quantidade_devolvida"] == 6
    assert rows["B"]["quantidade_devolvida"] == 2
    assert result["resumo"]["faturamento"] == 98
    result_pareto = pareto()
    assert [(row["sku"], row["rank"]) for row in result_pareto["todos_skus"]] == [
        ("MIX", 1), ("NEG", 2), ("479", None), ("B", None)]
    assert result_pareto["resumo"]["total_skus"] == 2
    assert result_pareto["resumo"]["componentes_sem_venda_direta"] == 2
    assert sum(row["registros"] for row in result_pareto["todos_skus"]) == 2
    assert sum(graph(sku="479")["quantidades_vendas"]) == 12


def test_graph_filters_account_and_business_unit_before_projection(sources):
    sources.insert([
        sale("a-full", "kit2", 3, 60, unit="Full"),
        sale("a-shop", "kit2", 2, 40),
        sale("b-full", "kit2", 1, 30, store="Loja B", store_id="store-b", unit="Full"),
    ])
    assert sum(graph(sku="479", loja="Loja A", unidade_negocio="Full")["quantidades_vendas"]) == 6
    assert sum(graph(sku="479", loja="Loja B", unidade_negocio="Full")["quantidades_vendas"]) == 10
    assert sum(graph(sku="479", unidade_negocio="Full")["quantidades_vendas"]) == 16
    assert sum(graph(sku="kit2", unidade_negocio="Full")["quantidades_vendas"]) == 4


def test_report_deduplicates_invoice_pair_before_expanding_kit(sources):
    sources.insert([sale("order", "KIT5", 3, 150, state="Em aberto", invoice="NF1"),
                    sale("invoice", "KIT5", 3, 150, invoice="NF1")])
    result = general()
    assert result["resumo"]["unidades_vendidas"] == 15
    assert result["resumo"]["faturamento"] == 150
    assert result["auditoria"]["vendas"]["duplicidades_pedido_nota"] == 1
    assert sum(row["registros_vendas"] for row in result["skus"]) == 1


def test_skus_without_sales_uses_latest_kit_sale_in_same_store(sources, monkeypatch):
    class FixedDate(datetime):
        @classmethod
        def now(cls, tz=None):
            return cls(2026, 3, 31, 12)

    monkeypatch.setattr(analytics, "datetime", FixedDate)
    (sources.root / "lojas_config.json").write_text(json.dumps(STORES), encoding="utf-8")
    products_path = sources.root / "produtos_compilado.csv"
    with products_path.open("w", encoding="utf-8", newline="") as handle:
        writer = csv.DictWriter(handle, ["sku", "store_id", "loja_sync", "saldo_loja", "nome_bling"])
        writer.writeheader()
        writer.writerow(dict(sku="479", store_id="store-a", loja_sync="Loja A", saldo_loja=59, nome_bling="Base"))
    monkeypatch.setattr(analytics, "get_vendas_dependencies", lambda: SimpleNamespace(paths=SimpleNamespace(legacy_products_csv=str(products_path))))
    monkeypatch.setattr(analytics, "_migrar_arquivo_legado_para_tenant", lambda *_: str(products_path))
    sources.insert([sale("a-kit", "kit2", 1, 20, day="2026-03-01"),
                    sale("b-kit", "kit2", 1, 20, day="2026-03-30", store="Loja B", store_id="store-b")])
    result = analytics.skus_sem_venda.__wrapped__(loja="Loja A", client_id="fixture")
    assert result["ultimos_30_dias"][0]["sku"] == "479"
    assert result["ultimos_30_dias"][0]["ultima_venda"] == "2026-03-01"
    assert result["ultimos_30_dias"][0]["estoque"] == 59


def test_component_only_rows_export_without_revenue_or_rank(sources):
    from openpyxl import load_workbook

    sources.insert([sale("nested", "MIX", 2, 100)])
    report = pareto()
    workbook = load_workbook(io.BytesIO(reports.build_pareto_xlsx(report)))
    rows = {row[1]: row for row in workbook["Todos os SKUs"].iter_rows(min_row=3, values_only=True)}
    assert rows["MIX"][0] == 1
    assert rows["MIX"][3:6] == (0, 1, 100)
    assert rows["479"][0] is None
    assert rows["479"][3:6] == (12, 0, 0)
    assert rows["B"][3:6] == (4, 0, 0)
    workbook.close()
    general_workbook = load_workbook(io.BytesIO(general_report.build_general_sku_xlsx(general())))
    sheet = general_workbook["Vendas e Estoque"]
    headers = [cell.value for cell in sheet[2]]
    general_rows = {row[1]: dict(zip(headers, row)) for row in sheet.iter_rows(min_row=3, values_only=True)}
    assert general_rows["479"]["2026-01"] == 12
    assert general_rows["479"]["Estoque Loja"] == 59
    assert general_rows["479"]["Faturamento"] == 0
    assert general_rows["MIX"]["Faturamento"] == 100
    general_workbook.close()
