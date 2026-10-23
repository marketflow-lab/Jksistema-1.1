"""External carton labels from the tenant's unchanged supplier Packing List."""

import copy
import io
import json
import re

import fitz
import openpyxl
import pytest
from fastapi import FastAPI, HTTPException
from fastapi.testclient import TestClient
from PIL import Image

from backend.routers.medias_compras import create_medias_compras_router
from backend.services import cadastro_compatibilidade, cadastro_importador_loja
from backend.services import medias_compras_common as common
from backend.services import medias_compras_excel as excel
from backend.services import medias_compras_packing as packing


ROTA = "/api/medias-compras/listas-pedidos/lista-1/etiquetas"


def _lista(**overrides):
    lista = {
        "id": "lista-1", "nome_lista": "Pedido sintético", "loja": "Loja A",
        "store_id": "store-a", "numero_invoice": "LABEL-2026",
        "supplier": "Fornecedor sintético", "itens": [{"SKU": "479", "Quantidade": 100}],
    }
    lista.update(overrides)
    return lista


def _grupo(caixas="1", ctns=1, skus=("SKU:479",), **overrides):
    grupo = {
        "caixas": caixas, "ctns": ctns, "tipo_embalagem": "ctns",
        "dimensoes": [10, 20, 30], "nw_ctn": 1, "gw_ctn": 2,
        "itens": [{"sku": sku, "quantidade": ctns * 10, "qty_ctn": 10} for sku in skus],
    }
    grupo.update(overrides)
    return grupo


def _dados_exemplo():
    # Synthetic fixture retaining only the example's carton-to-SKU assignment.
    return {"grupos": [_grupo("1-10", 10), _grupo("11", 1, ("SKU:479-1",))]}


def _xlsx(dados=None):
    dados = dados if dados is not None else _dados_exemplo()
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    for cell, value in {"A1": "CARTON NO.", "B1": "ITEM NO.", "D1": "QTY (PCS)",
                        "F1": "TOTAL CTNS"}.items():
        sheet[cell] = value
    row = 2
    for grupo in dados["grupos"]:
        inicio = row
        for item in grupo["itens"]:
            sheet.cell(row, 2, item["sku"])
            sheet.cell(row, 4, item["quantidade"])
            sheet.cell(row, 5, item["qty_ctn"])
            row += 1
        valores = {1: grupo["caixas"], 6: grupo["ctns"], 7: 10, 8: 20, 9: 30, 12: 1, 13: 2}
        if grupo.get("tipo_embalagem") == "trays":
            valores[6] = f'{grupo["ctns"]} Trays'
        for col, value in valores.items():
            sheet.cell(inicio, col, value)
            if row - inicio > 1:
                sheet.merge_cells(start_row=inicio, start_column=col, end_row=row - 1, end_column=col)
    buffer = io.BytesIO()
    workbook.save(buffer)
    workbook.close()
    return buffer.getvalue()


def _salvo(tenants, tenant="cliente-a"):
    return json.loads((tenants / tenant / "listas_pedidos.json").read_text(encoding="utf-8"))[0]


def _gravar(tenants, lista, tenant="cliente-a"):
    (tenants / tenant / "listas_pedidos.json").write_text(json.dumps([lista]), encoding="utf-8")


@pytest.fixture
def tenants(tmp_path, monkeypatch):
    root = tmp_path / "info"
    for tenant in ("cliente-a", "cliente-b"):
        pasta = root / tenant
        pasta.mkdir(parents=True)
        (pasta / "listas_pedidos.json").write_text(json.dumps([_lista()]), encoding="utf-8")
    monkeypatch.setattr(common, "_arquivo_listas_pedidos", lambda tenant: str(root / tenant / "listas_pedidos.json"))
    monkeypatch.setattr(common, "get_tenant_path", lambda tenant: str(root / tenant))
    monkeypatch.setattr(excel, "get_tenant_path", lambda tenant: str(root / tenant))
    monkeypatch.setattr(cadastro_compatibilidade, "visao_produtos_cadastro_contexto_loja",
                        lambda *_args, **_kwargs: {"store_id": "store-a", "scope": "store", "produtos": []})
    monkeypatch.setattr(cadastro_importador_loja, "ler_importador_loja_por_store_id", lambda *_args: {})
    monkeypatch.setattr(excel, "_cadastro_fornecedor_commercial_invoice", lambda *_args, **_kwargs: {})
    return root


def _textos_pdf(conteudo):
    with fitz.open(stream=conteudo, filetype="pdf") as documento:
        return [pagina.get_text() for pagina in documento]


def test_etiquetas_expande_exemplo_sem_alterar_packing():
    dados = _dados_exemplo()
    antes = copy.deepcopy(dados)
    caixas = packing._caixas_etiquetas_packing(dados)
    assert caixas == ([{"numero": numero, "skus": ["479"]} for numero in range(1, 11)]
                      + [{"numero": 11, "skus": ["479-1"]}])
    assert dados == antes


def test_etiquetas_pdf_tem_uma_pagina_por_caixa_no_modelo():
    lista, dados = _lista(), _dados_exemplo()
    antes = copy.deepcopy((lista, dados))
    pdf = packing._gerar_etiquetas_packing_bytes(lista, dados)
    assert pdf.startswith(b"%PDF-")
    with fitz.open(stream=pdf, filetype="pdf") as documento:
        assert len(documento) == 11
        for numero, pagina in enumerate(documento, 1):
            assert pagina.rect.width == pytest.approx(841.89, abs=.1)
            assert pagina.rect.height == pytest.approx(595.28, abs=.1)
            texto = pagina.get_text()
            assert "LABEL-2026" in texto
            assert re.search(rf"\b0*{numero}\s*/\s*0*11\b", texto)
            assert f'SKU {"479" if numero <= 10 else "479-1"}' in texto
            if numero <= 10:
                assert "479-1" not in texto
    assert (lista, dados) == antes


@pytest.mark.parametrize("caixas,ctns", [
    ("", 1), ("0", 1), ("-1", 1), ("2-1", 2), ("1-0", 1), ("1-3", 2),
    ("1.5", 1), ("1/2", 1), ("1-a", 1), ("1--2", 2), ("1,1", 2),
    ("1-2,2-3", 4), ("1-5001", 5001), ("1-999999999999", 999999999999),
])
def test_etiquetas_rejeita_numeracao_invalida_ou_divergente(caixas, ctns):
    with pytest.raises(HTTPException) as erro:
        packing._caixas_etiquetas_packing({"grupos": [_grupo(caixas, ctns)]})
    assert erro.value.status_code == 400


@pytest.mark.parametrize("grupos", [
    [_grupo("2", 1)],
    [_grupo("1", 1), _grupo("3", 1)],
    [_grupo("1-2", 2), _grupo("2-3", 2)],
    [_grupo("1", 1, ctns_origem="F2"), _grupo("1", 1, ctns_origem="F3")],
])
def test_etiquetas_rejeita_lacuna_ou_sobreposicao_entre_embalagens(grupos):
    with pytest.raises(HTTPException) as erro:
        packing._caixas_etiquetas_packing({"grupos": grupos})
    assert erro.value.status_code == 400


def test_etiquetas_aceita_5000_caixas_sem_gerar_pdf_e_ordena_numeracao():
    caixas = packing._caixas_etiquetas_packing({"grupos": [_grupo("2-5000", 4999), _grupo("1", 1)]})
    assert len(caixas) == 5000
    assert [caixas[0]["numero"], caixas[-1]["numero"]] == [1, 5000]


def test_etiquetas_lista_explicita_e_caixa_mista_preservam_todos_skus():
    dados = {"grupos": [_grupo("1, 2-3", 3, ("SKU:479", "SKU 479-1", "SKU:479"))]}
    caixas = packing._caixas_etiquetas_packing(dados)
    assert caixas == [{"numero": numero, "skus": ["479", "479-1"]} for numero in range(1, 4)]
    textos = _textos_pdf(packing._gerar_etiquetas_packing_bytes(_lista(), dados))
    assert len(textos) == 3
    for texto in textos:
        assert "SKU 479" in texto and "SKU 479-1" in texto
        assert "SKU SKU" not in texto


def test_etiquetas_mesma_ancora_de_embalagem_coalesce_skus_sem_duplicar_paginas():
    dados = {"grupos": [
        _grupo("1-2", 2, ("SKU:479",), ctns_origem="F2"),
        _grupo("1-2", 2, ("SKU:479-1",), ctns_origem="F2"),
    ]}
    assert packing._caixas_etiquetas_packing(dados) == [
        {"numero": 1, "skus": ["479", "479-1"]}, {"numero": 2, "skus": ["479", "479-1"]},
    ]


def test_etiquetas_trays_explicitos_preservam_numeracao_sem_inventar_caixas():
    dados = {"grupos": [_grupo("1-2", 2, tipo_embalagem="trays")]}
    antes = copy.deepcopy(dados)
    assert packing._caixas_etiquetas_packing(dados) == [
        {"numero": 1, "skus": ["479"]}, {"numero": 2, "skus": ["479"]},
    ]
    assert dados == antes
    dados["grupos"][0]["caixas"] = ""
    with pytest.raises(HTTPException) as erro:
        packing._caixas_etiquetas_packing(dados)
    assert erro.value.status_code == 400


def test_etiquetas_tipo_de_embalagem_divergente_na_mesma_ancora_e_rejeitado():
    dados = {"grupos": [
        _grupo("1", 1, ctns_origem="F2"),
        _grupo("1", 1, ctns_origem="F2", tipo_embalagem="trays"),
    ]}
    with pytest.raises(HTTPException) as erro:
        packing._caixas_etiquetas_packing(dados)
    assert erro.value.status_code == 400


@pytest.mark.parametrize("dados", [
    {"grupos": []},
    {"grupos": [_grupo(ctns=0)]},
    {"grupos": [_grupo(ctns=-1)]},
    {"grupos": [_grupo(ctns=1.5)]},
    {"grupos": [_grupo(skus=())]},
    {"grupos": [_grupo(skus=("SKU:",))]},
    {"grupos": [_grupo(tipo_embalagem="pallets")]},
    {"grupos": [_grupo("1", 1, ctns_origem="F2"), _grupo("2", 1, ctns_origem="F2")]},
])
def test_etiquetas_rejeita_volumes_sem_quantidade_produtos_ou_identidade_valida(dados):
    with pytest.raises(HTTPException) as erro:
        packing._caixas_etiquetas_packing(dados)
    assert erro.value.status_code == 400


@pytest.mark.parametrize("endereco,origem", [("Endereço sintético, Portugal", "PORTUGAL"), ("", "")])
def test_etiquetas_usa_cabecalho_e_destinatarios_atuais_e_origem_cadastrada(monkeypatch, endereco, origem):
    monkeypatch.setattr(excel, "_cadastro_fornecedor_commercial_invoice",
                        lambda *_args, **_kwargs: {"endereco_empresa": endereco})
    monkeypatch.setattr(excel, "_cabecalho_fornecedor_commercial_invoice",
                        lambda *_args, **_kwargs: ("EXPORTADOR <A> & B", 100))
    monkeypatch.setattr(excel, "_bloco_nacionalizador_commercial_invoice",
                        lambda *_args, **_kwargs: ("IMPORTER SYNTHETIC CURRENT", 100))
    monkeypatch.setattr(packing, "_texto_comprador_packing",
                        lambda *_args, **_kwargs: "BUYER SYNTHETIC CURRENT")
    pdf = packing._gerar_etiquetas_packing_bytes(_lista(), {"grupos": [_grupo()]}, client_id="cliente-a")
    texto = _textos_pdf(pdf)[0]
    for esperado in ("EXPORTADOR <A> & B", "IMPORTER SYNTHETIC CURRENT", "BUYER SYNTHETIC CURRENT"):
        assert esperado in texto
    assert "FROM" in texto
    if origem:
        assert origem in texto
    assert "CHINA" not in texto


def test_etiquetas_textos_longos_e_especiais_sao_literais_e_cabem_na_pagina():
    invoice = "INV <teste> & (A) " + "EXTENSO-" * 6
    sku = "PECA<1>&(A)"
    pdf = packing._gerar_etiquetas_packing_bytes(_lista(numero_invoice=invoice),
                                               {"grupos": [_grupo(skus=(f"SKU:{sku}",))]})
    with fitz.open(stream=pdf, filetype="pdf") as documento:
        assert len(documento) == 1
        pagina = documento[0]
        texto = "".join(pagina.get_text().split())
        assert "".join(invoice.split()) in texto
        assert sku in texto
        assert "&amp;" not in texto and "&lt;" not in texto
        spans = [span for bloco in pagina.get_text("dict")["blocks"] if "lines" in bloco
                 for linha in bloco["lines"] for span in linha["spans"] if span["text"].strip()]
        assert spans
        for span in spans:
            assert pagina.rect.contains(fitz.Rect(span["bbox"]))


def test_etiquetas_texto_que_excede_espaco_retorna_erro_sem_cortar():
    with pytest.raises(HTTPException) as erro:
        packing._gerar_etiquetas_packing_bytes(_lista(numero_invoice="SEM_ESPACO_" * 1000),
                                              {"grupos": [_grupo()]})
    assert erro.value.status_code == 400


def test_etiquetas_download_pdf_usa_invoice_atual_e_preserva_lista_e_original(tenants):
    original = _xlsx()
    packing._importar_packing_list_sync("lista-1", original, "packing.xlsx", "cliente-a")
    lista = _salvo(tenants)
    lista["numero_invoice"] = "CURRENT-2026"
    _gravar(tenants, lista)
    arquivo = tenants / "cliente-a" / "listas_pedidos.json"
    antes = arquivo.read_bytes()
    resposta = packing._baixar_etiquetas_packing_sync("lista-1", "cliente-a")
    assert resposta.media_type == "application/pdf"
    assert resposta.headers["content-disposition"] == 'attachment; filename="Etiquetas_Externas_CURRENT-2026.pdf"'
    assert resposta.headers["content-length"] == str(len(resposta.body))
    textos = _textos_pdf(resposta.body)
    assert len(textos) == 11
    assert all("CURRENT-2026" in texto for texto in textos)
    assert arquivo.read_bytes() == antes
    assert next((arquivo.parent / "packing_lists").glob("*.xlsx")).read_bytes() == original
    assert "packing_list" not in _salvo(tenants, "cliente-b")


@pytest.mark.parametrize("lista_id", ["lista-1", "inexistente", "../cliente-b/lista-1"])
def test_etiquetas_sem_documento_ou_lista_retorna_404_sem_gravar(tenants, lista_id):
    antes = {arquivo: arquivo.read_bytes() for arquivo in tenants.rglob("*.json")}
    with pytest.raises(HTTPException) as erro:
        packing._baixar_etiquetas_packing_sync(lista_id, "cliente-a")
    assert erro.value.status_code == 404
    assert {arquivo: arquivo.read_bytes() for arquivo in tenants.rglob("*.json")} == antes
    assert not list(tenants.rglob("*.pdf"))


def test_etiquetas_original_corrompido_nao_chega_ao_gerador(tenants, monkeypatch):
    packing._importar_packing_list_sync("lista-1", _xlsx(), "packing.xlsx", "cliente-a")
    original = next((tenants / "cliente-a" / "packing_lists").glob("*.xlsx"))
    original.write_bytes(b"Synthetic corrupted original")
    monkeypatch.setattr(packing, "_gerar_etiquetas_packing_bytes",
                        lambda *_args, **_kwargs: pytest.fail("A origem corrompida deve ser rejeitada antes da exportacao."))
    with pytest.raises(HTTPException) as erro:
        packing._baixar_etiquetas_packing_sync("lista-1", "cliente-a")
    assert erro.value.status_code == 409


def test_etiquetas_original_ausente_preserva_lista(tenants):
    packing._importar_packing_list_sync("lista-1", _xlsx(), "packing.xlsx", "cliente-a")
    arquivo = tenants / "cliente-a" / "listas_pedidos.json"
    antes = arquivo.read_bytes()
    next((arquivo.parent / "packing_lists").glob("*.xlsx")).unlink()
    with pytest.raises(HTTPException) as erro:
        packing._baixar_etiquetas_packing_sync("lista-1", "cliente-a")
    assert erro.value.status_code == 404
    assert arquivo.read_bytes() == antes


@pytest.mark.parametrize("traversal", [False, True])
def test_etiquetas_metadata_de_outro_tenant_nao_le_documento(tenants, monkeypatch, traversal):
    packing._importar_packing_list_sync("lista-1", _xlsx(), "packing.xlsx", "cliente-b")
    estrangeiro = _salvo(tenants, "cliente-b")["packing_list"]
    if traversal:
        estrangeiro["documento_id"] = "../cliente-b/packing_lists/" + estrangeiro["documento_id"]
    lista = _salvo(tenants)
    lista["packing_list"] = estrangeiro
    _gravar(tenants, lista)
    monkeypatch.setattr(packing, "_gerar_etiquetas_packing_bytes",
                        lambda *_args, **_kwargs: pytest.fail("Documento de outro cliente nao deve chegar ao gerador."))
    with pytest.raises(HTTPException) as erro:
        packing._baixar_etiquetas_packing_sync("lista-1", "cliente-a")
    assert erro.value.status_code == 404


def test_etiquetas_http_exige_autenticacao_e_usa_tenant_da_sessao(tenants, monkeypatch):
    packing._importar_packing_list_sync("lista-1", _xlsx(), "packing.xlsx", "cliente-a")

    def autenticar(_request, authorization):
        tenants_token = {"Bearer synthetic-a": "cliente-a", "Bearer synthetic-b": "cliente-b"}
        if authorization not in tenants_token:
            raise HTTPException(401, "Autenticacao necessaria.")
        return tenants_token[authorization]

    monkeypatch.setattr(common, "_get_tenant_id_fn", autenticar)
    app = FastAPI()
    app.include_router(create_medias_compras_router())
    with TestClient(app) as client:
        assert client.get(ROTA).status_code == 401
        assert client.get(ROTA, headers={"Authorization": "Bearer invalid"}).status_code == 401
        assert client.get(ROTA + "?client_id=cliente-a", headers={"Authorization": "Bearer synthetic-b"}).status_code == 404
        resposta = client.get(ROTA + "?client_id=cliente-b", headers={"Authorization": "Bearer synthetic-a"})
        assert resposta.status_code == 200, resposta.text
        assert resposta.headers["content-type"] == "application/pdf"
        assert len(_textos_pdf(resposta.content)) == 11


def _foto(cor, tamanho=(120, 60)):
    buffer = io.BytesIO()
    Image.new("RGB", tamanho, cor).save(buffer, format="PNG")
    return buffer.getvalue()


def _cores_pdf(conteudo):
    with fitz.open(stream=conteudo, filetype="pdf") as documento:
        cores = []
        for pagina in documento:
            for image in pagina.get_images(full=True):
                pixmap = fitz.Pixmap(documento, image[0])
                cores.append(tuple(pixmap.samples[:3]))
        return cores


def test_etiquetas_itens_preservam_zeros_case_e_agrupam_somente_sku_exato():
    dados = {"grupos": [_grupo(skus=("SKU:001", "SKU:1", "SKU:Ab", "SKU:ab", "SKU:001"))]}
    dados["grupos"][0]["itens"][0]["quantidade"] = 2
    dados["grupos"][0]["itens"][-1]["quantidade"] = 3
    antes = copy.deepcopy(dados)
    caixas = packing._caixas_etiquetas_packing(dados, incluir_itens=True)
    assert len(caixas) == 1
    assert caixas[0]["skus"] == ["001", "1", "Ab", "ab"]
    itens = caixas[0]["itens"]
    assert [item["sku"] for item in itens] == ["001", "1", "Ab", "ab"]
    assert itens[0]["quantidade"] == 5
    assert [item["quantidade"] for item in itens[1:]] == [10, 10, 10]
    assert dados == antes


@pytest.mark.parametrize("ctns,total,informado,individual,esperado", [
    (1, 7, None, False, 7),
    (1, 0, None, False, 0),
    (2, 20, 10, True, 10),
    (2, 25, 10, True, None),
    (2, 20, 10, False, None),
    (2, 20, None, True, None),
])
def test_etiquetas_quantidade_por_caixa_exige_fonte_individual_consistente(ctns, total, informado, individual, esperado):
    dados = {"grupos": [_grupo("1" if ctns == 1 else "1-2", ctns)]}
    item = dados["grupos"][0]["itens"][0]
    item.update(quantidade=total, qty_ctn_informado=informado, qty_ctn_individual=individual)
    antes = copy.deepcopy(dados)
    caixas = packing._caixas_etiquetas_packing(dados, incluir_itens=True)
    assert [caixa["itens"][0]["quantidade"] for caixa in caixas] == [esperado] * ctns
    assert dados == antes


def test_etiquetas_nao_inferem_quantidade_por_caixa_de_dados_legados():
    dados = {"grupos": [_grupo("1-2", 2)]}
    assert dados["grupos"][0]["itens"][0]["qty_ctn"] == 10
    caixas = packing._caixas_etiquetas_packing(dados, incluir_itens=True)
    assert [caixa["itens"][0]["quantidade"] for caixa in caixas] == [None, None]


def test_etiquetas_sku_repetido_com_quantidade_desconhecida_mantem_total_vazio():
    dados = {"grupos": [_grupo("1-2", 2, ("SKU:001", "SKU:001"))]}
    item = dados["grupos"][0]["itens"][0]
    item.update(qty_ctn_informado=10, qty_ctn_individual=True)
    caixas = packing._caixas_etiquetas_packing(dados, incluir_itens=True)
    assert all(len(caixa["itens"]) == 1 for caixa in caixas)
    assert all(caixa["itens"][0]["quantidade"] is None for caixa in caixas)


def test_etiquetas_parser_preserva_evidencia_de_quantidade_sem_mudar_calculo_packing():
    dados = {"grupos": [_grupo("1-2", 2, ("SKU:001", "SKU:002"))]}
    lido = packing._ler_packing_list_bytes(_xlsx(dados))
    itens = lido["grupos"][0]["itens"]
    assert [item["qty_ctn_informado"] for item in itens] == [10, 10]
    assert all(item["qty_ctn_individual"] is True for item in itens)
    assert [item["qty_ctn"] for item in itens] == [10, 10]

    workbook = openpyxl.load_workbook(io.BytesIO(_xlsx(dados)))
    workbook.active.merge_cells("E2:E3")
    buffer = io.BytesIO()
    workbook.save(buffer)
    workbook.close()
    lido_mesclado = packing._ler_packing_list_bytes(buffer.getvalue())
    itens_mesclados = lido_mesclado["grupos"][0]["itens"]
    assert [item["qty_ctn_informado"] for item in itens_mesclados] == [10, 10]
    assert all(item["qty_ctn_individual"] is False for item in itens_mesclados)
    assert [item["qty_ctn"] for item in itens_mesclados] == [10, 10]
    assert all(caixa["itens"][0]["quantidade"] is None
               for caixa in packing._caixas_etiquetas_packing(lido_mesclado, incluir_itens=True))


def test_etiquetas_fotos_do_cadastro_correspondem_ao_sku_e_tenant_atuais(tenants, monkeypatch):
    cores = {"cliente-a": ("red", "blue"), "cliente-b": ("green", "yellow")}
    for tenant, pares in cores.items():
        pasta = tenants / tenant / "cadastro_fotos"
        pasta.mkdir()
        for sku, cor in zip(("001", "002"), pares):
            (pasta / f"{sku}.png").write_bytes(_foto(cor))
    consultas = []

    def contexto(client_id, store_id, **_kwargs):
        consultas.append((client_id, store_id))
        return {"scope": "store", "store_id": store_id, "produtos": [
            {"sku": sku, "foto": f"{sku}.png", "cg_product description": f"Atual {sku}"}
            for sku in ("001", "002")
        ]}

    monkeypatch.setattr(cadastro_compatibilidade, "visao_produtos_cadastro_contexto_loja", contexto)
    resolver = lambda _client, _sku, foto="", *_args, **_kwargs: foto
    monkeypatch.setattr(common, "_resolver_foto_cadastro_sku", resolver)
    monkeypatch.setattr(excel, "_resolver_foto_cadastro_sku", resolver)
    dados = {"grupos": [_grupo(skus=("SKU:001", "SKU:002"))]}
    for item in dados["grupos"][0]["itens"]:
        item["foto"] = _foto("magenta")
    antes = copy.deepcopy(dados)
    for tenant, esperado in cores.items():
        pdf = packing._gerar_etiquetas_packing_bytes(_lista(), dados, client_id=tenant)
        assert set(_cores_pdf(pdf)) == {Image.new("RGB", (1, 1), cor).getpixel((0, 0)) for cor in esperado}
        assert "Sem foto" not in _textos_pdf(pdf)[0]
        assert (tenant, "store-a") in consultas
    assert dados == antes


def test_etiquetas_sem_foto_no_cadastro_nao_recuperam_foto_do_fornecedor(tenants):
    dados = {"grupos": [_grupo(skus=("SKU:001", "SKU:002"))]}
    for item, cor in zip(dados["grupos"][0]["itens"], ("red", "blue")):
        item["foto"] = _foto(cor)
    antes = copy.deepcopy(dados)
    pdf_tenant = packing._gerar_etiquetas_packing_bytes(_lista(), dados, client_id="cliente-a")
    assert _cores_pdf(pdf_tenant) == []
    assert _textos_pdf(pdf_tenant)[0].count("Sem foto") == 2
    pdf_legado = packing._gerar_etiquetas_packing_bytes(_lista(), dados)
    assert set(_cores_pdf(pdf_legado)) == {(255, 0, 0), (0, 0, 255)}
    assert dados == antes
