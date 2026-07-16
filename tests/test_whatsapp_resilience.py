from __future__ import annotations

import hashlib
import datetime
import json
import os
import sqlite3
import time
from pathlib import Path

from openpyxl import load_workbook

from backend.schemas.perguntas_pos_venda import PerguntasAprovacaoRequest
from backend.services import perguntas_pos_venda_endpoints, whatsapp_bridge, whatsapp_report_files
from backend.services.whatsapp_bridge_store import WhatsappBridgeStore


def test_sqlite_store_migrates_legacy_state_once_and_uses_wal(tmp_path: Path):
    legacy = tmp_path / "whatsapp_bridge_state.json"
    legacy.write_text(
        json.dumps(
            {
                "scheduled_report_deliveries": {"subject-1": {"weekly": "2026-W29"}},
                "dual_agent_conversations": {"conversation-1": {"thread_id": "thread-1"}},
                "question_approval_tokens": {
                    "ABCDEFGH": {
                        "approval_id": "approval-1",
                        "subject_id": "5511999999999",
                        "client_id": "cliente",
                        "username": "operador",
                        "used": False,
                        "created_at": time.time(),
                    }
                },
                "pending_messages": {
                    "wamid.1": {
                        "kind": "dual_worker",
                        "job_state": "running",
                        "client_id": "cliente",
                        "subject_id": "5511999999999",
                        "task_id": "task-1",
                        "deadline_at_epoch": time.time() + 120,
                    }
                },
            }
        ),
        encoding="utf-8",
    )
    database = tmp_path / "whatsapp_bridge.sqlite3"
    store = WhatsappBridgeStore(database, legacy)

    first = store.load_state()
    second = WhatsappBridgeStore(database, legacy).load_state()

    assert first == second
    assert first["scheduled_report_deliveries"]["subject-1"]["weekly"] == "2026-W29"
    assert legacy.with_suffix(".json.sqlite-migration.bak").exists()
    diagnostics = store.diagnostics()
    assert diagnostics["backend"] == "sqlite"
    assert diagnostics["journal_mode"].lower() == "wal"
    assert diagnostics["jobs"] == {"running": 1}
    assert diagnostics["approvals"] == {"pending": 1, "total": 1}


def test_sqlite_attempts_are_idempotent_and_audit_redacts_secrets(tmp_path: Path):
    database = tmp_path / "whatsapp_bridge.sqlite3"
    store = WhatsappBridgeStore(database)
    store.record_attempt("wamid.1", 1, "running", error_class="rate_limited", retryable=True)
    store.record_attempt("wamid.1", 1, "running", error_class="rate_limited", retryable=True)
    store.audit(
        "mercado_livre_approval_requested",
        message_id="wamid.1",
        subject_id="5511999999999",
        details={"access_token": "secret-token", "nested": {"authorization": "Bearer secret"}, "draft_hash": "abc"},
    )

    with sqlite3.connect(database) as connection:
        attempts = connection.execute("SELECT COUNT(*) FROM job_attempts").fetchone()[0]
        details = connection.execute("SELECT details_json FROM audit_events").fetchone()[0]

    assert attempts == 1
    assert "secret-token" not in details
    assert "Bearer secret" not in details
    assert "access_token" not in details
    assert "authorization" not in details
    assert "draft_hash" in details


def _tool_result() -> dict:
    return {
        "tool_id": "mercado_livre_orders",
        "source_label": "API do Mercado Livre",
        "manager_store": "JK Pecas",
        "success": True,
        "dados_suficientes": True,
        "coverage_complete": True,
        "data": [
            {"pedido": "2000001", "valor": 199.9, "status": "paid"},
            {"pedido": "2000002", "valor": 89.5, "status": "paid"},
        ],
    }


def test_report_documents_are_real_hashed_bounded_and_describe_coverage(tmp_path: Path):
    outcome = whatsapp_report_files.generate_report_documents(
        base_info_dir=tmp_path,
        client_id="cliente",
        task_id="task-report",
        prompt="Relatorio PDF & Excel <vendas>",
        tool_results=[_tool_result()],
        query_policy={"store": "JK Pecas", "period_start": "2026-07-01", "period_end": "2026-07-15"},
        formats={"pdf", "xlsx"},
    )

    assert outcome["status"] == "completed"
    assert len(outcome["artifacts"]) == 2
    by_type = {item["artifact_type"]: item for item in outcome["artifacts"]}
    pdf = Path(by_type["report_pdf"]["path"])
    xlsx = Path(by_type["report_xlsx"]["path"])
    assert pdf.read_bytes().startswith(b"%PDF-")
    assert xlsx.read_bytes().startswith(b"PK\x03\x04")
    assert load_workbook(xlsx, read_only=True)["Dados"].max_row == 3
    for artifact in outcome["artifacts"]:
        path = Path(artifact["path"])
        assert artifact["sha256"] == hashlib.sha256(path.read_bytes()).hexdigest()
        assert 0 < artifact["size"] <= 10 * 1024 * 1024
        assert artifact["stores"] == ["JK Pecas"]
        assert artifact["sources"] == ["API do Mercado Livre", "mercado_livre_orders"]
        assert artifact["record_count"] == 2
        assert artifact["coverage_complete"] is True


def test_report_cleanup_retains_seven_days_and_removes_only_expired_files(tmp_path: Path):
    outcome = whatsapp_report_files.generate_report_documents(
        base_info_dir=tmp_path,
        client_id="cliente",
        task_id="task-report-cleanup",
        prompt="Relatorio em PDF",
        tool_results=[_tool_result()],
        query_policy={"store": "JK Pecas"},
        formats={"pdf"},
    )
    path = Path(outcome["artifacts"][0]["path"])
    assert whatsapp_report_files.cleanup_stale_files(tmp_path) == 0
    old = time.time() - whatsapp_report_files.REPORT_FILE_TTL_SECONDS - 10
    os.utime(path, (old, old))
    assert whatsapp_report_files.cleanup_stale_files(tmp_path) == 1
    assert not path.exists()


def test_wait_notices_occur_only_at_15_45_105_and_then_every_60(monkeypatch):
    now = {"value": 1000.0}
    messages: list[str] = []
    saved: list[dict] = []
    monkeypatch.setattr(whatsapp_bridge.time, "time", lambda: now["value"])
    monkeypatch.setattr(whatsapp_bridge, "_pending_codex_tasks", lambda _pending: [{"status": "running"}])
    monkeypatch.setattr(
        whatsapp_bridge,
        "_post_proactive",
        lambda _config, payload: messages.append(payload["text"]) or {"success": True, "status": "sent"},
    )
    monkeypatch.setattr(whatsapp_bridge, "_save_pending", lambda _state, _message_id, pending: saved.append(dict(pending)))
    monkeypatch.setattr(whatsapp_bridge, "_update_pending_codex_tasks", lambda *_args, **_kwargs: None)
    pending = {"created_at_epoch": 1000.0, "subject_id": "subject", "job_group_id": "job"}
    task = {"task_id": "task", "status": "running"}
    config = whatsapp_bridge._default_config()

    for instant, expected in ((1014, 0), (1015, 1), (1044, 1), (1045, 2), (1104, 2), (1105, 3), (1164, 3), (1165, 4)):
        now["value"] = instant
        whatsapp_bridge._maybe_send_dual_conversation_tick(config, {}, "wamid.1", pending, task)
        assert len(messages) == expected

    assert pending["wait_notice_count"] == 4
    assert len(saved) == 4


def test_tool_result_contract_never_treats_empty_or_timeout_as_conclusive():
    empty = whatsapp_bridge._normalize_tool_result_contract({"tool_id": "bling_orders", "manager_store": "JK Pecas"})
    timeout = whatsapp_bridge._normalize_tool_result_contract(
        {"success": False, "error": "HTTP 504 timeout", "data": [], "tool_id": "mercado_livre_orders", "manager_store": "JK Pecas"},
    )

    for result in (empty, timeout):
        assert set(("success", "data", "sources", "dados_suficientes", "coverage_complete", "error_class", "retryable")).issubset(result)
        assert result["dados_suficientes"] is False
        assert result["coverage_complete"] is False
    assert timeout["retryable"] is True


def test_mercado_livre_approval_idempotency_replays_without_second_api_post(monkeypatch):
    response = "Sim, serve no modelo informado."
    draft_hash = hashlib.sha256(response.encode("utf-8")).hexdigest()
    key = hashlib.sha256(f"cliente|JK Pecas|question-1|{draft_hash}".encode("utf-8")).hexdigest()
    approvals = [{
        "id": "approval-1",
        "status": "pending",
        "loja": "JK Pecas",
        "question_id": "question-1",
        "resposta_sugerida": response,
    }]
    sends: list[tuple] = []
    monkeypatch.setattr(perguntas_pos_venda_endpoints, "dt", datetime, raising=False)
    monkeypatch.setattr(perguntas_pos_venda_endpoints, "_perguntas_ia_aprovacoes_carregar", lambda _client: approvals, raising=False)
    monkeypatch.setattr(perguntas_pos_venda_endpoints, "_perguntas_ia_aprovacoes_salvar", lambda _client, values: approvals.__setitem__(slice(None), values), raising=False)
    monkeypatch.setattr(perguntas_pos_venda_endpoints, "_obter_cfg_ml", lambda *_args: {}, raising=False)
    monkeypatch.setattr(perguntas_pos_venda_endpoints, "_perguntas_ia_limpar_resposta", lambda value: str(value).strip(), raising=False)
    monkeypatch.setattr(
        perguntas_pos_venda_endpoints,
        "_perguntas_ia_pergunta_respondida_ml",
        lambda *_args: (False, {"id": "question-1", "status": "UNANSWERED"}, {}), raising=False,
    )
    monkeypatch.setattr(
        perguntas_pos_venda_endpoints,
        "_perguntas_ia_enviar_resposta_ml",
        lambda *args: sends.append(args) or ({"status": "ANSWERED"}, {}), raising=False,
    )
    monkeypatch.setattr(perguntas_pos_venda_endpoints, "_perguntas_ia_state_carregar", lambda _client: {}, raising=False)
    monkeypatch.setattr(perguntas_pos_venda_endpoints, "_perguntas_ia_marcar_processada", lambda *_args: None, raising=False)
    monkeypatch.setattr(perguntas_pos_venda_endpoints, "_perguntas_ia_state_salvar", lambda *_args: None, raising=False)
    monkeypatch.setattr(perguntas_pos_venda_endpoints, "_perguntas_ia_memoria_registrar_resposta_aprovada", lambda *_args, **_kwargs: None, raising=False)

    request = PerguntasAprovacaoRequest(approval_id="approval-1", resposta=response, idempotency_key=key)
    first = perguntas_pos_venda_endpoints.ml_perguntas_aprovacoes_aprovar(request, "cliente")
    second = perguntas_pos_venda_endpoints.ml_perguntas_aprovacoes_aprovar(request, "cliente")

    assert first["approval"]["status"] == "sent"
    assert first["approval"]["send_idempotency_key"] == key
    assert second["idempotent_replay"] is True
    assert len(sends) == 1


def test_mercado_livre_preflight_blocks_question_answered_elsewhere(monkeypatch):
    response = "Resposta que nao pode ser repetida."
    approvals = [{
        "id": "approval-2",
        "status": "pending",
        "loja": "JK Pecas",
        "question_id": "question-2",
        "resposta_sugerida": response,
    }]
    monkeypatch.setattr(perguntas_pos_venda_endpoints, "dt", datetime, raising=False)
    monkeypatch.setattr(perguntas_pos_venda_endpoints, "_perguntas_ia_aprovacoes_carregar", lambda _client: approvals, raising=False)
    monkeypatch.setattr(perguntas_pos_venda_endpoints, "_perguntas_ia_aprovacoes_salvar", lambda _client, values: approvals.__setitem__(slice(None), values), raising=False)
    monkeypatch.setattr(perguntas_pos_venda_endpoints, "_obter_cfg_ml", lambda *_args: {}, raising=False)
    monkeypatch.setattr(perguntas_pos_venda_endpoints, "_perguntas_ia_limpar_resposta", lambda value: str(value).strip(), raising=False)
    monkeypatch.setattr(
        perguntas_pos_venda_endpoints,
        "_perguntas_ia_pergunta_respondida_ml",
        lambda *_args: (True, {"id": "question-2", "status": "ANSWERED", "answer": {"text": "Resposta de outro fluxo"}}, {}), raising=False,
    )
    monkeypatch.setattr(
        perguntas_pos_venda_endpoints,
        "_perguntas_ia_enviar_resposta_ml",
        lambda *_args: (_ for _ in ()).throw(AssertionError("answered question must not be posted again")), raising=False,
    )

    result = perguntas_pos_venda_endpoints.ml_perguntas_aprovacoes_aprovar(
        PerguntasAprovacaoRequest(approval_id="approval-2", resposta=response),
        "cliente",
    )

    assert result["approval"]["status"] == "answered_elsewhere"
    assert result["idempotent_replay"] is False
