from __future__ import annotations

import io
import json
import sqlite3
from pathlib import Path

import fitz
import pytest
from openpyxl import load_workbook

from backend.modules.vendas import general_report, reports
from backend.modules.vendas.errors import VendasDomainError
from backend.services import bling_vendas


SALES_SCHEMA = """
CREATE TABLE vendas (
    id_unico TEXT, data TEXT, loja_conta TEXT, canal TEXT, numero TEXT,
    situacao TEXT, devolucao INTEGER, sku TEXT, produto TEXT,
    quantidade REAL, valor REAL, comprador TEXT, unidade_negocio TEXT,
    nota_fiscal_id TEXT
);
CREATE TABLE notas_entrada_itens (
    id_unico TEXT, data_emissao TEXT, loja_conta TEXT, numero_nota TEXT,
    origem_codigo TEXT, sku TEXT, descricao TEXT, quantidade REAL,
    valor_total REAL, unidade_negocio_virtual TEXT, unidade_negocio TEXT,
    natureza_operacao TEXT, finalidade_operacao TEXT, devolucao INTEGER
);
"""


def _create_sales_db(path: Path) -> None:
    sales = [
        ("s-001-local", "2026-01-10", "JK Peças", "Bling", "1", "Concluída", 0, "001", "Produto 001", 10, 100, "Cliente", "JKPEÇAS LTDA", ""),
        ("s-001-full", "2026-02-10", "JK Peças", "Mercado Livre Full", "2", "Concluída", 0, "001", "Produto 001", 2, 30, "Cliente", "Mercado Livre Full", ""),
        ("order-open", "2026-03-03", "JK Peças", "Bling", "3", "Em aberto", 0, "ORDER", "Pedido antigo", 1, 40, "Cliente", "JKPEÇAS LTDA", "NF-1"),
        ("order-invoiced", "2026-03-03", "JK Peças", "Bling", "NF3", "Faturado", 0, "ORDER", "Pedido faturado", 1, 40, "Cliente", "JKPEÇAS LTDA", "NF-1"),
        ("cancel", "2026-01-11", "JK Peças", "Bling", "4", "Cancelada", 0, "CANCEL", "Cancelado", 1, 99, "Cliente", "JKPEÇAS LTDA", ""),
        ("return-line", "2026-01-11", "JK Peças", "Bling", "5", "Concluída", 1, "RETURN", "Devolução", 1, 99, "Cliente", "JKPEÇAS LTDA", ""),
        ("ebazar", "2026-01-11", "JK Peças", "Bling", "6", "Concluída", 0, "EBA", "Transferência", 1, 99, "EBAZAR", "JKPEÇAS LTDA", ""),
        ("blank", "2026-01-11", "JK Peças", "Bling", "7", "Concluída", 0, "", "Sem SKU", 1, 99, "Cliente", "JKPEÇAS LTDA", ""),
        ("fiscal", "2026-01-11", "JK Peças", "Bling", "8", "Concluída", 0, "ESTORNO DE CRÉDITO ICMS", "Ajuste", 1, 25, "Cliente", "JKPEÇAS LTDA", ""),
        ("other-store", "2026-01-11", "Uai Mineirinho", "Bling", "9", "Concluída", 0, "OUT", "Outra loja", 100, 9999, "Cliente", "LTDA MULTIMARCAS", ""),
    ]
    returns = [
        ("r-001-local", "2026-01-20", "JK Peças", "D1", "O1", "001", "Produto 001", 2, 20, "JKPEÇAS LTDA", "", "Devolução", "", 1),
        ("r-001-full", "2026-03-20", "JK Peças", "D2", "O2", "001", "Produto 001", 1, 12, "Devolução Full Estoque", "", "Compra para comercialização", "", 1),
        ("r-only", "2026-02-15", "JK Peças", "D3", "O3", "RET", "Somente devolvido", 4, 80, "JKPEÇAS LTDA", "", "Devolução", "", 1),
        ("r-001-local", "2026-01-20", "JK Peças", "D1", "O1", "001", "Produto 001", 2, 20, "JKPEÇAS LTDA", "", "Devolução", "", 1),
        ("r-other", "2026-01-20", "Carlos José", "D4", "O4", "OUT", "Outra loja", 50, 500, "Carlos Jose", "", "Devolução", "", 1),
    ]
    with sqlite3.connect(path) as connection:
        connection.executescript(SALES_SCHEMA)
        connection.executemany(
            "INSERT INTO vendas VALUES (?,?,?,?,?,?,?,?,?,?,?,?,?,?)",
            sales,
        )
        connection.executemany(
            "INSERT INTO notas_entrada_itens VALUES (?,?,?,?,?,?,?,?,?,?,?,?,?,?)",
            returns,
        )


def _create_stock_db(path: Path) -> None:
    with sqlite3.connect(path) as connection:
        connection.executescript(
            """
            CREATE TABLE estoque_sync_eventos (
                event_id TEXT PRIMARY KEY, recorded_at TEXT, data_ref TEXT,
                loja_sync TEXT, total_skus INTEGER, status TEXT
            );
            CREATE TABLE estoque_sync_evento_itens (
                id INTEGER PRIMARY KEY AUTOINCREMENT, event_id TEXT,
                sku TEXT, nome_bling TEXT, saldo_loja REAL, saldo_full REAL
            );
            """
        )
        connection.executemany(
            "INSERT INTO estoque_sync_eventos VALUES (?,?,?,?,?,?)",
            [
                ("old", "2026-03-01T09:00:00-03:00", "2026-03-01", "JK Peças", 1, "committed"),
                ("latest", "2026-03-31T10:00:00-03:00", "2026-03-31", "JK Peças", 4, "committed"),
                ("pending", "2026-04-01T10:00:00-03:00", "2026-04-01", "JK Peças", 1, "pending"),
                ("other", "2026-03-31T10:00:00-03:00", "2026-03-31", "Carlos José", 1, "committed"),
            ],
        )
        connection.executemany(
            "INSERT INTO estoque_sync_evento_itens (event_id,sku,nome_bling,saldo_loja,saldo_full) VALUES (?,?,?,?,?)",
            [
                ("old", "001", "Produto 001", 10, 900),
                ("latest", "001", "Produto 001", 5, 999),
                ("latest", "STOCK", "Somente estoque", 7, 700),
                ("latest", "ZERO", "Estoque zero", 0, 500),
                ("latest", "NEG", "Estoque negativo", -2, 400),
                ("pending", "001", "Produto 001", 999, 0),
                ("other", "OUT", "Outra loja", 500, 0),
            ],
        )


@pytest.fixture
def general_sources(tmp_path: Path, monkeypatch):
    tenant = tmp_path / "tenant-general"
    tenant.mkdir()
    sales_db = tenant / "vendas_historico_jk_pecas.db"
    stock_db = tenant / "estoque_historico.db"
    _create_sales_db(sales_db)
    _create_stock_db(stock_db)
    (tenant / "lojas_config.json").write_text(
        json.dumps([
            {"store_id": "store-jk", "nome": "JK Peças"},
            {"store_id": "store-carlos", "nome": "Carlos José"},
        ]),
        encoding="utf-8",
    )

    monkeypatch.setattr(reports, "_listar_bancos_vendas_tenant", lambda *_args: [str(sales_db)])
    monkeypatch.setattr(general_report, "_listar_bancos_vendas_tenant", lambda *_args: [str(sales_db)])
    monkeypatch.setattr(general_report, "get_tenant_path", lambda _client_id: str(tenant))
    monkeypatch.setattr(reports, "_deduplicar_vendas_consolidadas", bling_vendas._deduplicar_vendas_consolidadas)
    monkeypatch.setattr(
        reports,
        "_deve_excluir_venda_ebazar",
        lambda _return, buyer, channel: "EBAZAR" in f"{buyer} {channel}".upper(),
    )
    return tenant, sales_db, stock_db


def _generate() -> dict:
    return general_report.generate_general_sku_report.__wrapped__(
        client_id="tenant-general",
        loja="JK Peças",
        periodo="personalizado",
        data_inicio="2026-01-01",
        data_fim="2026-03-31",
    )


def test_all_relevant_skus_store_stock_and_shared_sales_rules(general_sources) -> None:
    payload = _generate()
    rows = {item["sku"]: item for item in payload["skus"]}
    no_movement_rows = {
        item["sku"]: item
        for group in payload["sem_movimentacao"]["grupos"]
        for item in group["itens"]
    }

    assert set(rows) == {"001", "ORDER", "RET"}
    assert set(no_movement_rows) == {"STOCK", "ZERO", "NEG"}
    assert payload["periodo"]["meses"] == ["2026-01", "2026-02", "2026-03"]
    assert payload["periodo"]["estoque_atualizado_em"] == "2026-03-31T10:00:00-03:00"
    assert rows["001"]["unidades_vendidas"] == 12
    assert rows["001"]["full_unidades"] == 2
    assert rows["001"]["media_mensal"] == 4
    assert rows["001"]["estoque_loja"] == 5
    assert rows["001"]["cobertura_meses"] == pytest.approx(1.25)
    assert rows["001"]["status_cobertura"] == "1 a 3 meses"
    assert rows["001"]["quantidade_devolvida"] == 3
    assert rows["001"]["taxa_devolucao"] == pytest.approx(0.25)
    assert rows["ORDER"]["unidades_vendidas"] == 1
    assert rows["ORDER"]["estoque_loja"] is None
    assert rows["ORDER"]["status_cobertura"] == "Estoque sem registro"
    assert no_movement_rows["STOCK"]["estoque_loja"] == 7
    assert no_movement_rows["ZERO"]["estoque_loja"] == 0
    assert no_movement_rows["NEG"]["estoque_loja"] == -2
    assert rows["RET"]["estoque_loja"] is None

    assert payload["resumo"]["total_skus"] == 3
    assert payload["resumo"]["total_skus_identificados"] == 6
    assert payload["resumo"]["skus_sem_movimentacao"] == 3
    assert payload["resumo"]["unidades_vendidas"] == 13
    assert payload["resumo"]["faturamento"] == 170
    assert payload["resumo"]["estoque_loja_total_conhecido"] == 5
    assert payload["resumo"]["quantidade_devolvida"] == 7
    assert payload["ajuste_fiscal"]["valor"] == 25
    assert payload["auditoria"]["vendas"]["duplicidades_pedido_nota"] == 1
    assert payload["auditoria"]["estoque"]["event_id"] == "latest"
    assert payload["auditoria"]["devolucoes_lidas"] == 3
    assert payload["top_20_devolucoes"][0]["sku"] == "RET"
    assert payload["top_20_devolucoes"][1]["predominancia"] == "Demais canais"
    assert "pode corresponder a venda anterior" in payload["top_20_devolucoes"][0]["analise"]
    assert payload["top_10_vendidos"][0]["sku"] == "001"
    assert payload["top_10_vendidos"][0]["unidades_vendidas"] == 12
    assert payload["top_10_vendidos"][0]["participacao_unidades"] == pytest.approx(12 / 13)
    assert payload["top_10_faturamento"][0]["sku"] == "001"
    assert payload["top_10_faturamento"][0]["faturamento"] == 130
    assert payload["top_10_faturamento"][0]["participacao_faturamento"] == pytest.approx(130 / 170)
    assert payload["mensal_conta"][1]["full_faturamento"] == 30
    assert payload["analise_executiva"]["participacao_full_faturamento"] == pytest.approx(30 / 170)
    assert payload["analise_executiva"]["sku_mais_vendido"] == "001"
    assert len(payload["analise_executiva"]["destaques"]) == 5
    assert {item["sku"] for item in payload["mensal_detalhado"]} == {
        "001",
        "ORDER",
        "RET",
    }


def test_account_isolation_all_store_block_and_top_20_limit(general_sources) -> None:
    _tenant, sales_db, _stock_db = general_sources
    with sqlite3.connect(sales_db) as connection:
        for index in range(25):
            connection.execute(
                "INSERT INTO notas_entrada_itens VALUES (?,?,?,?,?,?,?,?,?,?,?,?,?,?)",
                (
                    f"bulk-{index}", "2026-03-25", "JK Peças", f"B{index}", f"BO{index}",
                    f"B{index:02d}", f"Produto {index}", index + 1, (index + 1) * 10,
                    "JKPEÇAS LTDA", "", "Devolução", "", 1,
                ),
            )
    payload = _generate()
    assert len(payload["top_20_devolucoes"]) == 20
    assert payload["top_20_devolucoes"][0]["sku"] == "B24"
    assert all(item["sku"] != "OUT" for item in payload["skus"])
    with pytest.raises(VendasDomainError, match="conta específica"):
        general_report.generate_general_sku_report.__wrapped__(
            client_id="tenant-general", loja="__todas", periodo="12m"
        )


def test_homonymous_store_omits_local_stock_before_database(general_sources, monkeypatch) -> None:
    tenant, _sales_db, _stock_db = general_sources
    (tenant / "lojas_config.json").write_text(
        json.dumps([
            {"store_id": "store-a", "nome": "JK Peças"},
            {"store_id": "store-b", "nome": "jk peças"},
        ]),
        encoding="utf-8",
    )
    aberturas = []
    monkeypatch.setattr(
        general_report,
        "open_vendas_readonly",
        lambda *_args, **_kwargs: aberturas.append("db"),
    )

    snapshot = general_report._read_current_local_stock("tenant-general", "JK Peças")

    assert snapshot["available"] is False
    assert snapshot["error_code"] == "estoque_historico_loja_ambigua"
    assert "lojas homônimas" in snapshot["errors"][0]
    assert aberturas == []


def test_xlsx_pdf_and_export_contract(general_sources) -> None:
    payload = _generate()
    xlsx_bytes = general_report.build_general_sku_xlsx(payload)
    pdf_bytes = general_report.build_general_sku_pdf(payload)

    assert xlsx_bytes.startswith(b"PK")
    workbook = load_workbook(io.BytesIO(xlsx_bytes), data_only=False)
    assert workbook.sheetnames == [
        "Resumo",
        "Vendas e Estoque",
        "Mensal Detalhado",
        "Devoluções por SKU",
        "Top 20 Devoluções",
        "Fontes e Regras",
        "Sem movimentação",
    ]
    sku_sheet = workbook["Vendas e Estoque"]
    sku_row = next(row for row in sku_sheet.iter_rows(min_row=3) if row[1].value == "001")
    assert sku_row[1].number_format == "@"
    assert any(isinstance(cell.value, str) and cell.value.startswith("=SUM(") for cell in sku_row)
    assert all(row[1].value not in {"STOCK", "ZERO", "NEG"} for row in sku_sheet.iter_rows(min_row=3))
    assert len(workbook["Resumo"]._charts) == 5
    assert len(workbook["Top 20 Devoluções"]._charts) == 2
    assert workbook["Top 20 Devoluções"]["A1"].fill.fgColor.rgb.endswith("7F1D1D")
    assert workbook["Top 20 Devoluções"]["A2"].fill.fgColor.rgb.endswith("B91C1C")
    assert workbook["Top 20 Devoluções"]["A3"].fill.fgColor.rgb.endswith("FDECEC")
    assert workbook["Devoluções por SKU"]["A1"].fill.fgColor.rgb.endswith("7F1D1D")
    no_movement_sheet = workbook["Sem movimentação"]
    assert no_movement_sheet["A1"].value == "SKUs sem movimentação no período"
    grouped_skus = " ".join(str(no_movement_sheet.cell(row, 4).value or "") for row in range(6, 9))
    assert all(sku in grouped_skus for sku in ("STOCK", "ZERO", "NEG"))
    assert "Faturamento mensal por canal" in workbook["Resumo"]._charts[0].title.tx.rich.p[0].r[0].t
    assert "Top 10 mais vendidos" in workbook["Resumo"]._charts[3].title.tx.rich.p[0].r[0].t
    assert "Top 10 por valor vendido" in workbook["Resumo"]._charts[4].title.tx.rich.p[0].r[0].t
    formulas = [
        cell.value
        for sheet in workbook.worksheets
        for row in sheet.iter_rows()
        for cell in row
        if isinstance(cell.value, str) and cell.value.startswith("=")
    ]
    assert formulas
    assert not any("#REF!" in formula for formula in formulas)

    assert pdf_bytes.startswith(b"%PDF")
    document = fitz.open(stream=pdf_bytes, filetype="pdf")
    text = "\n".join(page.get_text() for page in document)
    assert len(document) >= 6
    assert "Relatório Geral de SKUs" in text
    assert "Leitura executiva" in text
    assert "Faturamento mensal por canal" in text
    assert "Top 10 mais vendidos por unidades" in text
    assert "Top 10 mais vendidos por unidades" in document[0].get_text()
    assert "Top 10 mais vendidos por valor" in document[0].get_text()
    assert "Taxa mensal de devolução" in text
    assert "20 SKUs mais devolvidos" in text
    assert "SKUs com movimentação" in text
    assert "Vendas mensais por SKU" in text
    assert "Devoluções de todos os SKUs" in text
    assert text.index("SKUs com movimentação - resumo") < text.rindex("Vendas mensais por SKU")
    top_returns_heading = text.index("20 SKUs mais devolvidos - análise")
    assert text.rindex("Vendas mensais por SKU") < top_returns_heading
    assert top_returns_heading < text.index("Devoluções de todos os SKUs")
    assert "001" in text
    last_page_text = document[-1].get_text()
    earlier_text = "\n".join(page.get_text() for page in document[:-1])
    assert "SKUs sem movimentação no período" in last_page_text
    assert all(sku in last_page_text for sku in ("STOCK", "ZERO", "NEG"))
    assert all(sku not in earlier_text for sku in ("STOCK", "ZERO", "NEG"))

    content, filename, media_type = general_report.export_general_sku_report(
        client_id="tenant-general",
        loja="JK Peças",
        formato="xlsx",
        periodo="personalizado",
        data_inicio="2026-01-01",
        data_fim="2026-03-31",
    )
    assert content.startswith(b"PK")
    assert filename == "relatorio_geral_skus_jk_pecas_2026-01-01_2026-03-31.xlsx"
    assert media_type.endswith("spreadsheetml.sheet")
    with pytest.raises(VendasDomainError, match="Formato inválido"):
        general_report.export_general_sku_report(
            client_id="tenant-general", loja="JK Peças", formato="csv"
        )
