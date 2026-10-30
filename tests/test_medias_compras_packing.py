"""Packing List transformation, tenant persistence and concurrency contracts."""

import asyncio
import copy
import hashlib
import io
import json
import os
import subprocess
import zipfile
from concurrent.futures import ThreadPoolExecutor
from threading import Barrier

import openpyxl
import pytest
from fastapi import FastAPI, HTTPException
from fastapi.testclient import TestClient
from openpyxl.drawing.image import Image as XLImage
from openpyxl.utils.units import pixels_to_EMU, points_to_pixels
from PIL import Image

from backend.services import medias_compras_common as common
from backend.services import medias_compras_excel as excel
from backend.services import medias_compras_packing as packing
from backend.services import cadastro_compatibilidade, cadastro_importador_loja
from backend.services import medias_compras_listas as listas_service
from backend.routers.medias_compras import create_medias_compras_router


def _foto(cor, tamanho=(120, 60)):
    buffer = io.BytesIO()
    Image.new("RGB", tamanho, cor).save(buffer, format="PNG")
    return buffer.getvalue()


def _modelo(*, alterar=None, fotos=True):
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.title = "Sheet1"
    for cell, value in {
        "A1": "箱号\nCARTON NO.", "B1": "款号\nITEM NO.",
        "C1": "图片\nPICTURE", "D1": "数量\nQTY（PCS）",
        "E1": "每件数量\nQTY/CTN", "F1": "总件数\nTOTAL CTNS",
        "G1": "每件尺寸\nCTN SIZE", "J1": "体积\nCBM",
        "L1": "净重\nN.W", "N1": "毛重\nG.W",
        "E2": "（PCS）", "F2": "（CTN）", "J2": "每件体积\nCBM/CTN",
        "K2": "总体积\nTOTAL CBM", "L2": "每件净重\nN.W/CTN",
        "M2": "每件毛重\nG.W/CTN", "N2": "总净重\nTOTAL N.W",
        "O2": "总毛重\nTOTAL G.W",
    }.items():
        sheet[cell] = value
    for interval in ("A1:A2", "B1:B2", "C1:C2", "D1:D2", "G1:I2", "J1:K1", "L1:M1", "N1:O1"):
        sheet.merge_cells(interval)
    rows = {
        3: ["1-2", "SKU:001 Sensor", None, 4, 2, 2, 10, 20, 30, None, None, 1.2, 1.5],
        4: ["3", "SKU:001 Sensor", None, 2, 5, 1, 20, 25, 40, None, None, 3.2, 3.6],
        5: [None, "SKU:002 Cable", None, 3],
        6: ["4", "SKU:999 Unknown", None, 10, 10, 1, 10, 10, 10, None, None, 2, 2.5],
    }
    for row, values in rows.items():
        for col, value in enumerate(values, 1):
            sheet.cell(row, col, value)
        if row != 5:
            for col, formula in {
                10: f"=G{row}*H{row}*I{row}/1000000", 11: f"=J{row}*F{row}",
                14: f"=L{row}*F{row}", 15: f"=M{row}*F{row}",
            }.items():
                sheet.cell(row, col, formula)
    for col in (1, *range(5, 16)):
        letter = openpyxl.utils.get_column_letter(col)
        sheet.merge_cells(f"{letter}4:{letter}5")
    sheet.merge_cells("A7:C7")
    sheet["A7"] = "TOTAL:"
    for col in ("D", "F", "K", "N", "O"):
        sheet[f"{col}7"] = f"=SUM({col}3:{col}6)"
    # Internal drawing collection order intentionally differs from worksheet order.
    if fotos:
        for row, cor in ((6, "blue"), (3, "red"), (5, "green"), (4, "yellow")):
            sheet.add_image(XLImage(io.BytesIO(_foto(cor))), f"C{row}")
    if alterar:
        alterar(sheet)
    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def _lista(**overrides):
    lista = {
        "id": "lista-1", "nome_lista": "Pedido teste", "loja": "Loja A",
        "store_id": "store-a", "numero_invoice": "TEST-14072026",
        "order_date": "2026-07-14", "status": "Analisando orçamento",
        "supplier": "Fornecedor teste", "currency": "USD", "incoterm": "EXW",
        "itens": [
            {"SKU": "001", "Título do produto em inglês": "Sensor description", "Color/side": "Unit",
             "Quantidade": 6, "Valor unidade": 8.25, "compra_aprovada": True,
             "compra_aprovada_em": "2026-07-01T10:00:00", "frete_internacional": 3.5},
            {"SKU": "002", "Título do produto em inglês": "Cable description", "Color/side": "Pair",
             "Quantidade": 3, "Valor unidade": 5, "compra_aprovada": False},
        ],
        "comentario": "Preserve this order note", "embarcado": False,
    }
    lista.update(overrides)
    return lista


def _sheet(conteudo, *, data_only=False):
    workbook = openpyxl.load_workbook(io.BytesIO(conteudo), data_only=data_only)
    assert len(workbook.sheetnames) == 1
    return workbook.worksheets[0]


@pytest.fixture
def tenants(tmp_path, monkeypatch):
    raiz = tmp_path / "info"
    for tenant in ("cliente-a", "cliente-b"):
        pasta = raiz / tenant
        pasta.mkdir(parents=True)
        (pasta / "listas_pedidos.json").write_text(json.dumps([_lista()]), encoding="utf-8")
    monkeypatch.setattr(common, "_arquivo_listas_pedidos", lambda tenant: str(raiz / tenant / "listas_pedidos.json"))
    monkeypatch.setattr(common, "get_tenant_path", lambda tenant: str(raiz / tenant))
    monkeypatch.setattr(excel, "get_tenant_path", lambda tenant: str(raiz / tenant))
    # These synthetic tenants have no store/importer registry. Do not inherit
    # the global store runtime configured by unrelated test modules.
    monkeypatch.setattr(cadastro_compatibilidade, "visao_produtos_cadastro_contexto_loja",
                        lambda *_args, **_kwargs: {"store_id": "store-a", "scope": "store", "produtos": [
                            {"sku": "001", "cg_product description": "Sensor description", "Unit/Pair": "Unit"},
                            {"sku": "002", "cg_product description": "Cable description", "Unit/Pair": "Pair"},
                        ]})
    monkeypatch.setattr(common, "_resolver_foto_cadastro_sku", lambda *_args, **_kwargs: "")
    monkeypatch.setattr(cadastro_importador_loja, "ler_importador_loja_por_store_id", lambda *_args: {})
    return raiz


def _salvo(tenants, tenant="cliente-a"):
    return json.loads((tenants / tenant / "listas_pedidos.json").read_text(encoding="utf-8"))[0]


def test_parser_preserva_grupos_quantidades_e_fotos_da_ancora_sem_cache():
    dados = packing._ler_packing_list_bytes(_modelo())
    grupos = dados["grupos"]
    assert [g["caixas"] for g in grupos] == ["1-2", "3", "4"]
    assert [g["ctns"] for g in grupos] == [2, 1, 1]
    assert [len(g["itens"]) for g in grupos] == [1, 2, 1]
    assert grupos[0]["itens"][0]["sku"] == "SKU:001 Sensor"
    assert [i["quantidade"] for g in grupos for i in g["itens"]] == [4, 2, 3, 10]
    assert [i["qty_ctn"] for g in grupos for i in g["itens"]] == [2, 2, 3, 10]
    assert grupos[0]["dimensoes"] == [10, 20, 30]
    assert sum(g["cbm_total"] for g in grupos) == pytest.approx(0.033)
    assert sum(g["nw_total"] for g in grupos) == pytest.approx(7.6)
    assert sum(g["gw_total"] for g in grupos) == pytest.approx(9.1)
    for item, cor in zip((i for g in grupos for i in g["itens"]), ("red", "yellow", "green", "blue")):
        image = Image.open(io.BytesIO(item["foto"]))
        assert image.getpixel((0, 0)) == Image.new("RGB", (1, 1), cor).getpixel((0, 0))


def test_parser_sem_fotos_nao_inventa_imagem():
    dados = packing._ler_packing_list_bytes(_modelo(fotos=False))
    assert all(i.get("foto") is None for g in dados["grupos"] for i in g["itens"])


@pytest.mark.parametrize("conteudo", [b"", b"not an Excel document", b"PK invalid zip"])
def test_parser_rejeita_documento_invalido(conteudo):
    with pytest.raises(HTTPException) as erro:
        packing._ler_packing_list_bytes(conteudo)
    assert erro.value.status_code in (400, 413)


@pytest.mark.parametrize("alterar", [
    lambda s: setattr(s["A1"], "value", "Unrelated spreadsheet"),
    lambda s: setattr(s["D3"], "value", -1),
    lambda s: setattr(s["F3"], "value", 0),
    lambda s: setattr(s["G3"], "value", "=WEBSERVICE(\"https://example.invalid\")"),
    lambda s: setattr(s["L3"], "value", -1),
])
def test_parser_rejeita_estrutura_ou_medidas_invalidas(alterar):
    with pytest.raises(HTTPException) as erro:
        packing._ler_packing_list_bytes(_modelo(alterar=alterar))
    assert erro.value.status_code == 400


def test_exportacao_enriquece_lista_e_conta_grupos_uma_vez():
    lista = _lista()
    antes = copy.deepcopy(lista)
    dados = packing._ler_packing_list_bytes(_modelo())
    dados_antes = copy.deepcopy(dados)
    conteudo = packing._gerar_packing_list_bytes(lista, dados)
    sheet = _sheet(conteudo)
    assert sheet.max_column == 17
    assert sheet["A2"].value == "PACKING LIST"
    assert "TEST-14072026" in sheet["A3"].value
    assert "July 14, 2026" in sheet["I3"].value
    assert [sheet[f"B{r}"].value for r in range(9, 13)] == ["SKU:001 Sensor", "SKU:001 Sensor", "SKU:002 Cable", "SKU:999 Unknown"]
    assert [sheet[f"D{r}"].value for r in range(9, 13)] == ["Sensor description", "Sensor description", "Cable description", None]
    assert [sheet[f"E{r}"].value for r in range(9, 13)] == ["Unit", "Unit", "Pair", None]
    assert [sheet[f"G{r}"].value for r in range(9, 13)] == [2, 2, 3, 10]
    assert sheet["F9"].value == 4
    assert sheet["F10"].value == 5
    assert sheet["F11"].value is None
    assert sheet["F12"].value == 10
    assert {"A10:A11", "F10:F11", "H10:H11", "Q10:Q11"}.issubset(set(map(str, sheet.merged_cells.ranges)))
    assert sheet["F13"].value == "=SUM(F9:F12)"
    assert sheet["H13"].value == "=SUM(H9:H12)"
    assert sheet["M13"].value == "=SUM(M9:M12)"
    assert sheet["O13"].value == "=SUM(O9:O12)"
    assert sheet["Q13"].value == "=SUM(Q9:Q12)"
    calculado = _sheet(conteudo, data_only=True)
    for endereco, esperado in {"F13": 19, "H13": 4, "M13": .033, "O13": 7.6, "Q13": 9.1}.items():
        assert calculado[endereco].data_type == "n"
        assert calculado[endereco].value == pytest.approx(esperado)
    assert sheet.parent.calculation.fullCalcOnLoad is True
    assert sum(sheet[f"F{r}"].value or 0 for r in range(9, 13)) == 19
    assert sum(sheet[f"H{r}"].value or 0 for r in range(9, 13)) == 4
    assert sum(sheet[f"M{r}"].value or 0 for r in range(9, 13)) == pytest.approx(.033)
    assert len(sheet._images) == 4
    assert sorted(i.anchor._from.row + 1 for i in sheet._images) == list(range(9, 13))
    assert sheet.page_setup.orientation == "landscape"
    assert "$A$1:$Q$14" in str(sheet.print_area)
    assert lista == antes
    assert dados == dados_antes


def test_exportacao_totais_somam_cbm_arredondado_das_caixas_mistas():
    dados = packing._ler_packing_list_bytes(_modelo(fotos=False))
    dados["grupos"] = dados["grupos"][:2]
    for grupo in dados["grupos"]:
        grupo["cbm_total"] = .0006
    conteudo = packing._gerar_packing_list_bytes(_lista(), dados)
    sheet = _sheet(conteudo, data_only=True)
    assert sheet["M9"].value == .001
    assert sheet["M10"].value == .001
    assert sheet["M11"].value is None
    assert sheet["M12"].value == .002
    assert sheet["F12"].value == 9
    assert sheet["H12"].value == 3
    assert _sheet(conteudo)["M12"].value == "=SUM(M9:M11)"


def test_exportacao_totais_zero_aparecem_sem_recalculo():
    dados = packing._ler_packing_list_bytes(_modelo(fotos=False))
    for grupo in dados["grupos"]:
        grupo.update(cbm_total=0, nw_total=0, gw_total=0)
        for item in grupo["itens"]:
            item["quantidade"] = 0
    conteudo = packing._gerar_packing_list_bytes(_lista(), dados)
    sheet = _sheet(conteudo, data_only=True)
    assert sheet["H13"].value == 4
    for endereco in ("F13", "M13", "O13", "Q13"):
        assert sheet[endereco].value == 0
        assert sheet[endereco].data_type == "n"


def test_exportacao_cache_preserva_imagens_relacionamentos_e_xml(monkeypatch):
    original = []
    preencher = packing._preencher_cache_totais_packing

    def capturar(conteudo, totais):
        original.append(conteudo)
        return preencher(conteudo, totais)

    monkeypatch.setattr(packing, "_preencher_cache_totais_packing", capturar)
    conteudo = packing._gerar_packing_list_bytes(_lista(), packing._ler_packing_list_bytes(_modelo()))
    with zipfile.ZipFile(io.BytesIO(original[0])) as antes, zipfile.ZipFile(io.BytesIO(conteudo)) as depois:
        assert depois.testzip() is None
        assert antes.namelist() == depois.namelist()
        assert antes.comment == depois.comment
        for info in antes.infolist():
            if info.filename == "xl/worksheets/sheet1.xml":
                esperado = antes.read(info)
                for endereco, valor in {"F13": "19", "H13": "4", "M13": "0.033", "O13": "7.6", "Q13": "9.1"}.items():
                    esperado = esperado.replace(
                        f'<f>SUM({endereco[0]}9:{endereco[0]}12)</f><v></v>'.encode(),
                        f'<f>SUM({endereco[0]}9:{endereco[0]}12)</f><v>{valor}</v>'.encode(),
                    )
                assert depois.read(info.filename) == esperado
            else:
                assert depois.read(info.filename) == antes.read(info)
            compressao_esperada = (zipfile.ZIP_STORED if info.filename.startswith("xl/media/")
                                  else info.compress_type)
            assert depois.getinfo(info.filename).compress_type == compressao_esperada
    assert len(_sheet(conteudo)._images) == 4


def test_exportacao_textos_que_parecem_formula_sao_literais():
    lista = _lista(numero_invoice="=1+1")
    lista["itens"][0]["Título do produto em inglês"] = '=HYPERLINK("https://example.invalid","literal")'
    lista["itens"][0]["Color/side"] = "=1+1"
    sheet = _sheet(packing._gerar_packing_list_bytes(lista, packing._ler_packing_list_bytes(_modelo())))
    assert sheet["D9"].value == lista["itens"][0]["Título do produto em inglês"]
    assert sheet["D9"].data_type == "s"
    assert sheet["E9"].value == "=1+1"
    assert sheet["E9"].data_type == "s"
    formulas = [c.coordinate for row in sheet for c in row if c.data_type == "f"]
    assert formulas == ["F13", "H13", "M13", "O13", "Q13"]


@pytest.mark.parametrize("tamanho", [(640, 80), (80, 640), (240, 240)])
def test_exportacao_fotos_cabem_sem_deformar(tamanho):
    dados = packing._ler_packing_list_bytes(_modelo(fotos=False))
    dados["grupos"][0]["itens"][0]["foto"] = _foto("red", tamanho)
    sheet = _sheet(packing._gerar_packing_list_bytes(_lista(), dados))
    anchor = sheet._images[0].anchor
    assert anchor._from.col == 2
    assert anchor._from.row == 8
    # Drawing extents are integer EMUs, so preserve proportion within quantization.
    proporcao = tamanho[0] / tamanho[1]
    assert abs(anchor.ext.cx - anchor.ext.cy * proporcao) <= 1 + proporcao
    assert anchor._from.colOff + anchor.ext.cx <= pixels_to_EMU(int(sheet.column_dimensions["C"].width * 7))
    assert anchor._from.rowOff + anchor.ext.cy <= pixels_to_EMU(points_to_pixels(sheet.row_dimensions[9].height))


def test_upload_persiste_original_e_preserva_dados_comerciais(tenants):
    conteudo = _modelo()
    antes = _salvo(tenants)
    packing._importar_packing_list_sync("lista-1", conteudo, "packing.xlsx", "cliente-a")
    depois = _salvo(tenants)
    for chave, valor in antes.items():
        assert depois[chave] == valor
    assert depois["packing_list"]["sha256"] == hashlib.sha256(conteudo).hexdigest()
    arquivos = list((tenants / "cliente-a" / "packing_lists").glob("*.xlsx"))
    assert len(arquivos) == 1
    assert arquivos[0].read_bytes() == conteudo
    assert "packing_list" not in _salvo(tenants, "cliente-b")


def test_upload_repetido_nao_substitui_original(tenants):
    packing._importar_packing_list_sync("lista-1", _modelo(), "first.xlsx", "cliente-a")
    arquivo = tenants / "cliente-a" / "listas_pedidos.json"
    antes = arquivo.read_bytes()
    docs = {p.name: p.read_bytes() for p in (arquivo.parent / "packing_lists").glob("*.xlsx")}
    with pytest.raises(HTTPException) as erro:
        packing._importar_packing_list_sync("lista-1", _modelo(fotos=False), "second.xlsx", "cliente-a")
    assert erro.value.status_code == 409
    assert arquivo.read_bytes() == antes
    assert {p.name: p.read_bytes() for p in (arquivo.parent / "packing_lists").glob("*.xlsx")} == docs


def test_upload_em_origem_unida_rejeita_antes_de_gravar_documento(tenants):
    arquivo = tenants / "cliente-a" / "listas_pedidos.json"
    lista = _salvo(tenants)
    lista["unida_em_lista_id"] = "lista-consolidada"
    arquivo.write_text(json.dumps([lista]), encoding="utf-8")
    antes = arquivo.read_bytes()
    with pytest.raises(HTTPException) as erro:
        packing._importar_packing_list_sync("lista-1", _modelo(), "packing.xlsx", "cliente-a")
    assert erro.value.status_code == 409
    assert arquivo.read_bytes() == antes
    assert not (arquivo.parent / "packing_lists").exists()


def test_upload_mesmo_id_em_tenants_distintos_fica_isolado(tenants):
    packing._importar_packing_list_sync("lista-1", _modelo(), "same.xlsx", "cliente-a")
    packing._importar_packing_list_sync("lista-1", _modelo(fotos=False), "same.xlsx", "cliente-b")
    a = list((tenants / "cliente-a" / "packing_lists").glob("*.xlsx"))
    b = list((tenants / "cliente-b" / "packing_lists").glob("*.xlsx"))
    assert len(a) == len(b) == 1
    assert a[0].name != b[0].name
    assert a[0].read_bytes() != b[0].read_bytes()


@pytest.mark.parametrize("lista_id", ["does-not-exist", "../cliente-b/lista-1"])
def test_upload_lista_inexistente_nao_grava_arquivo(tenants, lista_id):
    before = (tenants / "cliente-a" / "listas_pedidos.json").read_bytes()
    with pytest.raises(HTTPException) as erro:
        packing._importar_packing_list_sync(lista_id, _modelo(), "packing.xlsx", "cliente-a")
    assert erro.value.status_code == 404
    assert (tenants / "cliente-a" / "listas_pedidos.json").read_bytes() == before
    assert not list(tenants.rglob("*.xlsx"))


def test_upload_filename_nao_controla_destino(tenants):
    packing._importar_packing_list_sync("lista-1", _modelo(), "../../cliente-b/escape.xlsx", "cliente-a")
    docs = list(tenants.rglob("*.xlsx"))
    assert len(docs) == 1
    assert docs[0].parent == tenants / "cliente-a" / "packing_lists"
    assert docs[0].name != "escape.xlsx"


def test_uploads_concorrentes_mesma_lista_apenas_um_persiste(tenants):
    barrier = Barrier(2)
    contents = [_modelo(), _modelo(fotos=False)]
    def importar(indice):
        barrier.wait(timeout=10)
        try:
            packing._importar_packing_list_sync("lista-1", contents[indice], f"packing-{indice}.xlsx", "cliente-a")
            return indice, 200
        except HTTPException as error:
            return indice, error.status_code
    with ThreadPoolExecutor(max_workers=2) as executor:
        resultados = list(executor.map(importar, (0, 1)))
    assert sorted(code for _, code in resultados) == [200, 409]
    vencedor = next(indice for indice, code in resultados if code == 200)
    docs = list((tenants / "cliente-a" / "packing_lists").glob("*.xlsx"))
    assert len(docs) == 1
    assert docs[0].read_bytes() == contents[vencedor]
    assert _salvo(tenants)["packing_list"]["sha256"] == hashlib.sha256(contents[vencedor]).hexdigest()


def test_download_sem_documento_retora_404(tenants):
    with pytest.raises(HTTPException) as erro:
        packing._baixar_packing_list_sync("lista-1", "cliente-a")
    assert erro.value.status_code == 404


def test_documento_invalido_nao_altera_lista_nem_persiste(tenants):
    path = tenants / "cliente-a" / "listas_pedidos.json"
    before = path.read_bytes()
    with pytest.raises(HTTPException):
        packing._importar_packing_list_sync("lista-1", b"bad input", "packing.xlsx", "cliente-a")
    assert path.read_bytes() == before
    assert not list(tenants.rglob("*.xlsx"))


def test_exportacao_busca_foto_atual_por_sku_e_ignora_imagens_do_fornecedor(tenants, monkeypatch):
    pasta = tenants / "cliente-a" / "cadastro_fotos"
    pasta.mkdir()
    (pasta / "001.png").write_bytes(_foto("purple", (100, 40)))
    (pasta / "002.png").write_bytes(_foto("black", (40, 100)))
    produtos = [
        {"sku": "001", "cg_product description": "Current sensor, material and dimensions", "foto": "001.png"},
        {"sku": "002", "cg_product description": "Current cable, cardboard box", "foto": "002.png"},
    ]
    monkeypatch.setattr(cadastro_compatibilidade, "visao_produtos_cadastro_contexto_loja",
                        lambda *_args, **_kwargs: {"scope": "store", "store_id": "store-a", "produtos": produtos})
    resolver = lambda _client, _sku, foto="", *_args, **_kwargs: foto
    monkeypatch.setattr(common, "_resolver_foto_cadastro_sku", resolver)
    monkeypatch.setattr(excel, "_resolver_foto_cadastro_sku", resolver)
    dados = packing._ler_packing_list_bytes(_modelo())
    antes = copy.deepcopy(dados)
    sheet = _sheet(packing._gerar_packing_list_bytes(_lista(), dados, client_id="cliente-a"))
    assert [sheet[f"D{row}"].value for row in range(9, 13)] == [
        "Current sensor, material and dimensions", "Current sensor, material and dimensions",
        "Current cable, cardboard box", None,
    ]
    assert len(sheet._images) == 3
    assert [image.anchor._from.row + 1 for image in sheet._images] == [9, 10, 11]
    for image, cor in zip(sheet._images, ("purple", "purple", "black")):
        assert Image.open(io.BytesIO(image._data())).getpixel((0, 0)) == Image.new("RGB", (1, 1), cor).getpixel((0, 0))
    assert dados == antes
    assert sheet["F13"].value == "=SUM(F9:F12)"
    assert _sheet(packing._gerar_packing_list_bytes(_lista(), dados, client_id="cliente-a"), data_only=True)["F13"].value == 19
    produtos[0]["cg_product description"] = "New registered technical description"
    (pasta / "001.png").write_bytes(_foto("orange"))
    atualizado = _sheet(packing._gerar_packing_list_bytes(_lista(), dados, client_id="cliente-a"))
    assert atualizado["D9"].value == "New registered technical description"
    assert Image.open(io.BytesIO(atualizado._images[0]._data())).getpixel((0, 0)) == Image.new("RGB", (1, 1), "orange").getpixel((0, 0))


def test_exportacao_reconhece_sku_com_codigo_peca_apos_dois_pontos(tenants, monkeypatch):
    literal = "SKU:496-K:71198-T7A-000 71193-T7A-000"
    pasta = tenants / "cliente-a" / "cadastro_fotos"
    pasta.mkdir()
    (pasta / "496-K.png").write_bytes(_foto("green"))
    descricao = "Front Bumper Retainer Bracket (Pair), Plastic, cardboard box"
    produto = {"sku": "496-K", "store_id": "store-a", "cg_product description": descricao,
               "foto": "496-K.png", "Unit/Pair": "Pair"}

    def visao(client_id, store_id, *, skus, _normalizar_sku_consulta):
        assert (client_id, store_id) == ("cliente-a", "store-a")
        produtos = [produto] if _normalizar_sku_consulta(produto["sku"]) in skus else []
        return {"scope": "store", "store_id": store_id, "produtos": produtos}

    monkeypatch.setattr(cadastro_compatibilidade, "visao_produtos_cadastro_contexto_loja", visao)
    resolver = lambda _client, _sku, foto="", *_args, **_kwargs: foto
    monkeypatch.setattr(common, "_resolver_foto_cadastro_sku", resolver)
    monkeypatch.setattr(excel, "_resolver_foto_cadastro_sku", resolver)

    def alterar(sheet):
        sheet["A4"] = "16"
        sheet["B5"] = literal
        sheet["D5"] = 30
        sheet["E4"] = 32

    dados = packing._ler_packing_list_bytes(_modelo(alterar=alterar))
    antes = copy.deepcopy(dados)
    sheet = _sheet(packing._gerar_packing_list_bytes(_lista(), dados, client_id="cliente-a"))

    assert sheet["B11"].value == literal
    assert sheet["D11"].value == descricao
    assert sheet["E11"].value == "Pair"
    assert sheet["G11"].value == 30
    assert sheet["A10"].value == "16"
    assert sheet["F10"].value == 32
    assert sheet["H10"].value == 1
    assert {"A10:A11", "F10:F11", "H10:H11"}.issubset(set(map(str, sheet.merged_cells.ranges)))
    assert len(sheet._images) == 1
    assert sheet._images[0].anchor._from.row + 1 == 11
    assert Image.open(io.BytesIO(sheet._images[0]._data())).getpixel((0, 0)) == (0, 128, 0)
    assert dados == antes


def test_exportacao_sem_cadastro_nao_usa_descricao_ou_foto_externas(tenants, monkeypatch):
    monkeypatch.setattr(cadastro_compatibilidade, "visao_produtos_cadastro_contexto_loja",
                        lambda *_args, **_kwargs: {"scope": "unresolved", "store_id": "", "produtos": []})
    sheet = _sheet(packing._gerar_packing_list_bytes(_lista(), packing._ler_packing_list_bytes(_modelo()), client_id="cliente-a"))
    assert all(sheet[f"D{row}"].value is None for row in range(9, 13))
    assert len(sheet._images) == 0
    assert sheet["F9"].value == 4


def _response_bytes(response):
    if hasattr(response, "body"):
        return response.body

    async def ler():
        return b"".join([chunk async for chunk in response.body_iterator])
    return asyncio.run(ler())


def test_download_usa_invoice_e_enriquecimento_atuais_sem_mudar_original(tenants, monkeypatch):
    original = _modelo()
    packing._importar_packing_list_sync("lista-1", original, "packing.xlsx", "cliente-a")
    path = tenants / "cliente-a" / "listas_pedidos.json"
    lista = _salvo(tenants)
    lista["numero_invoice"] = "UPDATED-2026"
    lista["order_date"] = "2026-09-02"
    lista["itens"][0]["Título do produto em inglês"] = "Updated sensor description"
    lista["itens"][0]["Color/side"] = "Left"
    path.write_text(json.dumps([lista]), encoding="utf-8")
    monkeypatch.setattr(cadastro_compatibilidade, "visao_produtos_cadastro_contexto_loja",
                        lambda *_args, **_kwargs: {"store_id": "store-a", "scope": "store", "produtos": [
                            {"sku": "001", "cg_product description": "Current registered sensor", "Unit/Pair": "Unit"},
                            {"sku": "002", "cg_product description": "Cable description", "Unit/Pair": "Pair"},
                        ]})
    before = path.read_bytes()
    response = packing._baixar_packing_list_sync("lista-1", "cliente-a")
    conteudo = _response_bytes(response)
    sheet = _sheet(conteudo)
    calculado = _sheet(conteudo, data_only=True)
    for endereco, esperado in {"F13": 19, "H13": 4, "M13": .033, "O13": 7.6, "Q13": 9.1}.items():
        assert calculado[endereco].value == pytest.approx(esperado)
    assert "UPDATED-2026" in sheet["A3"].value
    assert "September 2, 2026" in sheet["I3"].value
    assert sheet["D9"].value == "Current registered sensor"
    assert sheet["E9"].value == "Unit"
    assert "UPDATED-2026" in response.headers["content-disposition"]
    assert path.read_bytes() == before
    assert next((path.parent / "packing_lists").glob("*.xlsx")).read_bytes() == original


def test_download_envia_xlsx_pronto_em_corpo_unico_com_tamanho_exato(tenants):
    packing._importar_packing_list_sync("lista-1", _modelo(), "packing.xlsx", "cliente-a")
    response = packing._baixar_packing_list_sync("lista-1", "cliente-a")
    conteudo = _response_bytes(response)
    mensagens = []

    async def receber():
        pytest.fail("O download pronto nao deve aguardar dados da requisicao.")

    async def enviar(mensagem):
        mensagens.append(mensagem)

    async def executar():
        await response({"type": "http", "asgi": {"spec_version": "2.4"}}, receber, enviar)

    asyncio.run(executar())
    assert [mensagem["type"] for mensagem in mensagens] == ["http.response.start", "http.response.body"]
    assert mensagens[1]["body"] == conteudo
    assert not mensagens[1].get("more_body", False)
    headers = dict(mensagens[0]["headers"])
    assert headers[b"content-length"] == str(len(conteudo)).encode("ascii")
    assert headers[b"content-type"] == b"application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    assert headers[b"content-disposition"] == b'attachment; filename="Packing_List_TEST-14072026.xlsx"'
    assert _sheet(conteudo)["A2"].value == "PACKING LIST"


def test_download_rejeita_documento_corrompido(tenants, monkeypatch):
    packing._importar_packing_list_sync("lista-1", _modelo(), "packing.xlsx", "cliente-a")
    path = next((tenants / "cliente-a" / "packing_lists").glob("*.xlsx"))
    path.write_bytes(b"Corrupted original")
    monkeypatch.setattr(packing, "_gerar_packing_list_bytes", lambda *_args, **_kwargs: pytest.fail("Corrupted bytes must not reach exporter"))
    with pytest.raises(HTTPException) as erro:
        packing._baixar_packing_list_sync("lista-1", "cliente-a")
    assert erro.value.status_code == 409


def test_download_sem_original_nao_altera_lista(tenants):
    packing._importar_packing_list_sync("lista-1", _modelo(), "packing.xlsx", "cliente-a")
    path = tenants / "cliente-a" / "listas_pedidos.json"
    before = path.read_bytes()
    next((path.parent / "packing_lists").glob("*.xlsx")).unlink()
    with pytest.raises(HTTPException) as erro:
        packing._baixar_packing_list_sync("lista-1", "cliente-a")
    assert erro.value.status_code == 404
    assert path.read_bytes() == before


def test_download_documento_id_nao_permite_outro_tenant(tenants):
    packing._importar_packing_list_sync("lista-1", _modelo(), "packing.xlsx", "cliente-b")
    estrangeiro = _salvo(tenants, "cliente-b")["packing_list"]
    lista = _salvo(tenants)
    lista["packing_list"] = {**estrangeiro, "documento_id": "../cliente-b/packing_lists/" + estrangeiro["documento_id"]}
    (tenants / "cliente-a" / "listas_pedidos.json").write_text(json.dumps([lista]), encoding="utf-8")
    with pytest.raises(HTTPException) as erro:
        packing._baixar_packing_list_sync("lista-1", "cliente-a")
    assert erro.value.status_code == 404
    assert len(list(tenants.rglob("*.xlsx"))) == 1


def test_upload_falha_antes_commit_limpa_blob_e_preserva_lista(tenants, monkeypatch):
    path = tenants / "cliente-a" / "listas_pedidos.json"
    before = path.read_bytes()
    def falhar(*_args, **_kwargs):
        raise OSError("Synthetic disk failure before commit")
    monkeypatch.setattr(common, "_salvar_listas_pedidos", falhar)
    with pytest.raises(OSError):
        packing._importar_packing_list_sync("lista-1", _modelo(), "packing.xlsx", "cliente-a")
    assert path.read_bytes() == before
    assert not list(tenants.rglob("*.xlsx"))
    assert not list(tenants.rglob(".packing-*.tmp"))


def test_upload_falha_apos_commit_preserva_original_referenciado_e_baixavel(tenants, monkeypatch):
    salvar = common._salvar_listas_pedidos
    def falhar_apos_commit(*args, **kwargs):
        salvar(*args, **kwargs)
        raise OSError("Synthetic failure after canonical commit")
    monkeypatch.setattr(common, "_salvar_listas_pedidos", falhar_apos_commit)
    original = _modelo()
    with pytest.raises(OSError):
        packing._importar_packing_list_sync("lista-1", original, "packing.xlsx", "cliente-a")
    salvo = _salvo(tenants)
    assert salvo["packing_list"]["sha256"] == hashlib.sha256(original).hexdigest()
    docs = list((tenants / "cliente-a" / "packing_lists").glob("*.xlsx"))
    assert len(docs) == 1
    assert docs[0].read_bytes() == original
    assert _sheet(_response_bytes(packing._baixar_packing_list_sync("lista-1", "cliente-a")))["A2"].value == "PACKING LIST"


@pytest.mark.parametrize("nome", ["packing.xls", "packing.xlsm", "packing.pdf", ""])
def test_upload_extensao_nao_suportada_nao_persiste(tenants, nome):
    with pytest.raises(HTTPException) as erro:
        packing._importar_packing_list_sync("lista-1", _modelo(), nome, "cliente-a")
    assert erro.value.status_code == 400
    assert not list(tenants.rglob("*.xlsx"))


def test_cabecalho_comprador_usa_identidade_exata_da_loja_e_tenant(monkeypatch):
    calls = []
    def contexto(client_id, referencia, **_kwargs):
        calls.append(("scope", client_id, referencia))
        return {"store_id": referencia, "scope": "store"}
    def comprador(client_id, store_id):
        calls.append(("buyer", client_id, store_id))
        return {"nome_empresa": f"Buyer {client_id} {store_id}", "tax_id": "000001"}
    monkeypatch.setattr(cadastro_compatibilidade, "visao_produtos_cadastro_contexto_loja", contexto)
    monkeypatch.setattr(cadastro_importador_loja, "ler_importador_loja_por_store_id", comprador)
    for tenant, store in (("cliente-a", "store-a"), ("cliente-b", "store-b")):
        lista = _lista(store_id=store, loja="Same display name")
        sheet = _sheet(packing._gerar_packing_list_bytes(lista, packing._ler_packing_list_bytes(_modelo()), client_id=tenant))
        assert sheet["I4"].value == f"Buyer: Buyer {tenant} {store}\nTAX ID: 000001"
    assert calls == [("scope", "cliente-a", "store-a"), ("buyer", "cliente-a", "store-a"),
                     ("scope", "cliente-b", "store-b"), ("buyer", "cliente-b", "store-b")]


def test_loja_nao_resolvida_nao_busca_outro_cadastro(monkeypatch):
    monkeypatch.setattr(cadastro_compatibilidade, "visao_produtos_cadastro_contexto_loja", lambda *_args, **_kwargs: {"store_id": "", "scope": "unresolved"})
    monkeypatch.setattr(cadastro_importador_loja, "ler_importador_loja_por_store_id", lambda *_args: pytest.fail("Unresolved store must not trigger buyer lookup"))
    sheet = _sheet(packing._gerar_packing_list_bytes(_lista(), packing._ler_packing_list_bytes(_modelo()), client_id="cliente-a"))
    assert sheet["I4"].value == "STORE: Loja A"


def _xlsx_com_xml_alterado(conteudo, antigo, novo):
    output = io.BytesIO()
    with zipfile.ZipFile(io.BytesIO(conteudo)) as source, zipfile.ZipFile(output, "w", zipfile.ZIP_DEFLATED) as target:
        for info in source.infolist():
            content = source.read(info.filename)
            if info.filename == "xl/worksheets/sheet1.xml":
                assert antigo in content
                content = content.replace(antigo, novo, 1)
            target.writestr(info, content)
    return output.getvalue()


@pytest.mark.parametrize("antigo,novo", [
    (b'ref="A4:A5"', b'ref="A4:A1048576"'),
    (b'ref="A1:O7"', b'ref="A1:XFD1048576"'),
])
def test_parser_limita_dimensao_e_mescla_antes_carregar_excel(monkeypatch, antigo, novo):
    conteudo = _xlsx_com_xml_alterado(_modelo(fotos=False), antigo, novo)
    monkeypatch.setattr(packing.openpyxl, "load_workbook", lambda *_args, **_kwargs: pytest.fail("Oversized XML must be rejected before workbook load"))
    with pytest.raises(HTTPException) as erro:
        packing._ler_packing_list_bytes(conteudo)
    assert erro.value.status_code == 400


def test_pasta_packing_symlink_nao_permite_acesso_a_outro_tenant(tenants):
    packing._importar_packing_list_sync("lista-1", _modelo(), "packing.xlsx", "cliente-b")
    target = tenants / "cliente-b" / "packing_lists"
    link = tenants / "cliente-a" / "packing_lists"
    try:
        os.symlink(target, link, target_is_directory=True)
    except OSError:
        if os.name != "nt":
            raise
        quote = lambda value: "'" + str(value).replace("'", "''") + "'"
        command = f"New-Item -ItemType Junction -Path {quote(link)} -Target {quote(target)} | Out-Null"
        subprocess.run(["powershell.exe", "-NoProfile", "-Command", command], check=True, capture_output=True)
    lista = _salvo(tenants)
    lista["packing_list"] = _salvo(tenants, "cliente-b")["packing_list"]
    path = tenants / "cliente-a" / "listas_pedidos.json"
    path.write_text(json.dumps([lista]), encoding="utf-8")
    estrangeiro = {p.name: p.read_bytes() for p in target.glob("*.xlsx")}
    with pytest.raises(HTTPException) as erro:
        packing._baixar_packing_list_sync("lista-1", "cliente-a")
    assert erro.value.status_code in (400, 404)
    # A fresh import must not use the linked foreign directory either.
    lista.pop("packing_list")
    path.write_text(json.dumps([lista]), encoding="utf-8")
    with pytest.raises(HTTPException) as erro:
        packing._importar_packing_list_sync("lista-1", _modelo(fotos=False), "packing.xlsx", "cliente-a")
    assert erro.value.status_code in (400, 404)
    assert {p.name: p.read_bytes() for p in target.glob("*.xlsx")} == estrangeiro


def test_http_fluxo_sem_documento_upload_resumo_e_download_atual(tenants, monkeypatch):
    monkeypatch.setattr(listas_service, "_construir_mapa_m3_sku", lambda *_args, **_kwargs: {})
    app = FastAPI()
    app.include_router(create_medias_compras_router())
    app.dependency_overrides[common.get_tenant_id] = lambda: "cliente-a"
    rota = "/api/medias-compras/listas-pedidos/lista-1/packing-list"
    itens = copy.deepcopy(_salvo(tenants)["itens"])
    with TestClient(app) as client:
        assert client.get(rota).status_code == 404
        before = client.get("/api/medias-compras/listas-pedidos")
        assert before.status_code == 200, before.text
        assert before.json()["listas"][0]["packing_list_ready"] is False
        uploaded = client.post(
            rota + "/importar-excel",
            files={"file": ("packing.xlsx", _modelo(), "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
        )
        assert uploaded.status_code == 200, uploaded.text
        assert uploaded.json()["packing_list_ready"] is True
        after = client.get("/api/medias-compras/listas-pedidos")
        assert after.status_code == 200, after.text
        assert after.json()["listas"][0]["packing_list_ready"] is True
        lista = _salvo(tenants)
        assert lista["itens"] == itens
        lista["numero_invoice"] = "HTTP-CURRENT-2026"
        (tenants / "cliente-a" / "listas_pedidos.json").write_text(json.dumps([lista]), encoding="utf-8")
        downloaded = client.get(rota)
        assert downloaded.status_code == 200, downloaded.text
        assert downloaded.headers["content-type"] == "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
        assert "HTTP-CURRENT-2026" in downloaded.headers["content-disposition"]
        sheet = _sheet(downloaded.content)
        assert "HTTP-CURRENT-2026" in sheet["A3"].value
        assert sheet["D9"].value == "Sensor description"
        assert sheet["E9"].value == "Unit"
        assert _salvo(tenants)["itens"] == itens


def _modelo_caixas_trays(*, primeira=57, segunda="3 Trays", terceira=1, separar_carton=False):
    def alterar(sheet):
        sheet["F3"] = primeira
        sheet["F4"] = segunda
        sheet["F6"] = terceira
        if separar_carton:
            sheet.unmerge_cells("A4:A5")

    return _modelo(fotos=False, alterar=alterar)


def test_packing_caixas_trays_parser_preserva_tipo_e_total_misto_recalculavel():
    dados = packing._ler_packing_list_bytes(_modelo_caixas_trays())
    assert [grupo["ctns"] for grupo in dados["grupos"]] == [57, 3, 1]
    assert [grupo["tipo_embalagem"] for grupo in dados["grupos"]] == ["ctns", "trays", "ctns"]
    original = copy.deepcopy(dados)
    conteudo = packing._gerar_packing_list_bytes(_lista(), dados)
    formulas = _sheet(conteudo)
    calculado = _sheet(conteudo, data_only=True)

    assert formulas["H9"].value == 57
    assert formulas["H10"].value == "3 Trays"
    assert formulas["H11"].value is None
    assert formulas["H12"].value == 1
    assert formulas["H13"].data_type == "f"
    assert formulas["H13"].value.ref == "H13"
    assert "H9" in formulas["H13"].value.text
    assert "H12" in formulas["H13"].value.text
    assert calculado["H13"].value == "58 CTNS + 3 Trays"
    assert calculado["H13"].data_type == "s"
    assert formulas.parent.calculation.fullCalcOnLoad is True
    assert calculado["F13"].value == 19
    assert dados == original


@pytest.mark.parametrize("texto", ["3 Tray", "3 trays", "  3   TRAYS  ", "3 TrAyS"])
def test_packing_caixas_trays_aceita_variacoes_de_tipo(texto):
    dados = packing._ler_packing_list_bytes(_modelo_caixas_trays(segunda=texto))
    grupo = dados["grupos"][1]
    assert grupo["ctns"] == 3
    assert grupo["tipo_embalagem"] == "trays"
    calculado = _sheet(packing._gerar_packing_list_bytes(_lista(), dados), data_only=True)
    assert calculado["H10"].value == "3 Trays"
    assert calculado["H13"].value == "58 CTNS + 3 Trays"


@pytest.mark.parametrize("texto", [
    "0 Trays", "-3 Trays", "1.5 Trays", "1,5 Trays", "Trays",
    "3 pallets", "abc3", "3 Trays extra", "=3 Trays", "2+1 Trays",
    "2+1", "=3", "3箱",
])
def test_packing_caixas_trays_rejeita_quantidade_ou_tipo_invalido(texto):
    with pytest.raises(HTTPException) as erro:
        packing._ler_packing_list_bytes(_modelo_caixas_trays(segunda=texto))
    assert erro.value.status_code == 400


@pytest.mark.parametrize("valor", [3, "3", "3.0", "3,0", "+3"])
def test_packing_caixas_trays_preserva_formatos_numericos_de_caixas(valor):
    dados = packing._ler_packing_list_bytes(_modelo_caixas_trays(segunda=valor))
    assert dados["grupos"][1]["ctns"] == 3
    assert dados["grupos"][1]["tipo_embalagem"] == "ctns"
    conteudo = packing._gerar_packing_list_bytes(_lista(), dados)
    assert _sheet(conteudo)["H13"].value == "=SUM(H9:H12)"
    assert _sheet(conteudo, data_only=True)["H13"].value == 61


def test_packing_caixas_trays_duas_celulas_distintas_contam_ambas():
    dados = packing._ler_packing_list_bytes(_modelo_caixas_trays(primeira=58, terceira="3 Trays"))
    trays = [grupo for grupo in dados["grupos"] if grupo["tipo_embalagem"] == "trays"]
    assert len(trays) == 2
    assert trays[0]["ctns_origem"] != trays[1]["ctns_origem"]
    conteudo = packing._gerar_packing_list_bytes(_lista(), dados)
    calculado = _sheet(conteudo, data_only=True)
    assert calculado["H10"].value == calculado["H12"].value == "3 Trays"
    assert calculado["H13"].value == "58 CTNS + 6 Trays"


def test_packing_caixas_trays_apenas_trays_preserva_total_textual():
    def alterar(sheet):
        sheet["F3"] = "3 Trays"
        for linha in (4, 5, 6):
            sheet.cell(linha, 2).value = None

    dados = packing._ler_packing_list_bytes(_modelo(fotos=False, alterar=alterar))
    assert len(dados["grupos"]) == 1
    conteudo = packing._gerar_packing_list_bytes(_lista(), dados)
    calculado = _sheet(conteudo, data_only=True)
    assert calculado["H9"].value == "3 Trays"
    assert calculado["H10"].value == "3 Trays"
    assert calculado["H10"].data_type == "s"
    assert "CTNS" not in calculado["H10"].value


@pytest.mark.parametrize("quantidade,total", [(1, 59), ("3 Trays", "58 CTNS + 3 Trays")])
def test_packing_caixas_trays_mesma_ancora_ctns_nao_conta_duas_vezes(quantidade, total):
    dados = packing._ler_packing_list_bytes(_modelo_caixas_trays(segunda=quantidade, separar_carton=True))
    assert len(dados["grupos"]) == 4
    assert dados["grupos"][1]["ctns_origem"] == dados["grupos"][2]["ctns_origem"]
    conteudo = packing._gerar_packing_list_bytes(_lista(), dados)
    calculado = _sheet(conteudo, data_only=True)
    assert calculado["H13"].value == total
    valores = [calculado[f"H{linha}"].value for linha in range(9, 13)]
    assert valores[1] == quantidade
    assert valores[2] is None
    assert calculado["F13"].value == 19


@pytest.mark.parametrize("segunda,total", [(1, "59 CTNS + 3 Trays"), ("3 Trays", "58 CTNS + 6 Trays")])
def test_packing_caixas_trays_mesmo_carton_ctns_distintos_contam_separadamente(segunda, total):
    def alterar(sheet):
        sheet["F3"] = 57
        sheet.unmerge_cells("F4:F5")
        sheet["F4"] = segunda
        sheet["F5"] = "3 Trays"

    dados = packing._ler_packing_list_bytes(_modelo(fotos=False, alterar=alterar))
    assert len(dados["grupos"]) == 4
    assert dados["grupos"][1]["caixas"] == dados["grupos"][2]["caixas"] == "3"
    assert dados["grupos"][1]["ctns_origem"] != dados["grupos"][2]["ctns_origem"]
    calculado = _sheet(packing._gerar_packing_list_bytes(_lista(), dados), data_only=True)
    assert calculado["H10"].value == segunda
    assert calculado["H11"].value == "3 Trays"
    assert calculado["H13"].value == total
    assert calculado["F13"].value == 19


def test_packing_caixas_trays_importacao_guarda_totais_separados_e_original(tenants):
    original = _modelo_caixas_trays(separar_carton=True)
    antes = _salvo(tenants)
    packing._importar_packing_list_sync("lista-1", original, "packing-trays.xlsx", "cliente-a")
    depois = _salvo(tenants)
    assert depois["packing_list"]["total_ctns"] == 58
    assert depois["packing_list"]["total_trays"] == 3
    for chave, valor in antes.items():
        assert depois[chave] == valor
    documento = next((tenants / "cliente-a" / "packing_lists").glob("*.xlsx"))
    assert documento.read_bytes() == original
    assert depois["packing_list"]["sha256"] == hashlib.sha256(original).hexdigest()


def test_packing_caixas_trays_download_historico_recalcula_sem_mudar_lista_ou_original(tenants):
    original = _modelo_caixas_trays()
    packing._importar_packing_list_sync("lista-1", original, "packing-antigo.xlsx", "cliente-a")
    path = tenants / "cliente-a" / "listas_pedidos.json"
    lista = _salvo(tenants)
    lista["packing_list"]["total_ctns"] = 61
    lista["packing_list"].pop("total_trays", None)
    path.write_text(json.dumps([lista]), encoding="utf-8")
    lista_antes = path.read_bytes()
    documento = next((path.parent / "packing_lists").glob("*.xlsx"))
    bytes_antes = documento.read_bytes()

    resposta = packing._baixar_packing_list_sync("lista-1", "cliente-a")
    calculado = _sheet(_response_bytes(resposta), data_only=True)

    assert calculado["H13"].value == "58 CTNS + 3 Trays"
    assert calculado["H10"].value == "3 Trays"
    assert path.read_bytes() == lista_antes
    assert documento.read_bytes() == bytes_antes == original


def test_packing_caixas_trays_totais_mesclados_com_ctns_distintos_contam_uma_vez(tenants):
    def alterar(sheet):
        sheet.unmerge_cells("F4:F5")
        sheet["F4"] = 1
        sheet["F5"] = "3 Trays"
        sheet["K4"] = .06
        sheet["N4"] = 3.2
        sheet["O4"] = 3.6

    original = _modelo(fotos=False, alterar=alterar)
    packing._importar_packing_list_sync("lista-1", original, "packing-totais-mesclados.xlsx", "cliente-a")
    path = tenants / "cliente-a" / "listas_pedidos.json"
    lista_antes = path.read_bytes()
    documento = next((path.parent / "packing_lists").glob("*.xlsx"))
    conteudo = _response_bytes(packing._baixar_packing_list_sync("lista-1", "cliente-a"))
    calculado = _sheet(conteudo, data_only=True)

    assert calculado["F13"].value == 19
    assert calculado["H13"].value == "4 CTNS + 3 Trays"
    assert calculado["M13"].value == pytest.approx(.073)
    assert calculado["O13"].value == pytest.approx(7.6)
    assert calculado["Q13"].value == pytest.approx(9.1)
    assert {"M10:M11", "O10:O11", "Q10:Q11"}.issubset(set(map(str, calculado.merged_cells.ranges)))
    assert path.read_bytes() == lista_antes
    assert documento.read_bytes() == original
