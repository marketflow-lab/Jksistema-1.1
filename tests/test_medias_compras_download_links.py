"""Saved purchase-list exports read current links without rewriting the order."""

import asyncio
import copy
import io
from pathlib import Path
from types import SimpleNamespace

import pytest
from fastapi import HTTPException
from openpyxl import load_workbook

from backend.services import cadastro_compatibilidade as cadastro
from backend.services import medias_compras_common as common
from backend.services import medias_compras_excel as excel
from backend.services import medias_compras_listas as listas


SAVED_LINK = "https://supplier.example/saved-product"
CURRENT_LINK = "https://supplier.example/current-product?variant=503-1&quantity=2"


def _lista(*, lista_id="saved-list", store_id="store-a", loja="Nome antigo da loja"):
    return {
        "id": lista_id,
        "nome_lista": "Pedido salvo",
        "store_id": store_id,
        "loja": loja,
        "created_at": "2026-10-01T09:00:00",
        "updated_at": "2026-10-20T10:00:00",
        "status": "Pedido Aprovado",
        "preferencias": {"colunas": ["SKU", "Link", "Quantidade"]},
        "itens": [{
            "SKU": "503-1",
            "Link": SAVED_LINK,
            "Foto": "foto-salva.png",
            "Título do produto em inglês": "Saved supplier description",
            "OEM": "OEM-503-1",
            "Color/side": "Left / black",
            "Quantidade": 7,
            "Valor unidade": 12.5,
            "Valor total": 87.5,
            "Estimed CBM": "0.012",
            "Estimed Weigh": "0.45",
            "Individual packaging": "Carton",
            "compra_aprovada": True,
            "compra_aprovada_em": "2026-10-19T14:00:00",
            "analise": {"aprovador": "Comprador", "observacoes": ["Manter quantidade"]},
        }],
    }


def _contexto(produtos, *, store_id="store-a", **overrides):
    return {
        "scope": "store",
        "loja_resolvida": True,
        "store_id": store_id,
        "produtos": produtos,
        **overrides,
    }


def _produto(link=CURRENT_LINK, *, sku="503-1", store_id="store-a", **extra):
    return {"sku": sku, "store_id": store_id, "link": link, **extra}


@pytest.fixture(params=[
    "_api_medias_compras_lista_pedido_download_sync",
    "_api_medias_compras_lista_pedido_gerar_download_sync",
], ids=["download", "gerar-download"])
def endpoint(request):
    return getattr(listas, request.param)


@pytest.fixture
def ambiente(tmp_path, monkeypatch):
    listas_por_cliente = {"tenant-a": [_lista()]}
    contextos = {("tenant-a", "store-a"): _contexto([_produto()])}
    consultas, geracoes, caminhos_cache = [], [], []
    memoria, temporarios = {}, {}
    arquivo_cache = common._arquivo_cache_lista_pedido

    monkeypatch.setattr(common, "get_tenant_path", lambda client_id: str(tmp_path / client_id))
    monkeypatch.setattr(listas, "_carregar_listas_pedidos", lambda client_id: listas_por_cliente[client_id])
    monkeypatch.setattr(
        listas, "_salvar_listas_pedidos",
        lambda *_args, **_kwargs: pytest.fail("Baixar a lista nao deve persistir alteracoes"),
    )
    monkeypatch.setattr(listas, "LISTA_PEDIDO_XLSX_CACHE", memoria)
    monkeypatch.setattr(listas, "LISTA_PEDIDO_XLSX_CACHE_MAX_ITENS", 100)
    monkeypatch.setattr(listas, "TEMP_FILES_STORAGE", temporarios)
    monkeypatch.setattr(listas, "TEMP_FILES_META", {})
    monkeypatch.setattr(excel, "_resolver_caminho_foto_cadastro_seguro", lambda *_args: None)

    def consultar(client_id, referencia, *, skus=None):
        # The real XLSX renderer also resolves photo scope, without a SKU filter.
        # Keep that existing lookup separate from the link-refresh contract.
        if skus is None:
            return {"store_id": "store-a", "scope": "store", "loja_resolvida": True}
        consultas.append((client_id, referencia, tuple(skus)))
        contexto = contextos[(client_id, referencia)]
        if isinstance(contexto, Exception):
            raise contexto
        return contexto

    def caminho(client_id, lista_id, versao):
        result = Path(arquivo_cache(client_id, lista_id, versao))
        caminhos_cache.append(result)
        return str(result)

    def gerar(nome_lista, itens, **kwargs):
        geracoes.append((nome_lista, itens, kwargs))
        return excel._gerar_excel_lista_pedido_bytes(nome_lista, itens, **kwargs)

    monkeypatch.setattr(cadastro, "visao_produtos_cadastro_contexto_loja", consultar)
    monkeypatch.setattr(listas, "_arquivo_cache_lista_pedido", caminho)
    monkeypatch.setattr(listas, "_gerar_excel_lista_pedido_bytes", gerar)
    return SimpleNamespace(
        root=tmp_path, listas=listas_por_cliente, contextos=contextos, consultas=consultas,
        geracoes=geracoes, caminhos_cache=caminhos_cache, memoria=memoria, temporarios=temporarios,
    )


def _baixar(ambiente, endpoint, *, client_id="tenant-a", lista_id="saved-list"):
    resposta = endpoint(lista_id, client_id=client_id)
    if isinstance(resposta, dict):
        assert resposta["success"] is True
        assert resposta["lista_id"] == lista_id
        if "file_id" in resposta:
            return ambiente.temporarios[resposta["file_id"]]
        assert resposta["download_url"] == f"/api/medias-compras/listas-pedidos/{lista_id}/download"
        # Read the exact disk cache selected by gerar-download, rather than
        # calling the other endpoint and potentially hiding its stale-cache bug.
        return ambiente.caminhos_cache[-1].read_bytes()

    async def ler():
        return b"".join([chunk async for chunk in resposta.body_iterator])

    return asyncio.run(ler())


def _planilha(conteudo):
    workbook = load_workbook(io.BytesIO(conteudo), data_only=False)
    assert workbook.sheetnames == ["Import 54"]
    return workbook.active


def _conferir_link(conteudo, esperado, celula="F2"):
    sheet = _planilha(conteudo)
    assert sheet[celula].value == esperado
    assert sheet[celula].hyperlink is not None
    assert sheet[celula].hyperlink.target == esperado
    return sheet


@pytest.mark.parametrize("cache_antigo", ["disco", "memoria", "ambos"])
def test_cadastro_atual_invalida_cache_sem_editar_lista(ambiente, endpoint, cache_antigo):
    lista = ambiente.listas["tenant-a"][0]
    original = copy.deepcopy(lista)
    itens_originais, item_original = lista["itens"], lista["itens"][0]
    bytes_antigos = excel._gerar_excel_lista_pedido_bytes(lista["nome_lista"], lista["itens"])
    versao_antiga = listas._versao_cache_lista_pedido(lista)
    chave_antiga = f"tenant-a:saved-list:{versao_antiga}"
    if cache_antigo in {"memoria", "ambos"}:
        ambiente.memoria[chave_antiga] = bytes_antigos
    if cache_antigo in {"disco", "ambos"}:
        common._salvar_bytes_cache_lista_pedido("tenant-a", "saved-list", versao_antiga, bytes_antigos)

    primeira = _baixar(ambiente, endpoint)
    sheet = _conferir_link(primeira, CURRENT_LINK)
    assert [sheet[cell].value for cell in ("A2", "B2", "C2", "D2", "E2", "G2", "H2", "I2")] == [
        "503-1", "foto-salva.png", "Saved supplier description", "OEM-503-1", "Left / black", 7, 12.5, "=G2*H2",
    ]
    primeiro_cache = ambiente.caminhos_cache[-1]
    _conferir_link(primeiro_cache.read_bytes(), CURRENT_LINK)
    assert primeira in ambiente.memoria.values()
    assert len(ambiente.geracoes) == 1
    assert _baixar(ambiente, endpoint) == primeira
    assert len(ambiente.geracoes) == 1

    novo_link = "https://supplier.example/changed-only-in-cadastro"
    ambiente.contextos[("tenant-a", "store-a")]["produtos"][0]["link"] = novo_link
    segunda = _baixar(ambiente, endpoint)
    _conferir_link(segunda, novo_link)
    segundo_cache = ambiente.caminhos_cache[-1]
    assert segundo_cache != primeiro_cache
    _conferir_link(segundo_cache.read_bytes(), novo_link)
    assert segunda in ambiente.memoria.values()
    assert primeira != segunda
    assert len(ambiente.geracoes) == 2
    assert lista == original
    assert lista["itens"] is itens_originais
    assert lista["itens"][0] is item_original
    assert lista["updated_at"] == original["updated_at"]
    for _, itens_exportados, kwargs in ambiente.geracoes:
        assert itens_exportados is not itens_originais
        assert itens_exportados[0] is not item_original
        assert {k: v for k, v in itens_exportados[0].items() if k != "Link"} == {
            k: v for k, v in item_original.items() if k != "Link"
        }
        assert kwargs == {"client_id": "tenant-a", "loja": "store-a"}
    assert all(call == ("tenant-a", "store-a", ("503-1",)) for call in ambiente.consultas)


def test_link_isolado_por_cliente_e_loja_mesmo_com_sku_e_lista_iguais(ambiente, endpoint):
    ambiente.listas["tenant-a"].append(_lista(lista_id="other-store", store_id="store-b"))
    ambiente.listas["tenant-b"] = [_lista()]
    casos = [
        ("tenant-a", "store-a", "saved-list"),
        ("tenant-a", "store-b", "other-store"),
        ("tenant-b", "store-a", "saved-list"),
    ]
    for client_id, store_id, _ in casos:
        link = f"https://supplier.example/{client_id}/{store_id}"
        ambiente.contextos[(client_id, store_id)] = _contexto([
            _produto(link, store_id=store_id),
            _produto("https://supplier.example/foreign-store", store_id="foreign-store"),
            _produto("https://supplier.example/deleted", store_id=store_id, deleted_at_utc="2026-10-21"),
        ], store_id=store_id)

    originais = copy.deepcopy(ambiente.listas)
    for client_id, store_id, lista_id in casos:
        _conferir_link(
            _baixar(ambiente, endpoint, client_id=client_id, lista_id=lista_id),
            f"https://supplier.example/{client_id}/{store_id}",
        )
    assert ambiente.consultas == [(client, store, ("503-1",)) for client, store, _ in casos]
    assert ambiente.listas == originais


def test_sku_preserva_pontuacao_e_normalizacao_existente(ambiente, endpoint):
    lista = ambiente.listas["tenant-a"][0]
    lista["itens"] = [{**copy.deepcopy(lista["itens"][0]), "SKU": sku}
                      for sku in ("503-1", "5031", "503-JAN", "abc-2")]
    lista["itens"][-1]["sku"] = lista["itens"][-1].pop("SKU")
    ambiente.contextos[("tenant-a", "store-a")] = _contexto([
        _produto("https://supplier.example/without-hyphen", sku="5031"),
        _produto("https://supplier.example/with-hyphen", sku="503-1"),
        _produto("https://supplier.example/case-normalized", sku="ABC-2"),
    ])
    original = copy.deepcopy(lista)
    conteudo = _baixar(ambiente, endpoint)
    for row, esperado in enumerate(("with-hyphen", "without-hyphen", "with-hyphen", "case-normalized"), start=2):
        sheet = _conferir_link(conteudo, f"https://supplier.example/{esperado}", f"F{row}")
        item_original = original["itens"][row - 2]
        assert sheet[f"A{row}"].value == (item_original.get("SKU") or item_original["sku"])
    assert ambiente.consultas == [("tenant-a", "store-a", ("503-1", "5031", "ABC-2"))]
    assert lista == original


def test_link_canonico_atual_tem_preferencia_sobre_aliases_antigos(ambiente, endpoint):
    produto = ambiente.contextos[("tenant-a", "store-a")]["produtos"][0]
    produto.update({
        "Link Aliexpress": "https://supplier.example/old-aliexpress-alias",
        "product link": "https://supplier.example/old-product-link",
        "url": "https://supplier.example/old-url-alias",
    })
    original = copy.deepcopy(produto)
    _conferir_link(_baixar(ambiente, endpoint), CURRENT_LINK)
    assert produto == original


@pytest.mark.parametrize("campo_link", ["link", "LINK", " Link "], ids=["minusculo", "maiusculo", "espacos"])
def test_link_atual_substitui_chave_legada_da_lista_sem_mutar_original(ambiente, endpoint, campo_link):
    lista = ambiente.listas["tenant-a"][0]
    itens_originais, item_original = lista["itens"], lista["itens"][0]
    item_original[campo_link] = item_original.pop("Link")
    original = copy.deepcopy(lista)

    _conferir_link(_baixar(ambiente, endpoint), CURRENT_LINK)

    assert lista == original
    assert lista["itens"] is itens_originais
    assert lista["itens"][0] is item_original
    assert "Link" not in item_original
    assert item_original[campo_link] == SAVED_LINK


@pytest.mark.parametrize("produto", [
    None,
    {"link": ""},
    {"link": "javascript:alert(1)"},
    {"link": "ftp://supplier.example/product"},
    {"link": "https://"},
    {"link": "https://supplier.example/invalid space"},
    {"link": "https://user:password@supplier.example/product"},
], ids=["ausente", "vazio", "javascript", "ftp", "sem-host", "espaco", "credenciais"])
def test_sem_link_http_valido_mantem_link_salvo(ambiente, endpoint, produto):
    produtos = [] if produto is None else [_produto(**produto)]
    ambiente.contextos[("tenant-a", "store-a")] = _contexto(produtos)
    original = copy.deepcopy(ambiente.listas)
    _conferir_link(_baixar(ambiente, endpoint), SAVED_LINK)
    assert ambiente.listas == original


@pytest.mark.parametrize("contexto", [
    _contexto([_produto()], scope="global_safe"),
    _contexto([_produto()], scope="unresolved", loja_resolvida=False),
    _contexto([_produto()], loja_resolvida=False),
    _contexto([_produto()], store_id=""),
    _contexto([_produto(store_id="store-b")], store_id="store-b"),
    RuntimeError("Cadastro indisponivel"),
    OSError("Arquivo de cadastro indisponivel"),
    HTTPException(status_code=409, detail="Loja ambigua"),
], ids=["global", "nao-resolvida", "sem-confirmacao", "sem-store-id", "outra-loja", "runtime", "arquivo", "http"])
def test_contexto_inseguro_ou_indisponivel_preserva_link_sem_fallback_global(ambiente, endpoint, contexto):
    ambiente.contextos[("tenant-a", "store-a")] = contexto
    _conferir_link(_baixar(ambiente, endpoint), SAVED_LINK)
    assert ambiente.consultas == [("tenant-a", "store-a", ("503-1",))]


@pytest.mark.parametrize("referencia", ["", "__todas", "Todas as lojas", "Não definida"])
def test_lista_sem_loja_especifica_nao_busca_links_globais(ambiente, endpoint, referencia):
    lista = ambiente.listas["tenant-a"][0]
    lista["store_id"], lista["loja"] = "", referencia
    _conferir_link(_baixar(ambiente, endpoint), SAVED_LINK)
    assert ambiente.consultas == []


def test_lista_legada_resolve_loja_pelo_nome_e_aceita_alias_de_link(ambiente, endpoint):
    lista = ambiente.listas["tenant-a"][0]
    lista.pop("store_id")
    lista["loja"] = "Loja Carlos"
    link = "http://supplier.example/legacy-field?sku=503-1"
    ambiente.contextos[("tenant-a", "Loja Carlos")] = _contexto([
        {"sku_normalizado": "503-1", "store_id": "store-a", "Link Aliexpress": link},
    ])
    _conferir_link(_baixar(ambiente, endpoint), link)
    assert ambiente.consultas == [("tenant-a", "Loja Carlos", ("503-1",))]
