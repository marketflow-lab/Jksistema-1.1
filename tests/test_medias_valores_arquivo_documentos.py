"""File-supplied amounts survive the real workbook and report consumers."""

import copy
import io

import pytest
from openpyxl import load_workbook

from backend.services import codex_reports_advanced as reports
from backend.services import medias_compras_excel as excel


def item(unit=10, total=20, **flags):
    return {
        "SKU": "001", "Quantidade": 2, "Valor unidade": unit,
        "Valor total": total, "Título do produto em inglês": "Synthetic item",
        **flags,
    }


@pytest.fixture(autouse=True)
def isolate_photos(monkeypatch):
    monkeypatch.setattr(excel, "_resolver_caminho_foto_cadastro_seguro", lambda *_a, **_kw: None)
    monkeypatch.setattr(excel, "_resolver_caminho_foto_commercial_invoice", lambda *_a, **_kw: None)


@pytest.mark.parametrize("source,unit,total", [
    (item(10, 0, valor_total_informado=True, valor_unitario_informado=True), 10, 0),
    (item(10, 7, valor_total_informado=True, valor_unitario_informado=True), 10, 7),
    (item(0, 12, valor_total_informado=True, valor_unitario_informado=True), 0, 12),
    (item(0, 0, valor_unitario_informado=True), 0, "=G2*H2"),
    (item(1.2345, 2.469, valor_total_informado=True, valor_unitario_informado=True), 1.2345, 2.469),
    (item(10, 7), 10, "=G2*H2"),
    (item(0, 12), None, "=G2*H2"),
], ids=["zero-total", "different-total", "zero-unit", "unit-only-zero", "precision", "legacy-formula", "legacy-blank-unit"])
def test_purchase_excel_preserves_file_amounts_and_legacy_formulas(source, unit, total):
    before = copy.deepcopy(source)
    content = excel._gerar_excel_lista_pedido_bytes("Synthetic order", [source])
    workbook = load_workbook(io.BytesIO(content), data_only=False)
    try:
        sheet = workbook.active
        assert sheet["A2"].value == "001"
        assert sheet["G2"].value == 2
        assert sheet["H2"].value == unit
        assert sheet["I2"].value == total
        assert sheet["I2"].data_type == ("f" if isinstance(total, str) else "n")
        assert source == before
        # Internal source flags must not become workbook columns or product text.
        assert not any("_informado" in str(cell.value) for row in sheet for cell in row)
    finally:
        workbook.close()


@pytest.mark.parametrize("source,unit,total", [
    (item(10, 0, valor_total_informado=True, valor_unitario_informado=True), 10, 0),
    (item(10, 7, valor_total_informado=True, valor_unitario_informado=True), 10, 7),
    (item(0, 12, valor_total_informado=True, valor_unitario_informado=True), 0, 12),
    (item(0, 0, valor_unitario_informado=True), 0, 0),
    (item(1.2345, 2.469, valor_total_informado=True, valor_unitario_informado=True), 1.2345, 2.469),
    (item(10, 7), 10, 20),
    (item(0, 12), 6, 12),
], ids=["zero-total", "different-total", "zero-unit", "unit-only-zero", "precision", "legacy-product", "legacy-derived-unit"])
def test_commercial_invoice_preserves_file_values_and_legacy_derivation(source, unit, total):
    order = {"id": "synthetic-order", "nome_lista": "Synthetic order", "itens": [source]}
    before = copy.deepcopy(order)
    content = excel._gerar_commercial_invoice_bytes(order)
    workbook = load_workbook(io.BytesIO(content), data_only=False)
    try:
        sheet = workbook["Commercial"]
        assert sheet["G9"].value == 2
        assert sheet["H9"].value == unit
        assert sheet["I9"].value == total
        assert sheet["I10"].value == total
        assert sheet["I9"].data_type == "n"
        assert order == before
        assert not any("_informado" in str(cell.value) for row in sheet for cell in row)
    finally:
        workbook.close()


@pytest.mark.parametrize("source,expected", [
    (item(10, 0, valor_total_informado=True), 0),
    (item(10, 7, valor_total_informado=True), 7),
    (item(0, 0, valor_unitario_informado=True, Cost=10), 0),
    (item(10, 0), 20),
], ids=["zero-total", "different-total", "zero-unit-legacy-alias", "legacy-fallback"])
def test_landed_cost_report_uses_supplied_amounts(source, expected):
    order = {"id": "synthetic-order", "loja": "Synthetic store", "dolar": 5, "itens": [source]}
    before = copy.deepcopy(order)
    result = reports._landed_cost_analysis("", order, reports.normalize_report_settings({}), [], 500000)
    assert result["fob_usd"] == expected
    assert result["items"][0]["fob_usd"] == expected
    assert order == before


@pytest.mark.parametrize("informed,expected_zero_variation", [(True, True), (False, False)])
def test_supplier_report_does_not_replace_explicit_zero_with_product(informed, expected_zero_variation):
    def order(number, source):
        return {
            "id": str(number), "supplier": "Synthetic supplier", "status": "Recebido",
            "order_date": "2026-09-01", "received_at": "2026-09-15",
            "promised_delivery_date": "2026-09-20", "itens": [source],
        }

    source = item(50, 0, valor_total_informado=informed)
    orders = [order(1, source), order(2, item()), order(3, item())]
    before = copy.deepcopy(orders)
    result = reports._supplier_performance(orders, "Synthetic supplier")
    assert result["status"] == "available"
    assert (result["cost_variation_pct"] == 0) is expected_zero_variation
    assert orders == before
