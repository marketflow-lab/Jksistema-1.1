import asyncio
import base64
import copy
import io
import json

import pytest
from fastapi import HTTPException
from openpyxl import load_workbook
from openpyxl.utils.units import pixels_to_EMU, points_to_pixels
from PIL import Image

from backend.routers.medias_compras import create_medias_compras_router
from backend.schemas.medias_compras import ListaPedidoUpdateRequest
from backend.services import medias_compras_excel as excel_service
from backend.services import medias_compras_listas as listas_service
from backend.services import medias_compras_documentos_cadastro as cadastro_documentos
from backend.services import cadastro_compatibilidade, cadastro_fornecedores, cadastro_importador_loja, cadastro_nacionalizadores


CONDICOES_COMERCIAIS_INVOICE = (
    "Commercial Use - Formal",
    "Within 45 days after full payment of Incoterms: - EXW",
    "Payment terms: 100% prepaid",
)


PNG_1X1 = base64.b64decode(
    "iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAQAAAC1HAwCAAAAC0lEQVR42mP8/x8AAusB9Y9Zc1sAAAAASUVORK5CYII="
)


def _lista_aprovada():
    return {
        "id": "lista-1",
        "nome_lista": "JK 48",
        "loja": "JK Pecas",
        "numero_invoice": "JK48-18082026",
        "supplier": "Fornecedor Exemplo Ltd.",
        "currency": "USD",
        "incoterm": "EXW",
        "itens": [
            {
                "SKU": "001",
                "NCM": "85365090",
                "Título do produto em inglês": "Temperature sensor switch",
                "Quantidade": 2,
                "Valor unidade": 3.5,
                "Foto": "001.png",
                "compra_aprovada": True,
                "compra_aprovada_em": "2026-08-18T09:00:00",
            },
            {
                "SKU": "SKU 002",
                "Código NCM": "8413.70.10",
                "Título em português": "=HYPERLINK(\"https://example.invalid\",\"unsafe\")",
                "Quantidade": 3,
                "Valor unidade": 0,
                "Valor total": 12,
                "compra_aprovada": True,
                "compra_aprovada_em": "2026-08-18T14:30:00",
            },
        ],
    }


def _carregar_commercial_invoice(conteudo: bytes):
    assert conteudo.startswith(b"PK")
    workbook = load_workbook(io.BytesIO(conteudo), data_only=False)
    assert workbook.sheetnames == ["Commercial"]
    return workbook["Commercial"]


def _cadastro_produtos_invoice(monkeypatch):
    produtos = {
        "001": {
            "sku": "001", "descricao": "Temperature sensor switch",
            "foto": "001.png", "unidade": "Unit",
        },
        "SKU 002": {
            "sku": "002",
            "descricao": '=HYPERLINK("https://example.invalid","unsafe")',
            "foto": "", "unidade": "Unit",
        },
    }
    monkeypatch.setattr(
        cadastro_documentos, "_contexto_cadastro_documentos",
        lambda *_args: {"store_id": "", "scope": "global_safe"},
    )
    monkeypatch.setattr(
        cadastro_documentos, "_dados_produto_documento",
        lambda _contexto, sku: produtos.get(sku, {}),
    )


def _formulas(sheet):
    return [
        cell.coordinate
        for row in sheet.iter_rows()
        for cell in row
        if cell.data_type == "f"
    ]


@pytest.mark.parametrize(
    ("data_pedido", "data_aprovacao", "esperado"),
    [
        ("2026-10-05", "2026-08-18", "October 5, 2026"),
        (" 2026-10-05 ", "2026-08-18", "October 5, 2026"),
        ("05/10/2026", "2026-08-18", "October 5, 2026"),
        ("2026-10-05T23:00:00Z", "2026-08-18", "October 5, 2026"),
        ("", "2026-08-18T14:30:00", "August 18, 2026"),
        (None, "2026-08-18", "August 18, 2026"),
        ("  ", "2026-08-18", "August 18, 2026"),
        ("", "", None),
        ("Data historica pendente", "2026-08-18", "Data historica pendente"),
        ("31/02/2026", "2026-08-18", "31/02/2026"),
        ("=1+1", "2026-08-18", "=1+1"),
    ],
)
def test_commercial_invoice_prioriza_data_pedido_e_preserva_fallback(
    data_pedido, data_aprovacao, esperado,
):
    sheet = _carregar_commercial_invoice(excel_service._gerar_commercial_invoice_bytes(
        {"order_date": data_pedido, "itens": []}, data_aprovacao=data_aprovacao,
    ))
    assert sheet["E3"].value == esperado
    assert _formulas(sheet) == []


@pytest.mark.parametrize("dimensoes", [(640, 80), (80, 640), (240, 240), (1, 1)])
def test_commercial_invoice_foto_cabe_na_celula_sem_deformar(tmp_path, monkeypatch, dimensoes):
    foto = tmp_path / "produto.png"
    Image.new("RGB", dimensoes, "gray").save(foto)
    monkeypatch.setattr(excel_service, "_resolver_caminho_foto_commercial_invoice", lambda *_args: str(foto))
    sheet = _carregar_commercial_invoice(excel_service._gerar_commercial_invoice_bytes({
        "itens": [{"SKU": "001", "Quantidade": 2, "Valor unidade": 3.5}],
    }))

    assert len(sheet._images) == 1
    ancora = sheet._images[0].anchor
    assert ancora._from.col == 5
    assert ancora._from.row == 8
    largura_celula = pixels_to_EMU(int(sheet.column_dimensions["F"].width * 7))
    altura_celula = pixels_to_EMU(points_to_pixels(sheet.row_dimensions[9].height))
    margem = pixels_to_EMU(4)
    assert ancora._from.colOff >= margem
    assert ancora._from.rowOff >= margem
    assert ancora._from.colOff + ancora.ext.cx <= largura_celula - margem
    assert ancora._from.rowOff + ancora.ext.cy <= altura_celula - margem
    assert abs(2 * ancora._from.colOff + ancora.ext.cx - largura_celula) <= 2
    assert abs(2 * ancora._from.rowOff + ancora.ext.cy - altura_celula) <= 2
    assert ancora.ext.cx / ancora.ext.cy == pytest.approx(dimensoes[0] / dimensoes[1], rel=1e-5)
    assert sheet["G9"].value == 2
    assert sheet["I9"].value == 7
    assert sheet["I10"].value == 7
    assert sheet.column_dimensions["F"].width == 18.25
    assert sheet.row_dimensions[9].height == 60


def test_gerador_commercial_invoice_reproduz_modelo_totais_e_foto(tmp_path, monkeypatch):
    _cadastro_produtos_invoice(monkeypatch)
    raiz_tenant = tmp_path / "tenant-a"
    foto = raiz_tenant / "cadastro_fotos" / "001.png"
    foto.parent.mkdir(parents=True)
    foto.write_bytes(PNG_1X1)
    monkeypatch.setattr(excel_service, "get_tenant_path", lambda _client_id: str(raiz_tenant))
    monkeypatch.setattr(
        cadastro_compatibilidade,
        "visao_produtos_cadastro_contexto_loja",
        lambda _client_id, _referencia, **_kwargs: {"store_id": "", "scope": "unresolved"},
    )
    monkeypatch.setattr(
        excel_service,
        "_resolver_foto_cadastro_sku",
        lambda _client_id, sku, _foto_ref: str(foto) if sku == "001" else "",
    )

    sheet = _carregar_commercial_invoice(
        excel_service._gerar_commercial_invoice_bytes(
            _lista_aprovada(),
            client_id="tenant-a",
            data_aprovacao="2026-08-18T14:30:00",
        )
    )

    assert sheet.freeze_panes is None
    assert sheet.sheet_view.showGridLines is True
    assert sheet.page_setup.orientation == "portrait"
    assert sheet.max_row == 15
    assert sheet.max_column == 9
    assert _formulas(sheet) == []
    assert sheet["A1"].value == "Fornecedor Exemplo Ltd."
    assert sheet["A2"].value == "COMMERCIAL INVOICE"
    assert sheet["A3"].value == "INV. No.:"
    assert sheet["C3"].value == "JK48-18082026"
    assert sheet["D3"].value == "DATE:"
    assert sheet["E3"].value == "August 18, 2026"
    assert sheet["F3"].value == "FROM:"
    assert sheet["H3"].value in (None, "")
    assert sheet["A4"].value == "LIST: JK 48"
    assert sheet["F4"].value == "STORE: JK Pecas"
    assert sheet["A7"].value == "Descriptions"
    assert sheet["C8"].value == "HSCODE"

    assert sheet["B9"].value == "SKU 001"
    assert sheet["C9"].value == "8536.50.90"
    assert sheet["D9"].value == "Temperature sensor switch"
    assert sheet["G9"].value == 2
    assert sheet["H9"].value == 3.5
    assert "US$" not in sheet["H9"].number_format
    assert sheet["I9"].value == 7
    assert "US$" in sheet["I9"].number_format
    assert sheet["B10"].value == "SKU 002"
    assert sheet["D10"].value.startswith("=HYPERLINK")
    assert sheet["D10"].data_type == "s"
    assert sheet["C10"].value == "8413.70.10"
    assert sheet["H10"].value == 4
    assert sheet["I10"].value == 12

    assert sheet["C11"].value == "TOTAL NET WEIGHT"
    assert sheet["F11"].value == "Units Total"
    assert sheet["G11"].value == 5
    assert sheet["H11"].value == "Subtotal"
    assert sheet["I11"].value == 19
    assert sheet["F12"].value == "Freight USD"
    assert sheet["I12"].value in (None, "")
    assert sheet["F13"].value == "Insurance USD"
    assert sheet["I13"].value in (None, "")
    assert sheet["F14"].value == "Shipping:"
    assert sheet["G14"].value in (None, "")
    assert sheet["H14"].value == "TOTAL:"
    assert sheet["I14"].value == 19
    assert sheet["A15"].value.splitlines()[0] == "Commercial Use - Formal"
    assert "Note - invoice JK48-18082026" in sheet["A15"].value.splitlines()
    assert "Incoterms: - EXW" in sheet["A15"].value
    assert "Payment Currency: USD" in sheet["A15"].value
    assert sheet["E15"].value in (None, "")
    assert sheet["I15"].value == "Company stamp and signature"

    merges = {str(intervalo) for intervalo in sheet.merged_cells.ranges}
    assert {
        "A1:I1",
        "A2:I2",
        "A3:B3",
        "F3:G3",
        "H3:I3",
        "A4:E6",
        "F4:I6",
        "A7:H7",
        "D8:E8",
        "D9:E9",
        "D10:E10",
        "A15:D15",
        "E15:H15",
    } <= merges
    assert sheet.column_dimensions["A"].width == pytest.approx(6)
    assert sheet.column_dimensions["B"].width == pytest.approx(15.63)
    assert sheet.column_dimensions["E"].width == pytest.approx(27.75)
    assert sheet.column_dimensions["I"].width == pytest.approx(12.38)
    assert sheet.row_dimensions[1].height == pytest.approx(93)
    assert sheet.row_dimensions[6].height == pytest.approx(147)
    assert sheet.row_dimensions[9].height == pytest.approx(60)
    assert sheet.row_dimensions[15].height == pytest.approx(101.25)
    assert len(sheet._images) == 1
    assert sheet._images[0].anchor._from.col == 5
    assert sheet._images[0].anchor._from.row == 8


def test_commercial_invoice_usa_todos_os_dados_da_loja_exata(tmp_path, monkeypatch):
    registros = {
        "store-a": {
            "nome_empresa": "Empresa A Ltda.",
            "tax_id": "00.000.000/0001-00",
            "telefone": "+5500000000000",
            "email": "a@example.test",
            "contato": "Pessoa A",
            "logradouro": "Rua A, 10",
            "bairro": "Centro",
            "cidade": "Cidade A",
            "cep": "00000-000",
            "estado": "Estado A",
            "pais": "Brazil",
        },
        "store-b": {"nome_empresa": "Empresa B Ltda.", "email": "b@example.test"},
    }
    (tmp_path / "cadastro_importador_loja.json").write_text(
        json.dumps({"schema": "jk.cadastro.importador_loja.v1", "lojas": registros}),
        encoding="utf-8",
    )
    monkeypatch.setattr(cadastro_importador_loja.lojas, "get_tenant_path", lambda _client_id: str(tmp_path))
    monkeypatch.setattr(
        cadastro_compatibilidade,
        "visao_produtos_cadastro_contexto_loja",
        lambda _client_id, referencia, **_kwargs: {"store_id": referencia if referencia in registros else ""},
    )

    lista = {"store_id": "store-a", "loja": "Loja A", "itens": []}
    sheet = _carregar_commercial_invoice(
        excel_service._gerar_commercial_invoice_bytes(lista, client_id="tenant-a")
    )
    assert sheet["F4"].value.splitlines() == [
        "Buyer: Empresa A Ltda.",
        "TAX ID: 00.000.000/0001-00",
        "Phone: +5500000000000",
        "E-mail: a@example.test",
        "Contact Personal: Pessoa A",
        "Address: Rua A, 10",
        "Neighborhood: Centro",
        "City: Cidade A",
        "ZIP CODE: 00000-000",
        "State: Estado A",
        "Country: Brazil",
    ]
    assert "Empresa B" not in sheet["F4"].value

    lista["store_id"] = "store-b"
    sheet = _carregar_commercial_invoice(
        excel_service._gerar_commercial_invoice_bytes(lista, client_id="tenant-a")
    )
    assert sheet["F4"].value == "Buyer: Empresa B Ltda.\nE-mail: b@example.test"

    lista["store_id"] = "unknown"
    sheet = _carregar_commercial_invoice(
        excel_service._gerar_commercial_invoice_bytes(lista, client_id="tenant-a")
    )
    assert sheet["F4"].value == "STORE: Loja A"


@pytest.fixture
def fornecedores_invoice(tmp_path, monkeypatch):
    fornecedor = {
        "id": "fornecedor-a",
        "nome_empresa": "Fornecedor Atual Co., Ltd.",
        "nome_contato": "Contato A",
        "endereco_empresa": "Room 702, Business Building\nNo. 1533, Zhejiang, China",
        "telefone": "+00 0000-0000",
        "email": "contato@example.test",
        "nome_beneficiario": "Beneficiario diferente da empresa",
        "conta_beneficiario": "dados-bancarios-nao-exportados",
    }
    registros = {
        "tenant-a": [fornecedor, {"id": "fornecedor-b", "nome_empresa": "Fornecedor B"}],
        "tenant-b": [{"id": "fornecedor-outro", "nome_empresa": "Empresa de outro cliente"}],
    }
    for client_id, fornecedores in registros.items():
        pasta = tmp_path / client_id
        pasta.mkdir()
        (pasta / "cadastro_fornecedores.json").write_text(
            json.dumps({"schema": "jk.cadastro.fornecedores.v1", "fornecedores": fornecedores}),
            encoding="utf-8",
        )
    monkeypatch.setattr(excel_service, "get_tenant_path", lambda client_id: str(tmp_path / client_id))
    monkeypatch.setattr(
        cadastro_compatibilidade,
        "visao_produtos_cadastro_contexto_loja",
        lambda *_args, **_kwargs: {"store_id": "", "scope": "unresolved"},
    )
    return tmp_path, fornecedor


def test_commercial_invoice_cabecalho_usa_cadastro_e_formato_da_imagem(fornecedores_invoice):
    lista = {"supplier": "Nome historico", "fornecedor_id": "fornecedor-a", "itens": []}
    conteudo = excel_service._gerar_commercial_invoice_bytes(lista, client_id="tenant-a")
    sheet = _carregar_commercial_invoice(conteudo)

    assert sheet["A1"].value.splitlines() == [
        "Fornecedor Atual Co., Ltd.",
        "Contact Name: Contato A",
        "Registered company address: Room 702, Business Building",
        "No. 1533, Zhejiang, China",
        "Company Tel: +00 0000-0000",
        "E-mail: contato@example.test",
    ]
    assert "dados-bancarios" not in sheet["A1"].value
    assert "Beneficiario" not in sheet["A1"].value
    assert sheet["A1"].alignment.horizontal == "center"
    assert sheet["A1"].alignment.vertical == "center"
    assert sheet["A1"].alignment.wrap_text is True
    assert "A1:I1" in {str(intervalo) for intervalo in sheet.merged_cells.ranges}
    rich_sheet = load_workbook(io.BytesIO(conteudo), rich_text=True)["Commercial"]
    blocos = rich_sheet["A1"].value
    assert blocos[0].font.rFont == "Arial"
    assert blocos[0].font.sz == 14
    assert blocos[0].font.b is True
    assert blocos[1].font.sz == 11
    assert blocos[1].font.b is True
    assert sheet["A2"].value == "COMMERCIAL INVOICE"
    assert _formulas(sheet) == []


def test_commercial_invoice_omite_campos_vazios_ao_trocar_fornecedor(fornecedores_invoice):
    lista = {"supplier": "Nome historico", "fornecedor_id": "fornecedor-b", "itens": []}
    sheet = _carregar_commercial_invoice(
        excel_service._gerar_commercial_invoice_bytes(lista, client_id="tenant-a")
    )
    assert sheet["A1"].value == "Fornecedor B"
    assert sheet.row_dimensions[1].height == 93


@pytest.mark.parametrize(
    ("client_id", "fornecedor_id"),
    [("tenant-a", ""), ("tenant-a", "removido"), ("tenant-a", "fornecedor-outro"),
     ("tenant-sem-cadastro", "fornecedor-a"), (None, "fornecedor-a")],
)
def test_commercial_invoice_preserva_nome_historico_sem_fornecedor_do_cliente(
    fornecedores_invoice, client_id, fornecedor_id,
):
    lista = {"supplier": "Fornecedor legado", "fornecedor_id": fornecedor_id, "itens": []}
    sheet = _carregar_commercial_invoice(
        excel_service._gerar_commercial_invoice_bytes(lista, client_id=client_id)
    )
    assert sheet["A1"].value == "Fornecedor legado"
    assert sheet.row_dimensions[1].height == 93


def test_commercial_invoice_cabecalho_extenso_literal_sem_alterar_cadastro(fornecedores_invoice):
    raiz, fornecedor = fornecedores_invoice
    fornecedor["nome_empresa"] = '=HYPERLINK("https://example.invalid","unsafe")'
    fornecedor["endereco_empresa"] = "\n".join(f"Endereco de exemplo {indice}" for indice in range(15))
    arquivo = raiz / "tenant-a" / "cadastro_fornecedores.json"
    arquivo.write_text(json.dumps({"fornecedores": [fornecedor]}), encoding="utf-8")
    cadastro_original = arquivo.read_bytes()
    lista = {"fornecedor_id": "fornecedor-a", "itens": []}

    sheet = _carregar_commercial_invoice(
        excel_service._gerar_commercial_invoice_bytes(lista, client_id="tenant-a")
    )
    assert sheet["A1"].value.startswith("=HYPERLINK")
    assert sheet["A1"].data_type == "s"
    assert "Endereco de exemplo 14" in sheet["A1"].value
    assert sheet.row_dimensions[1].height > 93
    assert sheet["A2"].value == "COMMERCIAL INVOICE"
    assert _formulas(sheet) == []
    assert arquivo.read_bytes() == cadastro_original


@pytest.fixture
def fornecedores_rodape_invoice(tmp_path, monkeypatch):
    fornecedor = {
        "id": "fornecedor-a",
        "nome_empresa": "Fornecedor Atual Co., Ltd.",
        "nome_contato": "Contato A",
        "endereco_empresa": "Industrial Park, Hubei Province, Tianmen City",
        "moeda_pagamento": "EUR",
        "conta_beneficiario": "000123456000",
        "swift": "TESTSGSG",
        "pais_regiao_beneficiario": "Singapore",
        "nome_beneficiario": "Beneficiario A Ltd.",
        "endereco_beneficiario": "Room 301, Commerce Park",
        "banco_beneficiario": "Banco de Exemplo",
        "endereco_banco": "Business Park, Singapore",
        "codigo_banco": "001",
        "codigo_agencia": "0007",
        "observacao_pagamento": "Notify the supplier after payment",
    }
    registros = {
        "tenant-a": [fornecedor, {"id": "fornecedor-b", "nome_empresa": "Fornecedor B"}],
        "tenant-b": [{**fornecedor, "nome_empresa": "Fornecedor do cliente B",
                      "endereco_empresa": "Toronto, Canada",
                      "moeda_pagamento": "GBP", "conta_beneficiario": "00009999",
                      "nome_beneficiario": "Beneficiario do cliente B"}],
    }
    for client_id, fornecedores in registros.items():
        pasta = tmp_path / client_id
        pasta.mkdir()
        (pasta / "cadastro_fornecedores.json").write_text(
            json.dumps({"schema": "jk.cadastro.fornecedores.v1", "fornecedores": fornecedores}),
            encoding="utf-8",
        )
    monkeypatch.setattr(excel_service, "get_tenant_path", lambda client_id: str(tmp_path / client_id))
    monkeypatch.setattr(excel_service, "_resolver_caminho_foto_commercial_invoice", lambda *_args: None)
    monkeypatch.setattr(
        cadastro_compatibilidade, "visao_produtos_cadastro_contexto_loja",
        lambda *_args, **_kwargs: {"store_id": "", "scope": "unresolved"},
    )
    monkeypatch.setattr(
        cadastro_fornecedores, "get_tenant_path",
        lambda *_args: pytest.fail("Invoice deve fornecer a pasta canonica do cliente"),
    )
    return tmp_path, fornecedor


@pytest.mark.parametrize(
    ("campos_incoterm", "esperado"),
    [
        pytest.param({}, "EXW", id="ausente"),
        pytest.param({"incoterm": None}, "EXW", id="nulo"),
        pytest.param({"incoterm": ""}, "EXW", id="vazio"),
        pytest.param({"incoterm": "   "}, "EXW", id="espacos"),
        pytest.param({"incoterm": "EXW"}, "EXW", id="exw-explicito"),
        pytest.param({"incoterm": "  FOB  "}, "FOB", id="fob-explicito-com-espacos"),
    ],
)
def test_commercial_invoice_rodape_garante_incoterm_sem_mutar_lista(
    fornecedores_rodape_invoice, campos_incoterm, esperado,
):
    lista = {"itens": [], **campos_incoterm}
    lista_original = copy.deepcopy(lista)

    sheet = _carregar_commercial_invoice(
        excel_service._gerar_commercial_invoice_bytes(lista, client_id="tenant-a")
    )

    assert lista == lista_original
    assert sheet["A13"].value.splitlines()[1] == (
        f"Within 45 days after full payment of Incoterms: - {esperado}"
    )


def test_commercial_invoice_rodape_exporta_dados_bancarios_do_fornecedor(fornecedores_rodape_invoice):
    raiz, _registro = fornecedores_rodape_invoice
    arquivo = raiz / "tenant-a" / "cadastro_fornecedores.json"
    bytes_originais = arquivo.read_bytes()
    lista = {**_lista_aprovada(), "fornecedor_id": "fornecedor-a"}
    lista_original = copy.deepcopy(lista)
    sheet = _carregar_commercial_invoice(excel_service._gerar_commercial_invoice_bytes(lista, client_id="tenant-a"))

    assert sheet["E15"].value == (
        "Payment currency: EUR. Beneficiary account number: 000123456000. "
        "Swift code: TESTSGSG. Beneficiary country/region: Singapore. "
        "Beneficiary name: Beneficiario A Ltd. Beneficiary address: Room 301, Commerce Park. "
        "Beneficiary bank: Banco de Exemplo. Beneficiary bank address: Business Park, Singapore. "
        "Bank code: 001. Branch code: 0007. Remark: Notify the supplier after payment. "
        "Please include the order invoice number JK48-18082026."
    )
    assert sheet["E15"].alignment.horizontal == "left"
    assert sheet["E15"].alignment.wrap_text is True
    assert sheet["E15"].data_type == "s"
    assert sheet["A15"].value.splitlines()[0] == "Commercial Use - Formal"
    assert sheet["A15"].value.splitlines() == [
        "Commercial Use - Formal",
        "Within 45 days after full payment of Incoterms: - EXW",
        "Payment terms: 100% prepaid",
        "Note - invoice JK48-18082026",
        "Payment Currency: EUR Payment Method: Bank transfer",
    ]
    assert sheet["I15"].value == "Company stamp and signature"
    assert sheet["G11"].value == 5
    assert sheet["I11"].value == sheet["I14"].value == 19
    assert sheet["H14"].value == "TOTAL:"
    assert {"A15:D15", "E15:H15"} <= {str(r) for r in sheet.merged_cells.ranges}
    for condicao in ("45 days", "100% prepaid", "Bank transfer"):
        assert condicao in sheet["A15"].value
    assert _formulas(sheet) == []
    assert lista == lista_original
    assert arquivo.read_bytes() == bytes_originais


@pytest.mark.parametrize(
    ("moeda_fornecedor", "moeda_lista", "esperado"),
    [("EUR", "USD", "EUR"), ("", "USD", "USD"), ("   ", "JPY", "JPY"), ("", "", "")],
)
def test_commercial_invoice_rodape_moeda_fornecedor_tem_prioridade_com_fallback(
    fornecedores_rodape_invoice, moeda_fornecedor, moeda_lista, esperado,
):
    raiz, registro = fornecedores_rodape_invoice
    registro["moeda_pagamento"] = moeda_fornecedor
    (raiz / "tenant-a" / "cadastro_fornecedores.json").write_text(
        json.dumps({"fornecedores": [registro]}), encoding="utf-8",
    )
    lista = {"fornecedor_id": "fornecedor-a", "currency": moeda_lista, "itens": []}
    sheet = _carregar_commercial_invoice(excel_service._gerar_commercial_invoice_bytes(lista, client_id="tenant-a"))
    assert sheet["A13"].value.splitlines() == [
        *CONDICOES_COMERCIAIS_INVOICE,
        f"Payment Currency: {esperado} Payment Method: Bank transfer" if esperado
        else "Payment Method: Bank transfer",
    ]
    assert "Note - invoice" not in sheet["A13"].value
    if moeda_fornecedor.strip():
        assert sheet["E13"].value.startswith(f"Payment currency: {moeda_fornecedor}. ")
    else:
        assert "Payment currency:" not in sheet["E13"].value


def test_commercial_invoice_rodape_troca_fornecedor_sem_reaproveitar_dados(fornecedores_rodape_invoice):
    lista = {"fornecedor_id": "fornecedor-a", "supplier": "Nome historico", "currency": "USD", "itens": []}
    primeira = _carregar_commercial_invoice(excel_service._gerar_commercial_invoice_bytes(lista, client_id="tenant-a"))
    assert "Beneficiary account number: 000123456000" in primeira["E13"].value
    lista["fornecedor_id"] = "fornecedor-b"
    segunda = _carregar_commercial_invoice(excel_service._gerar_commercial_invoice_bytes(lista, client_id="tenant-a"))
    assert segunda["A1"].value == "Fornecedor B"
    assert segunda["E13"].value in (None, "")
    assert segunda["A13"].value.splitlines() == [
        *CONDICOES_COMERCIAIS_INVOICE, "Payment Currency: USD Payment Method: Bank transfer",
    ]
    assert segunda.row_dimensions[13].height == 101.25


def test_commercial_invoice_rodape_omite_rotulos_vazios_e_preserva_pontuacao(fornecedores_rodape_invoice):
    raiz, registro = fornecedores_rodape_invoice
    registro = {"id": "fornecedor-a", "nome_empresa": registro["nome_empresa"],
                "swift": " ", "codigo_banco": "001", "codigo_agencia": "",
                "observacao_pagamento": "Referencia ja pontuada."}
    (raiz / "tenant-a" / "cadastro_fornecedores.json").write_text(
        json.dumps({"fornecedores": [registro]}), encoding="utf-8",
    )
    sheet = _carregar_commercial_invoice(excel_service._gerar_commercial_invoice_bytes(
        {"fornecedor_id": "fornecedor-a", "itens": []}, client_id="tenant-a",
    ))
    assert sheet["E13"].value == "Bank code: 001. Remark: Referencia ja pontuada."


@pytest.mark.parametrize(("client_id", "nome", "conta", "moeda"), [
    ("tenant-a", "Fornecedor Atual Co., Ltd.", "000123456000", "EUR"),
    ("tenant-b", "Fornecedor do cliente B", "00009999", "GBP"),
])
def test_commercial_invoice_rodape_id_igual_isolado_por_cliente(
    fornecedores_rodape_invoice, client_id, nome, conta, moeda,
):
    sheet = _carregar_commercial_invoice(excel_service._gerar_commercial_invoice_bytes(
        {"fornecedor_id": "fornecedor-a", "itens": []}, client_id=client_id,
    ))
    assert sheet["A1"].value.splitlines()[0] == nome
    assert f"Beneficiary account number: {conta}." in sheet["E13"].value
    assert sheet["A13"].value.splitlines() == [
        *CONDICOES_COMERCIAIS_INVOICE, f"Payment Currency: {moeda} Payment Method: Bank transfer",
    ]
    outra_conta = "00009999" if client_id == "tenant-a" else "000123456000"
    assert outra_conta not in sheet["E13"].value


@pytest.mark.parametrize(("client_id", "fornecedor_id"), [
    ("tenant-a", None), ("tenant-a", ""), ("tenant-a", "removido"),
    ("tenant-a", "FORNECEDOR-A"), ("tenant-sem-cadastro", "fornecedor-a"), (None, "fornecedor-a"),
])
def test_commercial_invoice_rodape_sem_vinculo_exato_nao_exporta_conta(
    fornecedores_rodape_invoice, client_id, fornecedor_id,
):
    lista = {"fornecedor_id": fornecedor_id, "supplier": "Fornecedor Atual Co., Ltd.",
             "conta_beneficiario": "conta-solta-na-lista", "currency": "USD",
             "numero_invoice": "INV-SEM-VINCULO-001", "itens": []}
    sheet = _carregar_commercial_invoice(excel_service._gerar_commercial_invoice_bytes(lista, client_id=client_id))
    assert sheet["A1"].value == lista["supplier"]
    assert sheet["E13"].value in (None, "")
    assert sheet["A13"].value.splitlines() == [
        *CONDICOES_COMERCIAIS_INVOICE, "Note - invoice INV-SEM-VINCULO-001",
        "Payment Currency: USD Payment Method: Bank transfer",
    ]


def test_commercial_invoice_rodape_sem_id_nao_consulta_cadastro(monkeypatch):
    monkeypatch.setattr(cadastro_fornecedores, "ler_fornecedor_por_id",
                        lambda *_args, **_kwargs: pytest.fail("Sem vinculo nao deve ler cadastro"))
    monkeypatch.setattr(cadastro_compatibilidade, "visao_produtos_cadastro_contexto_loja",
                        lambda *_args, **_kwargs: {"store_id": "", "scope": "unresolved"})
    sheet = _carregar_commercial_invoice(excel_service._gerar_commercial_invoice_bytes(
        {"supplier": "Fornecedor historico", "itens": []}, client_id="tenant-a",
    ))
    assert sheet["E13"].value in (None, "")


def test_commercial_invoice_cabecalho_e_rodape_usam_uma_leitura_sem_mutacao(fornecedores_rodape_invoice, monkeypatch):
    raiz, registro = fornecedores_rodape_invoice
    registro_original = copy.deepcopy(registro)
    arquivo = raiz / "tenant-a" / "cadastro_fornecedores.json"
    bytes_originais = arquivo.read_bytes()
    chamadas = []

    def ler(client_id, fornecedor_id, *, pasta_tenant):
        chamadas.append((client_id, fornecedor_id, pasta_tenant))
        assert len(chamadas) == 1, "Cabeçalho e rodape precisam refletir a mesma leitura"
        return registro

    monkeypatch.setattr(cadastro_fornecedores, "ler_fornecedor_por_id", ler)
    monkeypatch.setattr(cadastro_fornecedores, "_salvar_fornecedores",
                        lambda *_args: pytest.fail("Exportar nao deve gravar fornecedor"))
    lista = {"fornecedor_id": "fornecedor-a", "currency": "USD", "itens": []}
    lista_original = copy.deepcopy(lista)
    sheet = _carregar_commercial_invoice(excel_service._gerar_commercial_invoice_bytes(lista, client_id="tenant-a"))
    assert chamadas == [("tenant-a", "fornecedor-a", str(raiz / "tenant-a"))]
    assert sheet["A1"].value.startswith(registro["nome_empresa"])
    assert sheet["H3"].value == "CHINA"
    assert "Beneficiary account number: " + registro["conta_beneficiario"] in sheet["E13"].value
    assert registro == registro_original
    assert lista == lista_original
    assert arquivo.read_bytes() == bytes_originais


@pytest.mark.parametrize("client_id,fornecedor_id,esperado", [
    ("tenant-a", "fornecedor-a", "CHINA"),
    ("tenant-b", "fornecedor-a", "CANADA"),
    ("tenant-a", "fornecedor-b", None),
    ("tenant-a", "inexistente", None),
])
def test_commercial_invoice_origem_usa_fornecedor_exato_do_cliente(
    fornecedores_rodape_invoice, client_id, fornecedor_id, esperado,
):
    sheet = _carregar_commercial_invoice(excel_service._gerar_commercial_invoice_bytes(
        {"fornecedor_id": fornecedor_id, "itens": []}, client_id=client_id,
    ))
    assert sheet["H3"].value == esperado
    assert "H3:I3" in {str(intervalo) for intervalo in sheet.merged_cells.ranges}
    if fornecedor_id == "fornecedor-a":
        assert "Beneficiary country/region: Singapore" in sheet["E13"].value


def test_commercial_invoice_origem_acompanha_endereco_sem_usar_campo_manual_ou_bancario(fornecedores_rodape_invoice):
    raiz, registro = fornecedores_rodape_invoice
    arquivo = raiz / "tenant-a" / "cadastro_fornecedores.json"
    lista = {"fornecedor_id": "fornecedor-a", "itens": []}
    for endereco, pais in (("Industrial Park, Hubei Province", "CHINA"), ("Tokyo, Japan", "JAPAN"),
                            ("China / Vietnam", None), ("=1+1", None), ("", None)):
        registro["endereco_empresa"] = endereco
        registro["pais_origem"] = "VIETNAM"  # Deprecated field from the previous release is ignored.
        arquivo.write_text(json.dumps({"fornecedores": [registro]}), encoding="utf-8")
        original = arquivo.read_bytes()
        sheet = _carregar_commercial_invoice(excel_service._gerar_commercial_invoice_bytes(lista, client_id="tenant-a"))
        assert sheet["H3"].value == pais
        if pais:
            assert sheet["H3"].data_type == "s"
        assert "Beneficiary country/region: Singapore" in sheet["E13"].value
        assert arquivo.read_bytes() == original


def test_commercial_invoice_rodape_preserva_zeros_texto_literal_e_multilinhas(fornecedores_rodape_invoice):
    raiz, registro = fornecedores_rodape_invoice
    registro.update(conta_beneficiario="000000123456789", codigo_banco="0001", codigo_agencia="000007",
                    endereco_beneficiario="Room 301\nBuilding A,  2nd floor",
                    observacao_pagamento='=HYPERLINK("https://example.invalid","unsafe")\nReferência: 001 /  002!')
    (raiz / "tenant-a" / "cadastro_fornecedores.json").write_text(
        json.dumps({"fornecedores": [registro]}), encoding="utf-8",
    )
    sheet = _carregar_commercial_invoice(excel_service._gerar_commercial_invoice_bytes(
        {"fornecedor_id": "fornecedor-a", "itens": []}, client_id="tenant-a",
    ))
    texto = sheet["E13"].value
    for campo in ("conta_beneficiario", "codigo_banco", "codigo_agencia", "endereco_beneficiario", "observacao_pagamento"):
        assert registro[campo] in texto
    assert texto.endswith(registro["observacao_pagamento"])
    assert sheet["E13"].data_type == "s"
    assert _formulas(sheet) == []


@pytest.mark.parametrize(("numero_oficial", "numero_legado", "esperado"), [
    ("  INV-OFICIAL-0001  ", "INV-LEGADO-0002", "INV-OFICIAL-0001"),
    ("", "  INV-LEGADO-0002  ", "INV-LEGADO-0002"),
    (None, "INV-LEGADO-0003", "INV-LEGADO-0003"),
    ("=1+1", "INV-LEGADO-0004", "=1+1"),
    ("", "", ""),
    (None, None, ""),
])
def test_commercial_invoice_numero_atual_no_cabecalho_e_ambos_rodapes(
    fornecedores_rodape_invoice, numero_oficial, numero_legado, esperado,
):
    raiz, registro = fornecedores_rodape_invoice
    arquivo = raiz / "tenant-a" / "cadastro_fornecedores.json"
    bytes_originais = arquivo.read_bytes()
    lista = {"fornecedor_id": "fornecedor-a", "numero_invoice": numero_oficial,
             "invoice": numero_legado, "itens": []}
    lista_original = copy.deepcopy(lista)
    sheet = _carregar_commercial_invoice(excel_service._gerar_commercial_invoice_bytes(lista, client_id="tenant-a"))
    assert sheet["C3"].value == (esperado or None)
    assert registro["observacao_pagamento"] in sheet["E13"].value
    if esperado:
        assert f"Note - invoice {esperado}" in sheet["A13"].value.splitlines()
        assert sheet["E13"].value.endswith(f"Please include the order invoice number {esperado}.")
    else:
        assert "Note - invoice" not in sheet["A13"].value
        assert "Please include the order invoice number" not in sheet["E13"].value
        assert sheet["E13"].value.endswith("Remark: Notify the supplier after payment.")
    if numero_oficial:
        assert numero_legado not in sheet["A13"].value
        assert numero_legado not in sheet["E13"].value
    assert _formulas(sheet) == []
    assert lista == lista_original
    assert arquivo.read_bytes() == bytes_originais


def test_commercial_invoice_rodape_numero_sem_observacao_cadastrada(fornecedores_rodape_invoice):
    raiz, _registro = fornecedores_rodape_invoice
    arquivo = raiz / "tenant-a" / "cadastro_fornecedores.json"
    bytes_originais = arquivo.read_bytes()
    sheet = _carregar_commercial_invoice(excel_service._gerar_commercial_invoice_bytes(
        {"fornecedor_id": "fornecedor-b", "numero_invoice": "INV-SEM-OBS-0001", "itens": []}, client_id="tenant-a",
    ))
    assert sheet["A1"].value == "Fornecedor B"
    assert "Note - invoice INV-SEM-OBS-0001" in sheet["A13"].value.splitlines()
    assert sheet["E13"].value == "Remark: Please include the order invoice number INV-SEM-OBS-0001."
    assert arquivo.read_bytes() == bytes_originais


def test_commercial_invoice_duas_listas_usam_seu_numero_em_ambos_rodapes(fornecedores_rodape_invoice):
    numeros = ("INV-SINTETICA-A-001", "INV-SINTETICA-B-002")
    for indice, numero in enumerate(numeros):
        lista = {**_lista_aprovada(), "id": f"lista-{indice}", "fornecedor_id": "fornecedor-a",
                 "numero_invoice": numero}
        sheet = _carregar_commercial_invoice(excel_service._gerar_commercial_invoice_bytes(lista, client_id="tenant-a"))
        assert sheet["C3"].value == numero
        assert f"Note - invoice {numero}" in sheet["A15"].value.splitlines()
        assert sheet["E15"].value.endswith(f"Please include the order invoice number {numero}.")
        for numero_incorreto in (numeros[1 - indice], "JK55-18082026", "JK48-18082026"):
            assert numero_incorreto not in sheet["A15"].value
            assert numero_incorreto not in sheet["E15"].value


@pytest.mark.parametrize("multilinhas", [False, True])
def test_commercial_invoice_rodape_extenso_expande_linhas_sem_perder_dados_ou_totais(
    fornecedores_rodape_invoice, multilinhas,
):
    raiz, registro = fornecedores_rodape_invoice
    limites = {"moeda_pagamento": 12, "conta_beneficiario": 120, "swift": 32,
               "pais_regiao_beneficiario": 120, "nome_beneficiario": 240, "endereco_beneficiario": 1200,
               "banco_beneficiario": 240, "endereco_banco": 1200, "codigo_banco": 40,
               "codigo_agencia": 40, "observacao_pagamento": 1600}
    for campo, limite in limites.items():
        registro[campo] = ("Linha longa de exemplo\n" * limite)[:limite] if multilinhas else "W" * limite
    arquivo = raiz / "tenant-a" / "cadastro_fornecedores.json"
    arquivo.write_text(json.dumps({"fornecedores": [registro]}), encoding="utf-8")
    bytes_originais = arquivo.read_bytes()
    lista = {**_lista_aprovada(), "fornecedor_id": "fornecedor-a"}
    sheet = _carregar_commercial_invoice(excel_service._gerar_commercial_invoice_bytes(lista, client_id="tenant-a"))

    assert sheet.max_row > 15
    texto = sheet["E15"].value
    for campo in limites:
        assert registro[campo].strip() in texto
    alturas = [sheet.row_dimensions[linha].height for linha in range(15, sheet.max_row + 1)]
    assert sum(alturas) > 409.5
    assert all(0 < altura <= 409.5 for altura in alturas)
    assert {f"A15:D{sheet.max_row}", f"E15:H{sheet.max_row}", f"I15:I{sheet.max_row}"} <= {
        str(intervalo) for intervalo in sheet.merged_cells.ranges
    }
    assert sheet["I15"].value == "Company stamp and signature"
    assert f"$A$1:$I${sheet.max_row}" in str(sheet.print_area)
    assert sheet["B9"].value == "SKU 001"
    assert sheet["B10"].value == "SKU 002"
    assert sheet["C11"].value == "TOTAL NET WEIGHT"
    assert sheet["G11"].value == 5
    assert sheet["I11"].value == sheet["I14"].value == 19
    assert sheet["H14"].value == "TOTAL:"
    assert _formulas(sheet) == []
    assert arquivo.read_bytes() == bytes_originais


@pytest.mark.parametrize("conteudo", ["{", '{"fornecedores":{}}'])
def test_commercial_invoice_rodape_cadastro_invalido_propaga_erro(fornecedores_rodape_invoice, conteudo):
    raiz, _registro = fornecedores_rodape_invoice
    arquivo = raiz / "tenant-a" / "cadastro_fornecedores.json"
    arquivo.write_text(conteudo, encoding="utf-8")
    with pytest.raises(HTTPException) as erro:
        excel_service._gerar_commercial_invoice_bytes(
            {"fornecedor_id": "fornecedor-a", "supplier": "Nome historico", "itens": []}, client_id="tenant-a",
        )
    assert erro.value.status_code == 500
    assert arquivo.read_text(encoding="utf-8") == conteudo


def test_commercial_invoice_rodape_cadastro_ilegivel_propaga_erro(fornecedores_rodape_invoice, monkeypatch):
    import builtins

    raiz, _registro = fornecedores_rodape_invoice
    arquivo = raiz / "tenant-a" / "cadastro_fornecedores.json"
    abrir_original = builtins.open

    def abrir(caminho, *args, **kwargs):
        if str(caminho) == str(arquivo):
            raise PermissionError("arquivo indisponivel")
        return abrir_original(caminho, *args, **kwargs)

    monkeypatch.setattr(builtins, "open", abrir)
    with pytest.raises(HTTPException) as erro:
        excel_service._gerar_commercial_invoice_bytes(
            {"fornecedor_id": "fornecedor-a", "supplier": "Nome historico", "itens": []}, client_id="tenant-a",
        )
    assert erro.value.status_code == 500


@pytest.fixture
def nacionalizadores_invoice(tmp_path, monkeypatch):
    nacionalizador = {
        "id": "nacionalizador-a",
        "nome_empresa": "ACME BRASIL COMERCIAL IMPORTADORA E EXPORTADORA LTDA",
        "endereco_empresa": "RODOVIA DE EXEMPLO, 100, CENTRO, CIDADE EXEMPLO, BRAZIL",
        "nome_contato": "Contato de Exemplo",
        "telefone": "+00 0000-0000",
        "cnpj": "00.000.000/0001-00",
    }
    outro = {**nacionalizador, "nome_empresa": "Empresa do cliente B", "endereco_empresa": "Endereco B"}
    registros = {
        "tenant-a": [nacionalizador],
        "tenant-b": [outro, {**outro, "id": "nacionalizador-outro"}],
    }
    for client_id, empresas in registros.items():
        pasta = tmp_path / client_id
        pasta.mkdir()
        (pasta / "cadastro_nacionalizadores.json").write_text(
            json.dumps({"schema": "jk.cadastro.nacionalizadores.v1", "nacionalizadores": empresas}),
            encoding="utf-8",
        )
    monkeypatch.setattr(excel_service, "get_tenant_path", lambda client_id: str(tmp_path / client_id))
    monkeypatch.setattr(
        cadastro_nacionalizadores,
        "get_tenant_path",
        lambda *_args: pytest.fail("Invoice deve usar a pasta canonica fornecida pelo exportador"),
    )
    monkeypatch.setattr(
        cadastro_compatibilidade,
        "visao_produtos_cadastro_contexto_loja",
        lambda *_args, **_kwargs: {"store_id": "", "scope": "unresolved"},
    )
    return tmp_path, nacionalizador


def test_commercial_invoice_nacionalizador_preserva_layout_totais_foto_e_fontes(nacionalizadores_invoice, monkeypatch):
    _cadastro_produtos_invoice(monkeypatch)
    raiz, _registro = nacionalizadores_invoice
    arquivo = raiz / "tenant-a" / "cadastro_nacionalizadores.json"
    cadastro_original = arquivo.read_bytes()
    foto = raiz / "tenant-a" / "cadastro_fotos" / "001.png"
    foto.parent.mkdir()
    foto.write_bytes(PNG_1X1)
    monkeypatch.setattr(
        excel_service, "_resolver_foto_cadastro_sku",
        lambda _client_id, sku, _foto_ref: str(foto) if sku == "001" else "",
    )
    lista = {**_lista_aprovada(), "nacionalizador_id": "nacionalizador-a", "nacionalizador": "Nome historico"}
    lista_original = copy.deepcopy(lista)
    sheet = _carregar_commercial_invoice(
        excel_service._gerar_commercial_invoice_bytes(lista, client_id="tenant-a", data_aprovacao="2026-08-18")
    )

    assert sheet["A4"].value.splitlines() == [
        "IMPORTER: ACME BRASIL COMERCIAL IMPORTADORA E",
        "EXPORTADORA LTDA",
        "ADDRESS: RODOVIA DE EXEMPLO, 100, CENTRO, CIDADE EXEMPLO, BRAZIL",
        "CONTACT: Contato de Exemplo +00 0000-0000",
        "CNP.J. 00.000.000/0001-00",
    ]
    assert sheet["A4"].font.name == "Times New Roman"
    assert sheet["A4"].font.sz == 9
    assert sheet["A4"].font.bold is True
    assert sheet["A4"].alignment.horizontal == "center"
    assert sheet["A4"].alignment.vertical == "center"
    assert sheet["A4"].alignment.wrap_text is True
    assert sheet["A4"].data_type == "s"
    assert sheet["A1"].value == "Fornecedor Exemplo Ltd."
    assert sheet["C3"].value == "JK48-18082026"
    assert sheet["E3"].value == "August 18, 2026"
    assert sheet["F4"].value == "STORE: JK Pecas"
    assert {"A4:E6", "F4:I6", "D9:E9", "D10:E10"} <= {str(r) for r in sheet.merged_cells.ranges}
    assert sum(sheet.row_dimensions[r].height for r in (4, 5, 6)) == pytest.approx(171)
    assert sheet.row_dimensions[6].height == 147
    assert sheet.max_row == 15
    assert sheet["B9"].value == "SKU 001"
    assert sheet["G11"].value == 5
    assert sheet["I11"].value == sheet["I14"].value == 19
    assert len(sheet._images) == 1
    assert sheet._images[0].anchor._from.col == 5
    assert sheet._images[0].anchor._from.row == 8
    assert _formulas(sheet) == []
    assert lista == lista_original
    assert arquivo.read_bytes() == cadastro_original


@pytest.mark.parametrize("client_id", ["tenant-a", "tenant-b"])
def test_commercial_invoice_resolve_mesmo_id_exclusivamente_no_cliente(nacionalizadores_invoice, client_id):
    raiz, _registro = nacionalizadores_invoice
    lista = {"nacionalizador_id": "nacionalizador-a", "nacionalizador": "Nome historico", "itens": []}
    sheet = _carregar_commercial_invoice(excel_service._gerar_commercial_invoice_bytes(lista, client_id=client_id))
    registros = json.loads((raiz / client_id / "cadastro_nacionalizadores.json").read_text(encoding="utf-8"))
    cadastro = registros["nacionalizadores"][0]
    assert f"ADDRESS: {cadastro['endereco_empresa']}" in sheet["A4"].value
    assert "Nome historico" not in sheet["A4"].value
    assert ("Empresa do cliente B" in sheet["A4"].value) == (client_id == "tenant-b")


@pytest.mark.parametrize("store_id", ["store-a", "store-b"])
def test_commercial_invoice_nacionalizador_compartilhado_preserva_buyer_da_loja(nacionalizadores_invoice, monkeypatch, store_id):
    raiz, _registro = nacionalizadores_invoice
    compradores = {
        "store-a": {"nome_empresa": "Comprador A", "tax_id": "TAX-A"},
        "store-b": {"nome_empresa": "Comprador B", "tax_id": "TAX-B"},
    }
    (raiz / "tenant-a" / "cadastro_importador_loja.json").write_text(
        json.dumps({"schema": "jk.cadastro.importador_loja.v1", "lojas": compradores}), encoding="utf-8",
    )
    monkeypatch.setattr(cadastro_importador_loja.lojas, "get_tenant_path", lambda client_id: str(raiz / client_id))
    monkeypatch.setattr(
        cadastro_compatibilidade, "visao_produtos_cadastro_contexto_loja",
        lambda _client_id, referencia, **_kwargs: {"store_id": referencia},
    )
    lista = {"store_id": store_id, "loja": "Mesmo nome de loja", "nacionalizador_id": "nacionalizador-a", "itens": []}
    sheet = _carregar_commercial_invoice(excel_service._gerar_commercial_invoice_bytes(lista, client_id="tenant-a"))
    assert sheet["F4"].value == f"Buyer: {compradores[store_id]['nome_empresa']}\nTAX ID: {compradores[store_id]['tax_id']}"
    assert sheet["A4"].value.startswith("IMPORTER: ACME BRASIL")
    assert lista["store_id"] == store_id


@pytest.mark.parametrize(
    ("client_id", "nacionalizador_id"),
    [("tenant-a", ""), ("tenant-a", "removido"), ("tenant-a", "nacionalizador-outro"),
     ("tenant-sem-cadastro", "nacionalizador-a"), (None, "nacionalizador-a")],
)
def test_commercial_invoice_nacionalizador_ausente_usa_so_nome_historico(
    nacionalizadores_invoice, client_id, nacionalizador_id,
):
    lista = {"nacionalizador_id": nacionalizador_id, "nacionalizador": "Empresa historica", "itens": []}
    sheet = _carregar_commercial_invoice(excel_service._gerar_commercial_invoice_bytes(lista, client_id=client_id))
    assert sheet["A4"].value == "IMPORTER: Empresa historica"
    assert sheet.row_dimensions[6].height == 147


@pytest.mark.parametrize("nacionalizador_id", [None, "", "removido", "nacionalizador-outro"])
def test_commercial_invoice_sem_nome_nacionalizador_mantem_list(nacionalizadores_invoice, nacionalizador_id):
    lista = {"nome_lista": "Lista antiga", "nacionalizador_id": nacionalizador_id, "itens": []}
    sheet = _carregar_commercial_invoice(excel_service._gerar_commercial_invoice_bytes(lista, client_id="tenant-a"))
    assert sheet["A4"].value == "LIST: Lista antiga"


def test_commercial_invoice_sem_id_nao_consulta_cadastro_nacionalizadores(monkeypatch):
    monkeypatch.setattr(
        cadastro_nacionalizadores, "ler_nacionalizador_por_id",
        lambda *_args, **_kwargs: pytest.fail("Sem vinculo nao deve consultar cadastro"),
    )
    monkeypatch.setattr(
        cadastro_compatibilidade, "visao_produtos_cadastro_contexto_loja",
        lambda *_args, **_kwargs: {"store_id": "", "scope": "unresolved"},
    )
    sheet = _carregar_commercial_invoice(
        excel_service._gerar_commercial_invoice_bytes({"nome_lista": "Lista sem vinculo", "itens": []}, client_id="tenant-a")
    )
    assert sheet["A4"].value == "LIST: Lista sem vinculo"


def test_commercial_invoice_retirada_do_vinculo_volta_ao_nome_da_lista(nacionalizadores_invoice, monkeypatch):
    from backend.services import medias_compras_common

    raiz, _registro = nacionalizadores_invoice
    estado = [{"id": "lista-1", "nome_lista": "Lista desvinculada", "nacionalizador_id": "nacionalizador-a",
               "nacionalizador": "Empresa historica", "itens": []}]
    monkeypatch.setattr(medias_compras_common, "get_tenant_path", lambda client_id: str(raiz / client_id))
    monkeypatch.setattr(listas_service, "_carregar_listas_pedidos", lambda _client_id: copy.deepcopy(estado))
    monkeypatch.setattr(listas_service, "_salvar_listas_pedidos", lambda _client_id, listas: estado.__setitem__(slice(None), listas))
    monkeypatch.setattr(listas_service, "_limpar_cache_lista_pedido", lambda *_args, **_kwargs: None)
    asyncio.run(listas_service.api_medias_compras_lista_pedido_editar(
        "lista-1", ListaPedidoUpdateRequest(nacionalizador_id=""), client_id="tenant-a",
    ))
    assert estado[0]["nacionalizador_id"] == estado[0]["nacionalizador"] == ""
    sheet = _carregar_commercial_invoice(excel_service._gerar_commercial_invoice_bytes(estado[0], client_id="tenant-a"))
    assert sheet["A4"].value == "LIST: Lista desvinculada"


@pytest.mark.parametrize(("nome_contato", "telefone"), [("Pessoa", ""), ("", "+00 0000-0000"), ("", "")])
def test_commercial_invoice_nacionalizador_omite_campos_vazios(nacionalizadores_invoice, nome_contato, telefone):
    raiz, registro = nacionalizadores_invoice
    registro.update(nome_empresa="Empresa A", nome_contato=nome_contato, telefone=telefone, endereco_empresa="", cnpj="")
    (raiz / "tenant-a" / "cadastro_nacionalizadores.json").write_text(
        json.dumps({"schema": "jk.cadastro.nacionalizadores.v1", "nacionalizadores": [registro]}), encoding="utf-8",
    )
    lista = {"nacionalizador_id": "nacionalizador-a", "itens": []}
    sheet = _carregar_commercial_invoice(excel_service._gerar_commercial_invoice_bytes(lista, client_id="tenant-a"))
    esperado = "IMPORTER: Empresa A"
    contato = " ".join(filter(None, (nome_contato, telefone)))
    if contato:
        esperado += f"\nCONTACT: {contato}"
    assert sheet["A4"].value == esperado


def test_commercial_invoice_nacionalizador_preserva_token_longo_do_nome(nacionalizadores_invoice):
    raiz, registro = nacionalizadores_invoice
    nome = '=HYPERLINK("https://example.invalid/' + "a" * 60 + '","unsafe")'
    registro["nome_empresa"] = nome
    (raiz / "tenant-a" / "cadastro_nacionalizadores.json").write_text(
        json.dumps({"schema": "jk.cadastro.nacionalizadores.v1", "nacionalizadores": [registro]}), encoding="utf-8",
    )
    sheet = _carregar_commercial_invoice(excel_service._gerar_commercial_invoice_bytes(
        {"nacionalizador_id": "nacionalizador-a", "itens": []}, client_id="tenant-a",
    ))
    assert sheet["A4"].value.startswith("IMPORTER:\n" + nome + "\nADDRESS:")
    assert sheet["A4"].data_type == "s"
    assert _formulas(sheet) == []


@pytest.mark.parametrize("linhas_endereco", [20, 40, None])
def test_commercial_invoice_nacionalizador_extenso_literal_ajusta_altura_sem_mutacao(nacionalizadores_invoice, linhas_endereco):
    raiz, registro = nacionalizadores_invoice
    registro.update(
        nome_empresa='=HYPERLINK("https://example.invalid","unsafe")',
        endereco_empresa=(
            "W" * 1000 if linhas_endereco is None
            else "\n".join(f"Endereco exemplo {indice:02}" for indice in range(linhas_endereco))
        ),
        nome_contato="Contato de exemplo " * 8,
        telefone="+00 0000-0000",
    )
    arquivo = raiz / "tenant-a" / "cadastro_nacionalizadores.json"
    arquivo.write_text(json.dumps({"schema": "jk.cadastro.nacionalizadores.v1", "nacionalizadores": [registro]}), encoding="utf-8")
    bytes_originais = arquivo.read_bytes()
    lista = {"nacionalizador_id": "nacionalizador-a", "itens": []}
    sheet = _carregar_commercial_invoice(excel_service._gerar_commercial_invoice_bytes(lista, client_id="tenant-a"))
    assert '=HYPERLINK("https://example.invalid","unsafe")' in sheet["A4"].value
    if linhas_endereco is None:
        assert "ADDRESS: " + "W" * 1000 in sheet["A4"].value
        assert sum(sheet.row_dimensions[r].height for r in (4, 5, 6)) >= 20 * 12
    else:
        assert f"Endereco exemplo {linhas_endereco - 1:02}" in sheet["A4"].value
    assert sheet["A4"].data_type == "s"
    assert _formulas(sheet) == []
    _texto, altura = excel_service._bloco_nacionalizador_commercial_invoice(lista, "tenant-a")
    assert sum(sheet.row_dimensions[r].height for r in (4, 5, 6)) == pytest.approx(altura)
    assert all(sheet.row_dimensions[r].height <= 409.5 for r in (4, 5, 6))
    assert sheet.row_dimensions[6].height > 147
    if linhas_endereco == 40:
        assert sheet.row_dimensions[4].height > 12.75
    assert sheet["C9"].value == "TOTAL NET WEIGHT"
    assert sheet["F4"].value in (None, "")
    assert arquivo.read_bytes() == bytes_originais


@pytest.mark.parametrize("conteudo", ['{', '{"schema":"errado","nacionalizadores":[]}'])
def test_commercial_invoice_nacionalizador_invalido_nao_esconde_erro_com_historico(nacionalizadores_invoice, conteudo):
    raiz, _registro = nacionalizadores_invoice
    arquivo = raiz / "tenant-a" / "cadastro_nacionalizadores.json"
    arquivo.write_text(conteudo, encoding="utf-8")
    lista = {"nacionalizador_id": "nacionalizador-a", "nacionalizador": "Empresa historica", "itens": []}
    with pytest.raises(HTTPException) as erro:
        excel_service._gerar_commercial_invoice_bytes(lista, client_id="tenant-a")
    assert erro.value.status_code == 500
    assert arquivo.read_text(encoding="utf-8") == conteudo


def test_commercial_invoice_nacionalizador_ilegivel_nao_gera_bloco_parcial(nacionalizadores_invoice, monkeypatch):
    import builtins

    raiz, _registro = nacionalizadores_invoice
    arquivo = raiz / "tenant-a" / "cadastro_nacionalizadores.json"
    abrir_original = builtins.open

    def abrir(caminho, *args, **kwargs):
        if str(caminho) == str(arquivo):
            raise PermissionError("arquivo indisponivel")
        return abrir_original(caminho, *args, **kwargs)

    monkeypatch.setattr(builtins, "open", abrir)
    with pytest.raises(HTTPException) as erro:
        excel_service._gerar_commercial_invoice_bytes({"nacionalizador_id": "nacionalizador-a", "itens": []}, client_id="tenant-a")
    assert erro.value.status_code == 500


def test_gerador_inclui_pendentes_e_mantem_numericos_ausentes_em_branco():
    lista = {
        "id": "rascunho",
        "itens": [
            {"SKU": "001", "Título em português": "+nao executar", "compra_aprovada": False},
            {"SKU": "SKU 002", "compra_aprovada": None},
        ],
    }
    sheet = _carregar_commercial_invoice(
        excel_service._gerar_commercial_invoice_bytes(lista, data_aprovacao="")
    )

    for endereco in ("A1", "C3", "E3", "H3", "A4", "F4"):
        assert sheet[endereco].value in (None, "")
    assert sheet["A15"].value.splitlines() == [
        *CONDICOES_COMERCIAIS_INVOICE, "Payment Method: Bank transfer",
    ]
    assert sheet["E15"].value in (None, "")
    assert sheet["B9"].value == "SKU 001"
    assert sheet["D9"].value == "+nao executar"
    assert sheet["D9"].data_type == "s"
    assert sheet["B10"].value == "SKU 002"
    for endereco in ("G9", "H9", "I9", "G10", "H10", "I10"):
        assert sheet[endereco].value in (None, "")
    assert sheet["G11"].value == 0
    assert sheet["I11"].value == 0
    assert sheet["I14"].value == 0
    assert _formulas(sheet) == []


def test_gerador_commercial_invoice_aceita_lista_vazia():
    sheet = _carregar_commercial_invoice(
        excel_service._gerar_commercial_invoice_bytes({"id": "vazia", "itens": []})
    )

    assert sheet.max_row == 13
    assert sheet["C9"].value == "TOTAL NET WEIGHT"
    assert sheet["G9"].value == 0
    assert sheet["I9"].value == 0
    assert sheet["H12"].value == "TOTAL:"
    assert sheet["I12"].value == 0
    assert sheet["I13"].value == "Company stamp and signature"
    assert sheet["A13"].value.splitlines() == [
        *CONDICOES_COMERCIAIS_INVOICE, "Payment Method: Bank transfer",
    ]
    assert _formulas(sheet) == []


def test_gerador_posiciona_resumo_depois_de_93_itens():
    itens = [
        {
            "SKU": f"{indice:03d}",
            "NCM": "85365090",
            "Título em português": f"Produto {indice}",
            "Quantidade": 1,
            "Valor unidade": 2,
        }
        for indice in range(1, 94)
    ]

    sheet = _carregar_commercial_invoice(
        excel_service._gerar_commercial_invoice_bytes({"id": "jk-48", "itens": itens})
    )

    assert sheet.max_row == 106
    assert sheet["B101"].value == "SKU 093"
    assert sheet["C102"].value == "TOTAL NET WEIGHT"
    assert sheet["G102"].value == 93
    assert sheet["I102"].value == 186
    assert sheet["H105"].value == "TOTAL:"
    assert sheet["I105"].value == 186
    assert sheet["I106"].value == "Company stamp and signature"
    merges = {str(intervalo) for intervalo in sheet.merged_cells.ranges}
    assert "D101:E101" in merges
    assert "C102:D102" in merges
    assert "A106:D106" in merges
    assert "E106:H106" in merges


def test_foto_commercial_invoice_nao_pode_sair_do_tenant(tmp_path, monkeypatch):
    raiz_tenant = tmp_path / "tenant-a"
    (raiz_tenant / "cadastro_fotos").mkdir(parents=True)
    foto_outro_tenant = tmp_path / "tenant-b" / "cadastro_fotos" / "001.png"
    foto_outro_tenant.parent.mkdir(parents=True)
    foto_outro_tenant.write_bytes(PNG_1X1)
    monkeypatch.setattr(excel_service, "get_tenant_path", lambda _client_id: str(raiz_tenant))
    monkeypatch.setattr(excel_service, "PASTA_INFO", str(tmp_path / "default-info"))
    monkeypatch.setattr(excel_service, "_resolver_foto_cadastro_sku", lambda *_args: str(foto_outro_tenant))

    assert excel_service._resolver_caminho_foto_commercial_invoice("tenant-a", "", "001") is None


def test_endpoint_inclui_pendentes_e_usa_aprovacao_mais_recente(monkeypatch):
    lista = _lista_aprovada()
    lista["itens"][1]["compra_aprovada"] = False
    recebido = {}
    monkeypatch.setattr(listas_service, "_carregar_listas_pedidos", lambda _client_id: [lista])

    def gerar(lista_recebida, *, client_id, data_aprovacao):
        recebido.update(lista=copy.deepcopy(lista_recebida), client_id=client_id, data_aprovacao=data_aprovacao)
        return b"PK\x03\x04xlsx"

    monkeypatch.setattr(listas_service, "_gerar_commercial_invoice_bytes", gerar)
    resposta = asyncio.run(
        listas_service.api_medias_compras_lista_pedido_commercial_invoice("lista-1", client_id="tenant-a")
    )

    assert recebido["client_id"] == "tenant-a"
    assert recebido["data_aprovacao"] == "2026-08-18T14:30:00"
    assert len(recebido["lista"]["itens"]) == 2
    assert recebido["lista"]["itens"][1]["compra_aprovada"] is False
    assert resposta.media_type.endswith("spreadsheetml.sheet")
    assert resposta.headers["content-disposition"] == 'attachment; filename="commercial_invoice_JK48-18082026.xlsx"'


def test_endpoint_aceita_campos_datas_aprovacoes_e_lista_vazios(monkeypatch):
    lista = {"id": "lista-1", "nome_lista": "JK 48", "numero_invoice": "", "supplier": "", "itens": []}
    recebido = {}
    monkeypatch.setattr(listas_service, "_carregar_listas_pedidos", lambda _client_id: [lista])

    def gerar(lista_recebida, *, client_id, data_aprovacao):
        recebido.update(lista=copy.deepcopy(lista_recebida), data_aprovacao=data_aprovacao)
        return b"PK\x03\x04xlsx"

    monkeypatch.setattr(listas_service, "_gerar_commercial_invoice_bytes", gerar)
    resposta = asyncio.run(
        listas_service.api_medias_compras_lista_pedido_commercial_invoice("lista-1", client_id="tenant-a")
    )

    assert recebido["lista"]["itens"] == []
    assert recebido["data_aprovacao"] == ""
    assert resposta.headers["content-disposition"] == 'attachment; filename="commercial_invoice_JK_48.xlsx"'


@pytest.mark.parametrize(
    ("numero_invoice", "nome_lista", "lista_id", "esperado"),
    [
        ("INV 01", "JK 48", "lista-1", "commercial_invoice_INV_01.xlsx"),
        ("", "JK 48", "lista-1", "commercial_invoice_JK_48.xlsx"),
        ("", "", "lista-1", "commercial_invoice_lista-1.xlsx"),
    ],
)
def test_endpoint_nome_arquivo_prioriza_invoice_lista_e_id(monkeypatch, numero_invoice, nome_lista, lista_id, esperado):
    lista = {"id": lista_id, "numero_invoice": numero_invoice, "nome_lista": nome_lista, "itens": []}
    monkeypatch.setattr(listas_service, "_carregar_listas_pedidos", lambda _client_id: [lista])
    monkeypatch.setattr(listas_service, "_gerar_commercial_invoice_bytes", lambda *_args, **_kwargs: b"PK\x03\x04xlsx")

    resposta = asyncio.run(
        listas_service.api_medias_compras_lista_pedido_commercial_invoice(lista_id, client_id="tenant-a")
    )

    assert resposta.headers["content-disposition"] == f'attachment; filename="{esperado}"'


@pytest.fixture
def invoice_edit_state(monkeypatch):
    lista = _lista_aprovada()
    lista.update({
        "store_id": "store-a", "numero_invoice": "Invoice anterior", "invoice": "Invoice legado",
        "meta_extra": {"preservar": True},
    })
    outro_cliente = {**copy.deepcopy(lista), "store_id": "store-b", "numero_invoice": "Outro cliente"}
    estado = {"tenant-a": [lista], "tenant-b": [outro_cliente]}
    monkeypatch.setattr(listas_service, "_carregar_listas_pedidos", lambda client_id: copy.deepcopy(estado[client_id]))

    def salvar(client_id, listas):
        estado[client_id] = copy.deepcopy(listas)

    monkeypatch.setattr(listas_service, "_salvar_listas_pedidos", salvar)
    monkeypatch.setattr(listas_service, "_limpar_cache_lista_pedido", lambda *_args, **_kwargs: None)
    return estado


@pytest.mark.parametrize("numero_invoice", ["  JK49-20261005  ", "", None])
def test_numero_invoice_passa_a_ser_persistido_no_backend(invoice_edit_state, numero_invoice):
    estado = invoice_edit_state
    antes = copy.deepcopy(estado)
    resultado = asyncio.run(
        listas_service.api_medias_compras_lista_pedido_editar(
            "lista-1",
            ListaPedidoUpdateRequest(numero_invoice=numero_invoice),
            client_id="tenant-a",
        )
    )

    esperado = numero_invoice.strip() if numero_invoice is not None else None
    salva = estado["tenant-a"][0]
    assert salva["numero_invoice"] == esperado
    assert "invoice" not in salva, "Um alias antigo nao pode ressurgir quando o numero e apagado"
    assert salva["updated_at"]
    lista_esperada = {**antes["tenant-a"][0], "numero_invoice": esperado, "updated_at": salva["updated_at"]}
    lista_esperada.pop("invoice")
    assert salva == lista_esperada, "Editar Invoice preserva loja, store_id, itens e demais metadados"
    assert estado["tenant-b"] == antes["tenant-b"], "IDs iguais em outro cliente permanecem isolados"
    assert resultado["lista"]["numero_invoice"] == (esperado or "")
    assert resultado["lista"]["store_id"] == "store-a"
    assert resultado["lista"]["itens"] == antes["tenant-a"][0]["itens"]
    sheet = _carregar_commercial_invoice(excel_service._gerar_commercial_invoice_bytes(salva))
    assert sheet["C3"].value == (esperado or None), "Commercial Invoice usa o numero persistido sem alias antigo"
    if esperado:
        assert f"Note - invoice {esperado}" in sheet["A15"].value.splitlines()
    else:
        assert "Note - invoice" not in sheet["A15"].value


@pytest.mark.parametrize("payload", [{"nome_lista": "Pedido renomeado"}, {"currency": "BRL"}])
def test_edicao_de_outro_campo_preserva_alias_invoice_legado(invoice_edit_state, payload):
    estado = invoice_edit_state
    antes = copy.deepcopy(estado)
    resultado = asyncio.run(
        listas_service.api_medias_compras_lista_pedido_editar(
            "lista-1", ListaPedidoUpdateRequest(**payload), client_id="tenant-a",
        )
    )

    salva = estado["tenant-a"][0]
    assert salva == {**antes["tenant-a"][0], **payload, "updated_at": salva["updated_at"]}
    assert salva["invoice"] == "Invoice legado"
    assert salva["numero_invoice"] == "Invoice anterior"
    assert resultado["lista"]["numero_invoice"] == "Invoice anterior"
    assert estado["tenant-b"] == antes["tenant-b"]


def test_endpoint_entrega_bytes_completos_sem_fragmentar_por_linha(monkeypatch):
    lista = {"id": "lista-1", "numero_invoice": "INV-001", "itens": []}
    conteudo = b"PK\x03\x04" + b"foto\nproduto\r\n\x00" * 4096
    monkeypatch.setattr(listas_service, "_carregar_listas_pedidos", lambda _client: [lista])
    monkeypatch.setattr(
        listas_service, "_gerar_commercial_invoice_bytes",
        lambda *_args, **_kwargs: conteudo,
    )
    resposta = listas_service._api_medias_compras_lista_pedido_commercial_invoice_sync(
        "lista-1", client_id="tenant-a",
    )
    mensagens = []

    async def receber():
        pytest.fail("Documento pronto nao precisa aguardar entrada HTTP")

    async def enviar(mensagem):
        mensagens.append(mensagem)

    asyncio.run(resposta({"type": "http"}, receber, enviar))

    assert resposta.headers["content-length"] == str(len(conteudo))
    assert resposta.headers["content-type"] == "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    assert resposta.headers["content-disposition"] == 'attachment; filename="commercial_invoice_INV-001.xlsx"'
    assert [mensagem["type"] for mensagem in mensagens] == ["http.response.start", "http.response.body"]
    assert mensagens[1]["body"] == conteudo
    assert not mensagens[1].get("more_body", False)


def test_router_expoe_download_do_commercial_invoice():
    rota = next(
        (
            item
            for item in create_medias_compras_router().routes
            if item.path == "/api/medias-compras/listas-pedidos/{lista_id}/commercial-invoice"
        ),
        None,
    )
    assert rota is not None
    assert rota.methods == {"GET"}
