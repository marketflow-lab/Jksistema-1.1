from __future__ import annotations

import io
import sqlite3
from datetime import date
from pathlib import Path

import fitz
import pytest
from openpyxl import load_workbook

from backend.modules.vendas import reports
from backend.modules.vendas.errors import VendasDomainError
from backend.services import bling_vendas


SCHEMA = """
CREATE TABLE vendas (
    id_unico TEXT,
    data TEXT,
    loja_conta TEXT,
    canal TEXT,
    numero TEXT,
    situacao TEXT,
    devolucao INTEGER,
    sku TEXT,
    produto TEXT,
    quantidade REAL,
    valor REAL,
    comprador TEXT,
    unidade_negocio TEXT,
    nota_fiscal_id TEXT
)
"""


def _create_db(path: Path, rows: list[tuple]) -> None:
    with sqlite3.connect(path) as connection:
        connection.execute(SCHEMA)
        connection.executemany(
            """
            INSERT INTO vendas (
                id_unico, data, loja_conta, canal, numero, situacao, devolucao,
                sku, produto, quantidade, valor, comprador, unidade_negocio,
                nota_fiscal_id
            ) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
            """,
            rows,
        )


@pytest.fixture
def pareto_sources(tmp_path: Path, monkeypatch):
    main_db = tmp_path / "vendas_historico_jk_pecas.db"
    mirror_db = tmp_path / "vendas_historico.db"
    base_rows = [
        ("a1", "2026-07-01", "JK Peças", "Bling", "1", "Concluída", 0, "001", "Produto A & Especial", 1, 400.0045, "Cliente", "Loja", ""),
        ("b1", "2026-07-02", "JK Peças", "Mercado Livre Full", "2", "Concluída", 0, "B", "Produto B", 2, 300, "Cliente", "Mercado Livre - Full", ""),
        ("c1", "2026-06-03", "JK Peças", "Bling", "3", "Concluída", 0, "C", "Produto C", 1, 200, "Cliente", "Loja", ""),
        ("d1", "2026-05-04", "JK Peças", "Bling", "4", "Concluída", 0, "D", "Produto D", 1, 100, "Cliente", "Loja", ""),
        ("e-order", "2026-07-05", "JK Peças", "Bling", "5", "Em aberto", 0, "E", "Produto E antigo", 1, 120, "Cliente", "Loja", "NF-55"),
        ("e-invoice", "2026-07-05", "JK Peças", "Bling", "NF5", "Faturado", 0, "E", "Produto E", 1, 120, "Cliente", "Loja", "NF-55"),
        ("return", "2026-07-05", "JK Peças", "Bling", "6", "Concluída", 1, "RET", "Devolução", 1, 999, "Cliente", "Loja", ""),
        ("cancel", "2026-07-05", "JK Peças", "Bling", "7", "Cancelada", 0, "CAN", "Cancelado", 1, 888, "Cliente", "Loja", ""),
        ("ebazar", "2026-07-05", "JK Peças", "Bling", "8", "Concluída", 0, "EBA", "Transferência", 1, 777, "EBAZAR", "Loja", ""),
        ("blank", "2026-07-05", "JK Peças", "Bling", "9", "Concluída", 0, "  ", "Sem SKU", 1, 666, "Cliente", "Loja", ""),
        ("fiscal", "2026-07-05", "JK Peças", "Bling", "10", "Concluída", 0, "ESTORNO DE CRÉDITO ICMS", "Ajuste", 1, 50, "Cliente", "Loja", ""),
        ("other-store", "2026-07-05", "Uai Mineirinho", "Bling", "11", "Concluída", 0, "OUT", "Outra loja", 1, 5000, "Cliente", "Loja", ""),
    ]
    _create_db(main_db, base_rows)
    _create_db(mirror_db, [base_rows[0]])

    monkeypatch.setattr(reports, "_listar_bancos_vendas_tenant", lambda *_args: [str(main_db), str(mirror_db)])
    monkeypatch.setattr(reports, "_deduplicar_vendas_consolidadas", bling_vendas._deduplicar_vendas_consolidadas)
    monkeypatch.setattr(
        reports,
        "_deve_excluir_venda_ebazar",
        lambda _return, buyer, channel: "EBAZAR" in f"{buyer} {channel}".upper(),
    )
    return main_db, mirror_db


def _generate() -> dict:
    return reports.generate_pareto_report.__wrapped__(
        client_id="tenant-pareto",
        loja="JK Peças",
        periodo="personalizado",
        data_inicio="2026-05-01",
        data_fim="2026-07-31",
    )


def test_periods_are_calendar_based_and_all_stores_is_blocked() -> None:
    assert reports.resolve_pareto_period("6m", today=date(2026, 7, 28)) == (
        "6m",
        "2026-02-01",
        "2026-07-28",
    )
    assert reports.resolve_pareto_period("12m", today=date(2026, 7, 28)) == (
        "12m",
        "2025-08-01",
        "2026-07-28",
    )
    with pytest.raises(VendasDomainError, match="conta específica"):
        reports.generate_pareto_report.__wrapped__(
            client_id="tenant-pareto", loja="__todas", periodo="12m"
        )
    with pytest.raises(VendasDomainError, match="posterior"):
        reports.resolve_pareto_period(
            "personalizado", "2026-07-31", "2026-07-01"
        )


def test_pareto_business_rules_and_account_isolation(pareto_sources) -> None:
    payload = _generate()

    assert payload["loja"] == "JK Peças"
    assert payload["periodo"]["ultima_data_disponivel"] == "2026-07-05"
    assert payload["resumo"]["faturamento_produtos"] == pytest.approx(1120.0045)
    assert payload["resumo"]["total_skus"] == 5
    assert payload["resumo"]["pareto_skus"] == 3
    assert payload["resumo"]["sku_corte"] == "C"
    assert [row["sku"] for row in payload["pareto_skus"]] == ["001", "B", "C"]
    assert payload["todos_skus"][0]["faturamento"] == pytest.approx(400.0045)
    assert payload["todos_skus"][0]["sku"] == "001"

    full = next(item for item in payload["canais"] if item["canal"] == "Mercado Livre Full")
    other = next(item for item in payload["canais"] if item["canal"].startswith("Demais canais"))
    assert full["faturamento"] == pytest.approx(300)
    assert other["faturamento"] == pytest.approx(820.0045)
    assert payload["ajuste_fiscal"] == {
        "sku": "ESTORNO DE CRÉDITO ICMS",
        "linhas": 1,
        "valor": 50.0,
        "incluido_no_pareto": False,
    }

    audit = payload["auditoria"]
    assert audit["duplicidades_entre_bases"] == 1
    assert audit["duplicidades_pedido_nota"] == 1
    assert audit["devolucoes_excluidas"] == 1
    assert audit["cancelamentos_excluidos"] == 1
    assert audit["ebazar_excluidos"] == 1
    assert audit["sku_vazio_excluidos"] == 1
    assert audit["linhas_elegiveis"] == 6
    assert audit["bases_processadas"] == 2
    assert len(audit["assinatura_fontes"]) == 64
    assert len(payload["mensal"]) == 3
    assert payload["analise_executiva"]["participacao_full"] == pytest.approx(300 / 1120.0045)
    assert payload["analise_executiva"]["cauda_skus"] == 2
    assert len(payload["analise_executiva"]["destaques"]) == 4


def test_xlsx_and_pdf_exports_are_valid_and_auditable(pareto_sources) -> None:
    payload = _generate()
    xlsx_bytes = reports.build_pareto_xlsx(payload)
    pdf_bytes = reports.build_pareto_pdf(payload)

    assert xlsx_bytes.startswith(b"PK")
    workbook = load_workbook(io.BytesIO(xlsx_bytes), data_only=False)
    assert workbook.sheetnames == [
        "Resumo",
        "Pareto 80%",
        "Todos os SKUs",
        "Mensal e Canais",
        "Fontes e Regras",
    ]
    assert workbook["Pareto 80%"]["B3"].value == "001"
    assert workbook["Pareto 80%"]["B3"].number_format == "@"
    assert len(workbook["Pareto 80%"]._charts) == 2
    assert len(workbook["Pareto 80%"]._charts[1].series) == 2
    assert len(workbook["Mensal e Canais"]._charts) == 1
    formulas = [
        cell.value
        for row in workbook["Fontes e Regras"].iter_rows()
        for cell in row
        if isinstance(cell.value, str) and cell.value.startswith("=")
    ]
    assert any("SUM('Todos os SKUs'!F3:F" in formula for formula in formulas)
    assert not any("#REF!" in formula for formula in formulas)

    assert pdf_bytes.startswith(b"%PDF")
    document = fitz.open(stream=pdf_bytes, filetype="pdf")
    text = "\n".join(page.get_text() for page in document)
    assert len(document) >= 3
    assert "Relatório Pareto 80%" in text
    assert "Leitura executiva" in text
    assert "Curva Pareto" in text
    assert "Faturamento mensal por canal" in text
    assert "JK Peças" in text
    assert "SKU de corte" in text
    assert "001" in text
    assert "ESTORNO DE CRÉDITO ICMS" in text


def test_export_contract_and_invalid_format(pareto_sources) -> None:
    content, filename, media_type = reports.export_pareto_report(
        client_id="tenant-pareto",
        loja="JK Peças",
        formato="xlsx",
        periodo="personalizado",
        data_inicio="2026-05-01",
        data_fim="2026-07-31",
    )
    assert content.startswith(b"PK")
    assert filename == "relatorio_pareto_80_jk_pecas_2026-05-01_2026-07-31.xlsx"
    assert media_type.endswith("spreadsheetml.sheet")

    with pytest.raises(VendasDomainError, match="Formato inválido"):
        reports.export_pareto_report(
            client_id="tenant-pareto",
            loja="JK Peças",
            formato="csv",
        )
