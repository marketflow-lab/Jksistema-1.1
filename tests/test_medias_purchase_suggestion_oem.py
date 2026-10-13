"""Regression coverage for OEM selection through the suggestion XLSX pipeline."""

import copy
import csv
import io
import json

import openpyxl
import pandas as pd
import pytest
from PIL import Image

from backend.services import cadastro_fotos as fotos
from backend.services import cadastro_lojas_produtos as produtos
from backend.services import integracoes
from backend.services import medias_compras_common as common
from backend.services import medias_compras_excel as excel
from backend.services import medias_compras_fiscal as fiscal
from backend.services import medias_compras_sugestoes as sugestoes
from backend.services import medias_compras_visao as visao


TIMESTAMP = "2026-09-04T12:32:19.668610+00:00"


def _csv(path, rows):
    fields = list(dict.fromkeys(key for row in rows for key in row))
    with path.open("w", encoding="utf-8-sig", newline="") as stream:
        writer = csv.DictWriter(stream, fieldnames=fields)
        writer.writeheader()
        writer.writerows(rows)


@pytest.fixture
def ambiente(monkeypatch, tmp_path):
    root = tmp_path / "info"
    stores = {
        client: [{"store_id": store, "nome": f"{client} {store}"}
                 for store in ("store-a", "store-b")]
        for client in ("cliente-a", "cliente-b")
    }
    state = {"root": root, "stores": stores, "photo_bytes": {}, "load_calls": [], "view_calls": []}
    for client, client_stores in stores.items():
        tenant = root / client
        tenant.mkdir(parents=True)
        (tenant / fotos.CADASTRO_FOTOS_CONFIG_ARQUIVO).write_text(json.dumps({
            "schema": fotos.CADASTRO_FOTOS_CONFIG_SCHEMA,
            "strict_store_scope": True,
            "shared_groups": [],
        }), encoding="utf-8")
        rows = []
        for store in client_stores:
            store_id = store["store_id"]
            segment = fotos._cadastro_store_id_foto_segmento(store_id)
            for sku in ("001", "002", "003"):
                ref = f"cadastro_fotos/lojas/{segment}/{sku}.png"
                photo = tenant / ref
                photo.parent.mkdir(parents=True, exist_ok=True)
                color = (20 if client == "cliente-a" else 180,
                         30 if store_id == "store-a" else 200, int(sku))
                Image.new("RGB", (12, 8), color).save(photo)
                state["photo_bytes"][client, store_id, sku] = photo.read_bytes()
                rows.append({"store_id": store_id, "sku": sku, "sku_normalizado": sku,
                             "nome": f"Produto {sku}", "foto": ref, "row_version": "1",
                             "updated_at_utc": TIMESTAMP, "deleted_at_utc": ""})
        _csv(tenant / "cadastro_produtos_lojas.csv", rows)

    def tenant_path(client):
        assert client in stores
        return str(root / client)

    for module in (common, fiscal, sugestoes, excel, fotos, produtos):
        monkeypatch.setattr(module, "get_tenant_path", tenant_path)
        monkeypatch.setattr(module, "PASTA_INFO", str(root), raising=False)
    monkeypatch.setattr(integracoes, "_get_tenant_path", tenant_path)
    monkeypatch.setattr(integracoes, "PASTA_INFO", str(root))
    monkeypatch.setattr(integracoes, "ler_lojas", lambda client: copy.deepcopy(stores[client]))

    def catalog(client):
        state["load_calls"].append(client)
        path = root / client / "cadastro_produtos.csv"
        return pd.read_csv(path, dtype=str).fillna(""), str(path)

    def backups(client):
        path = root / client / "cadastro_produtos.backup_fixture.csv"
        return [str(path)] if path.exists() else []

    def configure(client, fields, source="principal"):
        rows = [{"sku": "001", "ncm": "87089990", "m3 individual": "0.125", **fields},
                {"sku": "002", "ncm": "87089990", "m3 individual": "0.125", "OEM": "OEM-002"},
                {"sku": "003", "ncm": "87089990", "m3 individual": "0.125", "OEM": "OEM-003"}]
        if source == "principal":
            _csv(root / client / "cadastro_produtos.csv", rows)
        else:
            _csv(root / client / "cadastro_produtos.csv", [
                {key: row[key] for key in ("sku", "ncm", "m3 individual")} for row in rows
            ])
            _csv(root / client / "cadastro_produtos.backup_fixture.csv", rows)

    def fake_visao(*, meses, loja, store_id, client_id):
        assert meses == 6
        assert loja == f"{client_id} {store_id}"
        assert store_id in {store["store_id"] for store in stores[client_id]}
        state["view_calls"].append((client_id, store_id))
        items = [{"sku": "001", "compra_sugerida": 5, "titulo_anuncio": "Produto 001"},
                 {"sku": "002", "compra_sugerida": 4, "titulo_anuncio": "Produto 002"},
                 {"sku": "003", "compra_sugerida": 0, "titulo_anuncio": "Produto 003"}]
        if state.get("bound_oem"):
            items[0]["oem"] = f"OEM-{client_id}-{store_id}"
        return {"itens": items, "loja": loja, "store_id": store_id}

    monkeypatch.setattr(fiscal, "_carregar_cadastro_para_impostos", catalog)
    monkeypatch.setattr(fiscal, "_caminhos_cadastros_meta", backups)
    monkeypatch.setattr(fiscal, "_indice_ncm_referencia_aliquotas", lambda _client: {
        "87089990": {"ii": 12.5, "ipi": 5, "pis": 2.1, "cofins": 9.65},
    })
    monkeypatch.setattr(fiscal, "_indice_ncm_referencia_por_ncm", lambda _client: {
        "87089990": {"descricao_completa": "Descricao fiscal de teste"},
    })
    monkeypatch.setattr(visao, "_api_medias_compras_visao_sync", fake_visao)
    monkeypatch.setattr(sugestoes, "TEMP_FILES_STORAGE", {})
    monkeypatch.setattr(sugestoes, "TEMP_FILES_META", {})
    state["configure"] = configure
    return state


def _gerar(client="cliente-a", store="store-a"):
    response = sugestoes._api_medias_compras_gerar_lista_sugestao_sync(
        meses=6, loja=f"{client} {store}", store_id=store,
        quantidades_sugeridas='{"001":7,"002":0,"003":3}', client_id=client,
    )
    workbook = openpyxl.load_workbook(io.BytesIO(sugestoes.TEMP_FILES_STORAGE[response["file_id"]]))
    return response, workbook.active


@pytest.mark.parametrize("source", ["principal", "backup"])
@pytest.mark.parametrize("fields, expected", [
    ({"monofasico_verificado_em": TIMESTAMP, "consultado_em_utc_bling": TIMESTAMP}, None),
    ({"monofasico_verificado_em": TIMESTAMP, "OEM/ Model": "93110-2D000"}, "93110-2D000"),
    ({"consultado_em_utc_bling": TIMESTAMP, "cg_OEM": "SKU:19.1"}, "SKU:19.1"),
    ({"consultado_em_utc": TIMESTAMP, "modelo": "Modelo 2008/2013"}, "Modelo 2008/2013"),
    ({"monofasico_verificado_em": TIMESTAMP, "OEM": "2026-09-04"}, "2026-09-04"),
])
def test_sugestao_exporta_oem_real_e_preserva_quantidades_e_planilha(ambiente, source, fields, expected):
    ambiente["configure"]("cliente-a", fields, source)
    response, sheet = _gerar()

    assert sheet["D1"].value == "OEM/ Model"
    assert sheet["D2"].value == expected
    if expected is not None:
        assert sheet["D2"].data_type == "s"
    assert sheet["D3"].value == "OEM-003"
    assert [(sheet.cell(row, 1).value, sheet.cell(row, 7).value)
            for row in range(2, sheet.max_row + 1)] == [("001", 7), ("003", 3)]
    assert response["total_itens"] == 2
    assert response["total_quantidades_ajustadas"] == 3
    assert (sheet["I2"].value, sheet["I3"].value) == ("=G2*H2", "=G3*H3")
    assert (sheet["J2"].value, sheet["J3"].value) == ("0.875", "0.375")
    assert sheet.freeze_panes == "A2"
    assert sheet.auto_filter.ref == "A1:L3"
    assert [picture._data() for picture in sheet._images] == [
        ambiente["photo_bytes"]["cliente-a", "store-a", sku] for sku in ("001", "003")
    ]
    assert ambiente["view_calls"] == [("cliente-a", "store-a")]
    assert set(ambiente["load_calls"]) == {"cliente-a"}


@pytest.mark.parametrize("client, store", [
    ("cliente-a", "store-a"), ("cliente-a", "store-b"),
    ("cliente-b", "store-a"), ("cliente-b", "store-b"),
])
def test_sugestao_preserva_oem_informado_e_fotos_no_escopo(ambiente, client, store):
    for tenant in ambiente["stores"]:
        ambiente["configure"](tenant, {"monofasico_verificado_em": TIMESTAMP})
    ambiente["bound_oem"] = True
    _, sheet = _gerar(client, store)

    assert sheet["D2"].value == f"OEM-{client}-{store}"
    assert [picture._data() for picture in sheet._images] == [
        ambiente["photo_bytes"][client, store, sku] for sku in ("001", "003")
    ]
    assert ambiente["view_calls"] == [(client, store)]
    assert set(ambiente["load_calls"]) == {client}


@pytest.mark.parametrize("source", ["principal", "backup"])
@pytest.mark.parametrize("fields, expected", [
    ({"mlb_principal": "MLB1234567890"}, None),
    ({"link": "MLB1234567890", "cg_Link": "https://pt.aliexpress.com/item/1005001234567890.html"},
     "https://pt.aliexpress.com/item/1005001234567890.html"),
    ({"url": "MLB1234567890", "Link Aliexpress": "https://pt.aliexpress.com/item/1005001234567890.html"},
     "https://pt.aliexpress.com/item/1005001234567890.html"),
    ({"url_foto": "https://example.test/foto.png"}, None),
    ({"link_ml": "https://produto.mercadolivre.com.br/MLB-1234567890-produto-teste-_JM"},
     "https://produto.mercadolivre.com.br/MLB-1234567890-produto-teste-_JM"),
])
def test_sugestao_exporta_somente_url_cadastrada_de_produto(ambiente, source, fields, expected):
    ambiente["configure"]("cliente-a", fields, source)
    _, sheet = _gerar()

    assert sheet["F1"].value == "Link Aliexpress"
    assert sheet["F2"].value == expected
    if expected is None:
        assert sheet["F2"].hyperlink is None
    else:
        assert sheet["F2"].hyperlink.target == expected
    assert sheet["G2"].value == 7
    assert sheet["I2"].value == "=G2*H2"


@pytest.mark.parametrize("source", ["principal", "backup"])
@pytest.mark.parametrize("client, store", [
    ("cliente-a", "store-a"), ("cliente-a", "store-b"),
    ("cliente-b", "store-a"), ("cliente-b", "store-b"),
])
def test_sugestao_prioriza_link_atual_da_loja_sobre_link_legado(ambiente, source, client, store):
    old_link = "https://pt.aliexpress.com/item/link-antigo.html"
    for tenant in ambiente["stores"]:
        ambiente["configure"](tenant, {"Link Aliexpress": old_link}, source)
        path = ambiente["root"] / tenant / "cadastro_produtos_lojas.csv"
        rows = pd.read_csv(path, dtype=str).fillna("").to_dict("records")
        for row in rows:
            row["link"] = f"https://www.alibaba.com/product-detail/{tenant}-{row['store_id']}-atual.html"
        _csv(path, rows)

    _, sheet = _gerar(client, store)
    expected = f"https://www.alibaba.com/product-detail/{client}-{store}-atual.html"
    assert sheet["F2"].value == expected
    assert sheet["F2"].hyperlink.target == expected
    assert sheet["G2"].value == 7
    assert sheet["I2"].value == "=G2*H2"
    assert ambiente["view_calls"] == [(client, store)]
    assert set(ambiente["load_calls"]) == {client}


@pytest.mark.parametrize("value", ["", "MLB1234567890", "javascript:alert(1)"])
def test_sugestao_preserva_fallback_quando_link_da_loja_nao_e_url(ambiente, value):
    expected = "https://pt.aliexpress.com/item/link-legado-valido.html"
    ambiente["configure"]("cliente-a", {"Link Aliexpress": expected}, "backup")
    path = ambiente["root"] / "cliente-a" / "cadastro_produtos_lojas.csv"
    rows = pd.read_csv(path, dtype=str).fillna("").to_dict("records")
    for row in rows:
        row["link"] = value
    _csv(path, rows)

    _, sheet = _gerar()
    assert sheet["F2"].value == expected
    assert sheet["F2"].hyperlink.target == expected


def test_enriquecimento_preserva_link_preenchido_na_lista(ambiente):
    ambiente["configure"]("cliente-a", {"Link Aliexpress": "https://pt.aliexpress.com/item/antigo.html"})
    path = ambiente["root"] / "cliente-a" / "cadastro_produtos_lojas.csv"
    rows = pd.read_csv(path, dtype=str).fillna("").to_dict("records")
    for row in rows:
        row["link"] = "https://www.alibaba.com/product-detail/cadastro-atual.html"
    _csv(path, rows)
    informed = "https://fornecedor.example.test/produto-escolhido.html"
    item = {"SKU": "001", "Link": informed, "Quantidade": 7}
    result = fiscal._enriquecer_itens_lista_pedido_com_impostos(
        "cliente-a", [item], incluir_impostos=False, loja="store-a",
    )
    assert result[0]["Link"] == informed
    assert result[0]["Quantidade"] == 7
    assert item == {"SKU": "001", "Link": informed, "Quantidade": 7}
