import streamlit as st
import pandas as pd
import openpyxl
import base64
from io import BytesIO

# ==============================================================================
# --- LÓGICA DE PROCESSAMENTO (BACKEND) ---
# ==============================================================================

def find_sheet_name(workbook):
    """Encontra o nome da aba de promoções ou retorna a primeira."""
    for name in workbook.sheetnames:
        if "promoções" in name.lower() or "promo" in name.lower():
            return name
    return workbook.sheetnames[0]

def find_common_key(df_antiga, df_nova):
    """Encontra uma coluna chave comum para unir os dataframes."""
    POSSIBLE_KEYS = ['ITEM_ID', 'SELLER_SKU', 'SKU', 'ID', 'CÓDIGO', 'MLB']
    for key in POSSIBLE_KEYS:
        if key in df_antiga.columns and key in df_nova.columns:
            return key
    return None

def find_col(df_columns, names):
    """Encontra o primeiro nome de coluna correspondente em uma lista de possibilidades."""
    upper_cols = [str(c).upper().strip() for c in df_columns]
    for name in names:
        target = str(name).upper().strip()
        if target in upper_cols:
            return df_columns[upper_cols.index(target)]
    return None

def extract_comparison_data(file_antiga, file_nova):
    """
    Lê as duas planilhas e retorna um dataframe contendo APENAS os itens da Planilha Nova (how='right').
    """
    try:
        file_antiga.seek(0)
        file_nova.seek(0)
        sheet_name = 'Promoções'
        
        try:
            df_antiga = pd.read_excel(file_antiga, sheet_name=sheet_name, header=0, engine='openpyxl')
            df_nova = pd.read_excel(file_nova, sheet_name=sheet_name, header=0, engine='openpyxl')
        except:
            file_antiga.seek(0)
            file_nova.seek(0)
            df_antiga = pd.read_excel(file_antiga, header=0, engine='openpyxl')
            df_nova = pd.read_excel(file_nova, header=0, engine='openpyxl')

        if len(df_antiga) > 4: df_antiga = df_antiga.iloc[4:].reset_index(drop=True)
        if len(df_nova) > 4: df_nova = df_nova.iloc[4:].reset_index(drop=True)
        
        df_antiga.columns = df_antiga.columns.astype(str).str.strip().str.upper()
        df_nova.columns = df_nova.columns.astype(str).str.strip().str.upper()

        # Estatísticas
        status_col_antiga = find_col(df_antiga.columns, ['STATUS', 'ESTADO'])
        total_ativos_antiga = 0
        if status_col_antiga:
            total_ativos_antiga = df_antiga[status_col_antiga].astype(str).str.strip().str.lower().isin(['ativo', 'active']).sum()
        
        total_elegiveis_nova = len(df_nova)
        
        stats = {
            'ativos_antiga': total_ativos_antiga,
            'elegiveis_nova': total_elegiveis_nova
        }

        key_col = find_common_key(df_antiga, df_nova)
        if not key_col:
            return None, None, None, "Não foi possível encontrar uma coluna chave comum (ITEM_ID, SKU, etc.) nas duas planilhas."

        DISCOUNT_NAMES = ['DISCOUNT_PERCENTAGE', 'DISCOUNT %', 'DESCONTO', 'DESC %', 'DISCOUNT']
        PRICE_NAMES = ['FINAL_PRICE', 'PREÇO FINAL', 'PRECO FINAL', 'PRICE', 'PREÇO']
        STATUS_NAMES = ['STATUS', 'ESTADO']

        def get_col_map(df, suffix):
            cols = {key_col: key_col}
            status_col = find_col(df.columns, STATUS_NAMES)
            if status_col: cols[status_col] = f"STATUS{suffix}"
            discount_col = find_col(df.columns, DISCOUNT_NAMES)
            if discount_col: cols[discount_col] = f"DESC %{suffix}"
            price_col = find_col(df.columns, PRICE_NAMES)
            if price_col: cols[price_col] = f"PREÇO FINAL{suffix}"
            return cols

        cols_antiga = get_col_map(df_antiga, "_ANTIGA")
        cols_nova = get_col_map(df_nova, "_NOVA")

        sku_col = find_col(df_antiga.columns, ['SKU', 'SELLER_SKU'])
        if sku_col and sku_col not in cols_antiga:
            cols_antiga[sku_col] = 'SKU'

        df_antiga_sel = df_antiga[list(cols_antiga.keys())].rename(columns=cols_antiga)
        df_nova_sel = df_nova[list(cols_nova.keys())].rename(columns=cols_nova)

        df_antiga_sel[key_col] = df_antiga_sel[key_col].astype(str).str.strip()
        df_nova_sel[key_col] = df_nova_sel[key_col].astype(str).str.strip()

        # --- ALTERAÇÃO PRINCIPAL: how='right' ---
        # Mantém apenas as chaves que existem na planilha NOVA (direita)
        df_merged = pd.merge(df_antiga_sel, df_nova_sel, on=key_col, how='right')
        
        # Filtra para manter apenas linhas onde a chave começa com MLB
        df_merged = df_merged[df_merged[key_col].astype(str).str.strip().str.upper().str.startswith('MLB')]

        if df_merged.empty:
            return None, None, None, "Nenhum item da planilha nova encontrado."
            
        def get_default_action(row):
            status = row.get('STATUS_ANTIGA', '')
            if pd.isna(status): return 'Participar' # Item novo na campanha tende a participar
            status_str = str(status).strip().lower()
            if status_str in ['ativo', 'programado']: return 'Participar'
            return 'Não participar'

        df_merged['AÇÃO MANUAL'] = df_merged.apply(get_default_action, axis=1)

        return df_merged, df_antiga, stats, "Planilhas comparadas com sucesso."

    except Exception as e:
        return None, None, None, f"Erro ao processar planilhas: {e}"


def generate_updated_sheet(original_file_nova, edited_data_df, df_antiga_full=None):
    """
    Aplica as ações manuais na planilha nova original.
    NÃO ADICIONA NENHUMA LINHA NOVA. Apenas edita o que já existe.
    """
    try:
        original_file_nova.seek(0)
        wb = openpyxl.load_workbook(original_file_nova)
        ws_name = find_sheet_name(wb)
        ws = wb[ws_name]

        header = [str(c.value).strip().upper() if c.value else "" for c in ws[1]]
        
        DISCOUNT_NAMES = ['DISCOUNT_PERCENTAGE', 'DISCOUNT %', 'DESCONTO', 'DESC %', 'DISCOUNT']
        PRICE_NAMES = ['FINAL_PRICE', 'PREÇO FINAL', 'PRECO FINAL', 'PRICE', 'PREÇO']
        ACTION_NAMES = ['ACTION', 'AÇÃO', 'ACAO']
        
        key_col_name_df = find_common_key(edited_data_df, edited_data_df)
        
        key_col_idx = -1
        if key_col_name_df and key_col_name_df in header:
             key_col_idx = header.index(key_col_name_df) + 1
        else:
            possible_keys = ['ITEM_ID', 'SELLER_SKU', 'SKU', 'ID', 'CÓDIGO', 'MLB']
            for pk in possible_keys:
                if pk in header:
                    key_col_idx = header.index(pk) + 1
                    key_col_name_df = pk 
                    break
        
        action_col_name = find_col(header, ACTION_NAMES)
        discount_col_name = find_col(header, DISCOUNT_NAMES)
        price_col_name = find_col(header, PRICE_NAMES)

        if key_col_idx == -1 or not action_col_name:
            return None, f"Coluna chave ou 'ACTION' não encontrada na planilha nova. Headers: {header}"
            
        action_col_idx = header.index(action_col_name) + 1
        discount_col_idx = header.index(discount_col_name) + 1 if discount_col_name else -1
        price_col_idx = header.index(price_col_name) + 1 if price_col_name else -1

        col_chave_df = next((k for k in ['ITEM_ID', 'SELLER_SKU', 'SKU', 'ID', 'CÓDIGO', 'MLB'] if k in edited_data_df.columns), None)
        if not col_chave_df:
             return None, "Coluna chave não encontrada nos dados editados."

        edited_data_df['TEMP_KEY'] = edited_data_df[col_chave_df].astype(str).str.strip()
        update_map = edited_data_df.set_index('TEMP_KEY').to_dict('index')

        # Atualiza APENAS linhas existentes
        for row in ws.iter_rows(min_row=2):
            cell_key = row[key_col_idx - 1].value
            if cell_key is None:
                continue
            key_val = str(cell_key).strip()
            
            if key_val in update_map:
                item_data = update_map[key_val]
                
                # Atualiza Ação
                if 'AÇÃO MANUAL' in item_data:
                    row[action_col_idx - 1].value = item_data['AÇÃO MANUAL']
                
                # Atualiza Desconto
                if discount_col_idx != -1 and 'DESC %_NOVA' in item_data:
                    new_discount = item_data['DESC %_NOVA']
                    if pd.notna(new_discount) and new_discount != "":
                        try: row[discount_col_idx - 1].value = float(new_discount)
                        except ValueError: pass 
                
                # Atualiza Preço
                if price_col_idx != -1 and 'PREÇO FINAL_NOVA' in item_data:
                    new_price = item_data['PREÇO FINAL_NOVA']
                    if pd.notna(new_price) and new_price != "":
                        try: row[price_col_idx - 1].value = float(new_price)
                        except ValueError: pass

        output = BytesIO()
        wb.save(output)
        output.seek(0)
        
        return output, "Planilha atualizada gerada com sucesso."

    except Exception as e:
        return None, f"Erro ao gerar planilha: {e}"


# ==============================================================================
# --- INTERFACE VISUAL (FRONTEND) ---
# ==============================================================================

def render_page(navegar_para_callback):
    """
    Renderiza a página interativa de Renovação de Campanha.
    """
    if st.button("⬅️ Voltar ao Submenu", key="btn_voltar_renovacao"):
        st.session_state.clear() 
        navegar_para_callback('submenu_promo')

    st.markdown("## Renovação de Campanha Fixa (Interativo)")

    if 'comparison_data' not in st.session_state: st.session_state['comparison_data'] = None
    if 'df_antiga_raw' not in st.session_state: st.session_state['df_antiga_raw'] = None
    if 'edited_data' not in st.session_state: st.session_state['edited_data'] = None
    if 'renovacao_buffer' not in st.session_state: st.session_state['renovacao_buffer'] = None
    if 'original_nova_file' not in st.session_state: st.session_state['original_nova_file'] = None
    if 'stats_comparacao' not in st.session_state: st.session_state['stats_comparacao'] = None
    if 'renovacao_nome' not in st.session_state: st.session_state['renovacao_nome'] = None

    if st.session_state.get('comparison_data') is None:
        with st.form("upload_files_form"):
            col1, col2 = st.columns(2)
            with col1:
                st.info("📄 **Passo 1:** Anexar campanha Antiga")
                arquivo_antiga = st.file_uploader("Selecione arquivo XLSX (Antigo)", type=["xlsx"], key="up_antiga")
            with col2:
                st.info("🆕 **Passo 2:** Anexar campanha Nova")
                arquivo_nova = st.file_uploader("Selecione arquivo XLSX (Novo)", type=["xlsx"], key="up_nova")

            submitted = st.form_submit_button("🔍 Comparar Planilhas", type="primary", use_container_width=True)

        if submitted:
            if arquivo_antiga and arquivo_nova:
                with st.spinner("Analisando e comparando dados..."):
                    df, df_antiga_full, stats, msg = extract_comparison_data(arquivo_antiga, arquivo_nova)
                    if df is not None:
                        st.session_state['comparison_data'] = df
                        st.session_state['df_antiga_raw'] = df_antiga_full 
                        st.session_state['edited_data'] = df 
                        st.session_state['original_nova_file'] = arquivo_nova 
                        st.session_state['stats_comparacao'] = stats
                        st.session_state['renovacao_buffer'] = None
                        st.success(msg)
                        st.rerun()
                    else:
                        st.error(msg)
            else:
                st.warning("Por favor, selecione ambos os arquivos antes de comparar.")
    
    st.markdown("---")

    stats = st.session_state.get('stats_comparacao')
    if stats:
        st.markdown("### 📊 Estatísticas da Campanha")
        kpi1, kpi2 = st.columns(2)
        kpi1.metric("Total Anúncios Ativos (Planilha Antiga)", stats.get('ativos_antiga', 0))
        kpi2.metric("Total Anúncios Elegíveis (Planilha Nova)", stats.get('elegiveis_nova', 0))
        st.markdown("---")

    df_display = st.session_state.get('edited_data') if st.session_state.get('edited_data') is not None else st.session_state.get('comparison_data')

    if df_display is not None:
        st.markdown("### 📝 **Passo 3:** Compare e Decida a Ação")
        st.info("Apenas itens presentes na planilha NOVA estão sendo exibidos e processados.")

        # --- Botão de Copiar (Simplificado para não adicionar itens) ---
        if st.button("🔁 Copiar Preços/Descontos da Antiga para a Nova", use_container_width=True):
            df_mod = df_display.copy()
            changes_count = 0
            
            if 'DESC %_ANTIGA' in df_mod.columns and 'DESC %_NOVA' in df_mod.columns:
                df_mod['DESC %_NOVA'] = df_mod['DESC %_ANTIGA']
                changes_count += 1
            if 'PREÇO FINAL_ANTIGA' in df_mod.columns and 'PREÇO FINAL_NOVA' in df_mod.columns:
                df_mod['PREÇO FINAL_NOVA'] = df_mod['PREÇO FINAL_ANTIGA']
                changes_count += 1
            
            if changes_count > 0:
                st.session_state['edited_data'] = df_mod
                st.success("Valores copiados para os itens correspondentes!")
                st.rerun()
            else:
                st.warning("Nada para copiar.")

        df_to_edit = st.session_state.get('edited_data')
        key_col = next((c for c in df_to_edit.columns if c in ['ITEM_ID', 'SELLER_SKU', 'SKU', 'ID', 'CÓDIGO', 'MLB']), df_to_edit.columns[0])

        start_cols = [key_col]
        if 'SKU' in df_to_edit.columns and 'SKU' not in start_cols:
            start_cols.append('SKU')
        end_cols = ['AÇÃO MANUAL']
        grouped_cols = ['STATUS_ANTIGA', 'STATUS_NOVA', 'DESC %_ANTIGA', 'DESC %_NOVA', 'PREÇO FINAL_ANTIGA', 'PREÇO FINAL_NOVA']
        middle_cols = [col for col in grouped_cols if col in df_to_edit.columns]
        other_cols = [c for c in df_to_edit.columns if c not in start_cols + middle_cols + end_cols and c != 'TEMP_KEY']
        display_cols = start_cols + other_cols + middle_cols + end_cols
        existing_display_cols = [col for col in display_cols if col in df_to_edit.columns]

        st.markdown("""<style>.stDataFrame { height: 60vh; }</style>""", unsafe_allow_html=True)

        with st.form("form_edicao_tabela"):
            edited_df = st.data_editor(
                df_to_edit,
                column_config={
                    "AÇÃO MANUAL": st.column_config.SelectboxColumn("Ação Manual", options=["Participar", "Não participar"], required=True),
                    key_col: st.column_config.Column(disabled=True),
                    "STATUS_ANTIGA": st.column_config.Column("Status (Antiga)", disabled=True),
                    "DESC %_ANTIGA": st.column_config.NumberColumn("Desc % (Antiga)", format="%.2f%%", disabled=True),
                    "DESC %_NOVA": st.column_config.NumberColumn("Desc % (Nova)", format="%.2f%%", disabled=False),
                    "PREÇO FINAL_ANTIGA": st.column_config.NumberColumn("Preço (Antiga)", format="R$ %.2f", disabled=True),
                    "PREÇO FINAL_NOVA": st.column_config.NumberColumn("Preço (Nova)", format="R$ %.2f", disabled=False),
                },
                use_container_width=True,
                column_order=existing_display_cols,
                hide_index=True,
                key='editor_renovacao'
            )
            
            confirmar_edicoes = st.form_submit_button("💾 Salvar e Baixar Tabela Atualizada", type="primary", use_container_width=True)

        if confirmar_edicoes:
            st.session_state['edited_data'] = edited_df
            
            if edited_df is not None and not edited_df.empty:
                with st.spinner("Gerando arquivo e iniciando download..."):
                    buffer, msg = generate_updated_sheet(
                        st.session_state['original_nova_file'], 
                        edited_df,
                        st.session_state.get('df_antiga_raw')
                    )
                    
                    if buffer:
                        st.session_state['renovacao_buffer'] = buffer
                        original_name = st.session_state['original_nova_file'].name
                        st.session_state['renovacao_nome'] = original_name 
                        
                        b64 = base64.b64encode(buffer.getvalue()).decode()
                        mime_type = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
                        
                        download_html = f"""
                            <a href="data:{mime_type};base64,{b64}" download="{original_name}" id="hidden_download_link" style="display:none;">Download</a>
                            <script>
                                document.getElementById("hidden_download_link").click();
                            </script>
                        """
                        st.markdown(download_html, unsafe_allow_html=True)
                        st.success("Tabela salva! O download foi iniciado automaticamente.")
                        
                        st.download_button(
                            label="📥 Clique aqui se o download não iniciar",
                            data=buffer,
                            file_name=original_name,
                            mime=mime_type
                        )
                    else:
                        st.error(msg)
        
        elif st.session_state.get('edited_data') is not None:
             edited_df = st.session_state['edited_data']
             
        if st.session_state.get('comparison_data') is not None:
            st.markdown("---")
            if st.button("🔄 Reiniciar Análise", use_container_width=True):
                keys = ['comparison_data', 'df_antiga_raw', 'edited_data', 'renovacao_buffer',
                        'original_nova_file', 'stats_comparacao', 'renovacao_nome']
                for key in keys:
                    if key in st.session_state:
                        del st.session_state[key]
                st.rerun()