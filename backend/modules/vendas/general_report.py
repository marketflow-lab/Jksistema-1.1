"""General SKU sales, local-stock and returns report for the Vendas module."""

from __future__ import annotations

import io
import os
import re
import sqlite3
from collections import Counter, defaultdict
from copy import copy
from datetime import datetime
from decimal import Decimal
from typing import Any
from xml.sax.saxutils import escape as xml_escape

from backend.services.estoque_historico import _estoque_historico_ambiguidade

from .errors import VendasDomainError
from .legacy import _listar_bancos_vendas_tenant
from .performance import (
    append_indexable_date_filter,
    cache_vendas_response,
    cached_table_columns,
    get_tenant_path,
    open_vendas_readonly,
    select_compatible_column,
    vendas_source_paths,
)
from .reports import (
    FISCAL_ADJUSTMENT_SKU,
    _decimal,
    _money,
    _month_keys,
    _normalize_sku,
    _normalize_text,
    _number,
    _operation_name,
    _read_sales_rows,
    _recent_period_comparison,
    _representative_description,
    _safe_filename,
    _source_hash,
    _trend_highlight,
    _validate_store,
    prepare_eligible_sales_rows,
    resolve_pareto_period,
)


AVERAGE_MONTH_DAYS = Decimal("30.4375")


def _general_report_cache_paths(arguments: dict[str, Any]) -> list[str]:
    client_id = arguments["client_id"]
    paths = _listar_bancos_vendas_tenant(client_id, arguments.get("loja"))
    return vendas_source_paths(client_id, paths, include_stock=True)


def _return_is_full(item: dict[str, Any]) -> bool:
    fields = (
        item.get("unidade_negocio_virtual"),
        item.get("unidade_negocio"),
        item.get("natureza_operacao"),
        item.get("finalidade_operacao"),
    )
    return any("FULL" in _normalize_text(value) for value in fields)


def _read_return_rows(
    client_id: str,
    store: str,
    start_iso: str,
    end_iso: str,
) -> tuple[list[dict[str, Any]], list[str], list[str], int]:
    rows: list[dict[str, Any]] = []
    processed: list[str] = []
    errors: list[str] = []
    blank_skus = 0
    select_specs = (
        ("id_unico", "''"),
        ("data_emissao", "''"),
        ("loja_conta", "''"),
        ("numero_nota", "''"),
        ("origem_codigo", "''"),
        ("sku", "''"),
        ("descricao", "''"),
        ("quantidade", "0"),
        ("valor_total", "0"),
        ("unidade_negocio_virtual", "''"),
        ("unidade_negocio", "''"),
        ("natureza_operacao", "''"),
        ("finalidade_operacao", "''"),
    )
    db_paths = _listar_bancos_vendas_tenant(client_id, store)
    for source_index, db_path in enumerate(db_paths):
        if not db_path or not os.path.isfile(db_path):
            continue
        connection = None
        try:
            connection = open_vendas_readonly(db_path, row_factory=sqlite3.Row)
            exists = connection.execute(
                "SELECT 1 FROM sqlite_master WHERE type='table' "
                "AND name='notas_entrada_itens' LIMIT 1"
            ).fetchone()
            if not exists:
                continue
            columns = cached_table_columns(db_path, "notas_entrada_itens", connection)
            if "devolucao" not in columns:
                continue
            query = "SELECT rowid AS __rowid, " + ", ".join(
                select_compatible_column(columns, name, default)
                for name, default in select_specs
            ) + " FROM notas_entrada_itens"
            conditions = ["coalesce(devolucao, 0) = 1"]
            params: list[Any] = []
            append_indexable_date_filter(
                conditions, params, "data_emissao", start_iso, end_iso
            )
            conditions.append("trim(coalesce(loja_conta, '')) = ? COLLATE NOCASE")
            params.append(store)
            query += " WHERE " + " AND ".join(conditions) + " ORDER BY rowid"
            for row in connection.execute(query, params).fetchall():
                item = dict(row)
                if not _normalize_sku(item.get("sku")):
                    blank_skus += 1
                    continue
                item["__source_index"] = source_index
                item["data"] = str(item.get("data_emissao") or "")
                item["produto"] = str(item.get("descricao") or "")
                rows.append(item)
            processed.append(os.path.basename(db_path))
        except Exception as exc:
            errors.append(f"{os.path.basename(db_path)}: {exc}")
        finally:
            if connection is not None:
                connection.close()

    deduplicated: dict[tuple[Any, ...], dict[str, Any]] = {}
    for item in rows:
        unique_id = str(item.get("id_unico") or "").strip()
        if unique_id:
            key: tuple[Any, ...] = ("id_unico", unique_id)
        else:
            key = (
                "fallback",
                str(item.get("data_emissao") or "")[:10],
                _normalize_text(item.get("loja_conta")),
                str(item.get("numero_nota") or "").strip(),
                str(item.get("origem_codigo") or "").strip(),
                _normalize_sku(item.get("sku")),
                str(_decimal(item.get("quantidade"))),
                str(_decimal(item.get("valor_total"))),
            )
        current = deduplicated.get(key)
        if current is None or len(str(item.get("descricao") or "")) > len(
            str(current.get("descricao") or "")
        ):
            deduplicated[key] = item
    return list(deduplicated.values()), processed, errors, len(rows) - len(deduplicated) + blank_skus


def _read_current_local_stock(
    client_id: str,
    store: str,
) -> dict[str, Any]:
    tenant_path = get_tenant_path(client_id)
    db_path = os.path.join(tenant_path, "estoque_historico.db")
    base = {
        "available": False,
        "source": None,
        "snapshot_at": None,
        "event_id": None,
        "total_skus_snapshot": 0,
        "items": {},
        "errors": [],
        "db_path": db_path,
    }
    ambiguidade = _estoque_historico_ambiguidade(
        client_id,
        store,
        tenant_path=tenant_path,
    )
    if ambiguidade:
        base["error_code"] = ambiguidade["code"]
        base["errors"].append(ambiguidade["message"])
        return base
    if not os.path.isfile(db_path):
        base["errors"].append("Histórico de estoque não encontrado.")
        return base

    connection = None
    try:
        connection = open_vendas_readonly(db_path, row_factory=sqlite3.Row)
        tables = {
            str(row[0])
            for row in connection.execute(
                "SELECT name FROM sqlite_master WHERE type='table'"
            ).fetchall()
        }
        if {"estoque_sync_eventos", "estoque_sync_evento_itens"}.issubset(tables):
            event = connection.execute(
                """
                SELECT event_id, recorded_at, data_ref, total_skus
                FROM estoque_sync_eventos
                WHERE status = 'committed'
                  AND trim(coalesce(loja_sync, '')) = ? COLLATE NOCASE
                ORDER BY recorded_at DESC, event_id DESC
                LIMIT 1
                """,
                (store,),
            ).fetchone()
            if event:
                items: dict[str, dict[str, Any]] = {}
                for row in connection.execute(
                    """
                    SELECT sku, nome_bling, saldo_loja
                    FROM estoque_sync_evento_itens
                    WHERE event_id = ?
                    ORDER BY id
                    """,
                    (event["event_id"],),
                ).fetchall():
                    sku = _normalize_sku(row["sku"])
                    if not sku or sku == FISCAL_ADJUSTMENT_SKU:
                        continue
                    items[sku] = {
                        "sku": str(row["sku"] or "").strip(),
                        "produto": str(row["nome_bling"] or "").strip(),
                        "saldo_loja": _number(_decimal(row["saldo_loja"])),
                    }
                return {
                    **base,
                    "available": True,
                    "source": "evento_confirmado",
                    "snapshot_at": str(event["recorded_at"] or event["data_ref"] or ""),
                    "event_id": str(event["event_id"]),
                    "total_skus_snapshot": int(event["total_skus"] or 0),
                    "items": items,
                }

        if "estoque_historico" in tables:
            snapshot = connection.execute(
                """
                SELECT data_ref, recorded_at
                FROM estoque_historico
                WHERE trim(coalesce(loja_sync, '')) = ? COLLATE NOCASE
                ORDER BY recorded_at DESC, data_ref DESC
                LIMIT 1
                """,
                (store,),
            ).fetchone()
            if snapshot:
                items = {}
                for row in connection.execute(
                    """
                    SELECT sku, nome_bling, saldo_loja
                    FROM estoque_historico
                    WHERE trim(coalesce(loja_sync, '')) = ? COLLATE NOCASE
                      AND data_ref = ?
                    ORDER BY id
                    """,
                    (store, snapshot["data_ref"]),
                ).fetchall():
                    sku = _normalize_sku(row["sku"])
                    if not sku or sku == FISCAL_ADJUSTMENT_SKU:
                        continue
                    items[sku] = {
                        "sku": str(row["sku"] or "").strip(),
                        "produto": str(row["nome_bling"] or "").strip(),
                        "saldo_loja": _number(_decimal(row["saldo_loja"])),
                    }
                return {
                    **base,
                    "available": True,
                    "source": "snapshot_legado",
                    "snapshot_at": str(snapshot["recorded_at"] or snapshot["data_ref"] or ""),
                    "total_skus_snapshot": len(items),
                    "items": items,
                }
        base["errors"].append("Nenhum snapshot confirmado para a conta selecionada.")
        return base
    except Exception as exc:
        base["errors"].append(f"Histórico de estoque indisponível: {exc}")
        return base
    finally:
        if connection is not None:
            connection.close()


def _coverage_status(stock: Decimal | None, average: Decimal) -> tuple[float | None, float | None, str]:
    if stock is None:
        return None, None, "Estoque sem registro"
    if average <= 0:
        if stock > 0:
            return None, None, "Sem vendas no período"
        return None, None, "Sem estoque e sem vendas"
    if stock <= 0:
        return 0.0, 0.0, "Ruptura"
    months = stock / average
    days = months * AVERAGE_MONTH_DAYS
    if months <= 1:
        status = "Até 1 mês"
    elif months <= 3:
        status = "1 a 3 meses"
    else:
        status = "Acima de 3 meses"
    return _number(months), _number(days), status


def _choose_display_sku(values: Counter[str], normalized: str) -> str:
    if not values:
        return normalized
    return max(values, key=lambda value: (values[value], value.casefold()))


@cache_vendas_response("vendas-estoque-devolucoes", _general_report_cache_paths)
def generate_general_sku_report(
    client_id: str,
    loja: str,
    periodo: str = "12m",
    data_inicio: str | None = None,
    data_fim: str | None = None,
) -> dict[str, Any]:
    store = _validate_store(loja)
    period_mode, start_iso, end_iso = resolve_pareto_period(
        periodo, data_inicio, data_fim
    )
    months = _month_keys(start_iso, end_iso)
    month_count = Decimal(len(months) or 1)

    raw_sales, sales_bases, sales_errors = _read_sales_rows(
        client_id, store, start_iso, end_iso
    )
    product_sales, fiscal_rows, sales_audit = prepare_eligible_sales_rows(raw_sales)
    return_rows, return_bases, return_errors, return_exclusions = _read_return_rows(
        client_id, store, start_iso, end_iso
    )
    stock_snapshot = _read_current_local_stock(client_id, store)

    display_skus: dict[str, Counter[str]] = defaultdict(Counter)
    sales_by_sku: dict[str, list[dict[str, Any]]] = defaultdict(list)
    returns_by_sku: dict[str, list[dict[str, Any]]] = defaultdict(list)
    monthly_sales: defaultdict[str, defaultdict[str, dict[str, Decimal]]] = defaultdict(
        lambda: defaultdict(
            lambda: {
                "unidades": Decimal("0"),
                "faturamento": Decimal("0"),
                "full_unidades": Decimal("0"),
                "full_faturamento": Decimal("0"),
                "demais_unidades": Decimal("0"),
                "demais_faturamento": Decimal("0"),
            }
        )
    )
    monthly_returns: defaultdict[str, defaultdict[str, dict[str, Decimal]]] = defaultdict(
        lambda: defaultdict(
            lambda: {
                "quantidade": Decimal("0"),
                "valor": Decimal("0"),
                "full_quantidade": Decimal("0"),
                "demais_quantidade": Decimal("0"),
            }
        )
    )

    for item in product_sales:
        sku = _normalize_sku(item.get("sku"))
        display_skus[sku][str(item.get("sku") or "").strip()] += 1
        sales_by_sku[sku].append(item)
        month = str(item.get("data") or "")[:7]
        if month not in months:
            continue
        quantity = _decimal(item.get("quantidade"))
        revenue = _decimal(item.get("valor"))
        bucket = monthly_sales[sku][month]
        bucket["unidades"] += quantity
        bucket["faturamento"] += revenue
        if _operation_name(item, store) == "Mercado Livre Full":
            bucket["full_unidades"] += quantity
            bucket["full_faturamento"] += revenue
        else:
            bucket["demais_unidades"] += quantity
            bucket["demais_faturamento"] += revenue

    for item in return_rows:
        sku = _normalize_sku(item.get("sku"))
        display_skus[sku][str(item.get("sku") or "").strip()] += 1
        returns_by_sku[sku].append(item)
        month = str(item.get("data_emissao") or "")[:7]
        if month not in months:
            continue
        quantity = _decimal(item.get("quantidade"))
        value = _decimal(item.get("valor_total"))
        bucket = monthly_returns[sku][month]
        bucket["quantidade"] += quantity
        bucket["valor"] += value
        if _return_is_full(item):
            bucket["full_quantidade"] += quantity
        else:
            bucket["demais_quantidade"] += quantity

    for sku, item in stock_snapshot["items"].items():
        display_skus[sku][str(item.get("sku") or sku)] += 1

    relevant_skus = set(sales_by_sku) | set(returns_by_sku) | set(stock_snapshot["items"])
    relevant_skus.discard(FISCAL_ADJUSTMENT_SKU)
    rows: list[dict[str, Any]] = []
    return_summaries: list[dict[str, Any]] = []
    total_sales_units = Decimal("0")
    total_revenue = Decimal("0")
    total_return_units = Decimal("0")
    total_return_value = Decimal("0")
    known_stock_total = Decimal("0")
    coverage_bands = Counter()
    recent_months = set(months[-3:])

    for normalized_sku in relevant_skus:
        sales_items = sales_by_sku.get(normalized_sku, [])
        returns_items = returns_by_sku.get(normalized_sku, [])
        stock_item = stock_snapshot["items"].get(normalized_sku)
        sales_units = sum(
            (_decimal(item.get("quantidade")) for item in sales_items), Decimal("0")
        )
        revenue = sum(
            (_decimal(item.get("valor")) for item in sales_items), Decimal("0")
        )
        full_units = sum(
            (
                _decimal(item.get("quantidade"))
                for item in sales_items
                if _operation_name(item, store) == "Mercado Livre Full"
            ),
            Decimal("0"),
        )
        full_revenue = sum(
            (
                _decimal(item.get("valor"))
                for item in sales_items
                if _operation_name(item, store) == "Mercado Livre Full"
            ),
            Decimal("0"),
        )
        return_units = sum(
            (_decimal(item.get("quantidade")) for item in returns_items), Decimal("0")
        )
        return_value = sum(
            (_decimal(item.get("valor_total")) for item in returns_items), Decimal("0")
        )
        full_return_units = sum(
            (
                _decimal(item.get("quantidade"))
                for item in returns_items
                if _return_is_full(item)
            ),
            Decimal("0"),
        )
        recent_return_units = sum(
            (
                _decimal(item.get("quantidade"))
                for item in returns_items
                if str(item.get("data_emissao") or "")[:7] in recent_months
            ),
            Decimal("0"),
        )
        average = sales_units / month_count
        stock = _decimal(stock_item.get("saldo_loja")) if stock_item is not None else None
        coverage_months, coverage_days, status = _coverage_status(stock, average)

        description_items = list(sales_items) + list(returns_items)
        product = _representative_description(description_items)
        if product == "-" and stock_item:
            product = str(stock_item.get("produto") or "-")
        display_sku = _choose_display_sku(display_skus[normalized_sku], normalized_sku)
        monthly_units = {
            month: _number(monthly_sales[normalized_sku][month]["unidades"])
            for month in months
        }
        row = {
            "sku": display_sku,
            "produto": product,
            "vendas_mensais": monthly_units,
            "unidades_vendidas": _number(sales_units),
            "faturamento": _number(revenue),
            "full_unidades": _number(full_units),
            "demais_unidades": _number(sales_units - full_units),
            "full_faturamento": _number(full_revenue),
            "demais_faturamento": _number(revenue - full_revenue),
            "media_mensal": _number(average),
            "estoque_loja": _number(stock) if stock is not None else None,
            "estoque_disponivel": stock is not None,
            "cobertura_meses": coverage_months,
            "cobertura_dias": coverage_days,
            "status_cobertura": status,
            "ocorrencias_devolucao": len(returns_items),
            "quantidade_devolvida": _number(return_units),
            "valor_devolvido": _number(return_value),
            "taxa_devolucao": _number(return_units / sales_units) if sales_units > 0 else None,
            "teve_movimentacao": bool(sales_items or returns_items),
        }
        rows.append(row)
        if returns_items:
            return_summaries.append(
                {
                    "sku": display_sku,
                    "produto": product,
                    "ocorrencias": len(returns_items),
                    "quantidade": _number(return_units),
                    "valor": _number(return_value),
                    "full_quantidade": _number(full_return_units),
                    "demais_quantidade": _number(return_units - full_return_units),
                    "unidades_vendidas": _number(sales_units),
                    "taxa_devolucao": _number(return_units / sales_units) if sales_units > 0 else None,
                    "quantidade_ultimos_3_meses": _number(recent_return_units),
                    "participacao_recente": _number(recent_return_units / return_units)
                    if return_units > 0
                    else 0.0,
                    "ultima_devolucao": max(
                        (str(item.get("data_emissao") or "")[:10] for item in returns_items),
                        default=None,
                    ),
                }
            )
        total_sales_units += sales_units
        total_revenue += revenue
        total_return_units += return_units
        total_return_value += return_value

    all_rows = rows
    rows = [item for item in all_rows if item["teve_movimentacao"]]
    no_movement_rows = [item for item in all_rows if not item["teve_movimentacao"]]
    for item in all_rows:
        item.pop("teve_movimentacao", None)

    coverage_bands = Counter(item["status_cobertura"] for item in rows)
    known_stock_total = sum(
        (
            _decimal(item["estoque_loja"])
            for item in rows
            if item["estoque_loja"] is not None
        ),
        Decimal("0"),
    )
    no_movement_group_specs = [
        ("Com estoque positivo", lambda item: item["estoque_loja"] is not None and _decimal(item["estoque_loja"]) > 0),
        ("Estoque zerado ou negativo", lambda item: item["estoque_loja"] is not None and _decimal(item["estoque_loja"]) <= 0),
        ("Sem registro de estoque", lambda item: item["estoque_loja"] is None),
    ]
    no_movement_groups = []
    for label, predicate in no_movement_group_specs:
        grouped_items = sorted(
            (item for item in no_movement_rows if predicate(item)),
            key=lambda item: _normalize_sku(item["sku"]),
        )
        if not grouped_items:
            continue
        known_group_stock = sum(
            (
                _decimal(item["estoque_loja"])
                for item in grouped_items
                if item["estoque_loja"] is not None
            ),
            Decimal("0"),
        )
        no_movement_groups.append(
            {
                "grupo": label,
                "quantidade_skus": len(grouped_items),
                "estoque_loja_total_conhecido": _number(known_group_stock),
                "estoque_disponivel": any(
                    item["estoque_loja"] is not None for item in grouped_items
                ),
                "itens": [
                    {
                        "sku": item["sku"],
                        "produto": item["produto"],
                        "estoque_loja": item["estoque_loja"],
                    }
                    for item in grouped_items
                ],
            }
        )

    rows.sort(
        key=lambda item: (
            -_decimal(item["faturamento"]),
            -_decimal(item["unidades_vendidas"]),
            _normalize_sku(item["sku"]),
        )
    )
    for rank, item in enumerate(rows, start=1):
        item["rank"] = rank

    top_sellers = []
    top_seller_rows = sorted(
        (item for item in rows if _decimal(item["unidades_vendidas"]) > 0),
        key=lambda item: (
            -_decimal(item["unidades_vendidas"]),
            -_decimal(item["faturamento"]),
            _normalize_sku(item["sku"]),
        ),
    )[:10]
    for sales_rank, item in enumerate(top_seller_rows, start=1):
        units = _decimal(item["unidades_vendidas"])
        top_sellers.append(
            {
                "rank": sales_rank,
                "sku": item["sku"],
                "produto": item["produto"],
                "unidades_vendidas": _number(units),
                "faturamento": item["faturamento"],
                "participacao_unidades": _number(units / total_sales_units)
                if total_sales_units
                else 0.0,
            }
        )

    top_revenue_sellers = []
    top_revenue_rows = sorted(
        (item for item in rows if _decimal(item["faturamento"]) > 0),
        key=lambda item: (
            -_decimal(item["faturamento"]),
            -_decimal(item["unidades_vendidas"]),
            _normalize_sku(item["sku"]),
        ),
    )[:10]
    for revenue_rank, item in enumerate(top_revenue_rows, start=1):
        revenue = _decimal(item["faturamento"])
        top_revenue_sellers.append(
            {
                "rank": revenue_rank,
                "sku": item["sku"],
                "produto": item["produto"],
                "unidades_vendidas": item["unidades_vendidas"],
                "faturamento": _number(revenue),
                "participacao_faturamento": _number(revenue / total_revenue)
                if total_revenue
                else 0.0,
            }
        )

    return_summaries.sort(
        key=lambda item: (
            -_decimal(item["quantidade"]),
            -_decimal(item["valor"]),
            _normalize_sku(item["sku"]),
        )
    )
    overall_return_rate = (
        total_return_units / total_sales_units if total_sales_units > 0 else None
    )
    top_returns = []
    for rank, item in enumerate(return_summaries[:20], start=1):
        quantity = _decimal(item["quantidade"])
        share = quantity / total_return_units if total_return_units > 0 else Decimal("0")
        full_quantity = _decimal(item["full_quantidade"])
        other_quantity = _decimal(item["demais_quantidade"])
        if full_quantity > other_quantity:
            predominance = "Mercado Livre Full"
        elif other_quantity > full_quantity:
            predominance = "Demais canais"
        else:
            predominance = "Equilibrada"
        rate_text = (
            f"taxa de {_number(_decimal(item['taxa_devolucao']) * 100):.1f}% sobre as vendas"
            if item["taxa_devolucao"] is not None
            else "sem venda elegível no período; pode corresponder a venda anterior"
        )
        analysis = (
            f"Representa {_number(share * 100):.1f}% das unidades devolvidas; "
            f"{rate_text}; {_number(_decimal(item['participacao_recente']) * 100):.1f}% "
            f"ocorreu nos últimos 3 meses; predominância: {predominance}."
        )
        top_returns.append(
            {
                **item,
                "rank": rank,
                "participacao_devolucoes": _number(share),
                "predominancia": predominance,
                "analise": analysis,
            }
        )

    detailed_monthly = []
    account_monthly = []
    for month in months:
        account_sales_units = Decimal("0")
        account_revenue = Decimal("0")
        account_full_units = Decimal("0")
        account_full_revenue = Decimal("0")
        account_other_units = Decimal("0")
        account_other_revenue = Decimal("0")
        account_return_units = Decimal("0")
        account_return_value = Decimal("0")
        for item in rows:
            normalized_sku = _normalize_sku(item["sku"])
            sale_bucket = monthly_sales[normalized_sku][month]
            return_bucket = monthly_returns[normalized_sku][month]
            detailed_monthly.append(
                {
                    "mes": month,
                    "sku": item["sku"],
                    "produto": item["produto"],
                    **{key: _number(value) for key, value in sale_bucket.items()},
                    "devolucoes_quantidade": _number(return_bucket["quantidade"]),
                    "devolucoes_valor": _number(return_bucket["valor"]),
                }
            )
            account_sales_units += sale_bucket["unidades"]
            account_revenue += sale_bucket["faturamento"]
            account_full_units += sale_bucket["full_unidades"]
            account_full_revenue += sale_bucket["full_faturamento"]
            account_other_units += sale_bucket["demais_unidades"]
            account_other_revenue += sale_bucket["demais_faturamento"]
            account_return_units += return_bucket["quantidade"]
            account_return_value += return_bucket["valor"]
        account_monthly.append(
            {
                "mes": month,
                "unidades_vendidas": _number(account_sales_units),
                "faturamento": _number(account_revenue),
                "full_unidades": _number(account_full_units),
                "full_faturamento": _number(account_full_revenue),
                "demais_unidades": _number(account_other_units),
                "demais_faturamento": _number(account_other_revenue),
                "quantidade_devolvida": _number(account_return_units),
                "valor_devolvido": _number(account_return_value),
                "taxa_devolucao": _number(account_return_units / account_sales_units)
                if account_sales_units > 0
                else None,
            }
        )

    revenue_comparison = _recent_period_comparison(account_monthly, "faturamento")
    comparison_months = int(revenue_comparison.get("meses") or 0)
    previous_months = (
        account_monthly[-2 * comparison_months : -comparison_months]
        if comparison_months
        else []
    )
    recent_month_rows = account_monthly[-comparison_months:] if comparison_months else []
    previous_sales_units = sum(
        (_decimal(item["unidades_vendidas"]) for item in previous_months), Decimal("0")
    )
    recent_sales_units = sum(
        (_decimal(item["unidades_vendidas"]) for item in recent_month_rows), Decimal("0")
    )
    previous_return_units = sum(
        (_decimal(item["quantidade_devolvida"]) for item in previous_months), Decimal("0")
    )
    recent_return_units = sum(
        (_decimal(item["quantidade_devolvida"]) for item in recent_month_rows), Decimal("0")
    )
    previous_return_rate = (
        previous_return_units / previous_sales_units if previous_sales_units else None
    )
    recent_return_rate = recent_return_units / recent_sales_units if recent_sales_units else None
    return_rate_delta = (
        recent_return_rate - previous_return_rate
        if recent_return_rate is not None and previous_return_rate is not None
        else None
    )
    total_full_revenue = sum(
        (_decimal(item["full_faturamento"]) for item in rows), Decimal("0")
    )
    full_revenue_share = total_full_revenue / total_revenue if total_revenue else Decimal("0")
    critical_rows = [
        item for item in rows if item["status_cobertura"] in {"Ruptura", "Até 1 mês"}
    ]
    critical_revenue = sum(
        (_decimal(item["faturamento"]) for item in critical_rows), Decimal("0")
    )
    critical_revenue_share = critical_revenue / total_revenue if total_revenue else Decimal("0")
    excess_rows = [item for item in rows if item["status_cobertura"] == "Acima de 3 meses"]
    top_20_return_units = sum(
        (_decimal(item["quantidade"]) for item in top_returns), Decimal("0")
    )
    top_5_return_units = sum(
        (_decimal(item["quantidade"]) for item in top_returns[:5]), Decimal("0")
    )
    top_return_item = top_returns[0] if top_returns else None
    top_seller_item = top_sellers[0] if top_sellers else None
    top_seller_units = sum(
        (_decimal(item["unidades_vendidas"]) for item in top_sellers), Decimal("0")
    )
    if return_rate_delta is None:
        return_rate_highlight = "A taxa recente de devolução não possui janela anterior comparável."
    else:
        direction = "subiu" if return_rate_delta > 0 else "caiu" if return_rate_delta < 0 else "ficou estável"
        delta_text = f" {abs(_number(return_rate_delta * 100)):.1f} p.p." if return_rate_delta else ""
        return_rate_highlight = (
            f"A taxa de devolução {direction}{delta_text} nos últimos {comparison_months} meses: "
            f"{_number((recent_return_rate or Decimal('0')) * 100):.1f}% versus "
            f"{_number((previous_return_rate or Decimal('0')) * 100):.1f}%."
        )
    highlights = [
        _trend_highlight("O faturamento", revenue_comparison),
        return_rate_highlight,
        (
            f"{len(critical_rows)} SKUs estão em ruptura ou têm até 1 mês de cobertura e "
            f"representam {_number(critical_revenue_share * 100):.1f}% do faturamento."
        ),
        (
            f"Mercado Livre Full responde por {_number(full_revenue_share * 100):.1f}% do faturamento; "
            f"os demais canais respondem por {_number((Decimal('1') - full_revenue_share) * 100):.1f}%."
        ),
        (
            f"Os 20 SKUs mais devolvidos concentram "
            f"{_number((top_20_return_units / total_return_units if total_return_units else Decimal('0')) * 100):.1f}% "
            f"das unidades devolvidas"
            + (f"; o SKU {top_return_item['sku']} lidera com {_number(_decimal(top_return_item['quantidade'])):.0f} unidade(s)." if top_return_item else ".")
        ),
    ]

    source_paths = vendas_source_paths(
        client_id,
        _listar_bancos_vendas_tenant(client_id, store),
        include_stock=True,
    )
    return {
        "success": True,
        "loja": store,
        "periodo": {
            "tipo": period_mode,
            "data_inicio": start_iso,
            "data_fim": end_iso,
            "meses": months,
            "quantidade_meses": len(months),
            "ultima_data_vendas": max(
                (str(item.get("data") or "")[:10] for item in raw_sales), default=None
            ),
            "ultima_data_devolucoes": max(
                (str(item.get("data_emissao") or "")[:10] for item in return_rows),
                default=None,
            ),
            "estoque_atualizado_em": stock_snapshot["snapshot_at"],
            "gerado_em": datetime.now().astimezone().isoformat(timespec="seconds"),
        },
        "resumo": {
            "total_skus": len(rows),
            "total_skus_identificados": len(all_rows),
            "skus_sem_movimentacao": len(no_movement_rows),
            "skus_com_venda": len(sales_by_sku),
            "skus_com_devolucao": len(return_summaries),
            "skus_com_estoque_registrado": sum(
                1 for item in rows if item["estoque_disponivel"]
            ),
            "skus_sem_registro_estoque": sum(
                1 for item in rows if not item["estoque_disponivel"]
            ),
            "unidades_vendidas": _number(total_sales_units),
            "faturamento": _number(total_revenue),
            "quantidade_devolvida": _number(total_return_units),
            "valor_devolvido": _number(total_return_value),
            "taxa_devolucao": _number(overall_return_rate)
            if overall_return_rate is not None
            else None,
            "estoque_loja_total_conhecido": _number(known_stock_total),
            "status_cobertura": dict(coverage_bands),
        },
        "skus": rows,
        "mensal_detalhado": detailed_monthly,
        "mensal_conta": account_monthly,
        "devolucoes": return_summaries,
        "top_10_vendidos": top_sellers,
        "top_10_faturamento": top_revenue_sellers,
        "top_20_devolucoes": top_returns,
        "sem_movimentacao": {
            "total_skus": len(no_movement_rows),
            "grupos": no_movement_groups,
        },
        "analise_executiva": {
            "destaques": highlights,
            "comparacao_faturamento": revenue_comparison,
            "taxa_devolucao_anterior": _number(previous_return_rate)
            if previous_return_rate is not None
            else None,
            "taxa_devolucao_recente": _number(recent_return_rate)
            if recent_return_rate is not None
            else None,
            "variacao_taxa_devolucao": _number(return_rate_delta)
            if return_rate_delta is not None
            else None,
            "participacao_full_faturamento": _number(full_revenue_share),
            "skus_cobertura_critica": len(critical_rows),
            "faturamento_cobertura_critica": _number(critical_revenue),
            "participacao_faturamento_cobertura_critica": _number(critical_revenue_share),
            "skus_acima_tres_meses": len(excess_rows),
            "concentracao_devolucoes_top_5": _number(top_5_return_units / total_return_units)
            if total_return_units
            else 0.0,
            "concentracao_devolucoes_top_20": _number(top_20_return_units / total_return_units)
            if total_return_units
            else 0.0,
            "sku_mais_vendido": top_seller_item["sku"] if top_seller_item else None,
            "unidades_sku_mais_vendido": top_seller_item["unidades_vendidas"]
            if top_seller_item
            else 0.0,
            "concentracao_unidades_top_10": _number(top_seller_units / total_sales_units)
            if total_sales_units
            else 0.0,
        },
        "ajuste_fiscal": {
            "sku": "ESTORNO DE CRÉDITO ICMS",
            "linhas": len(fiscal_rows),
            "valor": _number(
                sum((_decimal(item.get("valor")) for item in fiscal_rows), Decimal("0"))
            ),
            "incluido_nas_vendas": False,
        },
        "auditoria": {
            "vendas": sales_audit,
            "devolucoes_lidas": len(return_rows),
            "devolucoes_excluidas_ou_deduplicadas": return_exclusions,
            "bases_vendas": sales_bases,
            "bases_devolucoes": return_bases,
            "erros_vendas": sales_errors,
            "erros_devolucoes": return_errors,
            "estoque": {
                "disponivel": stock_snapshot["available"],
                "fonte": stock_snapshot["source"],
                "event_id": stock_snapshot["event_id"],
                "total_skus_snapshot": stock_snapshot["total_skus_snapshot"],
                "erros": stock_snapshot["errors"],
            },
            "assinatura_fontes": _source_hash(source_paths),
        },
    }


def build_general_sku_xlsx(report: dict[str, Any]) -> bytes:
    from openpyxl import Workbook
    from openpyxl.chart import BarChart, LineChart, Reference
    from openpyxl.chart.series import SeriesLabel
    from openpyxl.formatting.rule import ColorScaleRule, DataBarRule, FormulaRule
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter
    from openpyxl.worksheet.table import Table, TableStyleInfo

    workbook = Workbook()
    summary_sheet = workbook.active
    summary_sheet.title = "Resumo"
    sku_sheet = workbook.create_sheet("Vendas e Estoque")
    monthly_sheet = workbook.create_sheet("Mensal Detalhado")
    returns_sheet = workbook.create_sheet("Devoluções por SKU")
    top_returns_sheet = workbook.create_sheet("Top 20 Devoluções")
    rules_sheet = workbook.create_sheet("Fontes e Regras")
    no_movement_sheet = workbook.create_sheet("Sem movimentação")

    navy = "17365D"
    blue = "2F75B5"
    teal = "00A6A6"
    green = "70AD47"
    orange = "ED7D31"
    red = "C00000"
    dark_red = "7F1D1D"
    return_red = "B91C1C"
    return_light = "FDECEC"
    return_light_alt = "FFF7F7"
    return_border = "E7A9A9"
    light_blue = "D9EAF7"
    light_green = "E2F0D9"
    light_gray = "F2F4F7"
    light_orange = "FCE4D6"
    light_red = "F4CCCC"
    white = "FFFFFF"
    muted = "5B6573"
    thin = Side(style="thin", color="D5DCE5")

    def title(sheet, text: str, end_column: int, *, fill_color: str = navy) -> None:
        sheet.merge_cells(start_row=1, start_column=1, end_row=1, end_column=end_column)
        cell = sheet.cell(1, 1, text)
        cell.font = Font(name="Aptos Display", size=18, bold=True, color=white)
        cell.fill = PatternFill("solid", fgColor=fill_color)
        cell.alignment = Alignment(horizontal="left", vertical="center")
        sheet.row_dimensions[1].height = 30
        sheet.sheet_view.showGridLines = False

    def header(row, *, fill_color: str = blue) -> None:
        for cell in row:
            cell.font = Font(name="Aptos", size=10, bold=True, color=white)
            cell.fill = PatternFill("solid", fgColor=fill_color)
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            cell.border = Border(bottom=thin)

    def add_table(
        sheet,
        ref: str,
        name: str,
        *,
        style_name: str = "TableStyleMedium2",
        show_row_stripes: bool = True,
    ) -> None:
        table = Table(displayName=name, ref=ref)
        table.tableStyleInfo = TableStyleInfo(
            name=style_name,
            showFirstColumn=False,
            showLastColumn=False,
            showRowStripes=show_row_stripes,
            showColumnStripes=False,
        )
        sheet.add_table(table)

    def apply_returns_palette(sheet, last_row: int, last_column: int) -> None:
        palette_border = Side(style="thin", color=return_border)
        for cell in sheet[1]:
            if cell.column <= last_column:
                cell.fill = PatternFill("solid", fgColor=dark_red)
        for cell in sheet[2]:
            if cell.column <= last_column:
                cell.fill = PatternFill("solid", fgColor=return_red)
                cell.font = Font(name="Aptos", size=10, bold=True, color=white)
                cell.border = Border(bottom=palette_border)
        for row_index in range(3, last_row + 1):
            fill_color = return_light if row_index % 2 else return_light_alt
            for column in range(1, last_column + 1):
                cell = sheet.cell(row_index, column)
                cell.fill = PatternFill("solid", fgColor=fill_color)
                cell.border = Border(bottom=palette_border)

    summary = report["resumo"]
    period = report["periodo"]
    title(summary_sheet, "Relatório Geral de SKUs - Vendas, Estoque e Devoluções", 8)
    metadata = [
        ("Conta", report["loja"]),
        ("Período solicitado", f"{period['data_inicio']} a {period['data_fim']}"),
        ("Última venda disponível", period.get("ultima_data_vendas") or "Sem vendas"),
        ("Última devolução disponível", period.get("ultima_data_devolucoes") or "Sem devoluções"),
        ("Estoque da Loja atualizado em", period.get("estoque_atualizado_em") or "Indisponível"),
        ("Meses-calendário", period["quantidade_meses"]),
        ("Gerado em", period["gerado_em"]),
    ]
    for row_number, (label, value) in enumerate(metadata, start=3):
        summary_sheet.cell(row_number, 1, label).font = Font(bold=True, color=muted)
        summary_sheet.cell(row_number, 2, value)
    summary_sheet["B8"].number_format = "#,##0"

    kpis = [
        ("SKUs com movimentação", summary["total_skus"], "integer"),
        ("Unidades vendidas", summary["unidades_vendidas"], "number"),
        ("Faturamento", summary["faturamento"], "currency"),
        ("Estoque Loja conhecido", summary["estoque_loja_total_conhecido"], "number"),
        ("SKUs sem registro de estoque", summary["skus_sem_registro_estoque"], "integer"),
        ("Quantidade devolvida", summary["quantidade_devolvida"], "number"),
        ("Valor devolvido", summary["valor_devolvido"], "currency"),
        ("Taxa geral de devolução", summary["taxa_devolucao"], "percent"),
    ]
    for index, (label, value, kind) in enumerate(kpis):
        row = 12 + (index // 2) * 3
        column = 1 + (index % 2) * 3
        summary_sheet.merge_cells(start_row=row, start_column=column, end_row=row, end_column=column + 1)
        summary_sheet.merge_cells(start_row=row + 1, start_column=column, end_row=row + 1, end_column=column + 1)
        label_cell = summary_sheet.cell(row, column, label)
        value_cell = summary_sheet.cell(row + 1, column, value)
        label_cell.font = Font(bold=True, color=muted)
        value_cell.font = Font(size=15, bold=True, color=navy)
        fill = light_blue if index % 2 == 0 else light_green
        for target in (label_cell, value_cell):
            target.fill = PatternFill("solid", fgColor=fill)
            target.alignment = Alignment(vertical="center")
        if kind == "currency":
            value_cell.number_format = '"R$" #,##0.00'
        elif kind == "percent":
            value_cell.number_format = "0.00%"
        elif kind == "integer":
            value_cell.number_format = "#,##0"
        else:
            value_cell.number_format = "#,##0.00"

    analysis = report.get("analise_executiva") or {}
    analysis_start = 25
    summary_sheet.cell(analysis_start, 1, "Leitura executiva").font = Font(
        bold=True, color=navy, size=12
    )
    highlights = list(analysis.get("destaques") or [])
    for row_index, text in enumerate(highlights, start=analysis_start + 1):
        summary_sheet.merge_cells(
            start_row=row_index, start_column=1, end_row=row_index, end_column=8
        )
        cell = summary_sheet.cell(row_index, 1, f"• {text}")
        cell.fill = PatternFill("solid", fgColor=light_gray)
        cell.font = Font(color="303A46")
        cell.alignment = Alignment(vertical="center", wrap_text=True)
        summary_sheet.row_dimensions[row_index].height = 24

    monthly_start = analysis_start + max(1, len(highlights)) + 2
    summary_sheet.cell(monthly_start, 1, "Evolução mensal da conta").font = Font(
        bold=True, color=navy, size=12
    )
    monthly_headers = [
        "Mês",
        "Faturamento Full",
        "Faturamento demais",
        "Faturamento total",
        "Unidades vendidas",
        "Quantidade devolvida",
        "Taxa devolução",
        "Valor devolvido",
    ]
    for column, value in enumerate(monthly_headers, start=1):
        summary_sheet.cell(monthly_start + 1, column, value)
    header(summary_sheet[monthly_start + 1])
    for row_index, item in enumerate(report["mensal_conta"], start=monthly_start + 2):
        summary_sheet.cell(row_index, 1, item["mes"])
        summary_sheet.cell(row_index, 2, item["full_faturamento"])
        summary_sheet.cell(row_index, 3, item["demais_faturamento"])
        summary_sheet.cell(row_index, 4, item["faturamento"])
        summary_sheet.cell(row_index, 5, item["unidades_vendidas"])
        summary_sheet.cell(row_index, 6, item["quantidade_devolvida"])
        summary_sheet.cell(row_index, 7, item["taxa_devolucao"])
        summary_sheet.cell(row_index, 8, item["valor_devolvido"])
        for column in (5, 6):
            summary_sheet.cell(row_index, column).number_format = "#,##0.00"
        for column in (2, 3, 4, 8):
            summary_sheet.cell(row_index, column).number_format = '"R$" #,##0.00'
        summary_sheet.cell(row_index, 7).number_format = "0.00%"
    monthly_last = monthly_start + 1 + len(report["mensal_conta"])
    if report["mensal_conta"]:
        add_table(summary_sheet, f"A{monthly_start + 1}:H{monthly_last}", "ResumoMensalGeralTable")
        revenue_variation = (analysis.get("comparacao_faturamento") or {}).get("variacao")
        revenue_suffix = (
            f" - últimos meses {'+' if revenue_variation >= 0 else ''}{revenue_variation:.1%}"
            if revenue_variation is not None
            else ""
        )
        revenue_chart = BarChart()
        revenue_chart.type = "col"
        revenue_chart.style = 10
        revenue_chart.grouping = "stacked"
        revenue_chart.overlap = 100
        revenue_chart.title = f"Faturamento mensal por canal{revenue_suffix}"
        revenue_chart.y_axis.title = "Faturamento (R$)"
        revenue_chart.y_axis.numFmt = '"R$" #,##0'
        revenue_chart.x_axis.title = "Mês"
        revenue_chart.height = 9
        revenue_chart.width = 18
        revenue_chart.add_data(
            Reference(summary_sheet, min_col=2, max_col=3, min_row=monthly_start + 1, max_row=monthly_last),
            titles_from_data=True,
        )
        revenue_chart.set_categories(
            Reference(summary_sheet, min_col=1, min_row=monthly_start + 2, max_row=monthly_last)
        )
        revenue_chart.series[0].tx = SeriesLabel(v="Mercado Livre Full")
        revenue_chart.series[0].graphicalProperties.solidFill = teal
        revenue_chart.series[1].tx = SeriesLabel(v="Demais canais")
        revenue_chart.series[1].graphicalProperties.solidFill = blue
        summary_sheet.add_chart(revenue_chart, f"J{monthly_start}")

        return_chart = LineChart()
        return_chart.style = 13
        return_chart.title = "Taxa mensal de devolução sobre unidades vendidas"
        return_chart.y_axis.title = "Taxa de devolução"
        return_chart.y_axis.numFmt = "0.0%"
        return_chart.x_axis.title = "Mês"
        return_chart.height = 8
        return_chart.width = 18
        return_chart.add_data(
            Reference(summary_sheet, min_col=7, min_row=monthly_start + 1, max_row=monthly_last),
            titles_from_data=True,
        )
        return_chart.set_categories(
            Reference(summary_sheet, min_col=1, min_row=monthly_start + 2, max_row=monthly_last)
        )
        return_chart.series[0].tx = SeriesLabel(v="Taxa de devolução")
        return_chart.series[0].graphicalProperties.line.solidFill = orange
        return_chart.legend = None
        summary_sheet.add_chart(return_chart, f"J{monthly_start + 18}")

    coverage_start = monthly_last + 3
    summary_sheet.cell(coverage_start, 1, "Distribuição da cobertura").font = Font(
        bold=True, color=navy, size=12
    )
    summary_sheet.cell(coverage_start + 1, 1, "Status")
    summary_sheet.cell(coverage_start + 1, 2, "SKUs")
    header(summary_sheet[coverage_start + 1])
    coverage_order = [
        "Ruptura",
        "Até 1 mês",
        "1 a 3 meses",
        "Acima de 3 meses",
        "Sem vendas no período",
        "Sem estoque e sem vendas",
        "Estoque sem registro",
    ]
    coverage_items = [
        (status, summary["status_cobertura"].get(status, 0)) for status in coverage_order
    ]
    for row_index, (status, count) in enumerate(coverage_items, start=coverage_start + 2):
        summary_sheet.cell(row_index, 1, status)
        summary_sheet.cell(row_index, 2, count)
        summary_sheet.cell(row_index, 2).number_format = "#,##0"
    coverage_last = coverage_start + 1 + len(coverage_items)
    if coverage_items:
        add_table(summary_sheet, f"A{coverage_start + 1}:B{coverage_last}", "CoberturaGeralTable")
        coverage_chart = BarChart()
        coverage_chart.type = "bar"
        coverage_chart.style = 10
        coverage_chart.title = (
            f"Cobertura crítica: {analysis.get('skus_cobertura_critica', 0)} SKUs"
        )
        coverage_chart.height = 8
        coverage_chart.width = 16
        coverage_chart.add_data(
            Reference(summary_sheet, min_col=2, min_row=coverage_start + 1, max_row=coverage_last),
            titles_from_data=True,
        )
        coverage_chart.set_categories(
            Reference(summary_sheet, min_col=1, min_row=coverage_start + 2, max_row=coverage_last)
        )
        coverage_chart.series[0].graphicalProperties.solidFill = orange
        coverage_chart.legend = None
        summary_sheet.add_chart(coverage_chart, f"J{monthly_start + 35}")

    top_sales_start = coverage_last + 3
    summary_sheet.cell(top_sales_start, 1, "Top 10 mais vendidos por unidades").font = Font(
        bold=True, color=navy, size=12
    )
    top_sales_headers = [
        "Rank",
        "SKU",
        "Produto",
        "Unidades vendidas",
        "Faturamento",
        "Participa\u00e7\u00e3o nas unidades",
    ]
    for column, value in enumerate(top_sales_headers, start=1):
        summary_sheet.cell(top_sales_start + 1, column, value)
    header(summary_sheet[top_sales_start + 1])
    top_sales_rows = report.get("top_10_vendidos") or []
    for row_index, item in enumerate(top_sales_rows, start=top_sales_start + 2):
        summary_sheet.cell(row_index, 1, item["rank"])
        summary_sheet.cell(row_index, 2, item["sku"])
        summary_sheet.cell(row_index, 3, item["produto"])
        summary_sheet.cell(row_index, 4, item["unidades_vendidas"])
        summary_sheet.cell(row_index, 5, item["faturamento"])
        summary_sheet.cell(row_index, 6, item["participacao_unidades"])
        summary_sheet.cell(row_index, 2).number_format = "@"
        summary_sheet.cell(row_index, 4).number_format = "#,##0.00"
        summary_sheet.cell(row_index, 5).number_format = '"R$" #,##0.00'
        summary_sheet.cell(row_index, 6).number_format = "0.00%"
        summary_sheet.cell(row_index, 3).alignment = Alignment(
            wrap_text=True, vertical="top"
        )
        summary_sheet.row_dimensions[row_index].height = 30
    top_sales_last = top_sales_start + 1 + len(top_sales_rows)

    top_revenue_start = top_sales_last + 2
    summary_sheet.cell(top_revenue_start, 1, "Top 10 por valor vendido").font = Font(
        bold=True, color=navy, size=12
    )
    top_revenue_headers = [
        "Rank",
        "SKU",
        "Produto",
        "Unidades vendidas",
        "Faturamento",
        "Participa\u00e7\u00e3o no faturamento",
    ]
    for column, value in enumerate(top_revenue_headers, start=1):
        summary_sheet.cell(top_revenue_start + 1, column, value)
    header(summary_sheet[top_revenue_start + 1][:6])
    top_revenue_rows = report.get("top_10_faturamento") or []
    for row_index, item in enumerate(top_revenue_rows, start=top_revenue_start + 2):
        summary_sheet.cell(row_index, 1, item["rank"])
        summary_sheet.cell(row_index, 2, item["sku"])
        summary_sheet.cell(row_index, 3, item["produto"])
        summary_sheet.cell(row_index, 4, item["unidades_vendidas"])
        summary_sheet.cell(row_index, 5, item["faturamento"])
        summary_sheet.cell(row_index, 6, item["participacao_faturamento"])
        summary_sheet.cell(row_index, 2).number_format = "@"
        summary_sheet.cell(row_index, 4).number_format = "#,##0.00"
        summary_sheet.cell(row_index, 5).number_format = '"R$" #,##0.00'
        summary_sheet.cell(row_index, 6).number_format = "0.00%"
        summary_sheet.cell(row_index, 3).alignment = Alignment(
            wrap_text=True, vertical="top"
        )
        summary_sheet.row_dimensions[row_index].height = 30
    top_revenue_last = top_revenue_start + 1 + len(top_revenue_rows)
    if top_revenue_rows:
        add_table(
            summary_sheet,
            f"A{top_revenue_start + 1}:F{top_revenue_last}",
            "Top10FaturamentoGeralTable",
        )

    top_charts_start = max(top_sales_last, top_revenue_last) + 2
    if top_sales_rows:
        add_table(
            summary_sheet,
            f"A{top_sales_start + 1}:F{top_sales_last}",
            "Top10VendidosGeralTable",
        )
        top_sales_chart = BarChart()
        top_sales_chart.type = "bar"
        top_sales_chart.style = 10
        leader = top_sales_rows[0]
        leader_units_text = f"{float(leader['unidades_vendidas']):,.0f}".replace(",", ".")
        top_sales_chart.title = (
            f"Top 10 mais vendidos - SKU {leader['sku']} lidera com "
            f"{leader_units_text} un."
        )
        top_sales_chart.x_axis.title = "Unidades vendidas"
        top_sales_chart.height = 7
        top_sales_chart.width = 15
        top_sales_chart.add_data(
            Reference(
                summary_sheet,
                min_col=4,
                min_row=top_sales_start + 1,
                max_row=top_sales_last,
            ),
            titles_from_data=True,
        )
        top_sales_chart.set_categories(
            Reference(
                summary_sheet,
                min_col=2,
                min_row=top_sales_start + 2,
                max_row=top_sales_last,
            )
        )
        top_sales_chart.series[0].graphicalProperties.solidFill = blue
        top_sales_chart.legend = None
        summary_sheet.add_chart(top_sales_chart, f"A{top_charts_start}")

    if top_revenue_rows:
        top_revenue_chart = BarChart()
        top_revenue_chart.type = "bar"
        top_revenue_chart.style = 10
        leader = top_revenue_rows[0]
        top_revenue_chart.title = (
            f"Top 10 por valor vendido - SKU {leader['sku']} lidera"
        )
        top_revenue_chart.x_axis.title = "Faturamento (R$)"
        top_revenue_chart.height = 7
        top_revenue_chart.width = 15
        top_revenue_chart.add_data(
            Reference(
                summary_sheet,
                min_col=5,
                min_row=top_revenue_start + 1,
                max_row=top_revenue_last,
            ),
            titles_from_data=True,
        )
        top_revenue_chart.set_categories(
            Reference(
                summary_sheet,
                min_col=2,
                min_row=top_revenue_start + 2,
                max_row=top_revenue_last,
            )
        )
        top_revenue_chart.series[0].graphicalProperties.solidFill = teal
        top_revenue_chart.legend = None
        summary_sheet.add_chart(top_revenue_chart, f"N{top_charts_start}")

    months = period["meses"]
    base_headers = ["Rank", "SKU", "Produto"]
    trailing_headers = [
        "Total unidades",
        "Faturamento",
        "Unidades Full",
        "Demais canais",
        "Média mensal",
        "Estoque Loja",
        "Cobertura meses",
        "Cobertura dias",
        "Status cobertura",
        "Qtd. devolvida",
        "Taxa devolução",
    ]
    sku_headers = base_headers + months + trailing_headers
    title(sku_sheet, "Todos os SKUs - vendas mensais e estoque da Loja", len(sku_headers))
    for column, value in enumerate(sku_headers, start=1):
        sku_sheet.cell(2, column, value)
    header(sku_sheet[2])

    first_month_col = 4
    last_month_col = first_month_col + len(months) - 1
    total_col = last_month_col + 1
    revenue_col = total_col + 1
    full_col = total_col + 2
    other_col = total_col + 3
    average_col = total_col + 4
    stock_col = total_col + 5
    coverage_month_col = total_col + 6
    coverage_day_col = total_col + 7
    status_col = total_col + 8
    return_qty_col = total_col + 9
    return_rate_col = total_col + 10

    for row_index, item in enumerate(report["skus"], start=3):
        values = [item["rank"], str(item["sku"]), item["produto"]]
        values.extend(item["vendas_mensais"].get(month, 0) for month in months)
        for column, value in enumerate(values, start=1):
            sku_sheet.cell(row_index, column, value)
        month_start_letter = get_column_letter(first_month_col)
        month_end_letter = get_column_letter(last_month_col)
        total_letter = get_column_letter(total_col)
        average_letter = get_column_letter(average_col)
        stock_letter = get_column_letter(stock_col)
        coverage_letter = get_column_letter(coverage_month_col)
        return_qty_letter = get_column_letter(return_qty_col)
        sku_sheet.cell(row_index, total_col, f"=SUM({month_start_letter}{row_index}:{month_end_letter}{row_index})")
        sku_sheet.cell(row_index, revenue_col, item["faturamento"])
        sku_sheet.cell(row_index, full_col, item["full_unidades"])
        sku_sheet.cell(row_index, other_col, item["demais_unidades"])
        sku_sheet.cell(row_index, average_col, f"={total_letter}{row_index}/'Resumo'!$B$8")
        sku_sheet.cell(row_index, stock_col, item["estoque_loja"])
        sku_sheet.cell(
            row_index,
            coverage_month_col,
            f'=IF(OR(ISBLANK({stock_letter}{row_index}),{average_letter}{row_index}<=0),"",MAX({stock_letter}{row_index},0)/{average_letter}{row_index})',
        )
        sku_sheet.cell(
            row_index,
            coverage_day_col,
            f'=IF({coverage_letter}{row_index}="","",{coverage_letter}{row_index}*30.4375)',
        )
        sku_sheet.cell(
            row_index,
            status_col,
            f'=IF(ISBLANK({stock_letter}{row_index}),"Estoque sem registro",IF({average_letter}{row_index}<=0,IF({stock_letter}{row_index}>0,"Sem vendas no período","Sem estoque e sem vendas"),IF({stock_letter}{row_index}<=0,"Ruptura",IF({coverage_letter}{row_index}<=1,"Até 1 mês",IF({coverage_letter}{row_index}<=3,"1 a 3 meses","Acima de 3 meses")))))',
        )
        sku_sheet.cell(row_index, return_qty_col, item["quantidade_devolvida"])
        sku_sheet.cell(
            row_index,
            return_rate_col,
            f'=IF({total_letter}{row_index}>0,{return_qty_letter}{row_index}/{total_letter}{row_index},"")',
        )
        sku_sheet.cell(row_index, 2).number_format = "@"
        for column in range(first_month_col, last_month_col + 1):
            sku_sheet.cell(row_index, column).number_format = "#,##0.00"
        for column in (total_col, full_col, other_col, average_col, stock_col, coverage_month_col, coverage_day_col, return_qty_col):
            sku_sheet.cell(row_index, column).number_format = "#,##0.00"
        sku_sheet.cell(row_index, revenue_col).number_format = '"R$" #,##0.00'
        sku_sheet.cell(row_index, return_rate_col).number_format = "0.00%"

    sku_last = max(2, len(report["skus"]) + 2)
    if report["skus"]:
        add_table(
            sku_sheet,
            f"A2:{get_column_letter(len(sku_headers))}{sku_last}",
            "VendasEstoqueGeralTable",
        )
        sku_sheet.conditional_formatting.add(
            f"{get_column_letter(revenue_col)}3:{get_column_letter(revenue_col)}{sku_last}",
            DataBarRule(start_type="min", end_type="max", color=blue),
        )
        sku_sheet.conditional_formatting.add(
            f"{get_column_letter(coverage_month_col)}3:{get_column_letter(coverage_month_col)}{sku_last}",
            ColorScaleRule(
                start_type="min", start_color=light_red,
                mid_type="percentile", mid_value=50, mid_color=light_orange,
                end_type="max", end_color=light_green,
            ),
        )
        status_range = f"{get_column_letter(status_col)}3:{get_column_letter(status_col)}{sku_last}"
        sku_sheet.conditional_formatting.add(
            status_range,
            FormulaRule(
                formula=[f'${get_column_letter(status_col)}3="Ruptura"'],
                fill=PatternFill("solid", fgColor=light_red),
                font=Font(color=red, bold=True),
            ),
        )
    sku_sheet.freeze_panes = "D3"
    sku_sheet.auto_filter.ref = f"A2:{get_column_letter(len(sku_headers))}{sku_last}"
    widths = {"A": 8, "B": 18, "C": 42}
    for column, width in widths.items():
        sku_sheet.column_dimensions[column].width = width
    for column in range(first_month_col, last_month_col + 1):
        sku_sheet.column_dimensions[get_column_letter(column)].width = 12
    for column in range(total_col, return_rate_col + 1):
        sku_sheet.column_dimensions[get_column_letter(column)].width = 17
    sku_sheet.column_dimensions[get_column_letter(status_col)].width = 25

    monthly_headers = [
        "Mês", "SKU", "Produto", "Unidades", "Faturamento",
        "Unidades Full", "Faturamento Full", "Unidades demais", "Faturamento demais",
        "Qtd. devolvida", "Valor devolvido",
    ]
    title(monthly_sheet, "Detalhamento mensal por SKU e canal", len(monthly_headers))
    for column, value in enumerate(monthly_headers, start=1):
        monthly_sheet.cell(2, column, value)
    header(monthly_sheet[2])
    for row_index, item in enumerate(report["mensal_detalhado"], start=3):
        monthly_sheet.append(
            [
                item["mes"], str(item["sku"]), item["produto"], item["unidades"],
                item["faturamento"], item["full_unidades"], item["full_faturamento"],
                item["demais_unidades"], item["demais_faturamento"],
                item["devolucoes_quantidade"], item["devolucoes_valor"],
            ]
        )
        monthly_sheet.cell(row_index, 2).number_format = "@"
        for column in (4, 6, 8, 10):
            monthly_sheet.cell(row_index, column).number_format = "#,##0.00"
        for column in (5, 7, 9, 11):
            monthly_sheet.cell(row_index, column).number_format = '"R$" #,##0.00'
    monthly_last_detail = max(2, len(report["mensal_detalhado"]) + 2)
    if report["mensal_detalhado"]:
        add_table(monthly_sheet, f"A2:K{monthly_last_detail}", "MensalDetalhadoGeralTable")
    monthly_sheet.freeze_panes = "D3"
    for column, width in {"A": 13, "B": 18, "C": 42, "D": 14, "E": 18, "F": 16, "G": 19, "H": 18, "I": 21, "J": 17, "K": 18}.items():
        monthly_sheet.column_dimensions[column].width = width

    return_headers = [
        "Rank", "SKU", "Produto", "Ocorrências", "Quantidade", "Valor",
        "Qtd. Full", "Qtd. demais", "Unidades vendidas", "Taxa devolução",
        "Qtd. últimos 3 meses", "Participação recente", "Última devolução",
    ]
    title(
        returns_sheet,
        "Devoluções de todos os SKUs da conta",
        len(return_headers),
        fill_color=dark_red,
    )
    for column, value in enumerate(return_headers, start=1):
        returns_sheet.cell(2, column, value)
    header(returns_sheet[2], fill_color=return_red)
    for rank, item in enumerate(report["devolucoes"], start=1):
        row_index = rank + 2
        returns_sheet.append(
            [
                rank, str(item["sku"]), item["produto"], item["ocorrencias"],
                item["quantidade"], item["valor"], item["full_quantidade"],
                item["demais_quantidade"], item["unidades_vendidas"],
                item["taxa_devolucao"], item["quantidade_ultimos_3_meses"],
                item["participacao_recente"], item["ultima_devolucao"],
            ]
        )
        returns_sheet.cell(row_index, 2).number_format = "@"
        for column in (5, 7, 8, 9, 11):
            returns_sheet.cell(row_index, column).number_format = "#,##0.00"
        returns_sheet.cell(row_index, 6).number_format = '"R$" #,##0.00'
        for column in (10, 12):
            returns_sheet.cell(row_index, column).number_format = "0.00%"
    returns_last = max(2, len(report["devolucoes"]) + 2)
    if report["devolucoes"]:
        add_table(
            returns_sheet,
            f"A2:M{returns_last}",
            "DevolucoesTodosSkusTable",
            show_row_stripes=False,
        )
        returns_sheet.conditional_formatting.add(
            f"E3:E{returns_last}",
            DataBarRule(start_type="min", end_type="max", color=return_red),
        )
    apply_returns_palette(returns_sheet, returns_last, len(return_headers))
    returns_sheet.freeze_panes = "C3"
    for column, width in {"A": 8, "B": 18, "C": 42, "D": 13, "E": 14, "F": 18, "G": 13, "H": 14, "I": 17, "J": 16, "K": 20, "L": 18, "M": 17}.items():
        returns_sheet.column_dimensions[column].width = width

    top_headers = return_headers + ["Participação nas devoluções", "Predominância", "Análise"]
    title(
        top_returns_sheet,
        "20 SKUs mais devolvidos - análise",
        len(top_headers),
        fill_color=dark_red,
    )
    for column, value in enumerate(top_headers, start=1):
        top_returns_sheet.cell(2, column, value)
    header(top_returns_sheet[2], fill_color=return_red)
    for item in report["top_20_devolucoes"]:
        top_returns_sheet.append(
            [
                item["rank"], str(item["sku"]), item["produto"], item["ocorrencias"],
                item["quantidade"], item["valor"], item["full_quantidade"],
                item["demais_quantidade"], item["unidades_vendidas"],
                item["taxa_devolucao"], item["quantidade_ultimos_3_meses"],
                item["participacao_recente"], item["ultima_devolucao"],
                item["participacao_devolucoes"], item["predominancia"], item["analise"],
            ]
        )
    top_last = max(2, len(report["top_20_devolucoes"]) + 2)
    for row_index in range(3, top_last + 1):
        top_returns_sheet.cell(row_index, 2).number_format = "@"
        top_returns_sheet.cell(row_index, 6).number_format = '"R$" #,##0.00'
        for column in (5, 7, 8, 9, 11):
            top_returns_sheet.cell(row_index, column).number_format = "#,##0.00"
        for column in (10, 12, 14):
            top_returns_sheet.cell(row_index, column).number_format = "0.00%"
        top_returns_sheet.cell(row_index, 16).alignment = Alignment(wrap_text=True, vertical="top")
        top_returns_sheet.row_dimensions[row_index].height = 42
    if report["top_20_devolucoes"]:
        add_table(
            top_returns_sheet,
            f"A2:P{top_last}",
            "Top20DevolucoesGeralTable",
            show_row_stripes=False,
        )
        chart_limit = min(top_last, 22)
        chart = BarChart()
        chart.type = "bar"
        chart.style = 10
        chart.title = "Top 20 por quantidade devolvida"
        chart.height = 11
        chart.width = 18
        chart.add_data(
            Reference(top_returns_sheet, min_col=5, min_row=2, max_row=chart_limit),
            titles_from_data=True,
        )
        chart.set_categories(
            Reference(top_returns_sheet, min_col=2, min_row=3, max_row=chart_limit)
        )
        chart.series[0].tx = SeriesLabel(v="Quantidade devolvida")
        chart.series[0].graphicalProperties.solidFill = return_red
        chart.legend = None
        top_returns_sheet.add_chart(chart, "R2")

        rate_chart = BarChart()
        rate_chart.type = "bar"
        rate_chart.style = 10
        rate_chart.title = "Taxa de devolução dos Top 20 sobre vendas"
        rate_chart.x_axis.title = "Taxa de devolução"
        rate_chart.x_axis.numFmt = "0.0%"
        rate_chart.height = 11
        rate_chart.width = 18
        rate_chart.add_data(
            Reference(top_returns_sheet, min_col=10, min_row=2, max_row=chart_limit),
            titles_from_data=True,
        )
        rate_chart.set_categories(
            Reference(top_returns_sheet, min_col=2, min_row=3, max_row=chart_limit)
        )
        rate_chart.series[0].tx = SeriesLabel(v="Taxa de devolução")
        rate_chart.series[0].graphicalProperties.solidFill = red
        rate_chart.legend = None
        top_returns_sheet.add_chart(rate_chart, "R24")
    apply_returns_palette(top_returns_sheet, top_last, len(top_headers))
    top_returns_sheet.freeze_panes = "C3"
    for column in range(1, 17):
        top_returns_sheet.column_dimensions[get_column_letter(column)].width = 15
    top_returns_sheet.column_dimensions["B"].width = 18
    top_returns_sheet.column_dimensions["C"].width = 38
    top_returns_sheet.column_dimensions["P"].width = 70

    audit = report["auditoria"]
    title(rules_sheet, "Fontes, regras e controles de auditoria", 5)
    controls = [
        ("Conta", report["loja"]),
        ("Período", f"{period['data_inicio']} a {period['data_fim']}"),
        ("Assinatura das fontes", audit["assinatura_fontes"]),
        ("Bases de vendas", ", ".join(audit["bases_vendas"]) or "Nenhuma"),
        ("Bases de devoluções", ", ".join(audit["bases_devolucoes"]) or "Nenhuma"),
        ("Fonte do estoque", audit["estoque"].get("fonte") or "Indisponível"),
        ("Evento do estoque", audit["estoque"].get("event_id") or "Não aplicável"),
        ("SKUs no snapshot", audit["estoque"].get("total_skus_snapshot") or 0),
        ("Linhas de vendas lidas", audit["vendas"]["linhas_lidas"]),
        ("Duplicidades de vendas entre bases", audit["vendas"]["duplicidades_entre_bases"]),
        ("Duplicidades pedido/nota", audit["vendas"]["duplicidades_pedido_nota"]),
        ("Devoluções lidas", audit["devolucoes_lidas"]),
        ("Devoluções excluídas/deduplicadas", audit["devolucoes_excluidas_ou_deduplicadas"]),
        ("Ajuste fiscal segregado", f"{report['ajuste_fiscal']['linhas']} linha(s); {_money(report['ajuste_fiscal']['valor'])}"),
    ]
    rules_sheet.append(["Controle", "Valor"])
    header(rules_sheet[2])
    for label, value in controls:
        rules_sheet.append([label, value])
    rules_start = rules_sheet.max_row + 2
    rules_sheet.cell(rules_start, 1, "Regras aplicadas").font = Font(bold=True, color=navy, size=12)
    business_rules = [
        "A união de SKUs considera vendas, devoluções e o último estoque confirmado da conta.",
        "As tabelas principais exibem somente SKUs com venda ou devolução no período.",
        "SKUs sem movimentação aparecem apenas na aba final, agrupados pela situação do estoque.",
        "Vendas e devoluções incluem Mercado Livre Full e demais canais da conta exata.",
        "Estoque usa somente saldo_loja; saldo_full e eventos pendentes são ignorados.",
        "Ausência no snapshot é exibida como sem registro, nunca como estoque zero presumido.",
        "A média divide as unidades vendidas por todos os meses-calendário solicitados, inclusive meses zerados.",
        "O Top 20 usa quantidade devolvida, depois valor devolvido e SKU como desempates.",
    ]
    for offset, rule in enumerate(business_rules, start=1):
        row = rules_start + offset
        rules_sheet.cell(row, 1, f"{offset}. {rule}")
        rules_sheet.merge_cells(start_row=row, start_column=1, end_row=row, end_column=5)
        rules_sheet.cell(row, 1).alignment = Alignment(wrap_text=True, vertical="top")
    qc_row = rules_start + len(business_rules) + 3
    rules_sheet.cell(qc_row, 1, "Controle por fórmula")
    rules_sheet.cell(qc_row, 2, "Resultado")
    header(rules_sheet[qc_row])
    revenue_letter = get_column_letter(revenue_col)
    total_letter = get_column_letter(total_col)
    return_letter = get_column_letter(return_qty_col)
    rules_sheet.cell(qc_row + 1, 1, "Unidades vendidas na aba Vendas e Estoque")
    rules_sheet.cell(qc_row + 1, 2, f"=SUM('Vendas e Estoque'!{total_letter}3:{total_letter}{max(3, sku_last)})")
    rules_sheet.cell(qc_row + 2, 1, "Faturamento na aba Vendas e Estoque")
    rules_sheet.cell(qc_row + 2, 2, f"=SUM('Vendas e Estoque'!{revenue_letter}3:{revenue_letter}{max(3, sku_last)})")
    rules_sheet.cell(qc_row + 3, 1, "Quantidade devolvida na aba Vendas e Estoque")
    rules_sheet.cell(qc_row + 3, 2, f"=SUM('Vendas e Estoque'!{return_letter}3:{return_letter}{max(3, sku_last)})")
    rules_sheet.cell(qc_row + 1, 2).number_format = "#,##0.00"
    rules_sheet.cell(qc_row + 2, 2).number_format = '"R$" #,##0.00'
    rules_sheet.cell(qc_row + 3, 2).number_format = "#,##0.00"
    rules_sheet.column_dimensions["A"].width = 58
    rules_sheet.column_dimensions["B"].width = 80
    for column in ("C", "D", "E"):
        rules_sheet.column_dimensions[column].width = 16
    rules_sheet.freeze_panes = "A3"

    title(no_movement_sheet, "SKUs sem movimentação no período", 4)
    no_movement_sheet.merge_cells("A2:D3")
    no_movement_sheet["A2"] = (
        "Estes SKUs não tiveram venda nem devolução no período solicitado. "
        "Por isso, foram retirados das tabelas e dos gráficos principais e são citados "
        "somente aqui, de forma agrupada."
    )
    no_movement_sheet["A2"].fill = PatternFill("solid", fgColor=light_gray)
    no_movement_sheet["A2"].alignment = Alignment(wrap_text=True, vertical="center")
    no_movement_sheet.row_dimensions[2].height = 26
    no_movement_headers = [
        "Grupo",
        "Quantidade de SKUs",
        "Estoque Loja total conhecido",
        "SKUs agrupados",
    ]
    for column, value in enumerate(no_movement_headers, start=1):
        no_movement_sheet.cell(5, column, value)
    header(no_movement_sheet[5])
    no_movement_groups = (report.get("sem_movimentacao") or {}).get("grupos") or []
    if no_movement_groups:
        for row_index, group in enumerate(no_movement_groups, start=6):
            sku_text = ", ".join(str(item["sku"]) for item in group["itens"])
            no_movement_sheet.cell(row_index, 1, group["grupo"])
            no_movement_sheet.cell(row_index, 2, group["quantidade_skus"])
            no_movement_sheet.cell(
                row_index,
                3,
                group["estoque_loja_total_conhecido"]
                if group.get("estoque_disponivel")
                else "Indispon\u00edvel",
            )
            no_movement_sheet.cell(row_index, 4, sku_text)
            no_movement_sheet.cell(row_index, 2).number_format = "#,##0"
            if group.get("estoque_disponivel"):
                no_movement_sheet.cell(row_index, 3).number_format = "#,##0.00"
            no_movement_sheet.cell(row_index, 4).alignment = Alignment(
                wrap_text=True, vertical="top"
            )
            no_movement_sheet.row_dimensions[row_index].height = min(
                180, 30 + (len(sku_text) // 100) * 15
            )
        no_movement_last = 5 + len(no_movement_groups)
        add_table(
            no_movement_sheet,
            f"A5:D{no_movement_last}",
            "SkusSemMovimentacaoAgrupadosTable",
        )
    else:
        no_movement_sheet.cell(6, 1, "Nenhum SKU sem movimentação no período.")
        no_movement_sheet.merge_cells("A6:D6")
    no_movement_sheet.column_dimensions["A"].width = 32
    no_movement_sheet.column_dimensions["B"].width = 20
    no_movement_sheet.column_dimensions["C"].width = 28
    no_movement_sheet.column_dimensions["D"].width = 110
    no_movement_sheet.freeze_panes = "A6"

    summary_sheet.column_dimensions["A"].width = 30
    summary_sheet.column_dimensions["B"].width = 28
    summary_sheet.column_dimensions["C"].width = 22
    summary_sheet.column_dimensions["D"].width = 30
    summary_sheet.column_dimensions["E"].width = 24
    summary_sheet.column_dimensions["F"].width = 19
    summary_sheet.column_dimensions["G"].width = 17
    summary_sheet.column_dimensions["H"].width = 20
    summary_sheet.column_dimensions["I"].width = 4
    summary_sheet.freeze_panes = "A3"

    for sheet in workbook.worksheets:
        for row in sheet.iter_rows():
            for cell in row:
                if cell.value is not None and cell.row != 1:
                    font = copy(cell.font)
                    font.name = "Aptos"
                    cell.font = font
                    alignment = copy(cell.alignment)
                    alignment.vertical = alignment.vertical or "center"
                    cell.alignment = alignment
        sheet.page_setup.orientation = "landscape"
        sheet.page_setup.fitToWidth = 1
        sheet.sheet_properties.pageSetUpPr.fitToPage = True

    workbook.calculation.fullCalcOnLoad = True
    workbook.calculation.forceFullCalc = True
    workbook.calculation.calcMode = "auto"
    output = io.BytesIO()
    workbook.save(output)
    return output.getvalue()


def _pdf_general_monthly_chart(report: dict[str, Any]):
    from reportlab.graphics.charts.barcharts import VerticalBarChart
    from reportlab.graphics.shapes import Drawing, Rect, String
    from reportlab.lib import colors

    rows = report["mensal_conta"]
    comparison = (report.get("analise_executiva") or {}).get("comparacao_faturamento") or {}
    variation = comparison.get("variacao")
    suffix = f" ({'+' if variation >= 0 else ''}{variation * 100:.1f}% recente)" if variation is not None else ""
    drawing = Drawing(360, 165)
    drawing.add(
        String(
            10,
            150,
            f"Faturamento mensal por canal{suffix}",
            fontName="Helvetica-Bold",
            fontSize=10,
            fillColor=colors.HexColor("#17365D"),
        )
    )
    if not rows:
        drawing.add(String(20, 100, "Sem movimentação no período.", fontSize=9))
        return drawing
    chart = VerticalBarChart()
    chart.x = 40
    chart.y = 28
    chart.width = 300
    chart.height = 100
    chart.data = [
        [float(item["full_faturamento"]) for item in rows],
        [float(item["demais_faturamento"]) for item in rows],
    ]
    chart.categoryAxis.categoryNames = [item["mes"] for item in rows]
    chart.categoryAxis.labels.angle = 30
    chart.categoryAxis.labels.dy = -10
    chart.categoryAxis.labels.fontSize = 6.5
    chart.valueAxis.labels.fontSize = 6.5
    chart.valueAxis.valueMin = 0
    chart.valueAxis.labelTextFormat = lambda value: f"{value / 1000:.0f} mil" if value >= 1000 else f"{value:.0f}"
    chart.bars[0].fillColor = colors.HexColor("#00A6A6")
    chart.bars[1].fillColor = colors.HexColor("#2F75B5")
    chart.categoryAxis.style = "stacked"
    drawing.add(chart)
    drawing.add(Rect(220, 134, 6, 6, fillColor=colors.HexColor("#00A6A6"), strokeColor=None))
    drawing.add(String(229, 134, "Full", fontSize=6.5, fillColor=colors.HexColor("#17365D")))
    drawing.add(Rect(263, 134, 6, 6, fillColor=colors.HexColor("#2F75B5"), strokeColor=None))
    drawing.add(String(272, 134, "Demais canais", fontSize=6.5, fillColor=colors.HexColor("#17365D")))
    return drawing


def _pdf_general_return_rate_chart(report: dict[str, Any]):
    from reportlab.graphics.charts.linecharts import HorizontalLineChart
    from reportlab.graphics.shapes import Drawing, String
    from reportlab.lib import colors

    rows = report["mensal_conta"]
    drawing = Drawing(360, 165)
    drawing.add(
        String(
            10,
            150,
            "Taxa mensal de devolução sobre unidades vendidas",
            fontName="Helvetica-Bold",
            fontSize=10,
            fillColor=colors.HexColor("#17365D"),
        )
    )
    if not rows:
        drawing.add(String(20, 78, "Sem movimentação no período.", fontSize=9))
        return drawing
    chart = HorizontalLineChart()
    chart.x = 40
    chart.y = 30
    chart.width = 300
    chart.height = 100
    chart.data = [[float(item["taxa_devolucao"] or 0) for item in rows]]
    chart.categoryAxis.categoryNames = [item["mes"] for item in rows]
    chart.categoryAxis.labels.angle = 30
    chart.categoryAxis.labels.dy = -10
    chart.categoryAxis.labels.fontSize = 6.5
    chart.valueAxis.labels.fontSize = 6.5
    chart.valueAxis.valueMin = 0
    chart.valueAxis.labelTextFormat = lambda value: f"{value * 100:.0f}%"
    chart.lines[0].strokeColor = colors.HexColor("#C00000")
    drawing.add(chart)
    return drawing


def _pdf_general_coverage_chart(report: dict[str, Any]):
    from reportlab.graphics.charts.barcharts import HorizontalBarChart
    from reportlab.graphics.shapes import Drawing, String
    from reportlab.lib import colors

    order = [
        "Ruptura",
        "Até 1 mês",
        "1 a 3 meses",
        "Acima de 3 meses",
        "Sem vendas no período",
        "Sem estoque e sem vendas",
        "Estoque sem registro",
    ]
    counts = report["resumo"].get("status_cobertura") or {}
    critical = (report.get("analise_executiva") or {}).get("skus_cobertura_critica", 0)
    drawing = Drawing(360, 165)
    drawing.add(
        String(
            10,
            150,
            f"Cobertura de estoque - {critical} SKUs críticos",
            fontName="Helvetica-Bold",
            fontSize=10,
            fillColor=colors.HexColor("#17365D"),
        )
    )
    chart = HorizontalBarChart()
    chart.x = 115
    chart.y = 22
    chart.width = 220
    chart.height = 110
    chart.data = [[float(counts.get(status, 0)) for status in reversed(order)]]
    chart.categoryAxis.categoryNames = list(reversed(order))
    chart.valueAxis.valueMin = 0
    chart.categoryAxis.labels.fontSize = 6.2
    chart.valueAxis.labels.fontSize = 6.2
    chart.bars[0].fillColor = colors.HexColor("#ED7D31")
    drawing.add(chart)
    return drawing


def _pdf_general_top_sellers_chart(report: dict[str, Any]):
    from reportlab.graphics.charts.barcharts import HorizontalBarChart
    from reportlab.graphics.shapes import Drawing, String
    from reportlab.lib import colors

    rows = report.get("top_10_vendidos") or []
    drawing = Drawing(340, 98)
    if not rows:
        drawing.add(
            String(
                8,
                84,
                "Top 10 por unidades - sem vendas eleg\u00edveis",
                fontName="Helvetica-Bold",
                fontSize=9.3,
                fillColor=colors.HexColor("#17365D"),
            )
        )
        return drawing
    leader = rows[0]
    leader_units_text = f"{float(leader['unidades_vendidas']):,.0f}".replace(",", ".")
    drawing.add(
        String(
            8,
            86,
            "Top 10 mais vendidos por unidades",
            fontName="Helvetica-Bold",
            fontSize=9.3,
            fillColor=colors.HexColor("#17365D"),
        )
    )
    drawing.add(
        String(
            8,
            75,
            f"L\u00edder: SKU {leader['sku']} - {leader_units_text} un.",
            fontName="Helvetica",
            fontSize=7.2,
            fillColor=colors.HexColor("#5B6573"),
        )
    )
    chart = HorizontalBarChart()
    chart.x = 58
    chart.y = 6
    chart.width = 255
    chart.height = 58
    chart.data = [[float(item["unidades_vendidas"]) for item in reversed(rows)]]
    chart.categoryAxis.categoryNames = [str(item["sku"])[:18] for item in reversed(rows)]
    chart.valueAxis.valueMin = 0
    chart.valueAxis.labelTextFormat = lambda value: f"{value:,.0f}".replace(",", ".")
    chart.categoryAxis.labels.fontSize = 5.5
    chart.valueAxis.labels.fontSize = 5.5
    chart.bars[0].fillColor = colors.HexColor("#2F75B5")
    drawing.add(chart)
    return drawing


def _pdf_general_top_revenue_chart(report: dict[str, Any]):
    from reportlab.graphics.charts.barcharts import HorizontalBarChart
    from reportlab.graphics.shapes import Drawing, String
    from reportlab.lib import colors

    rows = report.get("top_10_faturamento") or []
    drawing = Drawing(340, 98)
    if not rows:
        drawing.add(
            String(
                8,
                84,
                "Top 10 por valor vendido - sem vendas eleg\u00edveis",
                fontName="Helvetica-Bold",
                fontSize=9.3,
                fillColor=colors.HexColor("#17365D"),
            )
        )
        return drawing
    leader = rows[0]
    drawing.add(
        String(
            8,
            86,
            "Top 10 mais vendidos por valor",
            fontName="Helvetica-Bold",
            fontSize=9.3,
            fillColor=colors.HexColor("#17365D"),
        )
    )
    drawing.add(
        String(
            8,
            75,
            f"L\u00edder: SKU {leader['sku']} - {_money(leader['faturamento'])}",
            fontName="Helvetica",
            fontSize=7.2,
            fillColor=colors.HexColor("#5B6573"),
        )
    )
    chart = HorizontalBarChart()
    chart.x = 58
    chart.y = 6
    chart.width = 255
    chart.height = 58
    chart.data = [[float(item["faturamento"]) for item in reversed(rows)]]
    chart.categoryAxis.categoryNames = [str(item["sku"])[:18] for item in reversed(rows)]
    chart.valueAxis.valueMin = 0
    chart.valueAxis.labelTextFormat = lambda value: (
        f"{value / 1_000_000:.1f} mi"
        if value >= 1_000_000
        else f"{value / 1000:.0f} mil"
        if value >= 1000
        else f"{value:.0f}"
    )
    chart.categoryAxis.labels.fontSize = 5.5
    chart.valueAxis.labels.fontSize = 5.5
    chart.bars[0].fillColor = colors.HexColor("#00A6A6")
    drawing.add(chart)
    return drawing


def _pdf_general_returns_chart(report: dict[str, Any]):
    from reportlab.graphics.charts.barcharts import HorizontalBarChart
    from reportlab.graphics.shapes import Drawing, String
    from reportlab.lib import colors

    rows = report["top_20_devolucoes"][:10]
    concentration = (report.get("analise_executiva") or {}).get(
        "concentracao_devolucoes_top_20", 0
    )
    drawing = Drawing(360, 165)
    drawing.add(
        String(
            10,
            150,
            f"Top 10 devolvidos - Top 20 concentram {concentration * 100:.1f}%",
            fontName="Helvetica-Bold",
            fontSize=10,
            fillColor=colors.HexColor("#17365D"),
        )
    )
    if not rows:
        drawing.add(String(20, 100, "Sem devoluções no período.", fontSize=9))
        return drawing
    chart = HorizontalBarChart()
    chart.x = 80
    chart.y = 20
    chart.width = 250
    chart.height = 112
    chart.data = [[float(item["quantidade"]) for item in reversed(rows)]]
    chart.categoryAxis.categoryNames = [str(item["sku"])[:18] for item in reversed(rows)]
    chart.valueAxis.valueMin = 0
    chart.categoryAxis.labels.fontSize = 6.5
    chart.valueAxis.labels.fontSize = 6.5
    chart.bars[0].fillColor = colors.HexColor("#B91C1C")
    drawing.add(chart)
    return drawing


def _pdf_value(value: Any, *, decimals: int = 2) -> str:
    if value is None:
        return "-"
    return f"{float(value):,.{decimals}f}".replace(",", "_").replace(".", ",").replace("_", ".")


def build_general_sku_pdf(report: dict[str, Any]) -> bytes:
    from reportlab.lib import colors
    from reportlab.lib.enums import TA_CENTER, TA_LEFT
    from reportlab.lib.pagesizes import A4, landscape
    from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
    from reportlab.lib.units import mm
    from reportlab.platypus import (
        KeepTogether,
        LongTable,
        PageBreak,
        Paragraph,
        SimpleDocTemplate,
        Spacer,
        Table,
        TableStyle,
    )

    output = io.BytesIO()
    page_size = landscape(A4)
    document = SimpleDocTemplate(
        output,
        pagesize=page_size,
        rightMargin=12 * mm,
        leftMargin=12 * mm,
        topMargin=17 * mm,
        bottomMargin=14 * mm,
        title="Relatório Geral de SKUs - Vendas, Estoque e Devoluções",
        author="JK Sistema",
    )
    styles = getSampleStyleSheet()
    title_style = ParagraphStyle(
        "GeneralTitle",
        parent=styles["Title"],
        fontName="Helvetica-Bold",
        fontSize=17,
        leading=20,
        textColor=colors.HexColor("#17365D"),
        alignment=TA_LEFT,
        spaceAfter=3,
    )
    heading = ParagraphStyle(
        "GeneralHeading",
        parent=styles["Heading2"],
        fontName="Helvetica-Bold",
        fontSize=12,
        leading=15,
        textColor=colors.HexColor("#17365D"),
        spaceBefore=5,
        spaceAfter=6,
    )
    returns_heading = ParagraphStyle(
        "GeneralReturnsHeading",
        parent=heading,
        textColor=colors.HexColor("#8B1E1E"),
    )
    body = ParagraphStyle(
        "GeneralBody",
        parent=styles["BodyText"],
        fontName="Helvetica",
        fontSize=8.5,
        leading=11,
        textColor=colors.HexColor("#303A46"),
    )
    metadata = ParagraphStyle("GeneralMetadata", parent=body, fontSize=7.2, leading=8.5)
    executive_heading = ParagraphStyle(
        "GeneralExecutiveHeading",
        parent=heading,
        fontSize=10.5,
        leading=12,
        spaceBefore=2,
        spaceAfter=3,
    )
    executive_body = ParagraphStyle(
        "GeneralExecutiveBody", parent=body, fontSize=7.2, leading=8.5
    )
    small = ParagraphStyle("GeneralSmall", parent=body, fontSize=6.8, leading=8.2)
    centered = ParagraphStyle("GeneralCentered", parent=small, alignment=TA_CENTER)
    table_header = ParagraphStyle(
        "GeneralTableHeader",
        parent=centered,
        fontName="Helvetica-Bold",
        textColor=colors.white,
    )

    def on_page(canvas, doc) -> None:
        canvas.saveState()
        width, height = page_size
        canvas.setStrokeColor(colors.HexColor("#D5DCE5"))
        canvas.line(12 * mm, height - 11 * mm, width - 12 * mm, height - 11 * mm)
        canvas.setFont("Helvetica", 7.2)
        canvas.setFillColor(colors.HexColor("#5B6573"))
        canvas.drawString(
            12 * mm,
            height - 8 * mm,
            f"JK Sistema - Relatório Geral de SKUs - {report['loja']}",
        )
        canvas.drawRightString(width - 12 * mm, 7 * mm, f"Página {doc.page}")
        canvas.drawString(12 * mm, 7 * mm, f"Gerado em {report['periodo']['gerado_em']}")
        canvas.restoreState()

    period = report["periodo"]
    summary = report["resumo"]
    stock_stamp = xml_escape(str(period.get("estoque_atualizado_em") or "indisponível"))
    story = [
        Paragraph("Relatório Geral de SKUs", title_style),
        Paragraph(
            f"<b>Conta:</b> {xml_escape(str(report['loja']))} &nbsp;&nbsp; "
            f"<b>Período:</b> {period['data_inicio']} a {period['data_fim']} &nbsp;&nbsp; "
            f"<b>Estoque da Loja atualizado em:</b> {stock_stamp}",
            metadata,
        ),
        Spacer(1, 4),
    ]
    kpi_data = [
        [
            Paragraph("SKUs com movimenta\u00e7\u00e3o", small),
            Paragraph("Unidades vendidas", small),
            Paragraph("Faturamento", small),
            Paragraph("Estoque Loja conhecido", small),
            Paragraph("Quantidade devolvida", small),
            Paragraph("Taxa de devolução", small),
        ],
        [
            Paragraph(f"<b>{summary['total_skus']}</b>", centered),
            Paragraph(f"<b>{_pdf_value(summary['unidades_vendidas'])}</b>", centered),
            Paragraph(f"<b>{_money(summary['faturamento'])}</b>", centered),
            Paragraph(f"<b>{_pdf_value(summary['estoque_loja_total_conhecido'])}</b>", centered),
            Paragraph(f"<b>{_pdf_value(summary['quantidade_devolvida'])}</b>", centered),
            Paragraph(
                f"<b>{_pdf_value(_decimal(summary['taxa_devolucao']) * 100) if summary['taxa_devolucao'] is not None else '-'}%</b>",
                centered,
            ),
        ],
    ]
    kpi_table = Table(kpi_data, colWidths=[39 * mm] * 6, rowHeights=[6 * mm, 8 * mm])
    kpi_table.setStyle(
        TableStyle(
            [
                ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#D9EAF7")),
                ("BACKGROUND", (0, 1), (-1, 1), colors.HexColor("#F7FAFC")),
                ("BOX", (0, 0), (-1, -1), 0.5, colors.HexColor("#B8C7D9")),
                ("INNERGRID", (0, 0), (-1, -1), 0.25, colors.HexColor("#D5DCE5")),
                ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
                ("ALIGN", (0, 0), (-1, -1), "CENTER"),
            ]
        )
    )
    insight_rows = [
        [Paragraph(f"• {xml_escape(str(text))}", executive_body)]
        for text in (report.get("analise_executiva") or {}).get("destaques", [])
    ]
    if not insight_rows:
        insight_rows = [[Paragraph("Sem destaques comparativos para o período.", executive_body)]]
    insight_table = Table(insight_rows, colWidths=[234 * mm])
    insight_table.setStyle(
        TableStyle(
            [
                ("BACKGROUND", (0, 0), (-1, -1), colors.HexColor("#F2F6FA")),
                ("BOX", (0, 0), (-1, -1), 0.4, colors.HexColor("#D5DCE5")),
                ("INNERGRID", (0, 0), (-1, -1), 0.2, colors.HexColor("#E4E9F0")),
                ("LEFTPADDING", (0, 0), (-1, -1), 7),
                ("RIGHTPADDING", (0, 0), (-1, -1), 7),
                ("TOPPADDING", (0, 0), (-1, -1), 2.2),
                ("BOTTOMPADDING", (0, 0), (-1, -1), 2.2),
            ]
        )
    )
    overview_charts = Table(
        [[_pdf_general_monthly_chart(report), _pdf_general_coverage_chart(report)]],
        colWidths=[120 * mm, 120 * mm],
    )
    rankings_charts = Table(
        [[_pdf_general_top_sellers_chart(report), _pdf_general_top_revenue_chart(report)]],
        colWidths=[120 * mm, 120 * mm],
    )
    for chart_table in (overview_charts, rankings_charts):
        chart_table.setStyle(
            TableStyle(
                [
                    ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
                    ("LEFTPADDING", (0, 0), (-1, -1), 0),
                    ("RIGHTPADDING", (0, 0), (-1, -1), 0),
                    ("TOPPADDING", (0, 0), (-1, -1), 0),
                    ("BOTTOMPADDING", (0, 0), (-1, -1), 0),
                ]
            )
        )
    story.extend(
        [
            kpi_table,
            Spacer(1, 3),
            Paragraph("Leitura executiva", executive_heading),
            insight_table,
            Spacer(1, 3),
            overview_charts,
            rankings_charts,
            PageBreak(),
        ]
    )

    top_data = [[
        Paragraph("Rank", table_header), Paragraph("SKU", table_header),
        Paragraph("Produto", table_header), Paragraph("Qtd.", table_header),
        Paragraph("Valor", table_header), Paragraph("Taxa", table_header),
        Paragraph("Análise", table_header),
    ]]
    for item in report["top_20_devolucoes"]:
        top_data.append(
            [
                item["rank"],
                Paragraph(xml_escape(str(item["sku"])), small),
                Paragraph(xml_escape(str(item["produto"])), small),
                _pdf_value(item["quantidade"]),
                _money(item["valor"]),
                f"{_pdf_value(_decimal(item['taxa_devolucao']) * 100)}%" if item["taxa_devolucao"] is not None else "-",
                Paragraph(xml_escape(str(item["analise"])), small),
            ]
        )
    if len(top_data) == 1:
        top_data.append(["-", "-", "Sem devoluções no período", "-", "-", "-", "-"])
    top_table = LongTable(
        top_data,
        repeatRows=1,
        colWidths=[10 * mm, 20 * mm, 48 * mm, 15 * mm, 24 * mm, 16 * mm, 108 * mm],
    )
    top_table.setStyle(_pdf_returns_table_style(colors))
    story.append(Paragraph("SKUs com movimentação - resumo", heading))

    sku_data = [[
        Paragraph("Rank", table_header), Paragraph("SKU", table_header),
        Paragraph("Produto", table_header), Paragraph("Vendidas", table_header),
        Paragraph("Faturamento", table_header), Paragraph("Média/mês", table_header),
        Paragraph("Estoque Loja", table_header), Paragraph("Cobertura", table_header),
        Paragraph("Devolvidas", table_header),
    ]]
    for item in report["skus"]:
        coverage = (
            f"{_pdf_value(item['cobertura_meses'])} mês(es)"
            if item["cobertura_meses"] is not None
            else item["status_cobertura"]
        )
        sku_data.append(
            [
                item["rank"], Paragraph(xml_escape(str(item["sku"])), small),
                Paragraph(xml_escape(str(item["produto"])), small),
                _pdf_value(item["unidades_vendidas"]), _money(item["faturamento"]),
                _pdf_value(item["media_mensal"]), _pdf_value(item["estoque_loja"]),
                Paragraph(xml_escape(str(coverage)), small),
                _pdf_value(item["quantidade_devolvida"]),
            ]
        )
    if len(sku_data) == 1:
        sku_data.append(["-", "-", "Sem SKUs no período", "-", "-", "-", "-", "-", "-"])
    sku_table = LongTable(
        sku_data,
        repeatRows=1,
        colWidths=[10 * mm, 20 * mm, 65 * mm, 20 * mm, 28 * mm, 22 * mm, 24 * mm, 34 * mm, 22 * mm],
    )
    sku_table.setStyle(_pdf_table_style(colors))
    story.extend([sku_table, PageBreak()])

    months = period["meses"]
    for block_start in range(0, len(months), 6):
        block = months[block_start:block_start + 6]
        story.append(
            Paragraph(
                f"Vendas mensais por SKU - {block[0]} a {block[-1]}",
                heading,
            )
        )
        monthly_data = [[
            Paragraph("SKU", table_header), Paragraph("Produto", table_header),
            *[Paragraph(month, table_header) for month in block],
            Paragraph("Total", table_header), Paragraph("Média", table_header),
            Paragraph("Estoque", table_header), Paragraph("Cobertura", table_header),
        ]]
        for item in report["skus"]:
            monthly_data.append(
                [
                    Paragraph(xml_escape(str(item["sku"])), small),
                    Paragraph(xml_escape(str(item["produto"])), small),
                    *[_pdf_value(item["vendas_mensais"].get(month, 0)) for month in block],
                    _pdf_value(item["unidades_vendidas"]),
                    _pdf_value(item["media_mensal"]),
                    _pdf_value(item["estoque_loja"]),
                    _pdf_value(item["cobertura_meses"]),
                ]
            )
        if len(monthly_data) == 1:
            monthly_data.append(["-", "Sem SKUs", *(["-"] * (len(block) + 4))])
        remaining = 241 - 20 - 47 - (len(block) * 18)
        monthly_table = LongTable(
            monthly_data,
            repeatRows=1,
            colWidths=[20 * mm, max(35, remaining) * mm / 1.0, *([18 * mm] * len(block)), 19 * mm, 19 * mm, 20 * mm, 20 * mm],
        )
        monthly_table.setStyle(_pdf_table_style(colors, font_size=6.2))
        story.extend([monthly_table, PageBreak()])

    story.append(Paragraph("20 SKUs mais devolvidos - análise", returns_heading))
    story.extend(
        [
            Table(
                [[_pdf_general_return_rate_chart(report), _pdf_general_returns_chart(report)]],
                colWidths=[120 * mm, 120 * mm],
            ),
            Spacer(1, 5),
            top_table,
            PageBreak(),
        ]
    )
    story.append(Paragraph("Devoluções de todos os SKUs", returns_heading))
    returns_data = [[
        Paragraph("Rank", table_header), Paragraph("SKU", table_header),
        Paragraph("Produto", table_header), Paragraph("Ocorr.", table_header),
        Paragraph("Quantidade", table_header), Paragraph("Valor", table_header),
        Paragraph("Full", table_header), Paragraph("Demais", table_header),
        Paragraph("Taxa", table_header), Paragraph("Última", table_header),
    ]]
    for rank, item in enumerate(report["devolucoes"], start=1):
        returns_data.append(
            [
                rank, Paragraph(xml_escape(str(item["sku"])), small),
                Paragraph(xml_escape(str(item["produto"])), small), item["ocorrencias"],
                _pdf_value(item["quantidade"]), _money(item["valor"]),
                _pdf_value(item["full_quantidade"]), _pdf_value(item["demais_quantidade"]),
                f"{_pdf_value(_decimal(item['taxa_devolucao']) * 100)}%" if item["taxa_devolucao"] is not None else "-",
                item["ultima_devolucao"] or "-",
            ]
        )
    if len(returns_data) == 1:
        returns_data.append(["-", "-", "Sem devoluções", "-", "-", "-", "-", "-", "-", "-"])
    returns_table = LongTable(
        returns_data,
        repeatRows=1,
        colWidths=[10 * mm, 20 * mm, 61 * mm, 16 * mm, 22 * mm, 27 * mm, 20 * mm, 20 * mm, 18 * mm, 23 * mm],
    )
    returns_table.setStyle(_pdf_returns_table_style(colors))
    story.extend([returns_table, PageBreak(), Paragraph("Metodologia e auditoria", heading)])

    audit = report["auditoria"]
    audit_rows = [
        ["Controle", "Resultado"],
        ["Assinatura das fontes", audit["assinatura_fontes"]],
        ["Linhas de vendas lidas", audit["vendas"]["linhas_lidas"]],
        ["Duplicidades de vendas entre bases", audit["vendas"]["duplicidades_entre_bases"]],
        ["Duplicidades pedido/nota", audit["vendas"]["duplicidades_pedido_nota"]],
        ["Devoluções lidas", audit["devolucoes_lidas"]],
        ["Devoluções excluídas/deduplicadas", audit["devolucoes_excluidas_ou_deduplicadas"]],
        ["Fonte do estoque", audit["estoque"].get("fonte") or "Indisponível"],
        ["Evento confirmado do estoque", audit["estoque"].get("event_id") or "Não aplicável"],
        ["Ajuste fiscal segregado", f"{report['ajuste_fiscal']['linhas']} linha(s) - {_money(report['ajuste_fiscal']['valor'])}"],
    ]
    audit_table = Table(audit_rows, colWidths=[72 * mm, 160 * mm], repeatRows=1)
    audit_table.setStyle(_pdf_table_style(colors, header_paragraphs=False, font_size=8))
    story.extend(
        [
            audit_table,
            Spacer(1, 10),
            KeepTogether(
                [
                    Paragraph("Regras aplicadas", heading),
                    Paragraph(
                        "1. O universo identificado reúne SKUs com venda, devolução ou registro no estoque atual.<br/>"
                        "2. As tabelas e gráficos principais exibem somente SKUs com venda ou devolução no período.<br/>"
                        "3. Os SKUs sem movimentação aparecem somente no agrupamento final.<br/>"
                        "4. Vendas e devoluções incluem Loja e Mercado Livre Full da conta selecionada.<br/>"
                        "5. O estoque usa somente saldo_loja do último evento confirmado; Full e eventos pendentes são ignorados.<br/>"
                        "6. SKU ausente do snapshot é tratado como estoque sem registro, nunca como zero.<br/>"
                        "7. A média mensal inclui todos os meses-calendário solicitados, mesmo quando não há venda.<br/>"
                        "8. O Top 20 é ordenado por quantidade devolvida, valor devolvido e SKU.",
                        body,
                    ),
                ]
            ),
        ]
    )
    warnings = (
        list(audit.get("erros_vendas") or [])
        + list(audit.get("erros_devolucoes") or [])
        + list(audit.get("estoque", {}).get("erros") or [])
    )
    if warnings:
        story.extend(
            [
                Spacer(1, 8),
                Paragraph("Avisos de cobertura", heading),
                Paragraph("<br/>".join(xml_escape(str(item)) for item in warnings), body),
            ]
        )

    story.extend(
        [
            PageBreak(),
            Paragraph("SKUs sem movimenta\u00e7\u00e3o no per\u00edodo", heading),
            Paragraph(
                "Estes SKUs n\u00e3o tiveram venda nem devolu\u00e7\u00e3o no per\u00edodo solicitado. "
                "Eles foram retirados das tabelas e dos gr\u00e1ficos principais e s\u00e3o citados "
                "somente nesta se\u00e7\u00e3o final, agrupados pela situa\u00e7\u00e3o do estoque.",
                body,
            ),
            Spacer(1, 8),
        ]
    )
    no_movement_groups = (report.get("sem_movimentacao") or {}).get("grupos") or []
    if no_movement_groups:
        for group in no_movement_groups:
            stock_text = (
                _pdf_value(group["estoque_loja_total_conhecido"])
                if group.get("estoque_disponivel")
                else "Indispon\u00edvel"
            )
            sku_text = ", ".join(str(item["sku"]) for item in group["itens"])
            story.extend(
                [
                    Paragraph(
                        f"<b>{xml_escape(str(group['grupo']))}</b> - "
                        f"{group['quantidade_skus']} SKU(s) - "
                        f"Estoque Loja total conhecido: {stock_text}",
                        body,
                    ),
                    Paragraph(xml_escape(sku_text), small),
                    Spacer(1, 7),
                ]
            )
    else:
        story.append(Paragraph("Nenhum SKU sem movimenta\u00e7\u00e3o no per\u00edodo.", body))

    document.build(story, onFirstPage=on_page, onLaterPages=on_page)
    return output.getvalue()


def _pdf_table_style(colors, *, font_size: float = 6.8, header_paragraphs: bool = True):
    from reportlab.platypus import TableStyle

    commands = [
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#17365D")),
        ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
        ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
        ("GRID", (0, 0), (-1, -1), 0.25, colors.HexColor("#D5DCE5")),
        ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#F2F6FA")]),
        ("FONTSIZE", (0, 0), (-1, -1), font_size),
        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
        ("TOPPADDING", (0, 0), (-1, -1), 2.5),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 2.5),
    ]
    if not header_paragraphs:
        commands.append(("ALIGN", (0, 0), (-1, 0), "CENTER"))
    return TableStyle(commands)


def _pdf_returns_table_style(colors, *, font_size: float = 6.8):
    from reportlab.platypus import TableStyle

    return TableStyle(
        [
            ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#7F1D1D")),
            ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
            ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
            ("GRID", (0, 0), (-1, -1), 0.25, colors.HexColor("#E7A9A9")),
            (
                "ROWBACKGROUNDS",
                (0, 1),
                (-1, -1),
                [colors.HexColor("#FFF7F7"), colors.HexColor("#FDECEC")],
            ),
            ("FONTSIZE", (0, 0), (-1, -1), font_size),
            ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
            ("TOPPADDING", (0, 0), (-1, -1), 2.5),
            ("BOTTOMPADDING", (0, 0), (-1, -1), 2.5),
        ]
    )


def export_general_sku_report(
    client_id: str,
    loja: str,
    formato: str,
    periodo: str = "12m",
    data_inicio: str | None = None,
    data_fim: str | None = None,
) -> tuple[bytes, str, str]:
    file_format = str(formato or "").strip().lower()
    if file_format not in {"xlsx", "pdf"}:
        raise VendasDomainError(400, "Formato inválido. Use xlsx ou pdf.")
    report = generate_general_sku_report(
        client_id=client_id,
        loja=loja,
        periodo=periodo,
        data_inicio=data_inicio,
        data_fim=data_fim,
    )
    suffix = f"{report['periodo']['data_inicio']}_{report['periodo']['data_fim']}"
    filename = (
        f"relatorio_geral_skus_{_safe_filename(report['loja'])}_{suffix}.{file_format}"
    )
    if file_format == "xlsx":
        return (
            build_general_sku_xlsx(report),
            filename,
            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        )
    return build_general_sku_pdf(report), filename, "application/pdf"


__all__ = [
    "build_general_sku_pdf",
    "build_general_sku_xlsx",
    "export_general_sku_report",
    "generate_general_sku_report",
]
