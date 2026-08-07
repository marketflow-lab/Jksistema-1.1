"""Pareto 80% sales report and in-memory XLSX/PDF exporters."""

from __future__ import annotations

import hashlib
import io
import os
import re
import sqlite3
import unicodedata
from collections import Counter, defaultdict
from copy import copy
from datetime import date, datetime
from decimal import Decimal, InvalidOperation
from typing import Any, Iterable
from xml.sax.saxutils import escape as xml_escape

from .errors import VendasDomainError
from .legacy import (
    _deduplicar_vendas_consolidadas,
    _deve_excluir_venda_ebazar,
    _listar_bancos_vendas_tenant,
)
from .performance import (
    append_indexable_date_filter,
    cache_vendas_response,
    cached_table_columns,
    open_vendas_readonly,
    select_compatible_column,
    vendas_source_paths,
    vendas_sources_signature,
)


PARETO_TARGET = Decimal("0.80")
FISCAL_ADJUSTMENT_SKU = "ESTORNO DE CREDITO ICMS"
PERIOD_OPTIONS = frozenset({"6m", "12m", "personalizado"})


def _normalize_text(value: Any) -> str:
    text = unicodedata.normalize("NFKD", str(value or ""))
    text = "".join(char for char in text if not unicodedata.combining(char))
    return re.sub(r"\s+", " ", text).strip().upper()


def _normalize_sku(value: Any) -> str:
    return _normalize_text(value)


def _decimal(value: Any) -> Decimal:
    if value in (None, ""):
        return Decimal("0")
    try:
        return Decimal(str(value))
    except (InvalidOperation, ValueError, TypeError):
        return Decimal("0")


def _number(value: Decimal) -> float:
    return float(value)


def _parse_iso_date(value: str | None, field: str) -> date:
    try:
        return date.fromisoformat(str(value or "")[:10])
    except ValueError as exc:
        raise VendasDomainError(400, f"{field} deve estar no formato AAAA-MM-DD.") from exc


def _shift_month_start(reference: date, months_back: int) -> date:
    absolute_month = reference.year * 12 + reference.month - 1 - months_back
    return date(absolute_month // 12, absolute_month % 12 + 1, 1)


def resolve_pareto_period(
    periodo: str,
    data_inicio: str | None = None,
    data_fim: str | None = None,
    *,
    today: date | None = None,
) -> tuple[str, str, str]:
    mode = str(periodo or "12m").strip().lower()
    if mode not in PERIOD_OPTIONS:
        raise VendasDomainError(400, "Período inválido. Use 6m, 12m ou personalizado.")

    current = today or date.today()
    if mode == "personalizado":
        start = _parse_iso_date(data_inicio, "data_inicio")
        end = _parse_iso_date(data_fim, "data_fim")
    else:
        start = _shift_month_start(current, 5 if mode == "6m" else 11)
        end = current

    if start > end:
        raise VendasDomainError(400, "data_inicio não pode ser posterior a data_fim.")
    return mode, start.isoformat(), end.isoformat()


def _validate_store(loja: str | None) -> str:
    store = str(loja or "").strip()
    if not store or store == "__todas":
        raise VendasDomainError(400, "Selecione uma conta específica para gerar o Pareto 80%.")
    return store


def _month_keys(start_iso: str, end_iso: str) -> list[str]:
    start = date.fromisoformat(start_iso[:10]).replace(day=1)
    end = date.fromisoformat(end_iso[:10]).replace(day=1)
    result: list[str] = []
    current = start
    while current <= end:
        result.append(current.strftime("%Y-%m"))
        if current.month == 12:
            current = date(current.year + 1, 1, 1)
        else:
            current = date(current.year, current.month + 1, 1)
    return result


def _operation_name(item: dict[str, Any], store: str) -> str:
    unit = _normalize_text(item.get("unidade_negocio"))
    channel = _normalize_text(item.get("canal"))
    if "FULL" in unit or "FULL" in channel:
        return "Mercado Livre Full"
    return f"Demais canais da conta {store}"


def _is_cancelled(item: dict[str, Any]) -> bool:
    return "CANCEL" in _normalize_text(item.get("situacao"))


def _is_return(item: dict[str, Any]) -> bool:
    try:
        return int(item.get("devolucao") or 0) == 1
    except (TypeError, ValueError):
        return False


def _is_invoiced(item: dict[str, Any]) -> bool:
    return "FATURAD" in _normalize_text(item.get("situacao"))


def _report_cache_paths(arguments: dict[str, Any]) -> list[str]:
    client_id = arguments["client_id"]
    paths = _listar_bancos_vendas_tenant(client_id, arguments.get("loja"))
    return vendas_source_paths(client_id, paths)


def _read_sales_rows(
    client_id: str,
    store: str,
    start_iso: str,
    end_iso: str,
) -> tuple[list[dict[str, Any]], list[str], list[str]]:
    db_paths = _listar_bancos_vendas_tenant(client_id, store)
    rows: list[dict[str, Any]] = []
    processed: list[str] = []
    errors: list[str] = []
    select_specs = (
        ("id_unico", "''"),
        ("data", "''"),
        ("loja_conta", "''"),
        ("canal", "''"),
        ("numero", "''"),
        ("situacao", "''"),
        ("devolucao", "0"),
        ("sku", "''"),
        ("produto", "''"),
        ("quantidade", "0"),
        ("valor", "0"),
        ("comprador", "''"),
        ("unidade_negocio", "''"),
        ("nota_fiscal_id", "''"),
    )

    for source_index, db_path in enumerate(db_paths):
        if not db_path or not os.path.isfile(db_path):
            continue
        connection = None
        try:
            connection = open_vendas_readonly(db_path, row_factory=sqlite3.Row)
            exists = connection.execute(
                "SELECT 1 FROM sqlite_master WHERE type='table' AND name='vendas' LIMIT 1"
            ).fetchone()
            if not exists:
                continue
            columns = cached_table_columns(db_path, "vendas", connection)
            query = "SELECT rowid AS __rowid, " + ", ".join(
                select_compatible_column(columns, name, default)
                for name, default in select_specs
            ) + " FROM vendas"
            conditions: list[str] = []
            params: list[Any] = []
            append_indexable_date_filter(conditions, params, "data", start_iso, end_iso)
            conditions.append("trim(coalesce(loja_conta, '')) = ? COLLATE NOCASE")
            params.append(store)
            if conditions:
                query += " WHERE " + " AND ".join(conditions)
            query += " ORDER BY rowid"

            for row in connection.execute(query, params).fetchall():
                item = dict(row)
                item["__source_index"] = source_index
                rows.append(item)
            processed.append(os.path.basename(db_path))
        except Exception as exc:
            errors.append(f"{os.path.basename(db_path)}: {exc}")
        finally:
            if connection is not None:
                connection.close()
    return rows, processed, errors


def _drop_invoice_order_pairs(rows: Iterable[dict[str, Any]]) -> tuple[list[dict[str, Any]], int]:
    materialized = list(rows)
    invoiced_keys = {
        (
            _normalize_text(item.get("loja_conta")),
            str(item.get("nota_fiscal_id") or "").strip(),
            _normalize_sku(item.get("sku")),
        )
        for item in materialized
        if str(item.get("nota_fiscal_id") or "").strip() and _is_invoiced(item)
    }
    kept: list[dict[str, Any]] = []
    removed = 0
    for item in materialized:
        note_id = str(item.get("nota_fiscal_id") or "").strip()
        key = (
            _normalize_text(item.get("loja_conta")),
            note_id,
            _normalize_sku(item.get("sku")),
        )
        if note_id and key in invoiced_keys and not _is_invoiced(item):
            removed += 1
            continue
        kept.append(item)
    return kept, removed


def prepare_eligible_sales_rows(
    raw_rows: Iterable[dict[str, Any]],
) -> tuple[list[dict[str, Any]], list[dict[str, Any]], dict[str, int]]:
    """Apply the shared audited sales rules used by every Vendas report."""

    materialized = list(raw_rows)
    primary_deduplicated = _deduplicar_vendas_consolidadas(materialized)
    counters = {
        "linhas_lidas": len(materialized),
        "duplicidades_entre_bases": len(materialized) - len(primary_deduplicated),
        "devolucoes_excluidas": 0,
        "cancelamentos_excluidos": 0,
        "ebazar_excluidos": 0,
        "sku_vazio_excluidos": 0,
        "duplicidades_pedido_nota": 0,
        "linhas_elegiveis": 0,
    }

    eligible: list[dict[str, Any]] = []
    for item in primary_deduplicated:
        if _is_return(item):
            counters["devolucoes_excluidas"] += 1
            continue
        if _is_cancelled(item):
            counters["cancelamentos_excluidos"] += 1
            continue
        if _deve_excluir_venda_ebazar(
            item.get("devolucao"), item.get("comprador"), item.get("canal")
        ):
            counters["ebazar_excluidos"] += 1
            continue
        if not _normalize_sku(item.get("sku")):
            counters["sku_vazio_excluidos"] += 1
            continue
        eligible.append(item)

    eligible, invoice_duplicates = _drop_invoice_order_pairs(eligible)
    counters["duplicidades_pedido_nota"] = invoice_duplicates
    counters["linhas_elegiveis"] = len(eligible)

    fiscal_rows: list[dict[str, Any]] = []
    product_rows: list[dict[str, Any]] = []
    for item in eligible:
        if _normalize_sku(item.get("sku")) == FISCAL_ADJUSTMENT_SKU:
            fiscal_rows.append(item)
        else:
            product_rows.append(item)
    return product_rows, fiscal_rows, counters


def _representative_description(items: Iterable[dict[str, Any]]) -> str:
    stats: dict[str, tuple[int, str]] = {}
    for item in items:
        description = re.sub(r"\s+", " ", str(item.get("produto") or "")).strip()
        if not description:
            continue
        count, latest = stats.get(description, (0, ""))
        item_date = str(item.get("data") or "")[:10]
        stats[description] = (count + 1, max(latest, item_date))
    if not stats:
        return "-"
    return max(stats, key=lambda value: (stats[value][0], stats[value][1], value.casefold()))


def _source_hash(paths: Iterable[str]) -> str:
    payload = repr(vendas_sources_signature(paths)).encode("utf-8", errors="replace")
    return hashlib.sha256(payload).hexdigest()


def _recent_period_comparison(
    rows: list[dict[str, Any]], field: str, *, max_window: int = 3
) -> dict[str, Any]:
    """Compare equal recent and prior monthly windows without inventing a baseline."""
    window = min(max_window, len(rows) // 2)
    if window <= 0:
        return {"meses": 0, "anterior": 0.0, "recente": 0.0, "variacao": None}
    previous = sum(
        (_decimal(item.get(field)) for item in rows[-2 * window : -window]),
        Decimal("0"),
    )
    recent = sum(
        (_decimal(item.get(field)) for item in rows[-window:]), Decimal("0")
    )
    variation = (recent - previous) / previous if previous else None
    return {
        "meses": window,
        "anterior": _number(previous),
        "recente": _number(recent),
        "variacao": _number(variation) if variation is not None else None,
    }


def _trend_highlight(label: str, comparison: dict[str, Any]) -> str:
    months = int(comparison.get("meses") or 0)
    variation = comparison.get("variacao")
    if not months or variation is None:
        return f"{label}: não há janela anterior comparável no período solicitado."
    rate = _decimal(variation)
    if abs(rate) < Decimal("0.005"):
        movement = "ficou praticamente estável"
    elif rate > 0:
        movement = f"cresceu {_number(abs(rate) * 100):.1f}%"
    else:
        movement = f"recuou {_number(abs(rate) * 100):.1f}%"
    return (
        f"{label} {movement} nos últimos {months} meses em relação "
        f"aos {months} meses anteriores."
    )


@cache_vendas_response("pareto-80", _report_cache_paths)
def generate_pareto_report(
    client_id: str,
    loja: str,
    periodo: str = "12m",
    data_inicio: str | None = None,
    data_fim: str | None = None,
) -> dict[str, Any]:
    store = _validate_store(loja)
    period_mode, start_iso, end_iso = resolve_pareto_period(periodo, data_inicio, data_fim)
    source_paths = _listar_bancos_vendas_tenant(client_id, store)
    raw_rows, processed_bases, read_errors = _read_sales_rows(client_id, store, start_iso, end_iso)
    if source_paths and not processed_bases and read_errors:
        raise VendasDomainError(503, "Não foi possível ler as bases de vendas da conta selecionada.")
    source_latest = max((str(item.get("data") or "")[:10] for item in raw_rows), default=None)

    product_rows, fiscal_rows, counters = prepare_eligible_sales_rows(raw_rows)

    rows_by_sku: dict[str, list[dict[str, Any]]] = defaultdict(list)
    display_skus: dict[str, Counter[str]] = defaultdict(Counter)
    for item in product_rows:
        normalized = _normalize_sku(item.get("sku"))
        rows_by_sku[normalized].append(item)
        display_skus[normalized][str(item.get("sku") or "").strip()] += 1

    sku_rows: list[dict[str, Any]] = []
    operation_totals: defaultdict[str, Decimal] = defaultdict(Decimal)
    monthly_totals: defaultdict[str, defaultdict[str, Decimal]] = defaultdict(
        lambda: defaultdict(Decimal)
    )
    total_revenue = Decimal("0")

    for normalized_sku, items in rows_by_sku.items():
        revenue = sum((_decimal(item.get("valor")) for item in items), Decimal("0"))
        quantity = sum((_decimal(item.get("quantidade")) for item in items), Decimal("0"))
        full_revenue = Decimal("0")
        other_revenue = Decimal("0")
        for item in items:
            amount = _decimal(item.get("valor"))
            operation = _operation_name(item, store)
            operation_totals[operation] += amount
            month = str(item.get("data") or "")[:7]
            if month:
                monthly_totals[month][operation] += amount
            if operation == "Mercado Livre Full":
                full_revenue += amount
            else:
                other_revenue += amount
        total_revenue += revenue
        display_sku = max(
            display_skus[normalized_sku],
            key=lambda value: (display_skus[normalized_sku][value], value.casefold()),
        )
        sku_rows.append(
            {
                "sku": display_sku,
                "produto": _representative_description(items),
                "quantidade": _number(quantity),
                "registros": len(items),
                "faturamento": _number(revenue),
                "mercado_livre_full": _number(full_revenue),
                "demais_canais": _number(other_revenue),
            }
        )

    sku_rows.sort(key=lambda item: (-_decimal(item["faturamento"]), _normalize_sku(item["sku"])))
    cumulative = Decimal("0")
    pareto_finished = total_revenue <= 0
    pareto_rows: list[dict[str, Any]] = []
    for rank, item in enumerate(sku_rows, start=1):
        amount = _decimal(item["faturamento"])
        participation = amount / total_revenue if total_revenue else Decimal("0")
        cumulative += participation
        include_pareto = not pareto_finished
        item.update(
            {
                "rank": rank,
                "participacao": _number(participation),
                "acumulado": _number(cumulative),
                "pareto": include_pareto,
            }
        )
        if include_pareto:
            pareto_rows.append(dict(item))
            if cumulative >= PARETO_TARGET:
                pareto_finished = True

    pareto_revenue = sum(
        (_decimal(item["faturamento"]) for item in pareto_rows), Decimal("0")
    )

    def top_share(limit: int) -> float:
        if not total_revenue:
            return 0.0
        value = sum(
            (_decimal(item["faturamento"]) for item in sku_rows[:limit]), Decimal("0")
        )
        return _number(value / total_revenue)

    other_name = f"Demais canais da conta {store}"
    monthly = []
    for month in _month_keys(start_iso, end_iso):
        full_value = monthly_totals[month]["Mercado Livre Full"]
        other_value = monthly_totals[month][other_name]
        monthly.append(
            {
                "mes": month,
                "mercado_livre_full": _number(full_value),
                "demais_canais": _number(other_value),
                "total": _number(full_value + other_value),
            }
        )

    fiscal_total = sum((_decimal(item.get("valor")) for item in fiscal_rows), Decimal("0"))
    channels = [
        {
            "canal": "Mercado Livre Full",
            "faturamento": _number(operation_totals["Mercado Livre Full"]),
            "participacao": _number(operation_totals["Mercado Livre Full"] / total_revenue)
            if total_revenue
            else 0.0,
        },
        {
            "canal": other_name,
            "faturamento": _number(operation_totals[other_name]),
            "participacao": _number(operation_totals[other_name] / total_revenue)
            if total_revenue
            else 0.0,
        },
    ]

    revenue_comparison = _recent_period_comparison(monthly, "total")
    full_share = (
        operation_totals["Mercado Livre Full"] / total_revenue
        if total_revenue
        else Decimal("0")
    )
    pareto_mix_share = (
        Decimal(len(pareto_rows)) / Decimal(len(sku_rows)) if sku_rows else Decimal("0")
    )
    tail_count = max(0, len(sku_rows) - len(pareto_rows))
    tail_revenue = max(Decimal("0"), total_revenue - pareto_revenue)
    pareto_average = (
        pareto_revenue / Decimal(len(pareto_rows)) if pareto_rows else Decimal("0")
    )
    tail_average = tail_revenue / Decimal(tail_count) if tail_count else Decimal("0")
    highlights = [
        _trend_highlight("O faturamento", revenue_comparison),
        (
            f"{len(pareto_rows)} SKUs ({_number(pareto_mix_share * 100):.1f}% do mix vendido) "
            f"concentram {_number((pareto_revenue / total_revenue if total_revenue else Decimal('0')) * 100):.1f}% "
            "do faturamento."
        ),
        (
            f"Mercado Livre Full responde por {_number(full_share * 100):.1f}% do faturamento; "
            f"os demais canais respondem por {_number((Decimal('1') - full_share) * 100):.1f}%."
        ),
        (
            f"A cauda fora do Pareto reúne {tail_count} SKUs e "
            f"{_number((tail_revenue / total_revenue if total_revenue else Decimal('0')) * 100):.1f}% do faturamento."
        ),
    ]

    return {
        "success": True,
        "loja": store,
        "periodo": {
            "tipo": period_mode,
            "data_inicio": start_iso,
            "data_fim": end_iso,
            "ultima_data_disponivel": source_latest,
            "gerado_em": datetime.now().astimezone().isoformat(timespec="seconds"),
        },
        "resumo": {
            "faturamento_produtos": _number(total_revenue),
            "total_skus": len(sku_rows),
            "pareto_skus": len(pareto_rows),
            "pareto_faturamento": _number(pareto_revenue),
            "pareto_participacao": _number(pareto_revenue / total_revenue)
            if total_revenue
            else 0.0,
            "sku_corte": pareto_rows[-1]["sku"] if pareto_rows else None,
            "top_1_participacao": top_share(1),
            "top_5_participacao": top_share(5),
            "top_10_participacao": top_share(10),
        },
        "pareto_skus": pareto_rows,
        "todos_skus": sku_rows,
        "mensal": monthly,
        "canais": channels,
        "analise_executiva": {
            "destaques": highlights,
            "comparacao_faturamento": revenue_comparison,
            "participacao_full": _number(full_share),
            "participacao_skus_pareto": _number(pareto_mix_share),
            "cauda_skus": tail_count,
            "cauda_faturamento": _number(tail_revenue),
            "cauda_participacao": _number(tail_revenue / total_revenue)
            if total_revenue
            else 0.0,
            "faturamento_medio_sku_pareto": _number(pareto_average),
            "faturamento_medio_sku_cauda": _number(tail_average),
        },
        "ajuste_fiscal": {
            "sku": "ESTORNO DE CRÉDITO ICMS",
            "linhas": len(fiscal_rows),
            "valor": _number(fiscal_total),
            "incluido_no_pareto": False,
        },
        "auditoria": {
            **counters,
            "bases_processadas": len(processed_bases),
            "nomes_bases": processed_bases,
            "erros_leitura": read_errors,
            "assinatura_fontes": _source_hash(source_paths),
        },
    }


def _money(value: Any) -> str:
    amount = _decimal(value).quantize(Decimal("0.01"))
    raw = f"{amount:,.2f}"
    return "R$ " + raw.replace(",", "_").replace(".", ",").replace("_", ".")


def _percentage(value: Any) -> str:
    return f"{_decimal(value) * 100:.2f}%".replace(".", ",")


def _safe_filename(value: str) -> str:
    normalized = unicodedata.normalize("NFKD", value)
    normalized = "".join(char for char in normalized if not unicodedata.combining(char))
    normalized = re.sub(r"[^A-Za-z0-9]+", "_", normalized).strip("_").lower()
    return normalized or "loja"


def build_pareto_xlsx(report: dict[str, Any]) -> bytes:
    from openpyxl import Workbook
    from openpyxl.chart import BarChart, LineChart, Reference
    from openpyxl.chart.series import SeriesLabel
    from openpyxl.formatting.rule import ColorScaleRule, DataBarRule
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.worksheet.table import Table, TableStyleInfo

    workbook = Workbook()
    summary_sheet = workbook.active
    summary_sheet.title = "Resumo"
    pareto_sheet = workbook.create_sheet("Pareto 80%")
    all_sheet = workbook.create_sheet("Todos os SKUs")
    monthly_sheet = workbook.create_sheet("Mensal e Canais")
    rules_sheet = workbook.create_sheet("Fontes e Regras")

    navy = "17365D"
    blue = "2F75B5"
    teal = "00A6A6"
    green = "70AD47"
    orange = "ED7D31"
    light_blue = "D9EAF7"
    light_green = "E2F0D9"
    light_gray = "F2F4F7"
    white = "FFFFFF"
    muted = "5B6573"
    thin = Side(style="thin", color="D5DCE5")

    def title(sheet, text: str, end_column: int) -> None:
        sheet.merge_cells(start_row=1, start_column=1, end_row=1, end_column=end_column)
        cell = sheet.cell(1, 1, text)
        cell.font = Font(name="Aptos Display", size=18, bold=True, color=white)
        cell.fill = PatternFill("solid", fgColor=navy)
        cell.alignment = Alignment(horizontal="left", vertical="center")
        sheet.row_dimensions[1].height = 30
        sheet.sheet_view.showGridLines = False

    def header(row) -> None:
        for cell in row:
            cell.font = Font(name="Aptos", size=10, bold=True, color=white)
            cell.fill = PatternFill("solid", fgColor=blue)
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            cell.border = Border(bottom=thin)

    def add_table(sheet, ref: str, name: str) -> None:
        table = Table(displayName=name, ref=ref)
        table.tableStyleInfo = TableStyleInfo(
            name="TableStyleMedium2",
            showFirstColumn=False,
            showLastColumn=False,
            showRowStripes=True,
            showColumnStripes=False,
        )
        sheet.add_table(table)

    summary = report["resumo"]
    period = report["periodo"]
    audit = report["auditoria"]
    title(summary_sheet, "Relatório Pareto 80% do Faturamento", 8)
    summary_sheet["A3"] = "Conta"
    summary_sheet["B3"] = report["loja"]
    summary_sheet["A4"] = "Período solicitado"
    summary_sheet["B4"] = f"{period['data_inicio']} a {period['data_fim']}"
    summary_sheet["A5"] = "Última data disponível"
    summary_sheet["B5"] = period.get("ultima_data_disponivel") or "Sem vendas no período"
    summary_sheet["A6"] = "Gerado em"
    summary_sheet["B6"] = period["gerado_em"]
    for row_number in range(3, 7):
        summary_sheet.cell(row_number, 1).font = Font(bold=True, color=muted)

    kpis = [
        ("Faturamento de produtos", summary["faturamento_produtos"], "currency"),
        ("SKUs vendidos", summary["total_skus"], "integer"),
        ("SKUs no Pareto", summary["pareto_skus"], "integer"),
        ("Faturamento do Pareto", summary["pareto_faturamento"], "currency"),
        ("Participação alcançada", summary["pareto_participacao"], "percent"),
        ("SKU de corte", summary.get("sku_corte") or "-", "text"),
        ("Concentração Top 5", summary["top_5_participacao"], "percent"),
        ("Ajuste fiscal segregado", report["ajuste_fiscal"]["valor"], "currency"),
    ]
    for index, (label, value, kind) in enumerate(kpis):
        row = 9 + (index // 2) * 3
        column = 1 + (index % 2) * 3
        summary_sheet.merge_cells(start_row=row, start_column=column, end_row=row, end_column=column + 1)
        summary_sheet.merge_cells(start_row=row + 1, start_column=column, end_row=row + 1, end_column=column + 1)
        label_cell = summary_sheet.cell(row, column, label)
        value_cell = summary_sheet.cell(row + 1, column, value)
        label_cell.font = Font(bold=True, color=muted)
        value_cell.font = Font(size=15, bold=True, color=navy)
        for target in (label_cell, value_cell):
            target.fill = PatternFill("solid", fgColor=light_blue if index % 2 == 0 else light_green)
            target.alignment = Alignment(vertical="center")
        if kind == "currency":
            value_cell.number_format = '"R$" #,##0.00'
        elif kind == "percent":
            value_cell.number_format = "0.00%"
        elif kind == "integer":
            value_cell.number_format = "#,##0"
        elif kind == "text":
            value_cell.number_format = "@"

    summary_sheet["A22"] = "Contribuição por operação"
    summary_sheet["A22"].font = Font(bold=True, color=navy, size=12)
    summary_sheet.append([])
    summary_sheet["A23"] = "Operação"
    summary_sheet["B23"] = "Faturamento"
    summary_sheet["C23"] = "Participação"
    header(summary_sheet[23])
    for channel in report["canais"]:
        summary_sheet.append(
            [channel["canal"], channel["faturamento"], channel["participacao"]]
        )
    for row in range(24, 24 + len(report["canais"])):
        summary_sheet.cell(row, 2).number_format = '"R$" #,##0.00'
        summary_sheet.cell(row, 3).number_format = "0.00%"

    analysis = report.get("analise_executiva") or {}
    summary_sheet["A28"] = "Leitura executiva"
    summary_sheet["A28"].font = Font(bold=True, color=navy, size=12)
    for row_index, text in enumerate(analysis.get("destaques") or [], start=29):
        summary_sheet.merge_cells(
            start_row=row_index, start_column=1, end_row=row_index, end_column=8
        )
        cell = summary_sheet.cell(row_index, 1, f"• {text}")
        cell.fill = PatternFill("solid", fgColor=light_gray)
        cell.font = Font(color="303A46")
        cell.alignment = Alignment(vertical="center", wrap_text=True)
        summary_sheet.row_dimensions[row_index].height = 24

    columns = [
        "Rank",
        "SKU",
        "Produto",
        "Quantidade",
        "Registros",
        "Faturamento",
        "Participação",
        "Acumulado",
        "Pareto",
        "Mercado Livre Full",
        "Demais canais",
        "Meta 80%",
    ]

    def populate_sku_sheet(sheet, sheet_title: str, rows: list[dict[str, Any]], table_name: str) -> None:
        title(sheet, sheet_title, len(columns))
        sheet.append(columns)
        header(sheet[2])
        for item in rows:
            sheet.append(
                [
                    item["rank"],
                    str(item["sku"]),
                    item["produto"],
                    item["quantidade"],
                    item["registros"],
                    item["faturamento"],
                    item["participacao"],
                    item["acumulado"],
                    "Sim" if item["pareto"] else "Não",
                    item["mercado_livre_full"],
                    item["demais_canais"],
                    0.8,
                ]
            )
        last_row = max(2, sheet.max_row)
        if rows:
            add_table(sheet, f"A2:L{last_row}", table_name)
            sheet.conditional_formatting.add(
                f"F3:F{last_row}",
                DataBarRule(start_type="min", end_type="max", color=blue),
            )
            sheet.conditional_formatting.add(
                f"H3:H{last_row}",
                ColorScaleRule(
                    start_type="min", start_color=light_green,
                    mid_type="percentile", mid_value=80, mid_color="FFF2CC",
                    end_type="max", end_color="F4CCCC",
                ),
            )
        for row in range(3, last_row + 1):
            sheet.cell(row, 2).number_format = "@"
            sheet.cell(row, 4).number_format = "#,##0.00"
            sheet.cell(row, 5).number_format = "#,##0"
            for column in (6, 10, 11):
                sheet.cell(row, column).number_format = '"R$" #,##0.00'
            for column in (7, 8):
                sheet.cell(row, column).number_format = "0.00%"
            sheet.cell(row, 12).number_format = "0.00%"
        sheet.freeze_panes = "A3"
        sheet.auto_filter.ref = f"A2:L{last_row}"
        widths = {"A": 9, "B": 18, "C": 55, "D": 14, "E": 12, "F": 18, "G": 14, "H": 14, "I": 11, "J": 20, "K": 20, "L": 12}
        for column, width in widths.items():
            sheet.column_dimensions[column].width = width
        sheet.column_dimensions["C"].width = 55

    populate_sku_sheet(pareto_sheet, "SKUs responsáveis por 80% do faturamento", report["pareto_skus"], "Pareto80Table")
    populate_sku_sheet(all_sheet, "Todos os SKUs elegíveis", report["todos_skus"], "TodosSkusTable")

    if report["pareto_skus"]:
        chart_limit = min(len(report["pareto_skus"]), 15) + 2
        bar = BarChart()
        bar.type = "bar"
        bar.style = 10
        bar.title = "Top 15 SKUs do Pareto por faturamento"
        bar.y_axis.title = "SKU"
        bar.x_axis.title = "Faturamento (R$)"
        bar.height = 10
        bar.width = 18
        bar.add_data(Reference(pareto_sheet, min_col=6, min_row=2, max_row=chart_limit), titles_from_data=True)
        bar.set_categories(Reference(pareto_sheet, min_col=2, min_row=3, max_row=chart_limit))
        bar.series[0].tx = SeriesLabel(v="Faturamento")
        bar.series[0].graphicalProperties.solidFill = blue
        bar.legend = None
        pareto_sheet.add_chart(bar, "M2")

        line = LineChart()
        line.style = 13
        line.title = f"Curva Pareto - 80% atingidos no rank {len(report['pareto_skus'])}"
        line.y_axis.title = "Participação acumulada"
        line.x_axis.title = "Rank"
        line.y_axis.numFmt = "0%"
        line.height = 9
        line.width = 18
        last_pareto = len(report["pareto_skus"]) + 2
        line.add_data(Reference(pareto_sheet, min_col=8, min_row=2, max_row=last_pareto), titles_from_data=True)
        line.add_data(Reference(pareto_sheet, min_col=12, min_row=2, max_row=last_pareto), titles_from_data=True)
        line.set_categories(Reference(pareto_sheet, min_col=1, min_row=3, max_row=last_pareto))
        line.series[0].tx = SeriesLabel(v="Acumulado")
        line.series[0].graphicalProperties.line.solidFill = orange
        line.series[1].tx = SeriesLabel(v="Meta 80%")
        line.series[1].graphicalProperties.line.solidFill = green
        pareto_sheet.add_chart(line, "M22")

    title(monthly_sheet, "Faturamento mensal por operação", 4)
    monthly_sheet.append(["Mês", "Mercado Livre Full", "Demais canais", "Total"])
    header(monthly_sheet[2])
    for index, item in enumerate(report["mensal"], start=3):
        monthly_sheet.cell(index, 1, item["mes"])
        monthly_sheet.cell(index, 2, item["mercado_livre_full"])
        monthly_sheet.cell(index, 3, item["demais_canais"])
        monthly_sheet.cell(index, 4, f"=SUM(B{index}:C{index})")
        for column in range(2, 5):
            monthly_sheet.cell(index, column).number_format = '"R$" #,##0.00'
    if report["mensal"]:
        monthly_last = len(report["mensal"]) + 2
        add_table(monthly_sheet, f"A2:D{monthly_last}", "MensalCanaisTable")
        chart = BarChart()
        chart.type = "col"
        chart.style = 10
        chart.grouping = "stacked"
        chart.overlap = 100
        revenue_variation = (analysis.get("comparacao_faturamento") or {}).get("variacao")
        variation_suffix = (
            f" - recente {'+' if revenue_variation >= 0 else ''}{revenue_variation:.1%}"
            if revenue_variation is not None
            else ""
        )
        chart.title = f"Faturamento mensal por canal{variation_suffix}"
        chart.y_axis.title = "Faturamento (R$)"
        chart.y_axis.numFmt = '"R$" #,##0'
        chart.x_axis.title = "Mês"
        chart.height = 10
        chart.width = 20
        chart.add_data(Reference(monthly_sheet, min_col=2, max_col=3, min_row=2, max_row=monthly_last), titles_from_data=True)
        chart.set_categories(Reference(monthly_sheet, min_col=1, min_row=3, max_row=monthly_last))
        chart.series[0].tx = SeriesLabel(v="Mercado Livre Full")
        chart.series[0].graphicalProperties.solidFill = teal
        chart.series[1].tx = SeriesLabel(v="Demais canais")
        chart.series[1].graphicalProperties.solidFill = blue
        monthly_sheet.add_chart(chart, "F2")
    monthly_sheet.freeze_panes = "A3"
    for column, width in {"A": 14, "B": 22, "C": 22, "D": 20}.items():
        monthly_sheet.column_dimensions[column].width = width

    title(rules_sheet, "Fontes, regras e controles de auditoria", 4)
    rules = [
        ("Conta", report["loja"]),
        ("Período solicitado", f"{period['data_inicio']} a {period['data_fim']}"),
        ("Última data disponível", period.get("ultima_data_disponivel") or "Sem vendas"),
        ("Assinatura das fontes", audit["assinatura_fontes"]),
        ("Bases processadas", audit["bases_processadas"]),
        ("Linhas lidas", audit["linhas_lidas"]),
        ("Duplicidades entre bases", audit["duplicidades_entre_bases"]),
        ("Duplicidades pedido/nota", audit["duplicidades_pedido_nota"]),
        ("Devoluções excluídas", audit["devolucoes_excluidas"]),
        ("Cancelamentos excluídos", audit["cancelamentos_excluidos"]),
        ("Transferências EBAZAR excluídas", audit["ebazar_excluidos"]),
        ("SKUs vazios excluídos", audit["sku_vazio_excluidos"]),
        ("Linhas elegíveis", audit["linhas_elegiveis"]),
        ("Ajuste fiscal", f"{report['ajuste_fiscal']['linhas']} linha(s); {_money(report['ajuste_fiscal']['valor'])}"),
    ]
    rules_sheet.append(["Controle", "Valor"])
    header(rules_sheet[2])
    for label, value in rules:
        rules_sheet.append([label, value])
    start_rules = rules_sheet.max_row + 2
    rules_sheet.cell(start_rules, 1, "Regras aplicadas")
    rules_sheet.cell(start_rules, 1).font = Font(bold=True, color=navy, size=12)
    business_rules = [
        "Somente a conta selecionada; o relatório não aceita Todas as lojas.",
        "Inclui Mercado Livre Full e demais canais pertencentes à conta.",
        "Exclui devoluções, cancelamentos, EBAZAR, SKU vazio e o ajuste fiscal do Pareto.",
        "Deduplica por ID/fallback e remove o pedido não faturado quando há linha Faturado para conta + NF + SKU.",
        "Ordena por faturamento e inclui o primeiro SKU que cruza 80%.",
    ]
    for offset, rule in enumerate(business_rules, start=1):
        rules_sheet.cell(start_rules + offset, 1, f"{offset}. {rule}")
        rules_sheet.merge_cells(start_row=start_rules + offset, start_column=1, end_row=start_rules + offset, end_column=4)
        rules_sheet.cell(start_rules + offset, 1).alignment = Alignment(wrap_text=True, vertical="top")
    qc_row = start_rules + len(business_rules) + 3
    all_last = max(3, len(report["todos_skus"]) + 2)
    rules_sheet.cell(qc_row, 1, "Controle por fórmula")
    rules_sheet.cell(qc_row, 2, "Resultado")
    header(rules_sheet[qc_row])
    rules_sheet.cell(qc_row + 1, 1, "Faturamento total da aba Todos os SKUs")
    rules_sheet.cell(qc_row + 1, 2, f"=SUM('Todos os SKUs'!F3:F{all_last})")
    rules_sheet.cell(qc_row + 2, 1, "Faturamento marcado como Pareto")
    rules_sheet.cell(qc_row + 2, 2, f'=SUMIF(\'Todos os SKUs\'!I3:I{all_last},"Sim",\'Todos os SKUs\'!F3:F{all_last})')
    rules_sheet.cell(qc_row + 3, 1, "Participação do Pareto")
    rules_sheet.cell(qc_row + 3, 2, f"=IFERROR(B{qc_row + 2}/B{qc_row + 1},0)")
    rules_sheet.cell(qc_row + 1, 2).number_format = '"R$" #,##0.00'
    rules_sheet.cell(qc_row + 2, 2).number_format = '"R$" #,##0.00'
    rules_sheet.cell(qc_row + 3, 2).number_format = "0.00%"
    rules_sheet.column_dimensions["A"].width = 54
    rules_sheet.column_dimensions["B"].width = 70
    rules_sheet.column_dimensions["C"].width = 18
    rules_sheet.column_dimensions["D"].width = 18
    rules_sheet.freeze_panes = "A3"

    summary_sheet.column_dimensions["A"].width = 28
    summary_sheet.column_dimensions["B"].width = 26
    summary_sheet.column_dimensions["C"].width = 16
    summary_sheet.column_dimensions["D"].width = 28
    summary_sheet.column_dimensions["E"].width = 26
    summary_sheet.column_dimensions["F"].width = 4
    summary_sheet.column_dimensions["G"].width = 16
    summary_sheet.column_dimensions["H"].width = 16
    summary_sheet.freeze_panes = "A3"

    for sheet in workbook.worksheets:
        for row in sheet.iter_rows():
            for cell in row:
                if cell.value is not None and cell.row != 1:
                    font = copy(cell.font)
                    font.name = "Aptos"
                    cell.font = font
                    alignment = copy(cell.alignment)
                    alignment.vertical = "center"
                    cell.alignment = alignment
        sheet.page_setup.orientation = "landscape"
        sheet.page_setup.fitToWidth = 1
        sheet.sheet_properties.pageSetUpPr.fitToPage = True

    workbook.calculation.fullCalcOnLoad = True
    workbook.calculation.forceFullCalc = True
    workbook.calculation.calcMode = "auto"
    output = io.BytesIO()
    workbook.save(output)
    return output.getvalue()


def _pdf_top_chart(report: dict[str, Any]):
    from reportlab.graphics.charts.barcharts import HorizontalBarChart
    from reportlab.graphics.shapes import Drawing, String
    from reportlab.lib import colors

    rows = report["pareto_skus"][:15]
    drawing = Drawing(740, 255)
    drawing.add(String(12, 235, "Top 15 SKUs do Pareto por faturamento (R$)", fontName="Helvetica-Bold", fontSize=12, fillColor=colors.HexColor("#17365D")))
    if not rows:
        drawing.add(String(20, 120, "Sem vendas elegíveis no período.", fontSize=11))
        return drawing
    chart = HorizontalBarChart()
    chart.x = 115
    chart.y = 20
    chart.width = 600
    chart.height = 200
    chart.data = [[float(item["faturamento"]) for item in reversed(rows)]]
    chart.categoryAxis.categoryNames = [str(item["sku"])[:22] for item in reversed(rows)]
    chart.valueAxis.valueMin = 0
    chart.valueAxis.labelTextFormat = lambda value: f"{value / 1000:.0f} mil" if value >= 1000 else f"{value:.0f}"
    chart.bars[0].fillColor = colors.HexColor("#2F75B5")
    chart.barWidth = 7
    chart.categoryAxis.labels.fontSize = 7
    chart.valueAxis.labels.fontSize = 7
    drawing.add(chart)
    return drawing


def _pdf_monthly_chart(report: dict[str, Any]):
    from reportlab.graphics.charts.barcharts import VerticalBarChart
    from reportlab.graphics.shapes import Drawing, Rect, String
    from reportlab.lib import colors

    rows = report["mensal"]
    comparison = (report.get("analise_executiva") or {}).get("comparacao_faturamento") or {}
    variation = comparison.get("variacao")
    suffix = f" ({'+' if variation >= 0 else ''}{variation * 100:.1f}% recente)" if variation is not None else ""
    drawing = Drawing(360, 180)
    drawing.add(String(10, 164, f"Faturamento mensal por canal{suffix}", fontName="Helvetica-Bold", fontSize=10, fillColor=colors.HexColor("#17365D")))
    if not rows:
        drawing.add(String(20, 120, "Sem meses no período.", fontSize=11))
        return drawing
    chart = VerticalBarChart()
    chart.x = 38
    chart.y = 30
    chart.width = 300
    chart.height = 112
    chart.data = [
        [float(item["mercado_livre_full"]) for item in rows],
        [float(item["demais_canais"]) for item in rows],
    ]
    chart.categoryAxis.categoryNames = [item["mes"] for item in rows]
    chart.categoryAxis.labels.angle = 30
    chart.categoryAxis.labels.dy = -12
    chart.categoryAxis.labels.fontSize = 7
    chart.valueAxis.labels.fontSize = 7
    chart.valueAxis.valueMin = 0
    chart.valueAxis.labelTextFormat = lambda value: f"{value / 1000:.0f} mil" if value >= 1000 else f"{value:.0f}"
    chart.bars[0].fillColor = colors.HexColor("#00A6A6")
    chart.bars[1].fillColor = colors.HexColor("#2F75B5")
    chart.categoryAxis.style = "stacked"
    drawing.add(chart)
    drawing.add(Rect(190, 147, 6, 6, fillColor=colors.HexColor("#00A6A6"), strokeColor=None))
    drawing.add(String(199, 147, "Mercado Livre Full", fontSize=6.5, fillColor=colors.HexColor("#17365D")))
    drawing.add(Rect(286, 147, 6, 6, fillColor=colors.HexColor("#2F75B5"), strokeColor=None))
    drawing.add(String(295, 147, "Demais", fontSize=6.5, fillColor=colors.HexColor("#17365D")))
    return drawing


def _pdf_pareto_curve_chart(report: dict[str, Any]):
    from reportlab.graphics.charts.linecharts import HorizontalLineChart
    from reportlab.graphics.shapes import Drawing, String
    from reportlab.lib import colors

    rows = report["pareto_skus"]
    drawing = Drawing(360, 180)
    drawing.add(
        String(
            10,
            164,
            f"Curva Pareto - corte no rank {len(rows)}",
            fontName="Helvetica-Bold",
            fontSize=10,
            fillColor=colors.HexColor("#17365D"),
        )
    )
    if not rows:
        drawing.add(String(20, 85, "Sem vendas elegíveis no período.", fontSize=9))
        return drawing
    sample_count = min(10, len(rows))
    indexes = sorted(
        {
            round(position * (len(rows) - 1) / max(1, sample_count - 1))
            for position in range(sample_count)
        }
    )
    sampled = [rows[index] for index in indexes]
    chart = HorizontalLineChart()
    chart.x = 42
    chart.y = 30
    chart.width = 295
    chart.height = 112
    chart.data = [
        [float(item["acumulado"]) for item in sampled],
        [0.8 for _item in sampled],
    ]
    chart.categoryAxis.categoryNames = [str(item["rank"]) for item in sampled]
    chart.categoryAxis.labels.fontSize = 6.5
    chart.valueAxis.labels.fontSize = 6.5
    chart.valueAxis.valueMin = 0
    chart.valueAxis.valueMax = 1
    chart.valueAxis.labelTextFormat = lambda value: f"{value * 100:.0f}%"
    chart.lines[0].strokeColor = colors.HexColor("#ED7D31")
    chart.lines[1].strokeColor = colors.HexColor("#70AD47")
    drawing.add(chart)
    drawing.add(String(85, 6, "Acumulado", fontSize=7, fillColor=colors.HexColor("#ED7D31")))
    drawing.add(String(170, 6, "Meta 80%", fontSize=7, fillColor=colors.HexColor("#70AD47")))
    return drawing


def build_pareto_pdf(report: dict[str, Any]) -> bytes:
    from reportlab.lib import colors
    from reportlab.lib.enums import TA_CENTER, TA_LEFT
    from reportlab.lib.pagesizes import A4, landscape
    from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
    from reportlab.lib.units import mm
    from reportlab.platypus import (
        KeepTogether,
        PageBreak,
        Paragraph,
        SimpleDocTemplate,
        Spacer,
        Table,
        TableStyle,
    )

    output = io.BytesIO()
    page_size = landscape(A4)
    document = SimpleDocTemplate(
        output,
        pagesize=page_size,
        rightMargin=14 * mm,
        leftMargin=14 * mm,
        topMargin=18 * mm,
        bottomMargin=15 * mm,
        title="Relatório Pareto 80% do Faturamento",
        author="JK Sistema",
    )
    styles = getSampleStyleSheet()
    title_style = ParagraphStyle(
        "ParetoTitle",
        parent=styles["Title"],
        fontName="Helvetica-Bold",
        fontSize=20,
        leading=24,
        textColor=colors.HexColor("#17365D"),
        alignment=TA_LEFT,
        spaceAfter=8,
    )
    heading = ParagraphStyle(
        "ParetoHeading",
        parent=styles["Heading2"],
        fontName="Helvetica-Bold",
        fontSize=13,
        leading=16,
        textColor=colors.HexColor("#17365D"),
        spaceBefore=6,
        spaceAfter=7,
    )
    body = ParagraphStyle(
        "ParetoBody",
        parent=styles["BodyText"],
        fontName="Helvetica",
        fontSize=9,
        leading=12,
        textColor=colors.HexColor("#303A46"),
    )
    small = ParagraphStyle(
        "ParetoSmall",
        parent=body,
        fontSize=7.5,
        leading=9,
    )
    centered = ParagraphStyle("ParetoCentered", parent=small, alignment=TA_CENTER)
    table_header = ParagraphStyle(
        "ParetoTableHeader",
        parent=centered,
        fontName="Helvetica-Bold",
        textColor=colors.white,
    )

    def on_page(canvas, doc) -> None:
        canvas.saveState()
        width, height = page_size
        canvas.setStrokeColor(colors.HexColor("#D5DCE5"))
        canvas.line(14 * mm, height - 12 * mm, width - 14 * mm, height - 12 * mm)
        canvas.setFont("Helvetica", 7.5)
        canvas.setFillColor(colors.HexColor("#5B6573"))
        canvas.drawString(14 * mm, height - 9 * mm, f"JK Sistema - Pareto 80% - {report['loja']}")
        canvas.drawRightString(width - 14 * mm, 8 * mm, f"Página {doc.page}")
        canvas.drawString(14 * mm, 8 * mm, f"Gerado em {report['periodo']['gerado_em']}")
        canvas.restoreState()

    summary = report["resumo"]
    period = report["periodo"]
    story = [
        Paragraph("Relatório Pareto 80% do Faturamento", title_style),
        Paragraph(
            f"<b>Conta:</b> {xml_escape(str(report['loja']))} &nbsp;&nbsp; <b>Período:</b> {period['data_inicio']} a {period['data_fim']} &nbsp;&nbsp; <b>Dados disponíveis até:</b> {xml_escape(str(period.get('ultima_data_disponivel') or 'sem vendas'))}",
            body,
        ),
        Spacer(1, 8),
    ]

    kpi_data = [
        [
            Paragraph("Faturamento de produtos", small),
            Paragraph("SKUs vendidos", small),
            Paragraph("SKUs no Pareto", small),
            Paragraph("Participação alcançada", small),
            Paragraph("SKU de corte", small),
        ],
        [
            Paragraph(f"<b>{_money(summary['faturamento_produtos'])}</b>", centered),
            Paragraph(f"<b>{summary['total_skus']}</b>", centered),
            Paragraph(f"<b>{summary['pareto_skus']}</b>", centered),
            Paragraph(f"<b>{_percentage(summary['pareto_participacao'])}</b>", centered),
            Paragraph(f"<b>{xml_escape(str(summary.get('sku_corte') or '-'))}</b>", centered),
        ],
    ]
    kpi_table = Table(kpi_data, colWidths=[47 * mm] * 5, rowHeights=[8 * mm, 11 * mm])
    kpi_table.setStyle(
        TableStyle(
            [
                ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#D9EAF7")),
                ("BACKGROUND", (0, 1), (-1, 1), colors.HexColor("#F7FAFC")),
                ("BOX", (0, 0), (-1, -1), 0.5, colors.HexColor("#B8C7D9")),
                ("INNERGRID", (0, 0), (-1, -1), 0.25, colors.HexColor("#D5DCE5")),
                ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
                ("ALIGN", (0, 0), (-1, -1), "CENTER"),
            ]
        )
    )
    insight_rows = [
        [Paragraph(f"• {xml_escape(str(text))}", body)]
        for text in (report.get("analise_executiva") or {}).get("destaques", [])
    ]
    if not insight_rows:
        insight_rows = [[Paragraph("Sem destaques comparativos para o período.", body)]]
    insight_table = Table(insight_rows, colWidths=[235 * mm])
    insight_table.setStyle(
        TableStyle(
            [
                ("BACKGROUND", (0, 0), (-1, -1), colors.HexColor("#F2F6FA")),
                ("BOX", (0, 0), (-1, -1), 0.4, colors.HexColor("#D5DCE5")),
                ("INNERGRID", (0, 0), (-1, -1), 0.2, colors.HexColor("#E4E9F0")),
                ("LEFTPADDING", (0, 0), (-1, -1), 7),
                ("RIGHTPADDING", (0, 0), (-1, -1), 7),
                ("TOPPADDING", (0, 0), (-1, -1), 4),
                ("BOTTOMPADDING", (0, 0), (-1, -1), 4),
            ]
        )
    )
    story.extend(
        [
            kpi_table,
            Spacer(1, 5),
            Paragraph("Leitura executiva", heading),
            insight_table,
            Spacer(1, 5),
            Table(
                [[_pdf_pareto_curve_chart(report), _pdf_monthly_chart(report)]],
                colWidths=[120 * mm, 120 * mm],
            ),
            PageBreak(),
            _pdf_top_chart(report),
            Spacer(1, 4),
        ]
    )

    story.append(Paragraph("SKUs responsáveis pelo Pareto 80%", heading))
    table_data = [[
        Paragraph("Rank", table_header),
        Paragraph("SKU", table_header),
        Paragraph("Produto", table_header),
        Paragraph("Faturamento", table_header),
        Paragraph("Part.", table_header),
        Paragraph("Acum.", table_header),
        Paragraph("Full", table_header),
        Paragraph("Demais canais", table_header),
    ]]
    for item in report["pareto_skus"]:
        table_data.append(
            [
                item["rank"],
                Paragraph(xml_escape(str(item["sku"])), small),
                Paragraph(xml_escape(str(item["produto"])), small),
                _money(item["faturamento"]),
                _percentage(item["participacao"]),
                _percentage(item["acumulado"]),
                _money(item["mercado_livre_full"]),
                _money(item["demais_canais"]),
            ]
        )
    pareto_table = Table(
        table_data,
        repeatRows=1,
        colWidths=[11 * mm, 22 * mm, 72 * mm, 28 * mm, 18 * mm, 18 * mm, 28 * mm, 31 * mm],
    )
    pareto_table.setStyle(
        TableStyle(
            [
                ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#17365D")),
                ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
                ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
                ("FONTSIZE", (0, 0), (-1, -1), 7.2),
                ("GRID", (0, 0), (-1, -1), 0.25, colors.HexColor("#D5DCE5")),
                ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#F2F6FA")]),
                ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
                ("ALIGN", (0, 1), (0, -1), "CENTER"),
                ("ALIGN", (3, 1), (-1, -1), "RIGHT"),
                ("TOPPADDING", (0, 0), (-1, -1), 3),
                ("BOTTOMPADDING", (0, 0), (-1, -1), 3),
            ]
        )
    )
    story.extend([pareto_table, PageBreak(), Paragraph("Metodologia e auditoria", heading)])

    audit = report["auditoria"]
    audit_rows = [
        ["Controle", "Resultado"],
        ["Linhas lidas", audit["linhas_lidas"]],
        ["Duplicidades entre bases", audit["duplicidades_entre_bases"]],
        ["Duplicidades pedido/nota", audit["duplicidades_pedido_nota"]],
        ["Devoluções excluídas", audit["devolucoes_excluidas"]],
        ["Cancelamentos excluídos", audit["cancelamentos_excluidos"]],
        ["Transferências EBAZAR excluídas", audit["ebazar_excluidos"]],
        ["SKUs vazios excluídos", audit["sku_vazio_excluidos"]],
        ["Linhas elegíveis", audit["linhas_elegiveis"]],
        ["Ajuste fiscal segregado", f"{report['ajuste_fiscal']['linhas']} linha(s) - {_money(report['ajuste_fiscal']['valor'])}"],
        ["Assinatura das fontes", audit["assinatura_fontes"]],
    ]
    audit_table = Table(audit_rows, colWidths=[70 * mm, 150 * mm], repeatRows=1)
    audit_table.setStyle(
        TableStyle(
            [
                ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#17365D")),
                ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
                ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
                ("GRID", (0, 0), (-1, -1), 0.3, colors.HexColor("#D5DCE5")),
                ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#F2F6FA")]),
                ("FONTNAME", (0, 1), (0, -1), "Helvetica-Bold"),
                ("FONTSIZE", (0, 0), (-1, -1), 8.5),
                ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
                ("TOPPADDING", (0, 0), (-1, -1), 5),
                ("BOTTOMPADDING", (0, 0), (-1, -1), 5),
            ]
        )
    )
    story.extend(
        [
            audit_table,
            Spacer(1, 12),
            KeepTogether(
                [
                    Paragraph("Regras aplicadas", heading),
                    Paragraph(
                        "1. A análise usa apenas a conta selecionada e reúne Mercado Livre Full e os demais canais da mesma conta.<br/>"
                        "2. Devoluções, cancelamentos, transferências EBAZAR e linhas sem SKU são excluídos.<br/>"
                        "3. O ajuste ESTORNO DE CRÉDITO ICMS é exibido para auditoria, mas não participa do faturamento de produtos.<br/>"
                        "4. Duplicidades são removidas por identificador/fallback e pela relação pedido não faturado versus linha Faturado da mesma conta, nota e SKU.<br/>"
                        "5. O Pareto inclui o primeiro SKU que faz o acumulado atingir ou ultrapassar 80%.",
                        body,
                    ),
                ]
            ),
        ]
    )
    if audit.get("erros_leitura"):
        story.extend(
            [
                Spacer(1, 10),
                Paragraph("Avisos de leitura", heading),
                Paragraph("<br/>".join(xml_escape(str(item)) for item in audit["erros_leitura"]), body),
            ]
        )

    document.build(story, onFirstPage=on_page, onLaterPages=on_page)
    return output.getvalue()


def export_pareto_report(
    client_id: str,
    loja: str,
    formato: str,
    periodo: str = "12m",
    data_inicio: str | None = None,
    data_fim: str | None = None,
) -> tuple[bytes, str, str]:
    file_format = str(formato or "").strip().lower()
    if file_format not in {"xlsx", "pdf"}:
        raise VendasDomainError(400, "Formato inválido. Use xlsx ou pdf.")
    report = generate_pareto_report(
        client_id=client_id,
        loja=loja,
        periodo=periodo,
        data_inicio=data_inicio,
        data_fim=data_fim,
    )
    suffix = f"{report['periodo']['data_inicio']}_{report['periodo']['data_fim']}"
    filename = f"relatorio_pareto_80_{_safe_filename(report['loja'])}_{suffix}.{file_format}"
    if file_format == "xlsx":
        return (
            build_pareto_xlsx(report),
            filename,
            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        )
    return build_pareto_pdf(report), filename, "application/pdf"


__all__ = [
    "build_pareto_pdf",
    "build_pareto_xlsx",
    "export_pareto_report",
    "generate_pareto_report",
    "prepare_eligible_sales_rows",
    "resolve_pareto_period",
]
