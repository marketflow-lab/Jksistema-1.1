"""Private, short-lived PDF and XLSX artifacts for WhatsApp reports."""

from __future__ import annotations

import hashlib
import html
import json
import re
import time
import unicodedata
from datetime import datetime, timezone
from pathlib import Path
from typing import Any


REPORT_FILE_DIRNAME = "whatsapp_report_files"
REPORT_FILE_TTL_SECONDS = 7 * 24 * 60 * 60
REPORT_DOCUMENT_MAX_BYTES = 10 * 1024 * 1024
FORMAT_MIMES = {
    "pdf": "application/pdf",
    "xlsx": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
}
SENSITIVE_KEYS = {
    "access_token", "refresh_token", "authorization", "token", "password", "secret",
    "credential", "credentials", "buyer_email", "buyer_phone", "address", "shipping_address",
}


def _key(value: Any) -> str:
    text = unicodedata.normalize("NFKD", str(value or ""))
    text = "".join(char for char in text if not unicodedata.combining(char))
    return re.sub(r"\s+", " ", re.sub(r"[^a-zA-Z0-9]+", " ", text).lower()).strip()


def _safe_client_id(value: Any) -> str:
    safe = re.sub(r"[^A-Za-z0-9._-]+", "_", str(value or "default").strip())[:80]
    return safe.strip("._-") or "default"


def requested_report_formats(prompt: Any) -> set[str]:
    text = _key(prompt)
    is_report = bool(re.search(r"\b(relatorio|resumo|analise|balanco|planilha|dashboard)\b", text))
    if not is_report:
        return set()
    if re.search(r"\b(todos os formatos|quatro formatos|texto imagem pdf (?:e )?excel|completo em formatos)\b", text):
        return {"text", "png", "pdf", "xlsx"}
    formats: set[str] = set()
    if re.search(r"\b(texto|mensagem)\b", text):
        formats.add("text")
    if re.search(r"\b(imagem|png|grafico|visual)\b", text):
        formats.add("png")
    if re.search(r"\bpdf\b", text):
        formats.add("pdf")
    if re.search(r"\b(excel|xlsx|planilha)\b", text):
        formats.add("xlsx")
    return formats


def report_requested(prompt: Any) -> bool:
    return bool(re.search(r"\b(relatorio|resumo|analise|balanco|planilha|dashboard)\b", _key(prompt)))


def report_offer_text(prompt: Any) -> str:
    if not report_requested(prompt) or requested_report_formats(prompt):
        return ""
    return "Posso enviar este relatorio tambem em Imagem, PDF ou Excel. Diga um ou mais formatos."


def output_dir(base_info_dir: str | Path, client_id: Any) -> Path:
    root = Path(base_info_dir).resolve()
    output = (root / _safe_client_id(client_id) / REPORT_FILE_DIRNAME).resolve()
    output.relative_to(root)
    output.mkdir(parents=True, exist_ok=True)
    return output


def cleanup_stale_files(base_info_dir: str | Path, max_age_seconds: int = REPORT_FILE_TTL_SECONDS) -> int:
    root = Path(base_info_dir).resolve()
    removed = 0
    if not root.exists():
        return 0
    cutoff = time.time() - max(3600, int(max_age_seconds or REPORT_FILE_TTL_SECONDS))
    for path in root.glob(f"*/{REPORT_FILE_DIRNAME}/*"):
        try:
            resolved = path.resolve()
            resolved.relative_to(root)
            if resolved.is_file() and resolved.stat().st_mtime < cutoff:
                resolved.unlink(missing_ok=True)
                removed += 1
        except (OSError, ValueError):
            continue
    return removed


def _safe_tree(value: Any, depth: int = 0) -> Any:
    if depth > 5:
        return str(value)[:500]
    if isinstance(value, dict):
        return {
            str(key)[:100]: _safe_tree(item, depth + 1)
            for key, item in value.items()
            if str(key).strip().lower() not in SENSITIVE_KEYS
        }
    if isinstance(value, list):
        return [_safe_tree(item, depth + 1) for item in value[:5000]]
    if isinstance(value, str):
        return value[:4000]
    return value


def _source_values(result: dict[str, Any]) -> list[str]:
    values = [result.get("source_label"), result.get("source"), result.get("tool_id")]
    values.extend(result.get("sources") if isinstance(result.get("sources"), list) else [])
    values.extend(result.get("sources_human") if isinstance(result.get("sources_human"), list) else [])
    output: list[str] = []
    for value in values:
        text = re.sub(r"\s+", " ", str(value or "")).strip()
        if text and text not in output:
            output.append(text[:300])
    return output


def _result_rows(result: dict[str, Any]) -> list[dict[str, Any]]:
    candidates = [result.get("data"), result.get("rows"), result.get("all_rows"), result.get("result")]
    for value in candidates:
        if isinstance(value, list):
            return [_safe_tree(item) if isinstance(item, dict) else {"valor": _safe_tree(item)} for item in value[:5000]]
        if isinstance(value, dict):
            for key in ("rows", "records", "items", "orders", "results", "data"):
                rows = value.get(key)
                if isinstance(rows, list):
                    return [_safe_tree(item) if isinstance(item, dict) else {"valor": _safe_tree(item)} for item in rows[:5000]]
            scalar = {key: item for key, item in _safe_tree(value).items() if not isinstance(item, (dict, list))}
            if scalar:
                return [scalar]
    return []


def build_report_dataset(query_policy: Any, tool_results: Any) -> dict[str, Any]:
    policy = query_policy if isinstance(query_policy, dict) else {}
    results = [item for item in (tool_results if isinstance(tool_results, list) else []) if isinstance(item, dict)]
    rows: list[dict[str, Any]] = []
    sources: list[str] = []
    stores: list[str] = []
    coverage_complete = True
    limitations: list[str] = []
    for result in results:
        store = str(result.get("manager_store") or result.get("loja") or result.get("store") or "").strip()
        if store and store not in stores:
            stores.append(store[:160])
        for source in _source_values(result):
            if source not in sources:
                sources.append(source)
        complete = bool(
            result.get("coverage_complete") is True
            or (
                result.get("dados_suficientes") is True
                and not (isinstance(result.get("paging"), dict) and result.get("paging", {}).get("has_more") is True)
            )
        )
        coverage_complete = coverage_complete and complete
        if not complete:
            reason = str(result.get("empty_reason") or result.get("error") or "cobertura parcial").strip()
            if reason and reason not in limitations:
                limitations.append(reason[:500])
        for row in _result_rows(result):
            rows.append({"loja": store, "fonte": str(result.get("tool_id") or ""), **row})
            if len(rows) >= 5000:
                break
        if len(rows) >= 5000:
            limitations.append("Arquivo limitado aos primeiros 5.000 registros.")
            coverage_complete = False
            break
    configured_stores = policy.get("stores") if isinstance(policy.get("stores"), list) else []
    for store in [*configured_stores, policy.get("store")]:
        text = str(store or "").strip()
        if text and text not in stores:
            stores.append(text[:160])
    period_start = str(policy.get("data_inicio") or policy.get("period_start") or "nao informado")
    period_end = str(policy.get("data_fim") or policy.get("period_end") or "nao informado")
    return {
        "generated_at": datetime.now(timezone.utc).isoformat(),
        "period_start": period_start,
        "period_end": period_end,
        "stores": stores,
        "sources": sources,
        "record_count": len(rows),
        "coverage_complete": bool(results and coverage_complete),
        "limitations": limitations,
        "rows": rows,
    }


def _write_xlsx(path: Path, title: str, dataset: dict[str, Any]) -> None:
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill
    from openpyxl.utils import get_column_letter

    workbook = Workbook()
    summary = workbook.active
    summary.title = "Resumo"
    summary.append([title])
    summary["A1"].font = Font(bold=True, size=14)
    metadata = [
        ("Periodo", f"{dataset['period_start']} a {dataset['period_end']}"),
        ("Lojas", ", ".join(dataset["stores"]) or "nao informadas"),
        ("Fontes", "; ".join(dataset["sources"]) or "nao informadas"),
        ("Registros", dataset["record_count"]),
        ("Cobertura completa", "Sim" if dataset["coverage_complete"] else "Nao"),
        ("Limitacoes", "; ".join(dataset["limitations"]) or "Nenhuma informada"),
        ("Gerado em", dataset["generated_at"]),
    ]
    for label, value in metadata:
        summary.append([label, value])
    summary.column_dimensions["A"].width = 24
    summary.column_dimensions["B"].width = 90
    rows = dataset["rows"]
    data_sheet = workbook.create_sheet("Dados")
    headers = list(dict.fromkeys(str(key) for row in rows for key in row.keys()))[:80] if rows else ["informacao"]
    data_sheet.append(headers)
    for cell in data_sheet[1]:
        cell.font = Font(bold=True)
        cell.fill = PatternFill("solid", fgColor="D9EAF7")
    for row in rows or [{"informacao": "Sem dados tabulares disponiveis"}]:
        values = []
        for header in headers:
            value = row.get(header, "")
            if isinstance(value, (dict, list)):
                value = json.dumps(value, ensure_ascii=False, default=str)[:32000]
            values.append(value)
        data_sheet.append(values)
    data_sheet.freeze_panes = "A2"
    data_sheet.auto_filter.ref = data_sheet.dimensions
    for index, header in enumerate(headers, start=1):
        sample = [str(header), *[str(data_sheet.cell(row=row, column=index).value or "") for row in range(2, min(data_sheet.max_row, 100) + 1)]]
        data_sheet.column_dimensions[get_column_letter(index)].width = min(60, max(10, max(map(len, sample)) + 2))
    workbook.save(path)


def _write_pdf(path: Path, title: str, dataset: dict[str, Any]) -> None:
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4, landscape
    from reportlab.lib.styles import getSampleStyleSheet
    from reportlab.lib.units import mm
    from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

    styles = getSampleStyleSheet()
    small = styles["BodyText"].clone("WhatsAppReportSmall")
    small.fontSize = 7
    small.leading = 9
    story: list[Any] = [Paragraph(html.escape(title[:180]), styles["Title"]), Spacer(1, 3 * mm)]
    metadata = (
        f"Periodo: {html.escape(str(dataset['period_start']))} a {html.escape(str(dataset['period_end']))}<br/>"
        f"Lojas: {html.escape(', '.join(dataset['stores']) or 'nao informadas')}<br/>"
        f"Fontes: {html.escape('; '.join(dataset['sources']) or 'nao informadas')}<br/>"
        f"Registros: {dataset['record_count']}<br/>"
        f"Cobertura completa: {'Sim' if dataset['coverage_complete'] else 'Nao'}<br/>"
        f"Limitacoes: {html.escape('; '.join(dataset['limitations']) or 'Nenhuma informada')}"
    )
    story.extend([Paragraph(metadata, styles["BodyText"]), Spacer(1, 4 * mm)])
    rows = dataset["rows"][:1000]
    if rows:
        headers = list(dict.fromkeys(str(key) for row in rows for key in row.keys()))[:10]
        data: list[list[Any]] = [[Paragraph(header[:60], small) for header in headers]]
        for row in rows:
            data.append([
                Paragraph(html.escape(str(row.get(header, ""))[:260]), small)
                for header in headers
            ])
        width = landscape(A4)[0] - 24 * mm
        table = Table(data, colWidths=[width / len(headers)] * len(headers), repeatRows=1)
        table.setStyle(TableStyle([
            ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#D9EAF7")),
            ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
            ("GRID", (0, 0), (-1, -1), 0.25, colors.grey),
            ("VALIGN", (0, 0), (-1, -1), "TOP"),
        ]))
        story.append(table)
    else:
        story.append(Paragraph("Sem dados tabulares disponiveis.", styles["BodyText"]))
    SimpleDocTemplate(
        str(path), pagesize=landscape(A4), rightMargin=12 * mm, leftMargin=12 * mm,
        topMargin=12 * mm, bottomMargin=12 * mm,
    ).build(story)


def generate_report_documents(
    *,
    base_info_dir: str | Path,
    client_id: Any,
    task_id: Any,
    prompt: Any,
    tool_results: Any,
    query_policy: Any,
    formats: set[str] | None = None,
) -> dict[str, Any]:
    requested = set(formats if formats is not None else requested_report_formats(prompt)) & {"pdf", "xlsx"}
    if not requested:
        return {"expected": False, "status": "not_requested", "artifacts": []}
    root = output_dir(base_info_dir, client_id)
    cleanup_stale_files(base_info_dir)
    safe_task = re.sub(r"[^A-Za-z0-9._-]+", "_", str(task_id or "report"))[:80] or "report"
    dataset = build_report_dataset(query_policy, tool_results)
    title = re.sub(r"\s+", " ", str(prompt or "Relatorio Black Jhon")).strip()[:180] or "Relatorio Black Jhon"
    artifacts: list[dict[str, Any]] = []
    errors: list[str] = []
    for fmt in ("pdf", "xlsx"):
        if fmt not in requested:
            continue
        path = root / f"{safe_task}.{fmt}"
        try:
            (_write_pdf if fmt == "pdf" else _write_xlsx)(path, title, dataset)
            size = path.stat().st_size
            if size < 1 or size > REPORT_DOCUMENT_MAX_BYTES:
                raise RuntimeError(f"report_{fmt}_size_limit")
            digest = hashlib.sha256(path.read_bytes()).hexdigest()
            artifacts.append(
                {
                    "kind": fmt,
                    "artifact_type": f"report_{fmt}",
                    "path": str(path),
                    "filename": path.name,
                    "mime_type": FORMAT_MIMES[fmt],
                    "size": size,
                    "sha256": digest,
                    "expires_at": int(time.time()) + REPORT_FILE_TTL_SECONDS,
                    "period_start": dataset["period_start"],
                    "period_end": dataset["period_end"],
                    "stores": dataset["stores"],
                    "sources": dataset["sources"],
                    "record_count": dataset["record_count"],
                    "coverage_complete": dataset["coverage_complete"],
                    "limitations": dataset["limitations"],
                }
            )
        except Exception as exc:
            path.unlink(missing_ok=True)
            errors.append(f"{fmt}: {str(exc)[:300]}")
    return {
        "expected": True,
        "status": "completed" if len(artifacts) == len(requested) else "partial" if artifacts else "generation_failed",
        "artifacts": artifacts,
        "errors": errors,
        "metadata": {key: value for key, value in dataset.items() if key != "rows"},
    }


__all__ = [
    "FORMAT_MIMES",
    "REPORT_DOCUMENT_MAX_BYTES",
    "REPORT_FILE_DIRNAME",
    "REPORT_FILE_TTL_SECONDS",
    "build_report_dataset",
    "cleanup_stale_files",
    "generate_report_documents",
    "output_dir",
    "report_offer_text",
    "report_requested",
    "requested_report_formats",
]
