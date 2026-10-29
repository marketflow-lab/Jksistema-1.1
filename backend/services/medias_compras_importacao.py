"""Excel import endpoints for Medias Compras pedido lists."""

from __future__ import annotations

import copy
import hashlib
import io
import json
import math
import re
import unicodedata
import uuid
from datetime import datetime

import openpyxl
from fastapi import Depends, File, Form, HTTPException, UploadFile

from backend.services import medias_compras_common as medias_common
from backend.services.cadastro_common import _normalizar_sku_mes
from backend.services.medias_compras_common import *
from backend.services.medias_compras_excel import *
from backend.services.medias_compras_fiscal import *
from backend.services.runtime_bridge import bind_runtime_globals

_RUNTIME_NAMES = (
    "logger",
    "get_tenant_path",
    "TEMP_FILES_STORAGE",
    "TEMP_FILES_META",
    "LISTA_PEDIDO_XLSX_CACHE",
    "LISTA_PEDIDO_XLSX_CACHE_MAX_ITENS",
    "PASTA_INFO",
)


def _sync_common_names() -> None:
    for name in medias_common.COMMON_EXPORTS:
        globals()[name] = getattr(medias_common, name)
    for name in _RUNTIME_NAMES:
        globals()[name] = getattr(medias_common, name)


def _configure_runtime_globals(runtime_module=None):
    runtime = medias_common.configure_medias_compras_common_runtime(runtime_module)
    bind_runtime_globals(globals(), runtime)
    _sync_common_names()
    return runtime


def _mapear_cabecalho_lista_pedido_excel(ws) -> tuple[dict[str, int], int]:
    max_row = int(ws.max_row or 0)
    max_col = int(ws.max_column or 0)
    if max_row <= 0 or max_col <= 0:
        return {}, 1

    linhas_para_testar = min(max_row, 15)
    melhores_headers: dict[str, int] = {}
    melhor_linha = 1
    melhor_pontuacao = -1
    colunas_relevantes = {
        "sku",
        "picture",
        "description",
        "oemmodel",
        "colorside",
        "linkaliexpress",
        "quantity",
        "cost",
        "subtotalusd",
        "valorunidade",
        "valorunitario",
        "precounitario",
        "preco",
        "valor",
        "custo",
        "valortotal",
        "subtotal",
    }

    for row in range(1, linhas_para_testar + 1):
        headers = {}
        for col in range(1, max_col + 1):
            chave = _normalizar_cabecalho_excel(ws.cell(row=row, column=col).value)
            if chave and chave not in headers:
                headers[chave] = col
        pontuacao = sum(1 for chave in headers if chave in colunas_relevantes)
        if "sku" in headers:
            pontuacao += 4
        if pontuacao > melhor_pontuacao:
            melhores_headers = headers
            melhor_linha = row
            melhor_pontuacao = pontuacao
        if "sku" in headers and pontuacao >= 6:
            return headers, row

    return melhores_headers, melhor_linha


def _normalizar_sku_lista_pedido_excel(sku: str) -> str:
    sku_txt = str(sku or "").strip()
    sku_txt = re.sub(r"(?i)^\s*sku\s*[:#-]?\s*", "", sku_txt).strip()
    return _normalizar_sku_mes(sku_txt)


def _opcao_importacao(valor) -> str:
    return valor.strip().lower() if isinstance(valor, str) else ""


def _token_atualizacao_arquivo(client_id, lista, conteudo, sincronizar):
    estado = json.dumps(
        [client_id, lista, bool(sincronizar)], ensure_ascii=False, sort_keys=True,
        separators=(",", ":"), default=str,
    ).encode("utf-8")
    return hashlib.sha256(estado + b"\0" + conteudo).hexdigest()


def _relatorio_atualizacao_arquivo(antes, depois, ausentes_mantidos):
    def mapa(lista):
        return {
            _normalizar_sku_lista_pedido_excel(item.get("SKU", "")): _normalizar_item_lista_pedido(item)
            for item in lista.get("itens") or []
        }

    antigos, novos = mapa(antes), mapa(depois)
    relatorio = {nome: [] for nome in ("atualizados", "incluidos", "excluidos", "mantidos", "inalterados", "lista")}
    for sku, novo in novos.items():
        if sku not in antigos:
            relatorio["incluidos"].append({"sku": sku, "item": copy.deepcopy(novo)})
            continue
        antigo = antigos[sku]
        campos = [
            {"campo": chave, "antes": antigo.get(chave), "depois": novo.get(chave)}
            for chave in sorted(set(antigo) | set(novo))
            if chave not in {"valor_total_informado", "valor_unitario_informado"} and antigo.get(chave) != novo.get(chave)
        ]
        if campos:
            relatorio["atualizados"].append({"sku": sku, "campos": campos})
        else:
            relatorio["inalterados"].append(sku)
        if sku in ausentes_mantidos:
            relatorio["mantidos"].append({"sku": sku, "item": copy.deepcopy(novo)})
    for sku, antigo in antigos.items():
        if sku not in novos:
            relatorio["excluidos"].append({"sku": sku, "item": copy.deepcopy(antigo)})
    for campo, titulo in (("status", "Status"), ("updated_at", "Atualizada em")):
        if antes.get(campo) != depois.get(campo):
            relatorio["lista"].append({"campo": titulo, "antes": antes.get(campo), "depois": depois.get(campo)})
    return relatorio

async def api_medias_compras_lista_pedido_importar_excel_precos(
    lista_id: str,
    file: UploadFile = File(...),
    confirmar_inclusoes: str = Form("0"),
    client_id: str = Depends(medias_common.get_tenant_id),
    previsualizar: str = Form("0"),
    sincronizar_arquivo: str = Form("0"),
    confirmar_exclusoes: str = Form(""),
    token_confirmacao: str = Form(""),
):
    from backend.services.blocking_workers import run_heavy

    conteudo = await file.read()
    return await run_heavy(
        _api_medias_compras_lista_pedido_importar_excel_precos_sync,
        lista_id=lista_id, conteudo=conteudo, confirmar_inclusoes=confirmar_inclusoes,
        client_id=client_id, previsualizar=previsualizar, sincronizar_arquivo=sincronizar_arquivo,
        confirmar_exclusoes=confirmar_exclusoes, token_confirmacao=token_confirmacao,
    )


@_retry_listas_pedidos
def _api_medias_compras_lista_pedido_importar_excel_precos_sync(
    lista_id: str,
    conteudo: bytes,
    confirmar_inclusoes: str = Form("0"),
    client_id: str = Depends(medias_common.get_tenant_id),
    previsualizar: str = "0",
    sincronizar_arquivo: str = "0",
    confirmar_exclusoes: str = "",
    token_confirmacao: str = "",
):
    listas = _carregar_listas_pedidos(client_id)
    idx = next((i for i, l in enumerate(listas) if str(l.get("id", "")) == str(lista_id)), -1)
    if idx < 0:
        raise HTTPException(status_code=404, detail="Lista de pedidos nÃ£o encontrada")

    medias_common._exigir_lista_pedido_editavel(listas[idx])
    antes = copy.deepcopy(listas[idx])
    verdadeiro = {"1", "true", "sim", "yes", "y"}
    previa = _opcao_importacao(previsualizar) in verdadeiro
    sincronizar = _opcao_importacao(sincronizar_arquivo) in verdadeiro
    decisao_exclusoes = _opcao_importacao(confirmar_exclusoes)
    if decisao_exclusoes not in {"", "0", "1"}:
        raise HTTPException(400, "Confirmação de exclusão inválida.")
    token_atual = _token_atualizacao_arquivo(client_id, antes, conteudo, sincronizar)
    token_recebido = _opcao_importacao(token_confirmacao)
    if not previa and token_recebido and token_recebido != token_atual:
        raise HTTPException(409, "A lista ou o arquivo mudou. Selecione o arquivo novamente e confira a nova prévia.")
    if not conteudo:
        raise HTTPException(status_code=400, detail="Arquivo Excel vazio")

    try:
        wb = openpyxl.load_workbook(io.BytesIO(conteudo), data_only=True)
        ws = wb.active
    except Exception:
        raise HTTPException(status_code=400, detail="Arquivo invalido. Envie um Excel .xlsx")

    mapa_headers, linha_cabecalho = _mapear_cabecalho_lista_pedido_excel(ws)

    col_sku = mapa_headers.get("sku")
    col_vu = (
        mapa_headers.get("valorunidade")
        or mapa_headers.get("valorunitario")
        or mapa_headers.get("precounitario")
        or mapa_headers.get("preco")
        or mapa_headers.get("valor")
        or mapa_headers.get("cost")
        or mapa_headers.get("custo")
    )
    col_vt = (
        mapa_headers.get("valortotal")
        or mapa_headers.get("subtotalusd")
        or mapa_headers.get("subtotal")
        or mapa_headers.get("subtotaldolar")
        or mapa_headers.get("subtotalusd")
    )
    col_qtd = (
        mapa_headers.get("quantidade")
        or mapa_headers.get("qtd")
        or mapa_headers.get("qtde")
        or mapa_headers.get("qty")
        or mapa_headers.get("quantity")
    )
    col_titulo = (
        mapa_headers.get("titulodoprodutoemingles")
        or mapa_headers.get("titulo")
        or mapa_headers.get("descricao")
        or mapa_headers.get("descricaodoproduto")
        or mapa_headers.get("produto")
        or mapa_headers.get("description")
    )
    col_foto = mapa_headers.get("foto") or mapa_headers.get("imagem") or mapa_headers.get("picture")
    col_oem = mapa_headers.get("oem") or mapa_headers.get("oemmodel") or mapa_headers.get("modelo")
    col_cor_lado = (
        mapa_headers.get("colorside")
        or mapa_headers.get("corside")
        or mapa_headers.get("corlado")
        or mapa_headers.get("ladocor")
        or mapa_headers.get("lado")
        or mapa_headers.get("cor")
    )
    col_link = mapa_headers.get("link") or mapa_headers.get("url") or mapa_headers.get("linkaliexpress")
    col_cbm = mapa_headers.get("estimedcbm") or mapa_headers.get("estimatedcbm") or mapa_headers.get("cbm")
    col_peso = mapa_headers.get("estimedweigh") or mapa_headers.get("estimatedweight") or mapa_headers.get("peso") or mapa_headers.get("weight")
    col_embalagem = mapa_headers.get("individualpackaging") or mapa_headers.get("embalagemindividual") or mapa_headers.get("embalagem")

    if not col_sku:
        raise HTTPException(status_code=400, detail="Excel sem coluna SKU")
    if not any((col_qtd, col_vu, col_vt)):
        raise HTTPException(400, "Excel sem coluna de quantidade ou valor do item.")

    confirmar = _opcao_importacao(confirmar_inclusoes) in verdadeiro

    def _normalizar_titulo_comparacao(texto: str) -> str:
        base = unicodedata.normalize("NFKD", str(texto or ""))
        base = "".join(ch for ch in base if not unicodedata.combining(ch))
        return re.sub(r"\s+", " ", base).strip().lower()

    def numero_arquivo(row, coluna, nome, sku):
        if not coluna:
            return None
        celula = ws.cell(row=row, column=coluna)
        bruto = celula.value
        if celula.data_type == "e":
            raise HTTPException(400, f"SKU {sku}: corrija o erro de {nome} na linha {row} do Excel.")
        if bruto is None or str(bruto).strip() == "":
            if ws_formulas.cell(row=row, column=coluna).data_type == "f":
                raise HTTPException(400, f"SKU {sku}: recalcule e salve a fórmula de {nome} no Excel antes de importar.")
            return None
        valor = _to_float_excel(bruto)
        if isinstance(bruto, bool) or valor is None or not math.isfinite(valor) or valor < 0:
            raise HTTPException(400, f"SKU {sku}: {nome} inválido na linha {row}.")
        if nome == "Quantidade" and not valor.is_integer():
            raise HTTPException(400, f"SKU {sku}: a quantidade deve ser inteira na linha {row}.")
        return valor

    wb_formulas = openpyxl.load_workbook(io.BytesIO(conteudo), data_only=False)
    ws_formulas = wb_formulas.active

    atualizacoes = {}
    for row in range(linha_cabecalho + 1, ws.max_row + 1):
        celula_sku = ws.cell(row=row, column=col_sku)
        if celula_sku.data_type == "e" or (
            celula_sku.value is None and ws_formulas.cell(row=row, column=col_sku).data_type == "f"
        ):
            raise HTTPException(400, f"SKU inválido na linha {row}: recalcule e salve as fórmulas no Excel antes de importar.")
        sku = str(celula_sku.value if celula_sku.value is not None else "").strip()
        if _linha_resumo_excel_sku(sku):
            continue
        sku_norm = _normalizar_sku_lista_pedido_excel(sku)
        if not sku_norm:
            continue
        if sku_norm in atualizacoes:
            raise HTTPException(400, f"SKU duplicado no arquivo: {sku_norm}. Mantenha uma linha por SKU.")
        vu = numero_arquivo(row, col_vu, "Valor unidade", sku_norm)
        vt = numero_arquivo(row, col_vt, "Valor total", sku_norm)
        qtd = numero_arquivo(row, col_qtd, "Quantidade", sku_norm)
        titulo = str(ws.cell(row=row, column=col_titulo).value or "").strip() if col_titulo else ""
        foto = str(ws.cell(row=row, column=col_foto).value or "").strip() if col_foto else ""
        oem = str(ws.cell(row=row, column=col_oem).value or "").strip() if col_oem else ""
        cor_lado = str(ws.cell(row=row, column=col_cor_lado).value or "").strip() if col_cor_lado else ""
        link = str(ws.cell(row=row, column=col_link).value or "").strip() if col_link else ""
        cbm = str(ws.cell(row=row, column=col_cbm).value or "").strip() if col_cbm else ""
        peso = str(ws.cell(row=row, column=col_peso).value or "").strip() if col_peso else ""
        embalagem = str(ws.cell(row=row, column=col_embalagem).value or "").strip() if col_embalagem else ""
        atualizacoes[sku_norm] = {
            "sku": sku_norm,
            "sku_original": sku,
            "vu": vu,
            "vt": vt,
            "qtd": qtd,
            "titulo": titulo,
            "foto": foto,
            "oem": oem,
            "cor_lado": cor_lado,
            "link": link,
            "cbm": cbm,
            "peso": peso,
            "embalagem": embalagem,
        }

    wb.close()
    wb_formulas.close()
    if not atualizacoes:
        raise HTTPException(400, "O arquivo não contém itens válidos. A lista foi preservada.")

    lista = listas[idx]
    itens = [_normalizar_item_lista_pedido(i) for i in (lista.get("itens") or [])]
    itens_por_sku = {}
    for item in itens:
        sku_norm_item = _normalizar_sku_lista_pedido_excel(item.get("SKU", ""))
        if sku_norm_item:
            if sku_norm_item in itens_por_sku:
                raise HTTPException(400, f"SKU duplicado na lista: {sku_norm_item}. Revise a lista antes de atualizar.")
            itens_por_sku[sku_norm_item] = item

    pendencias = []
    aplicaveis = []
    for sku_norm, novo in atualizacoes.items():
        item_existente = itens_por_sku.get(sku_norm)
        titulo_excel = str(novo.get("titulo") or "").strip()
        if not item_existente:
            if novo["qtd"] is None or (novo["vu"] is None and novo["vt"] is None):
                raise HTTPException(400, f"SKU novo {sku_norm}: preencha quantidade e valor unitário ou total.")
            pendencias.append({
                "tipo": "sku_novo",
                "sku": sku_norm,
                "titulo_arquivo": titulo_excel,
                "motivo": "SKU nÃ£o existia na lista original",
            })
            aplicaveis.append({"novo": novo, "item": None})
            continue

        titulo_atual = ""
        for chave_titulo in TITULO_PRODUTO_INGLES_KEYS_LEGADO:
            titulo_atual = str(item_existente.get(chave_titulo, "") or "").strip()
            if titulo_atual:
                break
        if titulo_excel and _normalizar_titulo_comparacao(titulo_excel) != _normalizar_titulo_comparacao(titulo_atual):
            pendencias.append({
                "tipo": "titulo_diferente",
                "sku": sku_norm,
                "titulo_atual": titulo_atual,
                "titulo_arquivo": titulo_excel,
                "motivo": "Título do arquivo é diferente do título atual da lista",
            })
        aplicaveis.append({"novo": novo, "item": item_existente})

    ausentes = set(itens_por_sku) - set(atualizacoes) if sincronizar else set()
    for sku in itens_por_sku:
        if sku in ausentes:
            pendencias.append({"tipo": "sku_ausente", "sku": sku,
                              "item": copy.deepcopy(itens_por_sku[sku]),
                              "motivo": "SKU existente na lista e ausente do arquivo"})

    for registro in aplicaveis:
        novo = registro.get("novo") or {}
        item = registro.get("item")

        vu = novo.get("vu")
        vt = novo.get("vt")
        qtd_nova = novo.get("qtd")
        titulo_novo = str(novo.get("titulo") or "").strip()
        foto_nova = str(novo.get("foto") or "").strip()
        oem_novo = str(novo.get("oem") or "").strip()
        cor_lado_novo = str(novo.get("cor_lado") or "").strip()
        link_novo = str(novo.get("link") or "").strip()
        cbm_novo = str(novo.get("cbm") or "").strip()
        peso_novo = str(novo.get("peso") or "").strip()
        embalagem_nova = str(novo.get("embalagem") or "").strip()

        if item is None:
            item = {
                "SKU": str(novo.get("sku") or "").strip(),
                "Foto": foto_nova,
                TITULO_PRODUTO_INGLES_KEY: titulo_novo,
                "OEM": oem_novo,
                COR_LADO_LISTA_PEDIDO_KEY: cor_lado_novo,
                "Link": link_novo,
                "Quantidade": 0,
                "Valor unidade": 0.0,
                "Valor total": 0.0,
                CBM_LISTA_PEDIDO_KEY: cbm_novo,
                PESO_LISTA_PEDIDO_KEY: peso_novo,
                EMBALAGEM_LISTA_PEDIDO_KEY: embalagem_nova,
            }
            itens.append(item)

        if qtd_nova is not None:
            item["Quantidade"] = int(qtd_nova)

        qtd_item = float(item.get("Quantidade", 0) or 0)
        if vu is not None:
            item["Valor unidade"] = vu
            item["valor_unitario_informado"] = True
        elif vt is not None and qtd_item > 0:
            item["Valor unidade"] = vt / qtd_item
            item.pop("valor_unitario_informado", None)
        elif vt is not None and vt > 0:
            raise HTTPException(400, f"SKU {novo['sku']}: informe quantidade maior que zero para calcular o valor unitário pelo total.")
        if vt is not None:
            item["Valor total"] = vt
            item["valor_total_informado"] = True
        elif vu is not None or qtd_nova is not None:
            item["Valor total"] = round(qtd_item * float(item.get("Valor unidade", 0) or 0), 2)
            item.pop("valor_total_informado", None)
        if not all(math.isfinite(float(item.get(campo, 0) or 0)) for campo in ("Quantidade", "Valor unidade", "Valor total")):
            raise HTTPException(400, f"SKU {novo['sku']}: quantidade ou valor excede o limite permitido.")

        if titulo_novo:
            item[TITULO_PRODUTO_INGLES_KEY] = titulo_novo
        if foto_nova:
            item["Foto"] = foto_nova
        if oem_novo:
            item["OEM"] = oem_novo
        if cor_lado_novo:
            item[COR_LADO_LISTA_PEDIDO_KEY] = cor_lado_novo
        if link_novo:
            item["Link"] = link_novo
        if cbm_novo:
            item[CBM_LISTA_PEDIDO_KEY] = cbm_novo
        if peso_novo:
            item[PESO_LISTA_PEDIDO_KEY] = peso_novo
        if embalagem_nova:
            item[EMBALAGEM_LISTA_PEDIDO_KEY] = embalagem_nova

    excluir = sincronizar and decisao_exclusoes == "1"
    propostos = [item for item in itens if not excluir or _normalizar_sku_lista_pedido_excel(item.get("SKU", "")) not in ausentes]
    precisa_confirmar = (
        any(p["tipo"] != "sku_ausente" for p in pendencias) and not confirmar
    ) or (bool(ausentes) and decisao_exclusoes == "")
    if previa or precisa_confirmar or (sincronizar and not token_recebido):
        # A preview never changes the loaded snapshot, status, cache or storage.
        proposta = {**antes, "itens": propostos}
        relatorio = _relatorio_atualizacao_arquivo(antes, proposta, ausentes if not excluir else set())
        if ausentes and decisao_exclusoes == "":
            relatorio["excluidos"] = [{"sku": sku, "item": copy.deepcopy(itens_por_sku[sku])}
                                      for sku in itens_por_sku if sku in ausentes]
            relatorio["mantidos"] = []
        return {
            "success": False,
            "previsualizacao": True,
            "requer_confirmacao": bool(pendencias),
            "token_confirmacao": token_atual,
            "mensagem": "Confira as alterações antes de atualizar a lista.",
            "pendencias": pendencias,
            "resumo_pendencias": {
                "total": len(pendencias),
                "total_sku_novo": sum(p["tipo"] == "sku_novo" for p in pendencias),
                "total_titulo_diferente": sum(p["tipo"] == "titulo_diferente" for p in pendencias),
                "total_sku_ausente": len(ausentes),
            },
            "total_skus_lidos": len(atualizacoes),
            "relatorio": relatorio,
        }

    if excluir:
        from backend.services.medias_compras_listas import _proteger_itens_aprovados_lista
        propostos = _proteger_itens_aprovados_lista(antes.get("itens"), propostos)

    lista["itens"] = _recalcular_frete_internacional_itens_lista(
        client_id,
        propostos,
        loja=str(lista.get("store_id") or lista.get("loja") or ""),
        preservar_cubagem=_cubagem_lista_historica(lista),
    )
    lista["status"] = _normalizar_status_lista_pedido("Analisando orçamento")
    lista["updated_at"] = datetime.now().isoformat(timespec="seconds")
    listas[idx] = lista

    relatorio = _relatorio_atualizacao_arquivo(antes, lista, ausentes if not excluir else set())
    skus_incluidos = [item["sku"] for item in relatorio["incluidos"]]
    skus_excluidos = [item["sku"] for item in relatorio["excluidos"]]

    resposta = {
        "success": True,
        "lista": {
            **_resumo_lista_pedido(lista, client_id=client_id),
            "itens": lista.get("itens") or [],
        },
        "total_skus_lidos": len(atualizacoes),
        "total_skus_atualizados": len(relatorio["atualizados"]),
        "total_quantidades_atualizadas": sum(any(campo["campo"] == "Quantidade" for campo in item["campos"]) for item in relatorio["atualizados"]),
        "total_skus_incluidos": len(skus_incluidos),
        "skus_incluidos": skus_incluidos,
        "total_skus_excluidos": len(skus_excluidos),
        "skus_excluidos": skus_excluidos,
        "total_skus_ignorados": 0,
        "skus_ignorados": [],
        "pendencias_confirmadas": len(pendencias) if confirmar else 0,
        "relatorio": relatorio,
    }
    # Finish response preparation before committing so a summary failure cannot
    # report an unsuccessful operation after the data has already been saved.
    _salvar_listas_pedidos(client_id, listas)
    try:
        _limpar_cache_lista_pedido(client_id, str(lista.get("id", "") or ""), manter_versao=lista.get("updated_at"))
    except OSError:
        pass  # Cached exports are versioned; cache cleanup cannot undo this save.
    return resposta


async def api_medias_compras_lista_pedido_importar_excel_nova_lista(
    file: UploadFile = File(...),
    nome_lista: str = Form(""),
    loja: str = Form("__todas"),
    store_id: str = Form(""),
    fornecedor_id: str = Form(""),
    client_id: str = Depends(medias_common.get_tenant_id),
):
    from backend.services.blocking_workers import run_heavy

    conteudo = await file.read()
    return await run_heavy(
        _api_medias_compras_lista_pedido_importar_excel_nova_lista_sync,
        conteudo=conteudo,
        nome_lista=nome_lista,
        loja=loja,
        store_id=store_id,
        fornecedor_id=fornecedor_id,
        client_id=client_id,
        filename=str(getattr(file, "filename", "") or ""),
    )


@_retry_listas_pedidos
def _api_medias_compras_lista_pedido_importar_excel_nova_lista_sync(
    conteudo: bytes,
    nome_lista: str = Form(""),
    loja: str = Form("__todas"),
    store_id: str = Form(""),
    fornecedor_id: str = Form(""),
    client_id: str = Depends(medias_common.get_tenant_id),
    filename: str = "",
):
    fornecedor = _resolver_fornecedor_lista(client_id, fornecedor_id)
    if not conteudo:
        raise HTTPException(status_code=400, detail="Arquivo Excel vazio")

    try:
        wb = openpyxl.load_workbook(io.BytesIO(conteudo), data_only=True)
        ws = wb.active
    except Exception:
        raise HTTPException(status_code=400, detail="Arquivo invalido. Envie um Excel .xlsx")

    mapa_headers, linha_cabecalho = _mapear_cabecalho_lista_pedido_excel(ws)

    col_sku = mapa_headers.get("sku")
    col_vu = (
        mapa_headers.get("valorunidade")
        or mapa_headers.get("valorunitario")
        or mapa_headers.get("precounitario")
        or mapa_headers.get("preco")
        or mapa_headers.get("valor")
        or mapa_headers.get("cost")
        or mapa_headers.get("custo")
    )
    col_vt = (
        mapa_headers.get("valortotal")
        or mapa_headers.get("subtotalusd")
        or mapa_headers.get("subtotal")
        or mapa_headers.get("subtotaldolar")
    )
    col_qtd = (
        mapa_headers.get("quantidade")
        or mapa_headers.get("qtd")
        or mapa_headers.get("qtde")
        or mapa_headers.get("qty")
        or mapa_headers.get("quantity")
    )
    col_titulo = (
        mapa_headers.get("titulodoprodutoemingles")
        or mapa_headers.get("titulo")
        or mapa_headers.get("descricao")
        or mapa_headers.get("descricaodoproduto")
        or mapa_headers.get("produto")
        or mapa_headers.get("description")
    )
    col_foto = mapa_headers.get("foto") or mapa_headers.get("imagem") or mapa_headers.get("picture")
    col_oem = mapa_headers.get("oem") or mapa_headers.get("oemmodel") or mapa_headers.get("modelo")
    col_cor_lado = (
        mapa_headers.get("colorside")
        or mapa_headers.get("corside")
        or mapa_headers.get("corlado")
        or mapa_headers.get("ladocor")
        or mapa_headers.get("lado")
        or mapa_headers.get("cor")
    )
    col_link = mapa_headers.get("link") or mapa_headers.get("url") or mapa_headers.get("linkaliexpress")
    col_cbm = mapa_headers.get("estimedcbm") or mapa_headers.get("estimatedcbm") or mapa_headers.get("cbm")
    col_peso = mapa_headers.get("estimedweigh") or mapa_headers.get("estimatedweight") or mapa_headers.get("peso") or mapa_headers.get("weight")
    col_embalagem = mapa_headers.get("individualpackaging") or mapa_headers.get("embalagemindividual") or mapa_headers.get("embalagem")

    if not col_sku:
        raise HTTPException(status_code=400, detail="Excel sem coluna SKU")
    itens_por_sku = {}

    for row in range(linha_cabecalho + 1, ws.max_row + 1):
        sku_raw = str(ws.cell(row=row, column=col_sku).value or "").strip()
        if _linha_resumo_excel_sku(sku_raw):
            continue
        sku = _normalizar_sku_lista_pedido_excel(sku_raw)
        if not sku:
            continue

        vu = _to_float_excel(ws.cell(row=row, column=col_vu).value) if col_vu else None
        vt = _to_float_excel(ws.cell(row=row, column=col_vt).value) if col_vt else None
        qtd = _to_float_excel(ws.cell(row=row, column=col_qtd).value) if col_qtd else None
        titulo = str(ws.cell(row=row, column=col_titulo).value or "").strip() if col_titulo else ""
        foto = str(ws.cell(row=row, column=col_foto).value or "").strip() if col_foto else ""
        oem = str(ws.cell(row=row, column=col_oem).value or "").strip() if col_oem else ""
        cor_lado = str(ws.cell(row=row, column=col_cor_lado).value or "").strip() if col_cor_lado else ""
        link = str(ws.cell(row=row, column=col_link).value or "").strip() if col_link else ""
        cbm = str(ws.cell(row=row, column=col_cbm).value or "").strip() if col_cbm else ""
        peso = str(ws.cell(row=row, column=col_peso).value or "").strip() if col_peso else ""
        embalagem = str(ws.cell(row=row, column=col_embalagem).value or "").strip() if col_embalagem else ""

        if vu is None and vt is None and qtd is None and not any([titulo, foto, oem, cor_lado, link, cbm, peso, embalagem]):
            continue

        qtd_int = int(max(0.0, round(float(qtd or 0))))
        valor_unidade = float(vu or 0)
        valor_total = float(vt or 0)

        if valor_unidade <= 0 and valor_total > 0 and qtd_int > 0:
            valor_unidade = valor_total / qtd_int
        if valor_total <= 0 and valor_unidade > 0 and qtd_int > 0:
            valor_total = valor_unidade * qtd_int

        itens_por_sku[sku] = {
            "SKU": sku,
            "Foto": foto,
            TITULO_PRODUTO_INGLES_KEY: titulo,
            "OEM": oem,
            COR_LADO_LISTA_PEDIDO_KEY: cor_lado,
            "Link": link,
            "Quantidade": qtd_int,
            "Valor unidade": round(max(0.0, valor_unidade), 2),
            "Valor total": round(max(0.0, valor_total), 2),
            CBM_LISTA_PEDIDO_KEY: cbm,
            PESO_LISTA_PEDIDO_KEY: peso,
            EMBALAGEM_LISTA_PEDIDO_KEY: embalagem,
        }

    itens_lidos = list(itens_por_sku.values())
    if not itens_lidos:
        raise HTTPException(status_code=400, detail="Nenhum SKU valido encontrado para importar")

    nome_lista_final = str(nome_lista or "").strip()
    if not nome_lista_final:
        nome_arquivo = str(filename or "").strip()
        nome_arquivo = re.sub(r"\.[A-Za-z0-9]{1,8}$", "", nome_arquivo).strip()
        nome_lista_final = nome_arquivo or f"Lista importada {datetime.now().strftime('%d/%m/%Y %H:%M')}"

    agora_iso = datetime.now().isoformat(timespec="seconds")
    lista_id = str(uuid.uuid4())
    escopo_loja = _resolver_escopo_loja_medias(
        client_id,
        loja,
        store_id,
        exigir_especifica=True,
    )
    loja_final = escopo_loja["loja"]
    store_id_cadastro = escopo_loja["store_id"]
    itens_norm = _recalcular_frete_internacional_itens_lista(
        client_id,
        itens_lidos,
        loja=store_id_cadastro or loja_final,
    )

    listas_salvas = _carregar_listas_pedidos(client_id)
    status_importacao = _normalizar_status_lista_pedido("Analisando orçamento")
    listas_salvas.insert(0, {
        "id": lista_id,
        "nome_lista": nome_lista_final,
        "loja": loja_final,
        "store_id": store_id_cadastro,
        **fornecedor,
        "status": status_importacao,
        "created_at": agora_iso,
        "updated_at": agora_iso,
        "origem": "importacao_excel",
        "itens": itens_norm,
    })
    _salvar_listas_pedidos(client_id, listas_salvas)
    _limpar_cache_lista_pedido(client_id, lista_id, manter_versao=agora_iso)

    return {
        "success": True,
        "mensagem": "Lista criada com sucesso a partir do Excel.",
        "lista": {
            "id": lista_id,
            "nome_lista": nome_lista_final,
            "loja": loja_final,
            "store_id": store_id_cadastro,
            **fornecedor,
            "status": status_importacao,
            "created_at": agora_iso,
            "updated_at": agora_iso,
            "itens": itens_norm,
            "total_itens": len(itens_norm),
            "total_quantidade": int(sum(float(i.get("Quantidade", 0) or 0) for i in itens_norm)),
        },
    }



def configure_medias_compras_importacao_runtime(runtime_module=None):
    return _configure_runtime_globals(runtime_module)


configure_medias_compras_importacao_runtime()

__all__ = [
    "configure_medias_compras_importacao_runtime",
    "api_medias_compras_lista_pedido_importar_excel_precos",
    "api_medias_compras_lista_pedido_importar_excel_nova_lista",
]
