"""Cadastro photo extraction, upload and serving helpers."""

from __future__ import annotations

import inspect
import json
import logging
from contextlib import contextmanager
from pathlib import Path
from typing import Any, Optional

from fastapi import Header, HTTPException, Request

from backend.services.path_coordination import path_locks_for
from backend.services.runtime_bridge import bind_runtime_globals
from backend.services.cadastro_tenant_trust import (
    CadastroTenantTrustErro,
    resolver_alias_tenant_registrado,
)
from backend.services.cadastro_fotos_coordenacao import (
    CadastroFotosCoordenacaoErro,
    bloquear_transicao_fotos_tenant,
)

logger = logging.getLogger("jk_sistema")
# FastAPI captures this module's dependency callable while routes are imported.
# Keep that callable stable and replace only its runtime target during startup.
_runtime_get_tenant_id = None

CADASTRO_FOTOS_CONFIG_ARQUIVO = "cadastro_fotos_config.json"
CADASTRO_FOTOS_CONFIG_SCHEMA = "jk.cadastro.fotos.v1"
CADASTRO_FOTOS_EXTENSOES = (".jpg", ".jpeg", ".png", ".webp", ".gif", ".bmp")


async def get_tenant_id(request: Request, authorization: Optional[str] = Header(default=None)):
    if not callable(_runtime_get_tenant_id):
        raise HTTPException(
            status_code=503,
            detail="Contexto de autenticacao do Cadastro ainda nao inicializado.",
        )
    resultado = _runtime_get_tenant_id(request, authorization)
    if inspect.isawaitable(resultado):
        return await resultado
    return resultado


def get_tenant_path(client_id: str):
    raise RuntimeError("Cadastro runtime was not configured.")


def _configure_runtime_globals(target_globals, runtime_module=None):
    runtime = bind_runtime_globals(target_globals, runtime_module)
    if runtime is not None:
        runtime_logger = getattr(runtime, "logger", None)
        if runtime_logger is not None:
            target_globals["logger"] = runtime_logger
        runtime_get_tenant_id = getattr(runtime, "get_tenant_id", None)
        if (
            callable(runtime_get_tenant_id)
            and runtime_get_tenant_id is not target_globals.get("get_tenant_id")
        ):
            target_globals["_runtime_get_tenant_id"] = runtime_get_tenant_id
        if hasattr(runtime, "get_tenant_path"):
            target_globals["get_tenant_path"] = getattr(runtime, "get_tenant_path")
    return runtime


import base64
import hashlib
import io
import os
import posixpath
import re
import tempfile
import unicodedata
from urllib.parse import quote_plus, unquote, urlsplit

import openpyxl
import requests
from fastapi import Depends, File, Form, UploadFile
from fastapi.responses import FileResponse
from google.oauth2.service_account import Credentials
from google.auth.transport.requests import Request as GoogleAuthRequest

from backend.services.cadastro_common import *


def configure_cadastro_fotos_runtime(runtime_module=None):
    configure_cadastro_common_runtime(runtime_module)
    return _configure_runtime_globals(globals(), runtime_module)


configure_cadastro_fotos_runtime()

def _exportar_planilha_xlsx_bytes(spreadsheet_id: str) -> bytes:
    scopes = ["https://www.googleapis.com/auth/spreadsheets.readonly"]
    creds = Credentials.from_service_account_file(CREDENTIALS_FILE, scopes=scopes)
    from google.auth.transport.requests import Request as GoogleAuthRequest
    creds.refresh(GoogleAuthRequest())
    headers = {"Authorization": f"Bearer {creds.token}"}
    url = f"https://docs.google.com/spreadsheets/d/{spreadsheet_id}/export"
    resp = requests.get(
        url,
        headers=headers,
        params={"format": "xlsx", "gid": str(SPREADSHEET_GID_FOTOS_SKU)},
        timeout=60,
    )
    resp.raise_for_status()
    return resp.content

def _extrair_imagens_planilha_por_sku() -> dict:
    """Extrai imagens ancoradas na coluna B e relaciona com o SKU da coluna A."""
    try:
        xlsx_bytes = _exportar_planilha_xlsx_bytes(SPREADSHEET_ID_FOTOS_SKU)
    except Exception as e:
        logger.warning(f"[CADASTRO] NÃƒÂ£o foi possÃƒÂ­vel exportar planilha para extrair imagens: {e}")
        return {}

    try:
        wb = openpyxl.load_workbook(io.BytesIO(xlsx_bytes))
        ws = None
        for nome in wb.sheetnames:
            if nome.strip().lower() == "dados":
                ws = wb[nome]
                break
        if ws is None:
            ws = wb[wb.sheetnames[0]]

        imagens_por_linha = {}
        for img in getattr(ws, "_images", []):
            anc = getattr(img, "anchor", None)
            from_marker = getattr(anc, "_from", None)
            if from_marker is None:
                continue
            row_1 = int(from_marker.row) + 1
            col_1 = int(from_marker.col) + 1
            if col_1 != 2 or row_1 < 2:
                continue
            try:
                payload = img._data()
            except Exception:
                continue
            ext = str(getattr(img, "format", "png") or "png").lower()
            if ext == "jpeg":
                ext = "jpg"
            imagens_por_linha[row_1] = {"bytes": payload, "ext": ext}

        if not imagens_por_linha:
            return {}

        mapa = {}
        for row_1, payload in imagens_por_linha.items():
            sku_raw = ws.cell(row=row_1, column=1).value
            sku_norm = _normalizar_sku_mes(str(sku_raw or "").strip())
            if sku_norm:
                atual = mapa.get(sku_norm)
                if atual is None or len(payload["bytes"]) > len(atual["bytes"]):
                    mapa[sku_norm] = payload
        return mapa
    except Exception as e:
        logger.warning(f"[CADASTRO] Falha ao processar imagens exportadas da planilha: {e}")
        return {}

def _nome_arquivo_foto_sku(sku: str, ext: str) -> str:
    sku_limpo = re.sub(r'[\\/:*?"<>|]+', '_', str(sku or '').strip())
    ext_limpa = re.sub(r'[^a-zA-Z0-9]+', '', str(ext or 'png').lower()) or 'png'
    return f"{sku_limpo}.{ext_limpa}"

def _cadastro_sku_sem_zeros_blocos(sku: str) -> str:
    sku_norm = _normalizar_sku_mes(str(sku or "").strip())
    partes = re.split(r"(\d+)", sku_norm.upper())
    return "".join(str(int(p)) if p.isdigit() else p for p in partes)


def _cadastro_store_id_foto_exato(store_id: str | None) -> str:
    valor = str(store_id or "").strip()
    if not valor:
        return ""
    return valor


def _cadastro_store_id_foto_segmento(store_id: str | None) -> str:
    valor = _cadastro_store_id_foto_exato(store_id)
    if not valor:
        return ""
    return "sid-" + hashlib.sha256(valor.encode("utf-8")).hexdigest()


def _cadastro_store_id_foto_seguro(store_id: str | None) -> str:
    """Compatibility alias for the new opaque, filesystem-safe segment."""

    return _cadastro_store_id_foto_segmento(store_id)


def _cadastro_store_id_foto_legado_sintaticamente_seguro(store_id: str | None) -> bool:
    return bool(re.fullmatch(r"[A-Za-z0-9_-]{1,128}", _cadastro_store_id_foto_exato(store_id)))


def _cadastro_caminho_e_link(caminho: os.PathLike[str] | str) -> bool:
    path = Path(caminho)
    if path.is_symlink():
        return True
    is_junction = getattr(path, "is_junction", None)
    return bool(callable(is_junction) and is_junction())


def _cadastro_raiz_lexical_sem_reparse(caminho: os.PathLike[str] | str) -> bool:
    """Rejeita redirecionamento no caminho ou em qualquer ancestral existente."""

    caminho_abs = os.path.abspath(os.fspath(caminho))
    caminho_real = os.path.realpath(caminho_abs)
    if os.path.normcase(os.path.normpath(caminho_real)) != os.path.normcase(
        os.path.normpath(caminho_abs)
    ):
        return False
    if os.path.lexists(caminho_abs):
        return not _cadastro_caminho_e_link(caminho_abs) and os.path.isdir(caminho_abs)
    return True


def _cadastro_diretorio_seguro_abaixo(
    raiz: os.PathLike[str] | str,
    diretorio: os.PathLike[str] | str,
) -> bool:
    """Valida contenção e rejeita reparse points abaixo da raiz autorizada."""

    raiz_abs = os.path.abspath(os.fspath(raiz))
    diretorio_abs = os.path.abspath(os.fspath(diretorio))
    try:
        if os.path.commonpath([raiz_abs, diretorio_abs]) != raiz_abs:
            return False
    except ValueError:
        return False
    relativo = os.path.relpath(diretorio_abs, raiz_abs)
    atual = raiz_abs
    if relativo != ".":
        for parte in relativo.split(os.sep):
            atual = os.path.join(atual, parte)
            if os.path.lexists(atual) and (
                _cadastro_caminho_e_link(atual) or not os.path.isdir(atual)
            ):
                return False
    raiz_real = os.path.realpath(raiz_abs)
    diretorio_real = os.path.realpath(diretorio_abs)
    try:
        return os.path.commonpath([raiz_real, diretorio_real]) == raiz_real
    except ValueError:
        return False


def _cadastro_tenant_path_fotos_seguro(client_id: str) -> str:
    identidade = str(client_id or "").strip()
    pasta_info = str(globals().get("PASTA_INFO") or "")
    info_root = os.path.abspath(pasta_info)
    tenant_path = os.path.abspath(str(get_tenant_path(client_id) or ""))
    tenant_lexico = os.path.abspath(os.path.join(info_root, identidade))
    if (
        not identidade
        or not pasta_info.strip()
        or not _cadastro_raiz_lexical_sem_reparse(info_root)
        or os.path.basename(os.path.normpath(tenant_path)) != identidade
        or os.path.normcase(os.path.normpath(tenant_path))
        != os.path.normcase(os.path.normpath(tenant_lexico))
    ):
        return ""
    if os.path.lexists(tenant_path):
        if not os.path.isdir(tenant_path):
            return ""
        if _cadastro_caminho_e_link(tenant_path):
            try:
                # A junction operacional so e aceita depois de provisionada no
                # registro local da raiz lexical. O resolvedor compara a
                # identidade fisica atual e devolve o destino canonico.
                destino_real = resolver_alias_tenant_registrado(pasta_info, identidade)
            except CadastroTenantTrustErro:
                return ""
            return destino_real
    return tenant_path


def _cadastro_pasta_fotos(client_id: str, store_id: str | None = None) -> tuple[str, str]:
    store_seguro = _cadastro_store_id_foto_segmento(store_id)
    tenant_path = _cadastro_tenant_path_fotos_seguro(client_id)
    if not tenant_path:
        raise HTTPException(status_code=409, detail="Diretorio do cliente inseguro.")
    pasta_base = os.path.join(tenant_path, "cadastro_fotos")
    if not store_seguro:
        pasta_destino = pasta_base
        relativo = "cadastro_fotos"
    else:
        pasta_lojas = os.path.join(pasta_base, "lojas")
        pasta_destino = os.path.join(pasta_lojas, store_seguro)
        relativo = f"cadastro_fotos/lojas/{store_seguro}"
    if not _cadastro_diretorio_seguro_abaixo(tenant_path, pasta_destino):
        raise HTTPException(status_code=409, detail="Diretorio de fotos inseguro.")
    return pasta_destino, relativo


def _preparar_foto_data_url_no_tenant(
    client_id: str,
    sku: str,
    foto_data_url: str,
    foto_filename: str = "",
    store_id: str | None = None,
) -> dict[str, Any]:
    """Valida uma foto e calcula seu destino sem alterar o sistema de arquivos."""
    raw = str(foto_data_url or "").strip()
    if not raw:
        return {}

    m = re.match(r"^data:image/([a-zA-Z0-9.+-]+);base64,(.+)$", raw, re.IGNORECASE | re.DOTALL)
    if not m:
        raise HTTPException(status_code=400, detail="Formato de imagem invalido para salvar.")

    mime_ext = m.group(1).lower()
    b64 = m.group(2).strip()
    ext_map = {
        "jpeg": "jpg",
        "jpg": "jpg",
        "png": "png",
        "webp": "webp",
        "gif": "gif",
        "bmp": "bmp",
    }
    ext = ext_map.get(mime_ext)
    if not ext:
        raise HTTPException(status_code=400, detail="Tipo de imagem nÃ£o suportado.")

    try:
        conteudo = base64.b64decode(b64, validate=True)
    except Exception:
        raise HTTPException(status_code=400, detail="ConteÃƒÂºdo da imagem invalido (base64).")

    if not conteudo:
        raise HTTPException(status_code=400, detail="Imagem vazia.")

    sku_norm = _normalizar_sku_mes(sku)
    if not sku_norm:
        raise HTTPException(status_code=400, detail="SKU invalido para salvar imagem.")

    pasta_fotos, prefixo_relativo = _cadastro_pasta_fotos(client_id, store_id)

    if store_id:
        # No cadastro por loja, o nome enviado pelo navegador e apenas metadado
        # de exibicao. A identidade persistida vem do SKU para que dois produtos
        # que anexem "imagem.png" nao sobrescrevam um ao outro na mesma loja.
        nome_arquivo = _nome_arquivo_foto_sku(sku_norm, ext)
        nome_base, nome_ext = os.path.splitext(nome_arquivo)
        if nome_base != sku_norm:
            digest = hashlib.sha256(sku_norm.encode("utf-8")).hexdigest()[:12]
            nome_arquivo = f"{nome_base}-{digest}{nome_ext}"
    elif foto_filename:
        # Mantem o contrato historico das rotas legadas fora do novo escopo.
        nome_base = os.path.splitext(os.path.basename(str(foto_filename)))[0]
        nome_base = re.sub(r'[\\/:*?"<>|]+', '_', nome_base).strip() or sku_norm
        nome_arquivo = f"{nome_base}.{ext}"
    else:
        nome_arquivo = _nome_arquivo_foto_sku(sku_norm, ext)

    return {
        "caminho": os.path.join(pasta_fotos, nome_arquivo),
        "relativo": f"{prefixo_relativo}/{nome_arquivo}",
        "conteudo": conteudo,
        "store_id": _cadastro_store_id_foto_exato(store_id),
        "sku": sku_norm,
    }


def _preparar_fotos_data_url_no_tenant(
    client_id: str,
    sku: str,
    foto_data_url: str,
    foto_filename: str = "",
    store_id: str | None = None,
) -> list[dict[str, Any]]:
    """Prepara o destino da loja e de todo o seu grupo compartilhado."""

    if not str(foto_data_url or "").strip():
        return []
    store_exato = _cadastro_store_id_foto_exato(store_id)
    if not store_exato and _cadastro_fotos_escopo_estrito(client_id):
        raise HTTPException(
            status_code=409,
            detail={
                "code": "store_id_required",
                "message": "Selecione a loja para salvar a foto do produto.",
            },
        )
    destinos = (
        _cadastro_fotos_store_ids_compartilhados(client_id, store_exato)
        if store_exato
        else [""]
    )
    preparadas = [
        _preparar_foto_data_url_no_tenant(
            client_id,
            sku,
            foto_data_url,
            foto_filename,
            store_id=destino or None,
        )
        for destino in destinos
    ]
    return [item for item in preparadas if item]


def _cadastro_caminhos_variantes_fotos_preparadas(
    preparadas: list[dict[str, Any]],
) -> list[str]:
    """Lista todos os formatos concorrentes do mesmo SKU em cada destino."""

    caminhos: list[str] = []
    vistos: set[str] = set()
    for preparada in preparadas:
        caminho = os.path.abspath(str((preparada or {}).get("caminho") or ""))
        base, extensao = os.path.splitext(caminho)
        if not caminho or extensao.casefold() not in CADASTRO_FOTOS_EXTENSOES:
            raise HTTPException(status_code=500, detail="Destino de imagem invalido.")
        for candidato in (caminho, *(f"{base}{ext}" for ext in CADASTRO_FOTOS_EXTENSOES)):
            chave = os.path.normcase(os.path.realpath(candidato))
            if chave not in vistos:
                vistos.add(chave)
                caminhos.append(candidato)
    return caminhos


def _remover_variantes_fotos_obsoletas(
    preparadas: list[dict[str, Any]],
    caminhos_variantes: list[str],
) -> None:
    alvos = {
        os.path.normcase(os.path.realpath(str(item.get("caminho") or "")))
        for item in preparadas
    }
    for caminho in caminhos_variantes:
        if os.path.normcase(os.path.realpath(caminho)) in alvos:
            continue
        if not os.path.lexists(caminho):
            continue
        if _cadastro_caminho_e_link(caminho) or not os.path.isfile(caminho):
            raise HTTPException(status_code=409, detail="Variante de imagem insegura.")
        os.unlink(caminho)


def _validar_caminhos_variantes_fotos(caminhos: list[str]) -> None:
    for caminho in caminhos:
        if os.path.lexists(caminho) and (
            _cadastro_caminho_e_link(caminho) or not os.path.isfile(caminho)
        ):
            raise HTTPException(status_code=409, detail="Destino de imagem invalido.")


def _salvar_foto_preparada_atomico(preparada: dict[str, Any]) -> str:
    caminho_arquivo = str((preparada or {}).get("caminho") or "")
    caminho_relativo = str((preparada or {}).get("relativo") or "")
    conteudo = (preparada or {}).get("conteudo")
    if not caminho_arquivo or not caminho_relativo or not isinstance(conteudo, bytes):
        raise HTTPException(status_code=500, detail="Imagem preparada invalida.")

    pasta_fotos = os.path.dirname(caminho_arquivo)
    os.makedirs(pasta_fotos, exist_ok=True)
    temporario = ""
    try:
        with tempfile.NamedTemporaryFile(
            mode="wb",
            prefix=".cadastro_foto.",
            suffix=".tmp",
            dir=pasta_fotos,
            delete=False,
        ) as arquivo:
            temporario = arquivo.name
            arquivo.write(conteudo)
            arquivo.flush()
            os.fsync(arquivo.fileno())
        os.replace(temporario, caminho_arquivo)
        temporario = ""
    finally:
        if temporario:
            try:
                os.unlink(temporario)
            except OSError:
                pass
    return caminho_relativo


def _salvar_fotos_preparadas_atomico(
    preparadas: list[dict[str, Any]],
) -> list[str]:
    """Grava um fan-out completo e restaura todos os destinos em caso de falha."""

    if not preparadas:
        return []
    caminhos_alvo = [str(item.get("caminho") or "") for item in preparadas]
    if any(not caminho for caminho in caminhos_alvo):
        raise HTTPException(status_code=500, detail="Destino de imagem invalido.")
    chaves = [os.path.normcase(os.path.realpath(caminho)) for caminho in caminhos_alvo]
    if len(set(chaves)) != len(chaves):
        raise HTTPException(
            status_code=400,
            detail="O compartilhamento gerou destinos de imagem duplicados.",
        )

    caminhos = _cadastro_caminhos_variantes_fotos_preparadas(preparadas)
    with path_locks_for(caminhos):
        _validar_caminhos_variantes_fotos(caminhos)
        estados: dict[str, bytes | None] = {}
        for caminho in caminhos:
            if os.path.isfile(caminho):
                with open(caminho, "rb") as arquivo:
                    estados[caminho] = arquivo.read()
            else:
                estados[caminho] = None
        try:
            relativos = [_salvar_foto_preparada_atomico(item) for item in preparadas]
            _remover_variantes_fotos_obsoletas(preparadas, caminhos)
            return relativos
        except BaseException:
            for caminho in reversed(caminhos):
                anterior = estados[caminho]
                if anterior is None:
                    try:
                        if os.path.isfile(caminho):
                            os.unlink(caminho)
                    except OSError:
                        logger.exception("[CADASTRO] Falha ao remover replica incompleta.")
                else:
                    try:
                        _salvar_foto_preparada_atomico(
                            {
                                "caminho": caminho,
                                "relativo": caminho,
                                "conteudo": anterior,
                            }
                        )
                    except Exception:
                        logger.exception("[CADASTRO] Falha ao restaurar foto apos rollback.")
            raise


def _salvar_foto_data_url_no_tenant(
    client_id: str,
    sku: str,
    foto_data_url: str,
    foto_filename: str = "",
    store_id: str | None = None,
) -> str:
    """Salva uma imagem enviada como data URL e retorna o caminho relativo para o campo foto."""
    preparadas = _preparar_fotos_data_url_no_tenant(
        client_id,
        sku,
        foto_data_url,
        foto_filename,
        store_id=store_id,
    )
    if not preparadas:
        return ""
    with _cadastro_fotos_bloquear_transicao(client_id):
        _cadastro_fotos_validar_preparadas_no_lock(client_id, preparadas)
        relativos = _salvar_fotos_preparadas_atomico(preparadas)
    return relativos[0] if relativos else ""


def _cadastro_mapa_fotos_locais(client_id: str, store_id: str | None = None) -> dict[str, str]:
    """Mapeia fotos jÃ¡ salvas em cadastro_fotos pelo SKU do nome do arquivo."""
    store_exato = _cadastro_store_id_foto_exato(store_id)
    store_seguro = _cadastro_store_id_foto_segmento(store_exato)
    escopo_estrito = _cadastro_fotos_escopo_estrito(client_id)
    tenant_path = _cadastro_tenant_path_fotos_seguro(client_id)
    if not tenant_path:
        return {}
    candidatos_dir: list[tuple[str, str, str]] = []
    if store_seguro:
        candidatos_dir.append((
            os.path.join(tenant_path, "cadastro_fotos", "lojas", store_seguro),
            f"cadastro_fotos/lojas/{store_seguro}",
            tenant_path,
        ))
    if not escopo_estrito and _cadastro_store_id_foto_legado_permitido(client_id, store_exato):
        candidatos_dir.append((
            os.path.join(tenant_path, "cadastro_fotos", "lojas", store_exato),
            f"cadastro_fotos/lojas/{store_exato}",
            tenant_path,
        ))
    if not escopo_estrito:
        candidatos_dir.extend([
            (os.path.join(tenant_path, "cadastro_fotos"), "cadastro_fotos", tenant_path),
            (
                os.path.join(PASTA_INFO, "default", "cadastro_fotos"),
                "cadastro_fotos",
                os.path.join(PASTA_INFO, "default"),
            ),
        ])
    mapa: dict[str, str] = {}
    extensoes = set(CADASTRO_FOTOS_EXTENSOES)

    for pasta, prefixo_relativo, raiz_autorizada in candidatos_dir:
        if (
            not _cadastro_diretorio_seguro_abaixo(raiz_autorizada, pasta)
            or not os.path.isdir(pasta)
        ):
            continue
        try:
            arquivos = sorted(
                os.listdir(pasta),
                key=lambda nome: (os.path.splitext(nome)[1].lower() != ".png", nome.lower()),
            )
        except Exception:
            continue

        for nome in arquivos:
            caminho_absoluto = os.path.join(pasta, nome)
            if _cadastro_caminho_e_link(caminho_absoluto) or not os.path.isfile(
                caminho_absoluto
            ):
                continue
            base, ext = os.path.splitext(nome)
            if ext.lower() not in extensoes:
                continue
            sku_norm = _normalizar_sku_mes(base)
            if not sku_norm:
                continue
            caminho_relativo = f"{prefixo_relativo}/{nome}"
            mapa.setdefault(f"exact:{sku_norm}", caminho_relativo)
            mapa.setdefault(f"soft:{_cadastro_sku_sem_zeros_blocos(sku_norm)}", caminho_relativo)

    return mapa

def _cadastro_resolver_foto_local(mapa_fotos: dict[str, str], sku: str) -> str:
    sku_norm = _normalizar_sku_mes(str(sku or "").strip())
    if not sku_norm:
        return ""
    return (
        mapa_fotos.get(f"exact:{sku_norm}")
        or mapa_fotos.get(f"soft:{_cadastro_sku_sem_zeros_blocos(sku_norm)}")
        or ""
    )

async def upload_foto_cadastro(
    sku: str = Form(...),
    arquivo: UploadFile = File(...),
    client_id: str = Depends(get_tenant_id)
):
    sku_norm = _normalizar_sku_mes(sku)
    if not sku_norm:
        raise HTTPException(status_code=400, detail="SKU invalido para upload da imagem.")
    if _cadastro_fotos_escopo_estrito(client_id):
        raise HTTPException(
            status_code=409,
            detail={
                "code": "store_id_required",
                "message": "Selecione a loja para salvar a foto do produto.",
            },
        )

    # Esta rota historica nao possui identidade de loja. Um upload para SKU ja
    # materializado por loja viraria fallback global e apareceria em todas as
    # lojas sem foto propria, portanto deve falhar antes de ler ou gravar bytes.
    from backend.services.cadastro_compatibilidade import (
        bloquear_mutacao_legada_sem_sku_controlado,
        exigir_mutacao_legada_sem_sku_controlado,
    )

    exigir_mutacao_legada_sem_sku_controlado(client_id, [sku_norm])

    if not arquivo:
        raise HTTPException(status_code=400, detail="Arquivo de imagem nÃ£o informado.")

    nome_original = str(getattr(arquivo, "filename", "") or "").strip()
    ext = os.path.splitext(nome_original)[1].lower()
    if not ext:
        ctype = str(getattr(arquivo, "content_type", "") or "").lower()
        if "png" in ctype:
            ext = ".png"
        elif "jpeg" in ctype or "jpg" in ctype:
            ext = ".jpg"
        elif "webp" in ctype:
            ext = ".webp"
        elif "gif" in ctype:
            ext = ".gif"

    if ext not in {".png", ".jpg", ".jpeg", ".webp", ".gif"}:
        raise HTTPException(status_code=400, detail="Formato invalido. Use PNG, JPG, WEBP ou GIF.")

    try:
        conteudo = await arquivo.read()
        if not conteudo:
            raise HTTPException(status_code=400, detail="Imagem vazia.")

        with (
            bloquear_mutacao_legada_sem_sku_controlado(client_id, [sku_norm]),
            _cadastro_fotos_bloquear_mutacao_global(client_id),
        ):
            pasta_fotos, _prefixo = _cadastro_pasta_fotos(client_id)
            ext_limpa = ext.replace(".", "")
            nome_arquivo = _nome_arquivo_foto_sku(sku_norm, ext_limpa)
            caminho_arquivo = os.path.join(pasta_fotos, nome_arquivo)
            caminho_relativo = f"cadastro_fotos/{nome_arquivo}"
            _salvar_fotos_preparadas_atomico(
                [
                    {
                        "caminho": caminho_arquivo,
                        "relativo": caminho_relativo,
                        "conteudo": conteudo,
                    }
                ]
            )
            return {
                "success": True,
                "foto": caminho_relativo,
                "url": f"/api/cadastro/foto/{quote_plus(str(client_id))}/{quote_plus(nome_arquivo)}"
            }
    except HTTPException:
        raise
    except Exception as e:
        raise HTTPException(status_code=500, detail=f"Erro ao salvar imagem do cadastro: {str(e)}")

def _cadastro_carregar_lojas_foto(client_id: str) -> list[dict[str, Any]]:
    # Importacao tardia evita o ciclo cadastro_lojas_produtos -> cadastro_fotos.
    from backend.services import integracoes

    return [
        dict(loja)
        for loja in (integracoes.carregar_lojas(client_id) or [])
        if isinstance(loja, dict)
    ]


def _cadastro_store_ids_foto_atuais(client_id: str) -> list[str]:
    ids: list[str] = []
    vistos: set[str] = set()
    for loja in _cadastro_carregar_lojas_foto(client_id):
        store_id = _cadastro_store_id_foto_exato(loja.get("store_id"))
        if store_id and store_id not in vistos:
            vistos.add(store_id)
            ids.append(store_id)
    return ids


class CadastroFotosConfigInvalida(ValueError):
    """A configuracao existe, mas nao pode autorizar fallback ou fan-out."""


def _cadastro_fotos_config_path(client_id: str) -> str:
    return os.path.join(get_tenant_path(client_id), CADASTRO_FOTOS_CONFIG_ARQUIVO)


def _cadastro_fotos_config_carregar(client_id: str) -> dict[str, Any]:
    tenant_path = _cadastro_tenant_path_fotos_seguro(client_id)
    if not tenant_path:
        raise CadastroFotosConfigInvalida("tenant_inseguro")
    caminho = os.path.join(tenant_path, CADASTRO_FOTOS_CONFIG_ARQUIVO)
    if not _cadastro_diretorio_seguro_abaixo(
        tenant_path,
        os.path.dirname(caminho),
    ):
        raise CadastroFotosConfigInvalida("diretorio_inseguro")
    if not os.path.lexists(caminho):
        return {
            "exists": False,
            "strict_store_scope": False,
            "shared_groups": [],
        }
    if _cadastro_caminho_e_link(caminho) or not os.path.isfile(caminho):
        raise CadastroFotosConfigInvalida("entrada_insegura")
    try:
        with path_locks_for([caminho]):
            if os.path.getsize(caminho) > 262_144:
                raise CadastroFotosConfigInvalida("arquivo_excede_limite")
            with open(caminho, "r", encoding="utf-8-sig") as arquivo:
                bruto = json.load(arquivo)
    except CadastroFotosConfigInvalida:
        raise
    except Exception as exc:
        raise CadastroFotosConfigInvalida("arquivo_invalido") from exc

    if not isinstance(bruto, dict) or bruto.get("schema") != CADASTRO_FOTOS_CONFIG_SCHEMA:
        raise CadastroFotosConfigInvalida("schema_invalido")
    estrito = bruto.get("strict_store_scope")
    grupos = bruto.get("shared_groups")
    if not isinstance(estrito, bool) or not isinstance(grupos, list):
        raise CadastroFotosConfigInvalida("campos_invalidos")

    store_ids_atuais = set(_cadastro_store_ids_foto_atuais(client_id))
    usados: set[str] = set()
    grupos_normalizados: list[dict[str, Any]] = []
    for indice, grupo in enumerate(grupos):
        if not isinstance(grupo, dict):
            raise CadastroFotosConfigInvalida("grupo_invalido")
        group_id = str(grupo.get("group_id") or "").strip()
        store_ids = [
            _cadastro_store_id_foto_exato(item)
            for item in (grupo.get("store_ids") or [])
        ] if isinstance(grupo.get("store_ids"), list) else []
        if (
            not group_id
            or len(store_ids) < 2
            or any(not item for item in store_ids)
            or len(set(store_ids)) != len(store_ids)
            or any(item not in store_ids_atuais for item in store_ids)
            or any(item in usados for item in store_ids)
        ):
            raise CadastroFotosConfigInvalida(f"grupo_{indice}_invalido")
        usados.update(store_ids)
        grupos_normalizados.append({"group_id": group_id, "store_ids": store_ids})

    return {
        "exists": True,
        "strict_store_scope": estrito,
        "shared_groups": grupos_normalizados,
    }


def _cadastro_fotos_escopo_estrito(client_id: str) -> bool:
    try:
        return bool(_cadastro_fotos_config_carregar(client_id)["strict_store_scope"])
    except CadastroFotosConfigInvalida:
        logger.warning("[CADASTRO] Configuracao de fotos invalida; fallback legado bloqueado.")
        return True


def _cadastro_foto_cabecalho_normalizar(valor: object) -> str:
    texto = unicodedata.normalize("NFKD", str(valor or "").strip())
    texto = "".join(ch for ch in texto if not unicodedata.combining(ch))
    return re.sub(r"[^a-z0-9]+", "_", texto.casefold()).strip("_")


def _cadastro_foto_coluna_candidata(valor: object) -> bool:
    cabecalho = _cadastro_foto_cabecalho_normalizar(valor)
    return any(
        marcador in cabecalho
        for marcador in (
            "foto",
            "imagem",
            "image",
            "picture",
            "thumbnail",
            "thumb",
        )
    )


def _cadastro_foto_referencia_limpar_wrappers(valor: object) -> str:
    texto = str(valor or "").strip()
    return re.sub(r'''^['"`‘’“”]+|['"`‘’“”]+$''', "", texto).strip()


def _cadastro_foto_referencia_local_cadastro(foto: object) -> bool:
    """Reconhece referencias locais de foto sem confundir URLs externas."""

    caminho = _cadastro_foto_referencia_limpar_wrappers(foto).replace("\\", "/")
    if not caminho or re.match(r"^data:", caminho, re.IGNORECASE):
        return False
    for _ in range(3):
        decodificado = unquote(caminho)
        if decodificado == caminho:
            break
        caminho = _cadastro_foto_referencia_limpar_wrappers(decodificado).replace(
            "\\", "/"
        )

    def _rota_api_local(valor: str) -> bool:
        normalizado = posixpath.normpath("/" + valor.replace("\\", "/").lstrip("/"))
        normalizado_fold = normalizado.casefold()
        return normalizado_fold.startswith("/api/cadastro/foto/") or normalizado_fold.startswith(
            "/api/cadastro/foto-arquivo/"
        )

    caminho_windows = bool(re.match(r"^[a-z]:/", caminho, re.IGNORECASE))
    if caminho.startswith("//"):
        try:
            url = urlsplit(f"https:{caminho}")
        except ValueError:
            return True
        return _rota_api_local(str(url.path or ""))
    esquema = re.match(r"^([a-z][a-z0-9+.-]*):", caminho, re.IGNORECASE)
    if esquema and not caminho_windows:
        nome_esquema = esquema.group(1).casefold()
        if nome_esquema in {"http", "https"}:
            if re.match(r"^https?://", caminho, re.IGNORECASE):
                try:
                    url = urlsplit(caminho)
                except ValueError:
                    return True
                return _rota_api_local(str(url.path or ""))
            return _rota_api_local(caminho.split(":", 1)[1])
        if nome_esquema != "file":
            return False
    caminho_file = caminho.casefold().startswith("file:")
    if caminho_file:
        try:
            caminho = str(urlsplit(caminho).path or "").replace("\\", "/")
        except ValueError:
            return True

    caminho_original = caminho.split("?", 1)[0].split("#", 1)[0]
    caminho = posixpath.normpath("/" + caminho_original.lstrip("/")).strip("/")
    caminho_fold = caminho.casefold()
    if caminho_fold.startswith("api/cadastro/foto/") or caminho_fold.startswith(
        "api/cadastro/foto-arquivo/"
    ):
        return True
    if caminho_fold.startswith("cadastro_fotos/") or caminho_fold.startswith("lojas/"):
        return True
    if "/cadastro_fotos/" in caminho_fold:
        return True
    nome = caminho.rsplit("/", 1)[-1]
    return os.path.splitext(nome)[1].casefold() in set(CADASTRO_FOTOS_EXTENSOES)


def _cadastro_fotos_exigir_mutacao_global_permitida(
    client_id: str,
    tenant_path: os.PathLike[str] | str | None = None,
) -> None:
    """Bloqueia qualquer writer legado quando a separacao por loja esta ativa."""

    bloqueado = False
    if tenant_path is not None:
        tenant_confiavel = _cadastro_tenant_path_fotos_seguro(client_id)
        tenant_recebido = os.path.abspath(os.fspath(tenant_path))
        bloqueado = not tenant_confiavel or os.path.normcase(
            os.path.normpath(os.path.realpath(tenant_confiavel))
        ) != os.path.normcase(os.path.normpath(os.path.realpath(tenant_recebido)))
    if bloqueado or _cadastro_fotos_escopo_estrito(client_id):
        raise HTTPException(
            status_code=409,
            detail={
                "code": "store_id_required",
                "message": (
                    "A gravacao global de fotos foi desativada. Informe o store_id."
                ),
                "skus": [],
            },
        )


@contextmanager
def _cadastro_fotos_bloquear_transicao(
    client_id: str,
    tenant_path: os.PathLike[str] | str | None = None,
    *,
    timeout_seconds: float = 10.0,
):
    tenant_confiavel = _cadastro_tenant_path_fotos_seguro(client_id)
    if not tenant_confiavel:
        raise HTTPException(status_code=409, detail="Diretorio do cliente inseguro.")
    tenant_lock_path = os.path.realpath(os.path.abspath(tenant_confiavel))
    if tenant_path is not None:
        tenant_recebido = os.path.abspath(os.fspath(tenant_path))
        tenant_runtime = os.path.abspath(str(get_tenant_path(client_id) or ""))
        chaves_lexicais_permitidas = {
            os.path.normcase(os.path.normpath(tenant_lock_path)),
            os.path.normcase(os.path.normpath(tenant_runtime)),
        }
        if (
            os.path.normcase(os.path.normpath(tenant_recebido))
            not in chaves_lexicais_permitidas
            or os.path.normcase(os.path.normpath(os.path.realpath(tenant_recebido)))
            != os.path.normcase(os.path.normpath(tenant_lock_path))
        ):
            raise HTTPException(status_code=409, detail="Diretorio do cliente inseguro.")
    try:
        with bloquear_transicao_fotos_tenant(
            tenant_lock_path,
            timeout_seconds=timeout_seconds,
        ):
            yield
    except CadastroFotosCoordenacaoErro as exc:
        indisponivel = exc.code != "locked"
        raise HTTPException(
            status_code=409,
            detail={
                "code": (
                    "cadastro_photo_transition_unavailable"
                    if indisponivel
                    else "cadastro_photo_transition_busy"
                ),
                "message": (
                    "Nao foi possivel coordenar as fotos deste cliente."
                    if indisponivel
                    else "As fotos deste cliente estao sendo atualizadas."
                ),
            },
        ) from exc


@contextmanager
def _cadastro_fotos_bloquear_mutacao_global(
    client_id: str,
    tenant_path: os.PathLike[str] | str | None = None,
    *,
    enabled: bool = True,
):
    if not enabled:
        yield
        return
    with _cadastro_fotos_bloquear_transicao(client_id, tenant_path):
        _cadastro_fotos_exigir_mutacao_global_permitida(client_id, tenant_path)
        yield


def _cadastro_fotos_validar_preparadas_no_lock(
    client_id: str,
    preparadas: list[dict[str, Any]],
) -> None:
    if not preparadas:
        return
    if any(not _cadastro_store_id_foto_exato(item.get("store_id")) for item in preparadas):
        _cadastro_fotos_exigir_mutacao_global_permitida(client_id)
        return

    grupos: dict[tuple[str, str], set[str]] = {}
    for item in preparadas:
        caminho = str(item.get("caminho") or "")
        conteudo = item.get("conteudo")
        store_id = _cadastro_store_id_foto_exato(item.get("store_id"))
        if not caminho or not isinstance(conteudo, bytes) or not store_id:
            raise HTTPException(status_code=500, detail="Imagem preparada invalida.")
        chave = (os.path.basename(caminho).casefold(), hashlib.sha256(conteudo).hexdigest())
        grupos.setdefault(chave, set()).add(store_id)

    for store_ids in grupos.values():
        origem = sorted(store_ids)[0]
        esperados = set(_cadastro_fotos_store_ids_compartilhados(client_id, origem))
        if store_ids != esperados:
            raise HTTPException(
                status_code=409,
                detail={
                    "code": "photo_scope_changed",
                    "message": "O compartilhamento de fotos mudou; tente salvar novamente.",
                },
            )


def _cadastro_fotos_store_ids_compartilhados(
    client_id: str,
    store_id: str,
) -> list[str]:
    store_exato = _cadastro_store_id_foto_exato(store_id)
    if not store_exato:
        return []
    try:
        config = _cadastro_fotos_config_carregar(client_id)
    except CadastroFotosConfigInvalida as exc:
        raise HTTPException(
            status_code=409,
            detail={
                "code": "cadastro_photo_config_invalid",
                "message": "A configuracao de fotos por loja e invalida.",
            },
        ) from exc
    if not config["exists"]:
        return [store_exato]
    for grupo in config["shared_groups"]:
        membros = list(grupo["store_ids"])
        if store_exato in membros:
            return [store_exato, *(item for item in membros if item != store_exato)]
    if store_exato not in set(_cadastro_store_ids_foto_atuais(client_id)):
        raise HTTPException(
            status_code=409,
            detail={
                "code": "store_config_changed",
                "message": "A identidade da loja nao existe mais.",
            },
        )
    return [store_exato]


def _cadastro_store_id_foto_legado_permitido(
    client_id: str,
    store_id: str | None,
    segmento: str | None = None,
) -> bool:
    store_exato = _cadastro_store_id_foto_exato(store_id)
    segmento_exato = _cadastro_store_id_foto_exato(segmento) if segmento is not None else store_exato
    if (
        not store_exato
        or segmento_exato != store_exato
        or not _cadastro_store_id_foto_legado_sintaticamente_seguro(store_exato)
    ):
        return False
    equivalentes = {
        atual
        for atual in _cadastro_store_ids_foto_atuais(client_id)
        if atual.casefold() == store_exato.casefold()
    }
    return store_exato in equivalentes and len(equivalentes) == 1


def _cadastro_store_id_por_segmento_foto(client_id: str, segmento: str) -> str:
    segmento_exato = str(segmento or "")
    store_ids = _cadastro_store_ids_foto_atuais(client_id)
    if re.fullmatch(r"sid-[a-f0-9]{64}", segmento_exato):
        correspondentes = [
            store_id
            for store_id in store_ids
            if _cadastro_store_id_foto_segmento(store_id) == segmento_exato
        ]
        return correspondentes[0] if len(correspondentes) == 1 else ""
    if not _cadastro_store_id_foto_legado_sintaticamente_seguro(segmento_exato):
        return ""
    return (
        segmento_exato
        if _cadastro_store_id_foto_legado_permitido(
            client_id,
            segmento_exato,
            segmento_exato,
        )
        else ""
    )


def _cadastro_foto_partes_loja_referencia_local(
    foto: str,
    client_id: str = "",
) -> list[str] | None:
    """Return local store-photo parts, an invalid sentinel, or ``None``.

    Persisted rows normally use ``cadastro_fotos/lojas/...``.  Older callers
    may submit one of the authenticated local API URLs; those URLs must be
    checked against the same exact store identity instead of being mistaken
    for an unrestricted external reference.
    """

    caminho = str(foto or "").strip().replace("\\", "/")
    if not caminho or re.match(r"^(?:data|blob):", caminho, re.IGNORECASE):
        return None

    # Decode repeatedly to catch local API paths hidden by one or two layers
    # of URL encoding.  Browsers/proxies may decode before routing, so treating
    # such a reference as an unrestricted external URL would cross store scope.
    for _ in range(3):
        decodificado = unquote(caminho)
        if decodificado == caminho:
            break
        caminho = decodificado.replace("\\", "/")

    if caminho.startswith("//") or re.match(r"^https?:", caminho, re.IGNORECASE):
        try:
            caminho_url = urlsplit(caminho if not caminho.startswith("//") else f"https:{caminho}")
        except ValueError:
            return None
        caminho_local = str(caminho_url.path or "").replace("\\", "/")
        while caminho_local.startswith("//"):
            caminho_local = caminho_local[1:]
        caminho_local = posixpath.normpath(caminho_local)
        caminho_local_fold = caminho_local.casefold()
        if not (
            caminho_local_fold.startswith("/api/cadastro/foto-arquivo/")
            or caminho_local_fold.startswith("/api/cadastro/foto/")
        ):
            return None
        caminho = caminho_local

    caminho = caminho.split("?", 1)[0].split("#", 1)[0]
    caminho_fold = caminho.casefold()
    prefixo_arquivo = "/api/cadastro/foto-arquivo/"
    prefixo_tenant = "/api/cadastro/foto/"
    if caminho_fold.startswith(prefixo_arquivo):
        caminho = caminho[len(prefixo_arquivo):]
    elif caminho_fold.startswith(prefixo_tenant):
        partes_api = caminho[len(prefixo_tenant):].strip("/").split("/")
        indice_lojas = next(
            (
                indice
                for indice, parte in enumerate(partes_api)
                if parte.casefold() == "lojas"
            ),
            -1,
        )
        # O endpoint autenticado carrega o tenant antes do caminho da foto.
        # Nao basta validar o hash/store da pasta: uma URL persistida por outro
        # tenant deve continuar invalida mesmo que use o mesmo store_id.
        if indice_lojas < 1:
            return [] if indice_lojas == 0 else None
        tenant_referencia = str(partes_api[0] or "")
        if str(client_id or "") and tenant_referencia != str(client_id):
            return []
        caminho = "/".join(partes_api[indice_lojas:])
    elif caminho.startswith("/"):
        return [] if "/lojas/" in caminho_fold else None

    caminho = caminho.strip("/")
    if caminho.casefold().startswith("cadastro_fotos/"):
        caminho = caminho.split("/", 1)[1]
    partes = caminho.split("/")
    if partes and partes[0].casefold() == "lojas":
        return partes
    if "cadastro_fotos/lojas/" in caminho.casefold():
        return []
    return None


def _cadastro_foto_store_id_referencia(client_id: str, foto: str) -> str:
    partes = _cadastro_foto_partes_loja_referencia_local(foto, client_id)
    if partes is None or len(partes) != 3:
        return ""
    if any(not parte or parte in {".", ".."} for parte in partes):
        return ""
    return _cadastro_store_id_por_segmento_foto(client_id, partes[1])


def _cadastro_foto_referencia_pertence_loja(
    client_id: str,
    store_id: str,
    foto: str,
) -> bool:
    partes = _cadastro_foto_partes_loja_referencia_local(foto, client_id)
    if partes is None:
        return True
    return (
        _cadastro_foto_store_id_referencia(client_id, foto)
        == _cadastro_store_id_foto_exato(store_id)
    )


def _cadastro_foto_referencias_candidatas(
    registro: object,
) -> list[tuple[str, object]]:
    itens = getattr(registro, "items", None)
    if not callable(itens):
        return []
    return [
        (str(chave or ""), valor)
        for chave, valor in itens()
        if _cadastro_foto_coluna_candidata(chave)
    ]


def _cadastro_foto_referencias_invalidas_loja(
    client_id: str,
    store_id: str,
    registro: object,
) -> list[str]:
    """Lista campos locais que escapariam do escopo exato da loja."""

    estrito = _cadastro_fotos_escopo_estrito(client_id)
    invalidas: list[str] = []
    for campo, referencia in _cadastro_foto_referencias_candidatas(registro):
        if not _cadastro_foto_referencia_local_cadastro(referencia):
            continue
        partes = _cadastro_foto_partes_loja_referencia_local(
            str(referencia or ""),
            client_id,
        )
        if partes is None:
            if estrito:
                invalidas.append(campo)
            continue
        if not _cadastro_foto_referencia_pertence_loja(
            client_id,
            store_id,
            str(referencia or ""),
        ):
            invalidas.append(campo)
    return invalidas


def _cadastro_partes_foto_seguras(
    client_id: str,
    filename: str,
    store_id_esperado: str = "",
) -> list[str]:
    nome_original = str(filename or "").replace("\\", "/").strip("/")
    if nome_original.lower().startswith("cadastro_fotos/"):
        nome_original = nome_original.split("/", 1)[1]

    partes = nome_original.split("/")
    if (
        not partes
        or any(not parte or parte in {".", ".."} for parte in partes)
        or any(
            ":" in parte
            or "\x00" in parte
            or any(ord(char) < 32 for char in parte)
            for parte in partes
        )
    ):
        raise HTTPException(status_code=404, detail="Arquivo de foto invalido")

    if len(partes) == 1:
        return partes
    if len(partes) != 3 or partes[0].lower() != "lojas":
        raise HTTPException(status_code=404, detail="Arquivo de foto invalido")

    store_id = _cadastro_store_id_por_segmento_foto(client_id, partes[1])
    if not store_id:
        raise HTTPException(status_code=404, detail="Loja nao encontrada para este cliente.")
    esperado = _cadastro_store_id_foto_exato(store_id_esperado)
    if esperado and store_id != esperado:
        raise HTTPException(status_code=404, detail="Foto nao pertence a loja solicitada.")
    return partes


def _cadastro_resolver_arquivo_foto(
    client_id: str,
    filename: str,
    store_id_esperado: str = "",
) -> str:
    partes = _cadastro_partes_foto_seguras(client_id, filename, store_id_esperado)
    if _cadastro_fotos_escopo_estrito(client_id) and partes[0].casefold() != "lojas":
        raise HTTPException(status_code=404, detail="Foto legada indisponivel neste cliente.")

    def _candidato_seguro(raiz_autorizada: str, base: str) -> str | None:
        candidato_lexico = os.path.join(base, *partes)
        if not _cadastro_diretorio_seguro_abaixo(
            raiz_autorizada,
            os.path.dirname(candidato_lexico),
        ):
            return None
        base_real = os.path.realpath(base)
        candidato = os.path.realpath(candidato_lexico)
        try:
            if os.path.commonpath([base_real, candidato]) != base_real:
                return None
        except ValueError:
            return None
        if _cadastro_caminho_e_link(candidato_lexico):
            return None
        return candidato

    tenant_path = _cadastro_tenant_path_fotos_seguro(client_id)
    if not tenant_path:
        raise HTTPException(status_code=404, detail="Diretorio do cliente inseguro.")
    tenant_fotos = os.path.join(tenant_path, "cadastro_fotos")
    candidatos = [
        _candidato_seguro(tenant_path, tenant_fotos),
    ]
    # O fallback historico em ``default`` so vale para fotos legadas na raiz.
    # Uma referencia ``lojas/<segmento>/<arquivo>`` ja declara identidade de
    # loja e nunca pode ser satisfeita por outro escopo quando o arquivo do
    # tenant estiver ausente.
    if partes[0].casefold() != "lojas":
        default_path = os.path.join(PASTA_INFO, "default")
        candidatos.append(
            _candidato_seguro(
                default_path,
                os.path.join(default_path, "cadastro_fotos"),
            )
        )

    caminho_arquivo = next((p for p in candidatos if p and os.path.isfile(p)), None)
    if not caminho_arquivo:
        raise HTTPException(status_code=404, detail="Foto nÃ£o encontrada")
    return caminho_arquivo


def _cadastro_resposta_foto(caminho_arquivo: str) -> FileResponse:
    if _cadastro_caminho_e_link(caminho_arquivo) or not os.path.isfile(caminho_arquivo):
        raise HTTPException(status_code=404, detail="Foto nao encontrada")
    response = FileResponse(caminho_arquivo)
    response.headers["Cache-Control"] = "no-cache, no-store, must-revalidate"
    response.headers["Pragma"] = "no-cache"
    response.headers["Expires"] = "0"
    return response


def _cadastro_tenant_foto_autenticado(client_id_url: str, client_id_sessao: str) -> str:
    tenant_sessao = str(client_id_sessao or "").strip()
    if not tenant_sessao:
        raise HTTPException(status_code=401, detail="Autenticacao necessaria para acessar a foto.")
    if str(client_id_url or "").strip() != tenant_sessao:
        raise HTTPException(status_code=403, detail="Acesso negado a foto de outro cliente.")
    return tenant_sessao


async def servir_foto_cadastro(
    client_id: str,
    filename: str,
    client_id_sessao: str = Depends(get_tenant_id),
    store_id: str = "",
):
    tenant_sessao = _cadastro_tenant_foto_autenticado(client_id, client_id_sessao)
    caminho_arquivo = _cadastro_resolver_arquivo_foto(
        tenant_sessao,
        filename,
        store_id,
    )
    return _cadastro_resposta_foto(caminho_arquivo)


async def servir_foto_cadastro_por_arquivo(
    filename: str,
    client_id_sessao: str = Depends(get_tenant_id),
):
    tenant_sessao = str(client_id_sessao or "").strip()
    if not tenant_sessao:
        raise HTTPException(status_code=401, detail="Autenticacao necessaria para acessar a foto.")
    caminho_arquivo = _cadastro_resolver_arquivo_foto(tenant_sessao, filename)
    return _cadastro_resposta_foto(caminho_arquivo)

__all__ = [
    "CADASTRO_FOTOS_CONFIG_ARQUIVO",
    "CADASTRO_FOTOS_CONFIG_SCHEMA",
    "CadastroFotosConfigInvalida",
    "_exportar_planilha_xlsx_bytes",
    "_extrair_imagens_planilha_por_sku",
    "_nome_arquivo_foto_sku",
    "_cadastro_sku_sem_zeros_blocos",
    "_cadastro_store_id_foto_exato",
    "_cadastro_store_id_foto_segmento",
    "_cadastro_store_id_foto_seguro",
    "_cadastro_store_id_foto_legado_sintaticamente_seguro",
    "_cadastro_store_id_foto_legado_permitido",
    "_cadastro_store_id_por_segmento_foto",
    "_cadastro_fotos_config_path",
    "_cadastro_fotos_config_carregar",
    "_cadastro_fotos_escopo_estrito",
    "_cadastro_foto_cabecalho_normalizar",
    "_cadastro_foto_coluna_candidata",
    "_cadastro_foto_referencia_limpar_wrappers",
    "_cadastro_foto_referencia_local_cadastro",
    "_cadastro_fotos_exigir_mutacao_global_permitida",
    "_cadastro_fotos_bloquear_transicao",
    "_cadastro_fotos_bloquear_mutacao_global",
    "_cadastro_fotos_validar_preparadas_no_lock",
    "_cadastro_fotos_store_ids_compartilhados",
    "_cadastro_foto_store_id_referencia",
    "_cadastro_foto_referencia_pertence_loja",
    "_cadastro_foto_referencias_candidatas",
    "_cadastro_foto_referencias_invalidas_loja",
    "_preparar_foto_data_url_no_tenant",
    "_preparar_fotos_data_url_no_tenant",
    "_cadastro_caminhos_variantes_fotos_preparadas",
    "_validar_caminhos_variantes_fotos",
    "_salvar_foto_preparada_atomico",
    "_salvar_fotos_preparadas_atomico",
    "_salvar_foto_data_url_no_tenant",
    "_cadastro_mapa_fotos_locais",
    "_cadastro_resolver_foto_local",
    "upload_foto_cadastro",
    "servir_foto_cadastro",
    "servir_foto_cadastro_por_arquivo",
    "configure_cadastro_fotos_runtime",
]
