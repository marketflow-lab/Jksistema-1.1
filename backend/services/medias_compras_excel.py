"""Excel helpers for Medias Compras pedido lists."""

from __future__ import annotations

import io
import os
import re
import unicodedata
from datetime import datetime

import openpyxl
from openpyxl.drawing.image import Image as XLImage
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

from backend.services import medias_compras_common as medias_common
from backend.services.medias_compras_common import *
from backend.services.runtime_bridge import bind_runtime_globals

_RUNTIME_NAMES = (
    "logger",
    "get_tenant_path",
    "TEMP_FILES_STORAGE",
    "TEMP_FILES_META",
    "LISTA_PEDIDO_XLSX_CACHE",
    "LISTA_PEDIDO_XLSX_CACHE_MAX_ITENS",
    "PASTA_INFO",
)


def _sync_common_names() -> None:
    for name in medias_common.COMMON_EXPORTS:
        globals()[name] = getattr(medias_common, name)
    for name in _RUNTIME_NAMES:
        globals()[name] = getattr(medias_common, name)


def _configure_runtime_globals(runtime_module=None):
    runtime = medias_common.configure_medias_compras_common_runtime(runtime_module)
    bind_runtime_globals(globals(), runtime)
    _sync_common_names()
    return runtime

def _normalizar_cabecalho_excel(texto: str) -> str:
    base = unicodedata.normalize("NFKD", str(texto or ""))
    base = "".join(ch for ch in base if not unicodedata.combining(ch))
    return re.sub(r"[^a-z0-9]+", "", base.lower())


def _to_float_excel(v) -> float | None:
    if v is None:
        return None
    if isinstance(v, (int, float)):
        return float(v)
    s = str(v).strip()
    if not s:
        return None
    try:
        s = s.replace("\u00a0", " ")
        s = re.sub(r"(?i)(r\$|us\$|usd|brl|\$)", "", s)
        s = re.sub(r"[^0-9,.\-]+", "", s.replace(" ", ""))
        if s in {"", "-", ".", ","}:
            return None
        if "," in s and "." in s:
            if s.rfind(",") > s.rfind("."):
                s = s.replace(".", "").replace(",", ".")
            else:
                s = s.replace(",", "")
        elif "," in s:
            s = s.replace(",", ".")
        return float(s)
    except Exception:
        return None


def _linha_resumo_excel_sku(sku: str) -> bool:
    txt = str(sku or "").strip()
    if not txt:
        return False
    base = unicodedata.normalize("NFKD", txt)
    base = "".join(ch for ch in base if not unicodedata.combining(ch))
    base = re.sub(r"[^a-z0-9]+", "", base.lower())
    return base in {"total", "totais", "subtotal", "subtotais"}


def _texto_excel_seguro(valor) -> str:
    texto = str(valor or "").strip()
    return texto


def _atribuir_texto_excel(celula, valor) -> None:
    celula.value = _texto_excel_seguro(valor)
    # Forca texto literal para impedir que dados externos iniciados por =, +,
    # - ou @ sejam interpretados como formulas ao abrir o arquivo.
    celula.data_type = "s"


def _formatar_ncm_commercial_invoice(valor) -> str:
    texto = str(valor or "").strip()
    digitos = re.sub(r"\D", "", texto)
    if len(digitos) == 8:
        return f"{digitos[:4]}.{digitos[4:6]}.{digitos[6:]}"
    return texto


def _formatar_data_commercial_invoice(valor) -> str:
    texto = str(valor or "").strip()
    if not texto:
        return ""
    try:
        data = datetime.fromisoformat(texto.replace("Z", "+00:00"))
        return data.strftime("%d/%m/%Y")
    except ValueError:
        return texto


def _resolver_caminho_foto_commercial_invoice(client_id: str | None, foto_ref: str, sku: str) -> str | None:
    diretorios = []
    if client_id:
        diretorios.append(os.path.join(get_tenant_path(client_id), "cadastro_fotos"))
    if PASTA_INFO:
        diretorios.append(os.path.join(PASTA_INFO, "default", "cadastro_fotos"))
    diretorios = [os.path.realpath(diretorio) for diretorio in diretorios if diretorio]

    def _caminho_permitido(caminho: str) -> bool:
        caminho_real = os.path.realpath(caminho)
        for diretorio in diretorios:
            try:
                if os.path.commonpath([caminho_real, diretorio]) == diretorio:
                    return True
            except ValueError:
                continue
        return False

    foto_txt = _resolver_foto_cadastro_sku(client_id, sku, foto_ref)
    foto_txt = str(foto_txt or "").replace("\\", "/").strip()
    if not foto_txt:
        return None
    if os.path.isabs(foto_txt):
        return foto_txt if os.path.isfile(foto_txt) and _caminho_permitido(foto_txt) else None

    nome = foto_txt.split("/", 1)[1] if foto_txt.lower().startswith("cadastro_fotos/") else foto_txt
    nome = os.path.basename(nome)
    if not nome:
        return None

    for diretorio in diretorios:
        candidato = os.path.join(diretorio, nome)
        if os.path.isfile(candidato):
            return candidato
    return None


def _gerar_commercial_invoice_bytes(
    lista: dict,
    *,
    client_id: str | None = None,
    data_aprovacao: str = "",
) -> bytes:
    lista = lista if isinstance(lista, dict) else {}
    itens = [_normalizar_item_lista_pedido(item) for item in (lista.get("itens") or []) if isinstance(item, dict)]

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Commercial Invoice"
    ws.sheet_view.zoomScale = 90
    ws.sheet_view.showGridLines = False
    ws.freeze_panes = "A7"
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.print_title_rows = "1:6"
    ws.sheet_properties.pageSetUpPr.fitToPage = True

    borda = Border(
        left=Side(style="thin", color="000000"),
        right=Side(style="thin", color="000000"),
        top=Side(style="thin", color="000000"),
        bottom=Side(style="thin", color="000000"),
    )
    fonte_base = Font(name="Times New Roman", size=9, color="000000")
    fonte_negrito = Font(name="Times New Roman", size=9, bold=True, color="000000")
    alinhamento_centro = Alignment(horizontal="center", vertical="center", wrap_text=True)

    ws.merge_cells("A1:I1")
    _atribuir_texto_excel(ws["A1"], f"SUPPLIER: {lista.get('supplier') or lista.get('fornecedor') or ''}")
    ws["A1"].font = Font(name="Times New Roman", size=12, bold=True)
    ws["A1"].alignment = alinhamento_centro
    ws.row_dimensions[1].height = 32

    ws.merge_cells("A2:I2")
    _atribuir_texto_excel(ws["A2"], "COMMERCIAL INVOICE")
    ws["A2"].font = Font(name="Times New Roman", size=11, bold=True)
    ws["A2"].alignment = alinhamento_centro
    ws.row_dimensions[2].height = 21

    ws.merge_cells("A3:B3")
    ws.merge_cells("C3:E3")
    ws.merge_cells("F3:G3")
    ws.merge_cells("H3:I3")
    _atribuir_texto_excel(ws["A3"], "INV. No.:")
    _atribuir_texto_excel(ws["C3"], lista.get("numero_invoice") or lista.get("invoice") or "")
    _atribuir_texto_excel(ws["F3"], "APPROVAL DATE:")
    _atribuir_texto_excel(ws["H3"], _formatar_data_commercial_invoice(data_aprovacao))

    ws.merge_cells("A4:E4")
    ws.merge_cells("F4:I4")
    _atribuir_texto_excel(ws["A4"], f"LIST: {lista.get('nome_lista') or ''}")
    loja = str(lista.get("loja") or "").strip()
    _atribuir_texto_excel(ws["F4"], f"STORE: {'' if loja == '__todas' else loja}")

    ws.merge_cells("A5:I5")
    _atribuir_texto_excel(ws["A5"], "Descriptions")
    ws["A5"].font = fonte_negrito
    ws["A5"].alignment = alinhamento_centro

    ws.merge_cells("D6:E6")
    cabecalhos = {
        "A6": "NO",
        "B6": "OE NO",
        "C6": "HSCODE / NCM",
        "D6": "PRODUCT DESCRIPTION",
        "F6": "PIC",
        "G6": "Quantity",
        "H6": "Price Unit",
        "I6": "Amount",
    }
    for endereco, texto in cabecalhos.items():
        _atribuir_texto_excel(ws[endereco], texto)
        ws[endereco].font = fonte_negrito
        ws[endereco].alignment = alinhamento_centro
    ws.row_dimensions[6].height = 26

    for linha in range(1, 7):
        for coluna in range(1, 10):
            celula = ws.cell(linha, coluna)
            celula.border = borda
            if linha in (3, 4):
                celula.font = fonte_negrito
                celula.alignment = alinhamento_centro

    primeira_linha_item = 7
    for indice, item in enumerate(itens, start=1):
        linha = primeira_linha_item + indice - 1
        ws.merge_cells(start_row=linha, start_column=4, end_row=linha, end_column=5)
        quantidade = int(max(0.0, _to_float(item.get("Quantidade"), 0.0)))
        preco = max(0.0, _to_float(item.get("Valor unidade"), 0.0))
        valor_total = max(0.0, _to_float(item.get("Valor total"), 0.0))
        if preco <= 0 and quantidade > 0 and valor_total > 0:
            preco = valor_total / quantidade
        descricao = str(item.get(TITULO_PRODUTO_INGLES_KEY) or "").strip() or _primeiro_texto_item(
            item,
            [
                "Título em português",
                "Titulo em portugues",
                "Título do produto em português",
                "Titulo do produto em portugues",
                "Descrição",
                "Descricao",
                "Descricao do NCM",
            ],
        )
        ncm = _primeiro_texto_item(item, ["NCM", "Código NCM", "Codigo NCM", "HS Code", "HSCODE"])

        _atribuir_texto_excel(ws.cell(linha, 1), indice)
        _atribuir_texto_excel(ws.cell(linha, 2), item.get("SKU", ""))
        _atribuir_texto_excel(ws.cell(linha, 3), _formatar_ncm_commercial_invoice(ncm))
        _atribuir_texto_excel(ws.cell(linha, 4), descricao)
        ws.cell(linha, 7).value = quantidade
        ws.cell(linha, 8).value = preco
        ws.cell(linha, 9).value = f"=G{linha}*H{linha}"

        for coluna in range(1, 10):
            celula = ws.cell(linha, coluna)
            celula.border = borda
            celula.font = fonte_negrito
            celula.alignment = alinhamento_centro
        ws.cell(linha, 4).alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        ws.cell(linha, 7).number_format = "#,##0"
        ws.cell(linha, 8).number_format = '"US$"#,##0.00'
        ws.cell(linha, 9).number_format = '"US$"#,##0.00'
        ws.row_dimensions[linha].height = 72

        caminho_foto = _resolver_caminho_foto_commercial_invoice(
            client_id,
            item.get("Foto", ""),
            item.get("SKU", ""),
        )
        if caminho_foto:
            try:
                imagem = XLImage(caminho_foto)
                largura = float(getattr(imagem, "width", 0) or 0)
                altura = float(getattr(imagem, "height", 0) or 0)
                altura_alvo = 66
                if largura > 0 and altura > 0:
                    imagem.height = altura_alvo
                    imagem.width = max(38, min(120, int(altura_alvo * (largura / altura))))
                else:
                    imagem.height = altura_alvo
                    imagem.width = 66
                ws.add_image(imagem, f"F{linha}")
            except Exception:
                pass

    ultima_linha_item = primeira_linha_item + len(itens) - 1
    linha_subtotal = ultima_linha_item + 1
    linha_total = linha_subtotal + 1
    ws.merge_cells(start_row=linha_subtotal, start_column=1, end_row=linha_subtotal, end_column=5)
    _atribuir_texto_excel(ws.cell(linha_subtotal, 1), "TOTALS")
    _atribuir_texto_excel(ws.cell(linha_subtotal, 6), "Units Total")
    ws.cell(linha_subtotal, 7).value = f"=SUM(G{primeira_linha_item}:G{ultima_linha_item})"
    _atribuir_texto_excel(ws.cell(linha_subtotal, 8), "Subtotal")
    ws.cell(linha_subtotal, 9).value = f"=SUM(I{primeira_linha_item}:I{ultima_linha_item})"

    ws.merge_cells(start_row=linha_total, start_column=1, end_row=linha_total, end_column=7)
    _atribuir_texto_excel(ws.cell(linha_total, 8), "TOTAL:")
    ws.cell(linha_total, 9).value = f"=I{linha_subtotal}"

    for linha in (linha_subtotal, linha_total):
        for coluna in range(1, 10):
            celula = ws.cell(linha, coluna)
            celula.border = borda
            celula.font = fonte_negrito
            celula.alignment = alinhamento_centro
        ws.cell(linha, 9).number_format = '"US$"#,##0.00'
    ws.cell(linha_subtotal, 7).number_format = "#,##0"
    ws.row_dimensions[linha_subtotal].height = 24
    ws.row_dimensions[linha_total].height = 24

    for coluna, largura in {
        "A": 7,
        "B": 17,
        "C": 17,
        "D": 24,
        "E": 24,
        "F": 20,
        "G": 14,
        "H": 15,
        "I": 17,
    }.items():
        ws.column_dimensions[coluna].width = largura

    if getattr(wb, "calculation", None) is not None:
        wb.calculation.fullCalcOnLoad = True
        wb.calculation.forceFullCalc = True
        wb.calculation.calcMode = "auto"

    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


def _gerar_excel_lista_pedido_bytes(nome_lista: str, itens: list[dict], client_id: str | None = None) -> bytes:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Import 54"
    ws.sheet_view.zoomScale = 110

    headers = list(C54_HEADERS_LISTA_PEDIDO)
    ws.append(headers)

    cabecalho_fill = PatternFill(fill_type="solid", fgColor="B4E4E8")
    cabecalho_font = Font(color="000000", bold=False, size=11)
    cabecalho_alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    borda_fina = Border(
        left=Side(style="thin", color="D9E1E2"),
        right=Side(style="thin", color="D9E1E2"),
        top=Side(style="thin", color="D9E1E2"),
        bottom=Side(style="thin", color="D9E1E2"),
    )

    for col_idx in range(1, len(headers) + 1):
        cell = ws.cell(row=1, column=col_idx)
        cell.fill = cabecalho_fill
        cell.font = cabecalho_font
        cell.alignment = cabecalho_alignment
        cell.border = borda_fina
    ws.row_dimensions[1].height = 25.5

    diretorios_fotos = []
    if client_id:
        diretorios_fotos.append(os.path.join(get_tenant_path(client_id), "cadastro_fotos"))
    diretorios_fotos.append(os.path.join(PASTA_INFO, "default", "cadastro_fotos"))

    cache_fotos: dict[str, str | None] = {}

    def _resolver_caminho_foto_lista(foto_ref: str, sku: str) -> str | None:
        foto_txt = _resolver_foto_cadastro_sku(client_id, sku, foto_ref)
        foto_txt = str(foto_txt or "").replace("\\", "/").strip()
        if not foto_txt:
            return None

        if os.path.isabs(foto_txt) and os.path.exists(foto_txt):
            return foto_txt

        nome = foto_txt
        if nome.lower().startswith("cadastro_fotos/"):
            nome = nome.split("/", 1)[1]
        nome = os.path.basename(nome)
        if not nome:
            return None

        if nome in cache_fotos:
            return cache_fotos[nome]

        caminho = None
        for pasta_fotos in diretorios_fotos:
            candidato = os.path.join(pasta_fotos, nome)
            if os.path.exists(candidato):
                caminho = candidato
                break
        cache_fotos[nome] = caminho
        return caminho

    itens_norm = [_normalizar_item_lista_pedido(i) for i in (itens or [])]
    for idx, item in enumerate(itens_norm, start=2):
        quantidade = int(max(0.0, _to_float(item.get("Quantidade", 0), 0.0)))
        valor_unidade = max(0.0, _to_float(item.get("Valor unidade", 0), 0.0))
        ws.append([
            item.get("SKU", ""),
            "",
            item.get(TITULO_PRODUTO_INGLES_KEY, ""),
            item.get("OEM", ""),
            item.get(COR_LADO_LISTA_PEDIDO_KEY, ""),
            item.get("Link", ""),
            quantidade,
            valor_unidade if valor_unidade > 0 else None,
            None,
            item.get(CBM_LISTA_PEDIDO_KEY, ""),
            item.get(PESO_LISTA_PEDIDO_KEY, ""),
            item.get(EMBALAGEM_LISTA_PEDIDO_KEY, ""),
        ])
        faixa_fill = PatternFill(fill_type="solid", fgColor=("F7FAFC" if idx % 2 == 0 else "FFFFFF"))
        for col_idx in range(1, len(headers) + 1):
            c = ws.cell(row=idx, column=col_idx)
            c.fill = faixa_fill
            c.border = borda_fina
            c.alignment = Alignment(
                vertical="center",
                horizontal=("center" if col_idx in (2, 7) else "left"),
                wrap_text=(col_idx in (3, 4, 5, 6, 12)),
            )

        ws.cell(row=idx, column=7).number_format = "#,##0"
        ws.cell(row=idx, column=8).number_format = '"$" #,##0.00'
        ws.cell(row=idx, column=9).number_format = '"$" #,##0.00'
        ws.cell(row=idx, column=10).number_format = "#,##0.000"
        ws.cell(row=idx, column=11).number_format = "#,##0.000"
        ws.row_dimensions[idx].height = max(float(ws.row_dimensions[idx].height or 0), 84.75)

        cel_sku = ws.cell(row=idx, column=1)
        cel_sku.font = Font(bold=True, color="0F2D52")
        cel_sku.fill = PatternFill(fill_type="solid", fgColor="EAF2FB")

        link_val = str(item.get("Link", "") or "").strip()
        if link_val:
            cel_link = ws.cell(row=idx, column=6)
            cel_link.value = link_val
            cel_link.hyperlink = link_val
            cel_link.style = "Hyperlink"
            cel_link.alignment = Alignment(vertical="center", horizontal="left")

        caminho_foto = _resolver_caminho_foto_lista(item.get("Foto", ""), item.get("SKU", ""))
        if caminho_foto:
            try:
                img = XLImage(caminho_foto)
                altura_alvo_px = 82
                largura_original = float(getattr(img, "width", 0) or 0)
                altura_original = float(getattr(img, "height", 0) or 0)

                if largura_original > 0 and altura_original > 0:
                    proporcao = largura_original / altura_original
                    img.height = altura_alvo_px
                    img.width = max(42, min(130, int(altura_alvo_px * proporcao)))
                else:
                    img.height = altura_alvo_px
                    img.width = 82

                ws.add_image(img, f"B{idx}")
            except Exception:
                ws[f"B{idx}"] = str(item.get("Foto", "") or "")
        else:
            ws[f"B{idx}"] = str(item.get("Foto", "") or "")

        ws[f"I{idx}"] = f"=G{idx}*H{idx}"

    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:L{max(1, len(itens_norm) + 1)}"
    for col, width in {
        "A": 12.57,
        "B": 17.43,
        "C": 34,
        "D": 18,
        "E": 16,
        "F": 24,
        "G": 13,
        "H": 13,
        "I": 16,
        "J": 13,
        "K": 13,
        "L": 18,
    }.items():
        ws.column_dimensions[col].width = width

    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()



def configure_medias_compras_excel_runtime(runtime_module=None):
    return _configure_runtime_globals(runtime_module)


configure_medias_compras_excel_runtime()

__all__ = [
    "configure_medias_compras_excel_runtime",
    "_normalizar_cabecalho_excel",
    "_to_float_excel",
    "_linha_resumo_excel_sku",
    "_gerar_excel_lista_pedido_bytes",
    "_formatar_ncm_commercial_invoice",
    "_gerar_commercial_invoice_bytes",
]
