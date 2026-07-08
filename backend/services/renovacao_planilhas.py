"""Spreadsheet analysis and export helpers for Renovacao."""

from __future__ import annotations

import io

import openpyxl
import pandas as pd

def find_sheet_name(workbook):
    """Encontra o nome da aba de promocoes ou retorna a primeira."""
    for name in workbook.sheetnames:
        if "promocoes" in name.lower() or "promo" in name.lower():
            return name
    return workbook.sheetnames[0]


def find_common_key(df_antiga, df_nova):
    """Encontra uma coluna chave comum para unir os dataframes."""
    possible_keys = ["ITEM_ID", "SELLER_SKU", "SKU", "ID", "CÃƒÆ’Ã¢â‚¬Å“DIGO", "MLB"]
    for key in possible_keys:
        if key in df_antiga.columns and key in df_nova.columns:
            return key
    return None


def find_col(df_columns, names):
    """Encontra o primeiro nome de coluna correspondente em uma lista de possibilidades."""
    upper_cols = [str(c).upper().strip() for c in df_columns]
    for name in names:
        target = str(name).upper().strip()
        if target in upper_cols:
            return df_columns[upper_cols.index(target)]
    return None


def processar_renovacao_logic(file_antiga_bytes, file_nova_bytes):
    try:
        with io.BytesIO(file_antiga_bytes) as b:
            try:
                wb = openpyxl.load_workbook(b, data_only=True, read_only=True)
                sheet_name = find_sheet_name(wb)
                wb.close()
                b.seek(0)
                df_antiga = pd.read_excel(b, sheet_name=sheet_name, header=0)
            except Exception:
                b.seek(0)
                df_antiga = pd.read_excel(b, header=0)

        cols_upper = [str(c).upper() for c in df_antiga.columns]
        if len(df_antiga) > 4 and not any(k in cols_upper for k in ["ITEM_ID", "SKU", "MLB", "STATUS"]):
            with io.BytesIO(file_antiga_bytes) as b:
                df_antiga = pd.read_excel(b, header=4)

        df_antiga.columns = df_antiga.columns.str.strip().str.upper()

        key_col = find_common_key(df_antiga, df_antiga)
        status_col = find_col(df_antiga.columns, ["STATUS", "ESTADO"])

        if not key_col or not status_col:
            return None, "Colunas chave ou STATUS nÃƒÆ’Ã‚Â£o encontradas na planilha antiga."

        df_antiga[key_col] = df_antiga[key_col].astype(str).str.strip()
        status_map = df_antiga.set_index(key_col)[status_col].to_dict()

        wb_nova = openpyxl.load_workbook(io.BytesIO(file_nova_bytes))
        ws_name = find_sheet_name(wb_nova)
        ws_nova = wb_nova[ws_name]

        header_row_idx = 1
        header_map = {}

        found_header = False
        for r in range(1, 11):
            row_vals = [str(c.value).strip().upper() if c.value else "" for c in ws_nova[r]]
            if (key_col in row_vals or any(k in row_vals for k in ["ITEM_ID", "SKU", "MLB"])) and any(
                x in row_vals for x in ["ACTION", "AÃƒÆ’Ã¢â‚¬Â¡ÃƒÆ’Ã†â€™O", "ACAO"]
            ):
                header_row_idx = r
                for idx, val in enumerate(row_vals):
                    if val:
                        header_map[val] = idx + 1
                found_header = True
                break

        if not found_header:
            return None, "CabeÃƒÆ’Ã‚Â§alho nÃƒÆ’Ã‚Â£o encontrado na planilha nova."

        idx_key_nova = header_map.get(key_col)
        if not idx_key_nova:
            for k in ["ITEM_ID", "SELLER_SKU", "SKU", "ID", "CÃƒÆ’Ã¢â‚¬Å“DIGO", "MLB"]:
                if k in header_map:
                    idx_key_nova = header_map[k]
                    break

        idx_action_nova = header_map.get("ACTION") or header_map.get("AÃƒÆ’Ã¢â‚¬Â¡ÃƒÆ’Ã†â€™O") or header_map.get("ACAO")

        if not idx_key_nova or not idx_action_nova:
            return None, "Colunas Chave ou ACTION nÃƒÆ’Ã‚Â£o mapeadas na planilha nova."

        for row in ws_nova.iter_rows(min_row=header_row_idx + 1):
            cell_key = row[idx_key_nova - 1].value
            if not cell_key:
                continue

            key_val = str(cell_key).strip()
            status_antigo = status_map.get(key_val)
            acao = "Participar"

            if status_antigo:
                st_str = str(status_antigo).strip().lower()
                if st_str not in ["ativo", "programado", "active"]:
                    acao = "NÃƒÆ’Ã‚Â£o participar"

            row[idx_action_nova - 1].value = acao

        output = io.BytesIO()
        wb_nova.save(output)
        output.seek(0)
        return output, "Sucesso"
    except Exception as e:
        return None, str(e)


def analisar_renovacao_logic(file_antiga_bytes, file_nova_bytes):
    try:
        with io.BytesIO(file_antiga_bytes) as b:
            try:
                wb = openpyxl.load_workbook(b, data_only=True, read_only=True)
                sheet_name_antiga = find_sheet_name(wb)
                wb.close()
                b.seek(0)
                df_antiga = pd.read_excel(b, sheet_name=sheet_name_antiga, header=0)
            except Exception:
                b.seek(0)
                df_antiga = pd.read_excel(b, header=0)

        cols_antiga_upper = [str(c).upper() for c in df_antiga.columns]
        if len(df_antiga) > 4 and not any(k in cols_antiga_upper for k in ["ITEM_ID", "SKU", "MLB", "STATUS"]):
            with io.BytesIO(file_antiga_bytes) as b:
                df_antiga = pd.read_excel(b, sheet_name=sheet_name_antiga if "sheet_name_antiga" in locals() else 0, header=4)

        df_antiga.columns = df_antiga.columns.str.strip().str.upper()

        with io.BytesIO(file_nova_bytes) as b:
            try:
                wb = openpyxl.load_workbook(b, data_only=True, read_only=True)
                sheet_name_nova = find_sheet_name(wb)
                wb.close()
                b.seek(0)
                df_nova = pd.read_excel(b, sheet_name=sheet_name_nova, header=0)
            except Exception:
                b.seek(0)
                df_nova = pd.read_excel(b, header=0)

        cols_nova_upper = [str(c).upper() for c in df_nova.columns]
        if len(df_nova) > 4 and not any(k in cols_nova_upper for k in ["ITEM_ID", "SKU", "MLB", "ACTION", "AÃƒÆ’Ã¢â‚¬Â¡ÃƒÆ’Ã†â€™O"]):
            with io.BytesIO(file_nova_bytes) as b:
                df_nova = pd.read_excel(b, sheet_name=sheet_name_nova if "sheet_name_nova" in locals() else 0, header=4)

        df_nova.columns = df_nova.columns.str.strip().str.upper()

        key_col = find_common_key(df_antiga, df_antiga)
        status_col = find_col(df_antiga.columns, ["STATUS", "ESTADO"])

        if not key_col:
            return None, "NÃƒÆ’Ã‚Â£o foi possÃƒÆ’Ã‚Â­vel encontrar uma coluna chave comum (ITEM_ID, SKU, etc.)."

        discount_names = ["DISCOUNT_PERCENTAGE", "DISCOUNT %", "DESCONTO", "DESC %", "DISCOUNT"]
        price_names = ["FINAL_PRICE", "PREÃƒÆ’Ã¢â‚¬Â¡O FINAL", "PRECO FINAL", "PRICE", "PREÃƒÆ’Ã¢â‚¬Â¡O"]

        df_antiga[key_col] = df_antiga[key_col].astype(str).str.strip()

        disc_col_ant = find_col(df_antiga.columns, discount_names)
        price_col_ant = find_col(df_antiga.columns, price_names)

        antiga_map = {}

        def get_clean_val(row, col):
            if not col:
                return ""
            val = row[col]
            if pd.isna(val):
                return ""
            return str(val).strip()

        def parse_pct(val):
            if not val:
                return 0.0
            try:
                s = str(val).replace("%", "").strip()
                if "," in s and "." in s:
                    s = s.replace(".", "").replace(",", ".")
                elif "," in s:
                    s = s.replace(",", ".")
                return float(s)
            except Exception:
                return 0.0

        for _, row in df_antiga.iterrows():
            k = str(row[key_col]).strip()
            antiga_map[k] = {
                "STATUS": get_clean_val(row, status_col),
                "DESC": get_clean_val(row, disc_col_ant),
                "PRICE": get_clean_val(row, price_col_ant),
            }

        df_nova[key_col] = df_nova[key_col].astype(str).str.strip()

        title_col = find_col(df_nova.columns, ["TITLE", "TITULO", "TÃƒÆ’Ã‚ÂTULO"])
        disc_col_nova = find_col(df_nova.columns, discount_names)
        price_col_nova = find_col(df_nova.columns, price_names)

        dados_analise = []
        for _, row in df_nova.iterrows():
            key_val = str(row[key_col])

            if not key_val.strip().upper().startswith("MLB"):
                continue

            antiga_data = antiga_map.get(key_val, {})
            status_antigo = antiga_data.get("STATUS", "")

            if not status_antigo:
                acao = "Participar"
            else:
                st_str = str(status_antigo).strip().lower()
                if "elegÃƒÆ’Ã‚Â­vel" in st_str or "elegivel" in st_str:
                    acao = "NÃƒÆ’Ã‚Â£o participar"
                elif st_str in ["ativo", "programado"]:
                    acao = "Participar"
                else:
                    acao = "NÃƒÆ’Ã‚Â£o participar"

                desc_antiga_val = parse_pct(antiga_data.get("DESC", ""))
                desc_nova_val = parse_pct(str(row[disc_col_nova]) if disc_col_nova else "")
                if desc_nova_val > desc_antiga_val + 1e-9:
                    acao = "NÃƒÆ’Ã‚Â£o participar"

            dados_analise.append({
                "KEY": key_val,
                "TITLE": str(row[title_col]) if title_col else "",
                "STATUS_ANTIGO": status_antigo,
                "DESC_ANTIGA": antiga_data.get("DESC", ""),
                "DESC_NOVA": str(row[disc_col_nova]) if disc_col_nova else "",
                "PRICE_ANTIGA": antiga_data.get("PRICE", ""),
                "PRICE_NOVA": str(row[price_col_nova]) if price_col_nova else "",
                "ACTION": acao,
            })

        return dados_analise, "Sucesso"
    except Exception as e:
        return None, f"Erro interno: {str(e)}"


def gerar_excel_renovacao(file_nova_bytes, decisoes_list):
    try:
        wb = openpyxl.load_workbook(io.BytesIO(file_nova_bytes))
        ws_name = find_sheet_name(wb)
        ws = wb[ws_name]

        decision_map = {str(d["KEY"]).strip(): d for d in decisoes_list}

        header_row_idx = 1
        col_key_idx = None
        col_action_idx = None
        col_desc_idx = None
        col_price_idx = None

        discount_names = ["DISCOUNT_PERCENTAGE", "DISCOUNT %", "DESCONTO", "DESC %", "DISCOUNT"]
        price_names = ["FINAL_PRICE", "PREÃƒÆ’Ã¢â‚¬Â¡O FINAL", "PRECO FINAL", "PRICE", "PREÃƒÆ’Ã¢â‚¬Â¡O"]

        for r in range(1, 21):
            row_vals = [str(c.value).strip().upper() if c.value else "" for c in ws[r]]
            if any(x in row_vals for x in ["ACTION", "AÃƒÆ’Ã¢â‚¬Â¡ÃƒÆ’Ã†â€™O", "ACAO"]):
                header_row_idx = r

                for idx, val in enumerate(row_vals):
                    if val in ["ACTION", "AÃƒÆ’Ã¢â‚¬Â¡ÃƒÆ’Ã†â€™O", "ACAO"]:
                        col_action_idx = idx + 1
                    if val in discount_names:
                        col_desc_idx = idx + 1
                    if val in price_names:
                        col_price_idx = idx + 1

                if "ITEM_ID" in row_vals:
                    col_key_idx = row_vals.index("ITEM_ID") + 1
                elif "MLB" in row_vals:
                    col_key_idx = row_vals.index("MLB") + 1
                else:
                    for idx, val in enumerate(row_vals):
                        if val in ["SKU", "SELLER_SKU", "ID", "CÃƒÆ’Ã¢â‚¬Å“DIGO"]:
                            col_key_idx = idx + 1
                            break

                if col_key_idx and col_action_idx:
                    break

        if not col_key_idx or not col_action_idx:
            return None, "Colunas nÃƒÆ’Ã‚Â£o encontradas para exportaÃƒÆ’Ã‚Â§ÃƒÆ’Ã‚Â£o."

        for row in ws.iter_rows(min_row=header_row_idx + 1):
            cell_key = row[col_key_idx - 1].value
            if cell_key:
                k = str(cell_key).strip()
                if k in decision_map:
                    data = decision_map[k]
                    row[col_action_idx - 1].value = data.get("ACTION", "")

                    if col_desc_idx and "DESC_NOVA" in data:
                        try:
                            val = data["DESC_NOVA"]
                            if val and val != "nan" and val != "None":
                                row[col_desc_idx - 1].value = float(val) if str(val).replace(".", "", 1).isdigit() else val
                        except Exception:
                            row[col_desc_idx - 1].value = data["DESC_NOVA"]

                    if col_price_idx and "PRICE_NOVA" in data:
                        try:
                            val = data["PRICE_NOVA"]
                            if val and val != "nan" and val != "None":
                                row[col_price_idx - 1].value = float(val) if str(val).replace(".", "", 1).isdigit() else val
                        except Exception:
                            row[col_price_idx - 1].value = data["PRICE_NOVA"]

        output = io.BytesIO()
        wb.save(output)
        output.seek(0)
        return output, "Sucesso"
    except Exception as e:
        return None, str(e)
