"""Purchase-list suggestion endpoints for Medias Compras."""

from __future__ import annotations

import io
import json
import math
import os
import re
import sqlite3
import unicodedata
import uuid
from datetime import datetime

import openpyxl
import pandas as pd
from fastapi import Depends, HTTPException
from openpyxl.drawing.image import Image as XLImage
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

from backend.schemas import ListaCompraRequest
from backend.services import medias_compras_common as medias_common
from backend.services.medias_compras_common import *
from backend.services.medias_compras_excel import (
    _gerar_excel_lista_pedido_bytes,
    _resolver_caminho_foto_cadastro_seguro,
)
from backend.services.medias_compras_fiscal import *
from backend.services.medias_compras_visao import api_medias_compras_visao
from backend.services.runtime_bridge import bind_runtime_globals
from medias_compras import calcular_reposicao_periodo


_RUNTIME_NAMES = (
    "logger",
    "get_tenant_path",
    "TEMP_FILES_STORAGE",
    "TEMP_FILES_META",
    "LISTA_PEDIDO_XLSX_CACHE",
    "LISTA_PEDIDO_XLSX_CACHE_MAX_ITENS",
)


def _sync_common_names() -> None:
    for name in medias_common.COMMON_EXPORTS:
        globals()[name] = getattr(medias_common, name)
    for name in _RUNTIME_NAMES:
        globals()[name] = getattr(medias_common, name)


def _configure_runtime_globals(runtime_module=None):
    runtime = medias_common.configure_medias_compras_common_runtime(runtime_module)
    bind_runtime_globals(globals(), runtime)
    _sync_common_names()
    return runtime


def _exigir_loja_especifica_para_lista(loja: str | None) -> str:
    loja_final = str(loja or "").strip()
    if not loja_final or loja_final == "__todas":
        raise HTTPException(
            status_code=400,
            detail="Selecione uma loja especifica para gerar a lista.",
        )
    return loja_final


def _normalizar_quantidades_sugeridas(valor) -> dict[str, int]:
    if valor is None or valor == "":
        return {}

    dados = valor
    if isinstance(valor, str):
        try:
            dados = json.loads(valor)
        except (TypeError, ValueError) as exc:
            raise HTTPException(
                status_code=400,
                detail="Quantidades sugeridas devem ser enviadas em JSON válido.",
            ) from exc

    if not isinstance(dados, dict):
        raise HTTPException(
            status_code=400,
            detail="Quantidades sugeridas devem ser um objeto por SKU.",
        )
    if len(dados) > 10000:
        raise HTTPException(
            status_code=400,
            detail="Quantidade de ajustes de compra excede o limite permitido.",
        )

    resultado: dict[str, int] = {}
    for sku_original, quantidade_original in dados.items():
        sku = _normalizar_sku_mes(str(sku_original or "").strip()).upper()
        if not sku:
            raise HTTPException(status_code=400, detail="SKU inválido nas quantidades sugeridas.")
        if isinstance(quantidade_original, bool):
            raise HTTPException(
                status_code=400,
                detail=f"Quantidade sugerida inválida para o SKU {sku}.",
            )
        try:
            quantidade_numero = float(quantidade_original)
        except (TypeError, ValueError) as exc:
            raise HTTPException(
                status_code=400,
                detail=f"Quantidade sugerida inválida para o SKU {sku}.",
            ) from exc
        if (
            not math.isfinite(quantidade_numero)
            or quantidade_numero < 0
            or not quantidade_numero.is_integer()
        ):
            raise HTTPException(
                status_code=400,
                detail=f"Quantidade sugerida do SKU {sku} deve ser um inteiro maior ou igual a zero.",
            )
        resultado[sku] = int(quantidade_numero)
    return resultado


def _resolver_quantidade_sugerida(
    sku: str,
    quantidade_calculada,
    quantidades_sugeridas: dict[str, int],
    quantidades_aplicadas: set[str],
) -> int:
    sku_cmp = _normalizar_sku_mes(str(sku or "").strip()).upper()
    if sku_cmp in quantidades_sugeridas:
        quantidades_aplicadas.add(sku_cmp)
        return quantidades_sugeridas[sku_cmp]
    return int(quantidade_calculada or 0)


async def api_medias_compras_gerar_lista_compra(
    req: ListaCompraRequest,
    client_id: str = Depends(medias_common.get_tenant_id)
):
    try:
        opcao = str(req.opcao or "").strip().lower()
        if opcao not in {"media_6m", "crescimento"}:
            raise HTTPException(status_code=400, detail="Opção inválida para geração da lista de compra")

        crescimento_percent = float(req.crescimento_percent or 0)
        if opcao == "crescimento":
            if crescimento_percent < 0 or crescimento_percent > 100:
                raise HTTPException(status_code=400, detail="Percentual de crescimento deve estar entre 0 e 100")
        else:
            crescimento_percent = 0

        hidden_skus = {
            _normalizar_sku_mes(str(sku or "").strip()).upper()
            for sku in (req.hidden_skus or [])
            if str(sku or "").strip()
        }
        quantidades_sugeridas = _normalizar_quantidades_sugeridas(req.quantidades_sugeridas)

        nome_lista = str(req.nome_lista or "").strip()
        if not nome_lista:
            raise HTTPException(status_code=400, detail="Informe o nome da lista antes de gerar o pedido")

        escopo_loja = _resolver_escopo_loja_medias(
            client_id,
            req.loja,
            req.store_id,
            exigir_especifica=True,
        )
        loja_sel = escopo_loja["loja"]
        store_id_escopo = escopo_loja["store_id"]
        loja_sel_norm = loja_sel.lower()
        periodo_meses = int(req.periodo_meses or 6)
        if periodo_meses not in (3, 6, 12):
            periodo_meses = 6
        lead_time_meses = 6.0

        def _safe_float(valor, default=0.0):
            if valor is None:
                return default
            if isinstance(valor, (int, float)):
                return float(valor)
            txt = str(valor).strip().replace(".", "").replace(",", ".")
            if not txt:
                return default
            try:
                return float(txt)
            except Exception:
                return default

        def _meses_referencia(qtd_meses: int) -> list[str]:
            agora = datetime.now()
            ano = agora.year
            mes = agora.month
            refs = []
            for i in range(qtd_meses - 1, -1, -1):
                m = mes - i
                a = ano
                while m <= 0:
                    m += 12
                    a -= 1
                refs.append(f"{a:04d}-{m:02d}")
            return refs

        def _norm_col(col: str) -> str:
            txt = str(col or "").strip().lower().replace("_", " ")
            txt = unicodedata.normalize("NFKD", txt)
            txt = "".join(ch for ch in txt if not unicodedata.combining(ch))
            txt = re.sub(r"[^a-z0-9\s]", " ", txt)
            txt = re.sub(r"\s+", " ", txt).strip()
            if txt.startswith("cg "):
                txt = txt[3:].strip()
            return txt

        def _pick_valor(row: dict, candidatos_norm: list[str]) -> str:
            candidatos = [_norm_col(c) for c in (candidatos_norm or []) if str(c or "").strip()]
            for key, valor in row.items():
                key_norm = _norm_col(key)
                if any((key_norm == c or c in key_norm) for c in candidatos):
                    txt = str(valor or "").strip()
                    if txt:
                        return txt
            return ""

        def _sku_sort_key(sku_valor: str):
            sku_txt = str(sku_valor or "").strip().upper()
            primeiro_bloco = re.split(r"[.-]", sku_txt)[0] if sku_txt else ""
            somente_digitos = "".join(ch for ch in primeiro_bloco if ch.isdigit())
            if somente_digitos:
                return (0, int(somente_digitos), sku_txt)
            return (1, sku_txt)

        meses_ref = _meses_referencia(periodo_meses)
        inicio_periodo = f"{meses_ref[0]}-01"
        fim_periodo = datetime.now().strftime("%Y-%m-%d")

        # Vendas totais do perÃƒÂ­odo selecionado por SKU, consolidadas por tenant e filtradas pela loja.
        vendas_total_periodo = {}
        db_paths = _listar_bancos_vendas_tenant(client_id, loja_sel)
        ids_vendas_processados = set()
        for db_vendas in db_paths:
            if not (db_vendas and os.path.exists(db_vendas)):
                continue

            conn = sqlite3.connect(db_vendas)
            try:
                conn.row_factory = sqlite3.Row
                cur = conn.cursor()
                cur.execute("CREATE INDEX IF NOT EXISTS idx_vendas_sku_data ON vendas(sku, data)")

                query = (
                    """
                    SELECT id_unico, sku, COALESCE(quantidade, 0) AS qtd
                    FROM vendas
                    WHERE sku IS NOT NULL
                      AND trim(sku) != ''
                      AND date(data) BETWEEN ? AND ?
                    """
                )
                params = [inicio_periodo, fim_periodo]
                filtro_loja_sql, filtro_loja_params = _sql_filtro_loja_vendas(loja_sel)
                if filtro_loja_sql:
                    query += filtro_loja_sql
                    params.extend(filtro_loja_params)

                rows = cur.execute(query, params).fetchall()
                for r in rows:
                    sku = _normalizar_sku_mes(str(r["sku"] or "").strip())
                    if not sku:
                        continue
                    id_unico = str(r["id_unico"] or "").strip()
                    qtd = float(r["qtd"] or 0)
                    chave_unica = id_unico or f"{os.path.basename(db_vendas)}::{sku}::{qtd}"
                    if chave_unica in ids_vendas_processados:
                        continue
                    ids_vendas_processados.add(chave_unica)
                    vendas_total_periodo[sku] = vendas_total_periodo.get(sku, 0.0) + qtd
            finally:
                conn.close()

        # Estoque atual por SKU considerando a loja selecionada.
        saldo_por_sku = {}
        arquivo_estoque = _migrar_arquivo_legado_para_tenant(client_id, "produtos_compilado.csv", ARQUIVO_DB_PRODUTOS)
        if arquivo_estoque and os.path.exists(arquivo_estoque):
            try:
                df_est = pd.read_csv(arquivo_estoque).fillna("")
                if not df_est.empty:
                    colunas_lower = {str(c).strip().lower(): c for c in df_est.columns}
                    col_sku = colunas_lower.get("sku") or next((c for c in df_est.columns if str(c).strip().lower() == "sku"), None)
                    col_loja_sync = colunas_lower.get("loja_sync")
                    if col_loja_sync is not None and loja_sel_norm != "__todas":
                        df_est = df_est[
                            df_est[col_loja_sync].astype(str).str.strip().str.lower() == loja_sel_norm
                        ].copy()
                    if col_sku is not None:
                        for _, row in df_est.iterrows():
                            sku_raw = str(row.get(col_sku, "") or "").strip()
                            if not sku_raw:
                                continue
                            sku = _normalizar_sku_mes(sku_raw)
                            saldo_loja = _safe_float(row.get("saldo_loja", 0), 0.0)
                            # Compatibilidade com bases antigas sem coluna saldo_loja.
                            saldo_total = saldo_loja if saldo_loja != 0 else _safe_float(row.get("saldo", row.get("saldo_total", 0)), 0.0)
                            saldo_por_sku[sku] = saldo_por_sku.get(sku, 0.0) + saldo_total
            except Exception:
                saldo_por_sku = saldo_por_sku or {}

        # Estoque em trÃƒÂ¢nsito vindo das listas com status Pedido Aprovado.
        transito_por_sku = _mapa_estoque_em_transito_por_sku(client_id, loja_sel)

        # Cadastro por SKU para foto, tÃƒÂ­tulo em inglÃƒÂªs, OEM e link.
        cadastro_por_sku = {}
        arquivo_cadastro = _migrar_arquivo_legado_para_tenant(client_id, "cadastro_produtos.csv", ARQUIVO_DB_CADASTRO_PRODUTOS)
        linhas_cadastro_legado = []
        if arquivo_cadastro and os.path.exists(arquivo_cadastro):
            try:
                df_cad = pd.read_csv(arquivo_cadastro, dtype=str).fillna("")
                if not df_cad.empty:
                    df_cad.columns = [str(c).strip().lower() for c in df_cad.columns]
                    linhas_cadastro_legado = df_cad.to_dict(orient="records")
            except Exception:
                linhas_cadastro_legado = []

        from backend.services.cadastro_compatibilidade import (
            mesclar_produtos_legados_com_contexto_loja,
        )

        try:
            contexto_cadastro = mesclar_produtos_legados_com_contexto_loja(
                client_id,
                linhas_cadastro_legado,
                store_id_escopo,
            )
        except RuntimeError:
            contexto_cadastro = {
                "produtos": [],
                "store_id": "",
                "loja_resolvida": False,
                "scope": "unavailable",
            }
        store_id_cadastro = (
            str(contexto_cadastro.get("store_id") or "").strip()
            or store_id_escopo
        )
        for row_dict in contexto_cadastro.get("produtos") or []:
            if not isinstance(row_dict, dict):
                continue
            sku_raw = str(row_dict.get("sku", "") or "").strip()
            if not sku_raw:
                continue
            sku = _normalizar_sku_mes(sku_raw)
            titulo_ingles = _pick_valor(row_dict, [
                "titulo em ingles", "titulo ingles", "title english", "english title", "nome em ingles", "nome ingles",
                "product name", "product description", "description"
            ])
            if not titulo_ingles:
                titulo_ingles = _pick_valor(row_dict, ["titulos anuncios mlb", "nome", "produto", "cg product name"])
            oem = _pick_valor(row_dict, ["oem", "codigo oem", "part number", "partnumber", "oem model"])
            cor_lado = _pick_valor(row_dict, ["color side", "color/side", "cor lado", "cor/lado", "lado cor", "lado/cor", "lado", "cor", "color", "side"])
            link = _pick_valor(row_dict, ["link", "url", "url produto", "link produto", "product link", "link aliexpress", "aliexpress"])
            foto = _pick_valor(row_dict, ["foto", "imagem", "image", "url foto", "foto produto"])
            foto = _resolver_foto_cadastro_sku(
                client_id,
                sku,
                foto,
                store_id_cadastro or None,
            )
            cadastro_por_sku[sku] = {
                "foto": foto,
                "titulo_ingles": titulo_ingles,
                "oem": oem,
                "cor_lado": cor_lado,
                "link": link,
            }

        todos_skus = set(vendas_total_periodo.keys()) | set(saldo_por_sku.keys())
        itens_lista = []
        fator_crescimento = 1 + (crescimento_percent / 100.0)
        quantidades_sugeridas_aplicadas = set()

        for sku in sorted(todos_skus, key=_sku_sort_key):
            sku_cmp = _normalizar_sku_mes(sku).upper()
            if sku_cmp in hidden_skus:
                continue

            total_periodo = float(vendas_total_periodo.get(sku, 0) or 0)
            estoque_atual = float(saldo_por_sku.get(sku, 0) or 0)
            estoque_em_transito = float(transito_por_sku.get(sku, 0) or 0)
            fator_reposicao = 1.0 if opcao == "media_6m" else fator_crescimento
            reposicao = calcular_reposicao_periodo(
                total_vendido_periodo=total_periodo,
                saldo_atual=estoque_atual,
                estoque_em_transito=estoque_em_transito,
                periodo_meses=periodo_meses,
                lead_time_meses=lead_time_meses,
                ciclo_compra_meses=3,
                margem_seguranca_meses=1,
                fator_crescimento=fator_reposicao,
            )
            quantidade_compra = _resolver_quantidade_sugerida(
                sku,
                reposicao.get("compra_sugerida", 0),
                quantidades_sugeridas,
                quantidades_sugeridas_aplicadas,
            )

            if quantidade_compra <= 0:
                continue

            cad = cadastro_por_sku.get(sku, {})
            itens_lista.append({
                "SKU": sku,
                "Foto": str(cad.get("foto", "") or ""),
                TITULO_PRODUTO_INGLES_KEY: str(cad.get("titulo_ingles", "") or ""),
                "OEM": str(cad.get("oem", "") or ""),
                COR_LADO_LISTA_PEDIDO_KEY: str(cad.get("cor_lado", "") or ""),
                "Link": str(cad.get("link", "") or ""),
                "Quantidade": quantidade_compra,
                "Valor unidade": "",
                "Valor total": "",
            })

        def _resolver_caminho_foto_excel(foto_ref: str, sku: str) -> str | None:
            # Esta exportacao historicamente so tenta incorporar uma referencia
            # explicita. Evita acionar o runtime de fotos quando a celula esta
            # vazia e preserva o fallback textual do XLSX.
            if not str(foto_ref or "").strip():
                return None
            return _resolver_caminho_foto_cadastro_seguro(
                client_id,
                foto_ref,
                sku,
                store_id_cadastro or None,
            )

        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Lista Compra"
        ws.sheet_view.zoomScale = 110

        headers = [
            "SKU", "Foto", TITULO_PRODUTO_INGLES_KEY, "OEM", "Link",
            "Quantidade", "Valor unidade", "Valor total"
        ]
        ws.append(headers)

        # Estilo moderno para cabeÃƒÂ§alho.
        cabecalho_fill = PatternFill(fill_type="solid", fgColor="1F4E78")
        cabecalho_font = Font(color="FFFFFF", bold=True, size=11)
        cabecalho_alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        borda_fina = Border(
            left=Side(style="thin", color="D9E1F2"),
            right=Side(style="thin", color="D9E1F2"),
            top=Side(style="thin", color="D9E1F2"),
            bottom=Side(style="thin", color="D9E1F2"),
        )

        for col_idx in range(1, len(headers) + 1):
            cell = ws.cell(row=1, column=col_idx)
            cell.fill = cabecalho_fill
            cell.font = cabecalho_font
            cell.alignment = cabecalho_alignment
            cell.border = borda_fina

        ws.row_dimensions[1].height = 26

        for idx, item in enumerate(itens_lista, start=2):
            ws.append([
                item["SKU"],
                "",
                item[TITULO_PRODUTO_INGLES_KEY],
                item["OEM"],
                item["Link"],
                item["Quantidade"],
                item["Valor unidade"],
                None,
            ])

            # Zebra style e acabamento das cÃƒÂ©lulas de dados.
            faixa_fill = PatternFill(fill_type="solid", fgColor=("F7FAFC" if idx % 2 == 0 else "FFFFFF"))
            for col_idx in range(1, 9):
                c = ws.cell(row=idx, column=col_idx)
                c.fill = faixa_fill
                c.border = borda_fina
                c.alignment = Alignment(vertical="center", horizontal=("center" if col_idx in (2, 6) else "left"), wrap_text=(col_idx in (3, 4)))

            ws.cell(row=idx, column=6).number_format = "#,##0"
            ws.cell(row=idx, column=7).number_format = '"R$" #,##0.00'
            ws.cell(row=idx, column=8).number_format = '"R$" #,##0.00'

            # SKU com destaque visual.
            cel_sku = ws.cell(row=idx, column=1)
            cel_sku.font = Font(bold=True, color="0F2D52")
            cel_sku.fill = PatternFill(fill_type="solid", fgColor="EAF2FB")

            link_val = str(item.get("Link", "") or "").strip()
            if link_val:
                cel_link = ws.cell(row=idx, column=5)
                cel_link.value = link_val
                cel_link.hyperlink = link_val
                cel_link.style = "Hyperlink"
                cel_link.alignment = Alignment(vertical="center", horizontal="left")

            # Insere a foto real do produto na coluna B quando o arquivo existir.
            caminho_foto = _resolver_caminho_foto_excel(
                item.get("Foto", ""),
                item.get("SKU", ""),
            )
            if caminho_foto:
                try:
                    img = XLImage(caminho_foto)
                    # Ajuste visual para caber na cÃƒÂ©lula sem distorcer excessivamente.
                    altura_alvo_px = 56
                    largura_original = float(getattr(img, "width", 0) or 0)
                    altura_original = float(getattr(img, "height", 0) or 0)

                    if largura_original > 0 and altura_original > 0:
                        proporcao = largura_original / altura_original
                        img.height = altura_alvo_px
                        img.width = max(26, min(120, int(altura_alvo_px * proporcao)))
                    else:
                        img.height = altura_alvo_px
                        img.width = 56

                    ws.add_image(img, f"B{idx}")
                    ws.row_dimensions[idx].height = max(float(ws.row_dimensions[idx].height or 0), 44)
                except Exception:
                    # Fallback para manter rastreabilidade se falhar leitura da imagem.
                    ws[f"B{idx}"] = str(item.get("Foto", "") or "")
            else:
                ws[f"B{idx}"] = str(item.get("Foto", "") or "")

            ws[f"H{idx}"] = f"=F{idx}*G{idx}"

        # Linha de totais ao final da lista.
        if itens_lista:
            total_row_num = len(itens_lista) + 2
            last_data_row = len(itens_lista) + 1
            for col_idx in range(1, 9):
                c = ws.cell(row=total_row_num, column=col_idx)
                c.fill = PatternFill(fill_type="solid", fgColor="1F4E78")
                c.font = Font(bold=True, color="FFFFFF", size=11)
                c.border = borda_fina
                c.alignment = Alignment(horizontal="center", vertical="center")
            ws.cell(row=total_row_num, column=1).value = "TOTAL"
            ws.cell(row=total_row_num, column=6).value = f"=SUM(F2:F{last_data_row})"
            ws.cell(row=total_row_num, column=6).number_format = "#,##0"
            ws.cell(row=total_row_num, column=8).value = f"=SUM(H2:H{last_data_row})"
            ws.cell(row=total_row_num, column=8).number_format = '"R$" #,##0.00'
            ws.cell(row=total_row_num, column=8).alignment = Alignment(horizontal="right", vertical="center")
            ws.row_dimensions[total_row_num].height = 22

        ws.freeze_panes = "A2"
        ws.auto_filter.ref = f"A1:H{max(1, len(itens_lista) + 1)}"

        largura_colunas = {
            "A": 16,
            "B": 14,
            "C": 52,
            "D": 24,
            "E": 20,
            "F": 14,
            "G": 14,
            "H": 14,
        }
        for col, width in largura_colunas.items():
            ws.column_dimensions[col].width = width

        nome_arquivo = f"lista_compra_{datetime.now().strftime('%Y%m%d_%H%M%S')}.xlsx"
        buffer = io.BytesIO()
        wb.save(buffer)
        file_bytes = buffer.getvalue()

        lista_id = str(uuid.uuid4())
        agora_iso = datetime.now().isoformat(timespec="seconds")
        listas_salvas = _carregar_listas_pedidos(client_id)
        itens_lista_norm = _recalcular_frete_internacional_itens_lista(
            client_id,
            itens_lista,
            loja=store_id_cadastro,
        )
        file_bytes = _gerar_excel_lista_pedido_bytes(
            nome_lista,
            itens_lista_norm,
            client_id=client_id,
            loja=store_id_cadastro,
        )

        listas_salvas.insert(0, {
            "id": lista_id,
            "nome_lista": nome_lista,
            "loja": loja_sel,
            "store_id": store_id_cadastro,
            "status": "Lista gerada",
            "created_at": agora_iso,
            "updated_at": agora_iso,
            "opcao": opcao,
            "crescimento_percent": crescimento_percent,
            "itens": itens_lista_norm,
        })
        _salvar_listas_pedidos(client_id, listas_salvas)
        _salvar_bytes_cache_lista_pedido(client_id, lista_id, agora_iso, file_bytes)

        file_id = str(uuid.uuid4())
        TEMP_FILES_STORAGE[file_id] = file_bytes
        TEMP_FILES_META[file_id] = {
            "mime": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            "filename": nome_arquivo,
        }

        return {
            "success": True,
            "file_id": file_id,
            "lista_id": lista_id,
            "nome_lista": nome_lista,
            "loja": loja_sel,
            "filename": nome_arquivo,
            "total_itens": len(itens_lista),
            "total_skus_ocultados": len(hidden_skus),
            "opcao": opcao,
            "crescimento_percent": crescimento_percent,
            "periodo_meses": periodo_meses,
            "lead_time_meses": lead_time_meses,
            "total_quantidades_ajustadas": len(quantidades_sugeridas_aplicadas),
        }
    except HTTPException:
        raise
    except Exception as e:
        logger.exception(f"Erro ao gerar lista de compra: {e}")
        raise HTTPException(status_code=500, detail="Erro ao gerar lista de compra")


async def api_medias_compras_gerar_lista_compra_get(
    opcao: str,
    crescimento_percent: float | None = None,
    hidden_skus: str = "",
    nome_lista: str = "",
    loja: str = "__todas",
    store_id: str = "",
    periodo_meses: int = 6,
    quantidades_sugeridas: str = "",
    client_id: str = Depends(medias_common.get_tenant_id)
):
    req = ListaCompraRequest(
        opcao=opcao,
        crescimento_percent=crescimento_percent,
        hidden_skus=[s for s in (hidden_skus or "").split(",") if str(s).strip()],
        nome_lista=nome_lista,
        loja=loja,
        store_id=store_id,
        periodo_meses=periodo_meses,
        quantidades_sugeridas=_normalizar_quantidades_sugeridas(quantidades_sugeridas),
    )
    return await api_medias_compras_gerar_lista_compra(req, client_id)


async def api_medias_compras_gerar_lista_sugestao(
    meses: int = 12,
    loja: str = "__todas",
    store_id: str = "",
    quantidades_sugeridas: str = "",
    client_id: str = Depends(medias_common.get_tenant_id)
):
    escopo_loja = _resolver_escopo_loja_medias(
        client_id,
        loja,
        store_id,
        exigir_especifica=True,
    )
    loja = escopo_loja["loja"]
    store_id = escopo_loja["store_id"]
    data = await api_medias_compras_visao(
        meses=meses,
        loja=loja,
        store_id=store_id,
        client_id=client_id,
    )
    quantidades_ajustadas = _normalizar_quantidades_sugeridas(quantidades_sugeridas)
    quantidades_ajustadas_aplicadas = set()
    itens = []
    for item_original in (data.get("itens") or []):
        item = dict(item_original)
        item["compra_sugerida"] = _resolver_quantidade_sugerida(
            item.get("sku", ""),
            item.get("compra_sugerida", 0),
            quantidades_ajustadas,
            quantidades_ajustadas_aplicadas,
        )
        if int(item.get("compra_sugerida", 0) or 0) > 0:
            itens.append(item)
    itens_lista = [
        _normalizar_item_lista_pedido({
            "SKU": str(item.get("sku", "") or ""),
            "Foto": str(item.get("foto", "") or ""),
            TITULO_PRODUTO_INGLES_KEY: str(item.get("titulo_anuncio", "") or ""),
            "OEM": str(item.get("oem", "") or ""),
            COR_LADO_LISTA_PEDIDO_KEY: str(item.get(COR_LADO_LISTA_PEDIDO_KEY, item.get("color_side", "")) or ""),
            "Link": str(item.get("link", "") or ""),
            "Quantidade": int(item.get("compra_sugerida", 0) or 0),
            "Valor unidade": "",
            "Valor total": "",
        })
        for item in itens
    ]
    itens_lista = _recalcular_frete_internacional_itens_lista(
        client_id,
        itens_lista,
        loja=store_id,
    )

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Sugestões"
    headers = [
        "SKU",
        "Nome do Produto",
        "VMM",
        "Posição de Estoque (Físico + Trânsito)",
        "Quantidade Sugerida de Compra",
    ]
    ws.append(headers)

    cabecalho_fill = PatternFill(fill_type="solid", fgColor="1F4E78")
    cabecalho_font = Font(color="FFFFFF", bold=True, size=11)
    for idx in range(1, len(headers) + 1):
        c = ws.cell(row=1, column=idx)
        c.fill = cabecalho_fill
        c.font = cabecalho_font
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    for item in itens:
        ws.append([
            str(item.get("sku", "") or ""),
            str(item.get("titulo_anuncio", "") or ""),
            float(item.get("media_mensal", 0) or 0),
            float(item.get("posicao_estoque", 0) or 0),
            int(item.get("compra_sugerida", 0) or 0),
        ])

    for col, width in {"A": 16, "B": 52, "C": 12, "D": 20, "E": 20}.items():
        ws.column_dimensions[col].width = width
    ws.freeze_panes = "A2"

    buffer = io.BytesIO()
    wb.save(buffer)
    file_bytes = buffer.getvalue()
    file_bytes = _gerar_excel_lista_pedido_bytes(
        "sugestao_compra",
        itens_lista,
        client_id=client_id,
        loja=store_id,
    )
    file_id = str(uuid.uuid4())
    filename = f"sugestao_compra_{datetime.now().strftime('%Y%m%d_%H%M%S')}.xlsx"
    TEMP_FILES_STORAGE[file_id] = file_bytes
    TEMP_FILES_META[file_id] = {
        "mime": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        "filename": filename,
    }

    return {
        "success": True,
        "file_id": file_id,
        "filename": filename,
        "total_itens": len(itens),
        "loja": str(data.get("loja") or loja),
        "store_id": store_id,
        "periodo_meses": meses,
        "total_quantidades_ajustadas": len(quantidades_ajustadas_aplicadas),
    }


async def api_medias_compras_produtos_sem_venda(
    faixa: str = "3m",
    order: str = "desc",
    loja: str = "__todas",
    store_id: str = "",
    client_id: str = Depends(medias_common.get_tenant_id)
):
    faixa_sel = str(faixa or "3m").strip().lower()
    if faixa_sel not in {"3m", "6m", "6m+"}:
        faixa_sel = "3m"

    order_sel = "asc" if str(order or "desc").strip().lower() == "asc" else "desc"
    data = await api_medias_compras_visao(
        meses=12,
        loja=loja,
        store_id=store_id,
        client_id=client_id,
    )

    itens_base = []
    count_3m = 0
    count_6m = 0
    count_6m_mais = 0

    for item in (data.get("itens") or []):
        meses_sem_vender = int(item.get("meses_sem_vender", 0) or 0)
        saldo = float(item.get("saldo_atual_estoque", 0) or 0)
        if saldo <= 0:
            continue

        if meses_sem_vender >= 3 and meses_sem_vender < 6:
            count_3m += 1
        elif meses_sem_vender == 6:
            count_6m += 1
        elif meses_sem_vender > 6:
            count_6m_mais += 1

        incluir = False
        if faixa_sel == "3m":
            incluir = meses_sem_vender >= 3 and meses_sem_vender < 6
        elif faixa_sel == "6m":
            incluir = meses_sem_vender == 6
        elif faixa_sel == "6m+":
            incluir = meses_sem_vender > 6

        if not incluir:
            continue

        itens_base.append({
            "sku": str(item.get("sku", "") or ""),
            "foto": str(item.get("foto", "") or ""),
            "titulo_anuncio": str(item.get("titulo_anuncio", "") or ""),
            "ultima_venda": str(item.get("ultima_venda", "") or ""),
            "meses_sem_vender": meses_sem_vender,
            "saldo_atual_estoque": round(saldo, 2),
            "aviso_sem_venda": str(item.get("aviso_sem_venda", "") or ""),
        })

    itens_base.sort(
        key=lambda i: (
            int(i.get("meses_sem_vender", 0) or 0),
            float(i.get("saldo_atual_estoque", 0) or 0),
            str(i.get("sku", "") or "")
        ),
        reverse=(order_sel == "desc")
    )

    return {
        "success": True,
        "faixa": faixa_sel,
        "order": order_sel,
        "loja": str(data.get("loja") or loja),
        "store_id": str(data.get("store_id") or ""),
        "items": itens_base,
        "resumo": {
            "total": len(itens_base),
            "faixa_3m": count_3m,
            "faixa_6m": count_6m,
            "faixa_6m_mais": count_6m_mais,
        },
    }



def configure_medias_compras_sugestoes_runtime(runtime_module=None):
    return _configure_runtime_globals(runtime_module)


configure_medias_compras_sugestoes_runtime()

__all__ = [
    "configure_medias_compras_sugestoes_runtime",
    "api_medias_compras_gerar_lista_compra",
    "api_medias_compras_gerar_lista_compra_get",
    "api_medias_compras_gerar_lista_sugestao",
    "api_medias_compras_produtos_sem_venda",
]
