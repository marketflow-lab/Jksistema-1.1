"""Supplier Packing List ingestion and list-scoped XLSX generation."""

from __future__ import annotations

import copy
import hashlib
import io
import math
import os
import re
import tempfile
import textwrap
import uuid
import zipfile
import xml.etree.ElementTree as ET
from datetime import datetime
from xml.sax.saxutils import escape

import openpyxl
from fastapi import Depends, File, HTTPException, UploadFile
from fastapi.responses import StreamingResponse
from openpyxl.compat import safe_string
from openpyxl.drawing.image import Image as XLImage
from openpyxl.drawing.spreadsheet_drawing import AnchorMarker, OneCellAnchor
from openpyxl.drawing.xdr import XDRPositiveSize2D
from openpyxl.styles import Alignment, Border, Font, Side
from openpyxl.utils import get_column_letter
from openpyxl.utils.cell import range_boundaries
from openpyxl.utils.units import pixels_to_EMU, points_to_pixels
from openpyxl.worksheet.formula import ArrayFormula
from backend.services.medias_compras_documentos_cadastro import (
    _contexto_cadastro_documentos, _dados_produto_documento,
)

from backend.services import medias_compras_common as common
from backend.services import medias_compras_excel as excel


MAX_PACKING_BYTES = 20 * 1024 * 1024
PACKING_SCHEMA = "jk.packing-list.v1"


def _numero_packing(valor, *, obrigatorio=False):
    if valor is None or (isinstance(valor, str) and not valor.strip()):
        if obrigatorio:
            raise ValueError("Preencha os dados de quantidade, caixas, medidas e pesos.")
        return None
    numero = excel._to_float_excel(valor)
    if isinstance(valor, bool) or numero is None or not math.isfinite(numero) or numero < 0:
        raise ValueError("O Packing List contém um valor numérico inválido.")
    return int(numero) if numero.is_integer() else numero


def _embalagem_packing(valor):
    """Keep trays distinct from cartons instead of dropping the unit label."""
    tipo = "ctns"
    if isinstance(valor, str):
        texto = valor.strip()
        trays = re.fullmatch(r"(?i)([0-9]+)\s+trays?", texto)
        if trays:
            valor, tipo = trays.group(1), "trays"
        elif not re.fullmatch(r"\+?[0-9]+(?:[.,][0-9]+)?", texto):
            raise ValueError("Informe uma quantidade inteira e positiva de caixas ou Trays, como 3 Trays.")
    quantidade = _numero_packing(valor, obrigatorio=True)
    if quantidade <= 0 or quantidade != int(quantidade):
        raise ValueError("Informe uma quantidade inteira e positiva de caixas ou Trays.")
    return int(quantidade), tipo


def _totais_embalagens_packing(grupos):
    totais = {"ctns": 0, "trays": 0}
    origens = set()
    for indice, grupo in enumerate(grupos):
        origem = grupo.get("ctns_origem") or f"grupo-{indice}"
        if origem not in origens:
            origens.add(origem)
            totais[grupo.get("tipo_embalagem", "ctns")] += grupo["ctns"]
    return totais


def _ler_packing_list_bytes(conteudo: bytes) -> dict:
    """Read the supplier's 15-column layout, keeping mixed cartons together."""
    if not conteudo or len(conteudo) > MAX_PACKING_BYTES:
        raise HTTPException(400, "Envie um Packing List .xlsx de até 20 MB.")
    wb = None
    try:
        with zipfile.ZipFile(io.BytesIO(conteudo)) as archive:
            entries = archive.infolist()
            if (len(entries) > 2500 or sum(entry.file_size for entry in entries) > 100 * 1024 * 1024
                    or any(entry.flag_bits & 1 for entry in entries)):
                raise ValueError("Arquivo Excel excede o tamanho permitido.")
            # Validate XML coordinates before openpyxl expands merged cells.
            for entry in entries:
                if not entry.filename.startswith("xl/worksheets/") or not entry.filename.endswith(".xml"):
                    continue
                total_cells = total_merged = 0
                with archive.open(entry) as stream:
                    for _event, node in ET.iterparse(stream, events=("end",)):
                        tag = node.tag.rsplit("}", 1)[-1]
                        if tag in {"c", "mergeCell", "dimension"}:
                            reference = node.get("r") if tag == "c" else node.get("ref")
                            if reference:
                                min_col, min_row, max_col, max_row = range_boundaries(reference)
                                if (None in (min_col, min_row, max_col, max_row) or min_col < 1 or min_row < 1
                                        or max_row > 5000 or max_col > 100):
                                    raise ValueError("O Packing List excede o limite de linhas ou colunas.")
                                if tag == "mergeCell":
                                    total_merged += (max_row - min_row + 1) * (max_col - min_col + 1)
                                elif tag == "c":
                                    total_cells += 1
                                if total_cells > 100000 or total_merged > 100000:
                                    raise ValueError("O Packing List excede o limite de células.")
                        node.clear()
        wb = openpyxl.load_workbook(io.BytesIO(conteudo), data_only=True, keep_links=False)
        ws = None
        cabecalho = None
        for candidato in wb.worksheets:
            for linha in range(1, min(candidato.max_row, 12) + 1):
                a, b, d, f = (excel._normalizar_cabecalho_excel(candidato.cell(linha, coluna).value)
                              for coluna in (1, 2, 4, 6))
                if "cartonno" in a and "itemno" in b and "qty" in d and "ctns" in f:
                    ws, cabecalho = candidato, linha
                    break
            if ws is not None:
                break
        if ws is None or cabecalho is None:
            raise ValueError("Use o modelo do fornecedor com Carton No, Item No, QTY e Total CTNS.")
        if ws.max_row > 5000 or ws.max_column > 100:
            raise ValueError("O Packing List excede o limite de linhas ou colunas.")

        merged = {}
        grupos_linhas = {}
        for intervalo in ws.merged_cells.ranges:
            if intervalo.min_col == 1 and intervalo.max_col == 1:
                for linha in range(intervalo.min_row, intervalo.max_row + 1):
                    grupos_linhas[linha] = (intervalo.min_row, intervalo.max_row)
            for linha in range(intervalo.min_row, intervalo.max_row + 1):
                for coluna in range(intervalo.min_col, min(intervalo.max_col, 15) + 1):
                    merged[(linha, coluna)] = (intervalo.min_row, intervalo.min_col)

        def valor(linha, coluna):
            inicio = merged.get((linha, coluna), (linha, coluna))
            return ws.cell(*inicio).value

        fotos = {}
        for foto in ws._images:
            marker = getattr(foto.anchor, "_from", None)
            if marker is not None and marker.col == 2:
                fotos.setdefault(marker.row + 1, foto._data())

        grupos = []
        por_inicio = {}
        for linha in range(cabecalho + 1, ws.max_row + 1):
            sku = ws.cell(linha, 2).value
            if sku is None or not str(sku).strip():
                continue
            if excel._linha_resumo_excel_sku(str(sku)):
                break
            caixa = valor(linha, 1)
            if caixa is not None and excel._linha_resumo_excel_sku(str(caixa)):
                break
            # The second, bilingual header row is not an item.
            if "itemno" in excel._normalizar_cabecalho_excel(sku):
                continue
            inicio, _fim = grupos_linhas.get(linha, (linha, linha))
            origem_ctns = merged.get((linha, 6), (linha, 6))
            chave_grupo = (inicio, origem_ctns)
            if chave_grupo not in por_inicio:
                ctns, tipo_embalagem = _embalagem_packing(valor(linha, 6))
                dimensoes = [_numero_packing(valor(linha, coluna), obrigatorio=True) for coluna in (7, 8, 9)]
                if any(numero <= 0 for numero in dimensoes):
                    raise ValueError("Informe as três medidas positivas da caixa em centímetros.")
                cbm_ctn = _numero_packing(valor(linha, 10))
                if cbm_ctn is None:
                    cbm_ctn = math.prod(dimensoes) / 1_000_000
                cbm_total = _numero_packing(valor(linha, 11))
                nw_ctn = _numero_packing(valor(linha, 12), obrigatorio=True)
                gw_ctn = _numero_packing(valor(linha, 13), obrigatorio=True)
                nw_total = _numero_packing(valor(linha, 14))
                gw_total = _numero_packing(valor(linha, 15))
                origens_totais = {}
                for campo, coluna, informado in (("cbm_total", 11, cbm_total),
                                                 ("nw_total", 14, nw_total), ("gw_total", 15, gw_total)):
                    if informado is not None:
                        origem = merged.get((linha, coluna), (linha, coluna))
                        origens_totais[campo] = f"{get_column_letter(origem[1])}{origem[0]}"
                grupo = {
                    "caixas": str(caixa or "").strip(), "ctns": int(ctns), "dimensoes": dimensoes,
                    "tipo_embalagem": tipo_embalagem, "ctns_origem": f"{get_column_letter(origem_ctns[1])}{origem_ctns[0]}",
                    "totais_origem": origens_totais,
                    "cbm_ctn": cbm_ctn, "cbm_total": cbm_total if cbm_total is not None else cbm_ctn * ctns,
                    "nw_ctn": nw_ctn, "nw_total": nw_total if nw_total is not None else nw_ctn * ctns,
                    "gw_ctn": gw_ctn, "gw_total": gw_total if gw_total is not None else gw_ctn * ctns,
                    "itens": [],
                }
                grupos.append(grupo)
                por_inicio[chave_grupo] = grupo
            grupo = por_inicio[chave_grupo]
            quantidade = _numero_packing(ws.cell(linha, 4).value, obrigatorio=True)
            qty_ctn = _numero_packing(valor(linha, 5))
            grupo["itens"].append({
                "sku": str(sku).strip(), "quantidade": quantidade,
                "qty_ctn": qty_ctn if qty_ctn is not None else quantidade / grupo["ctns"],
                "foto": fotos.get(linha),
            })
        if not grupos or sum(len(grupo["itens"]) for grupo in grupos) > 2000:
            raise ValueError("O Packing List precisa conter de 1 a 2.000 linhas de produtos.")
        for grupo in grupos:
            if len(grupo["itens"]) > 1:
                for item in grupo["itens"]:
                    item["qty_ctn"] = item["quantidade"] / grupo["ctns"]
        return {"grupos": grupos}
    except HTTPException:
        raise
    except Exception as error:
        detail = str(error) if isinstance(error, ValueError) else "Arquivo inválido. Envie um Packing List Excel .xlsx."
        raise HTTPException(400, detail) from None
    finally:
        if wb is not None:
            wb.close()


def _chave_sku_packing(valor):
    return re.sub(r"(?i)^sku[\s:]+", "", str(valor or "").strip()).strip().casefold()


def _texto_comprador_packing(lista, client_id, *, contexto_cadastro=None):
    loja = str(lista.get("loja") or "").strip()
    texto = f"STORE: {loja}" if loja and loja != "__todas" else ""
    if not client_id:
        return texto
    from backend.services.cadastro_compatibilidade import visao_produtos_cadastro_contexto_loja
    from backend.services.cadastro_importador_loja import ler_importador_loja_por_store_id
    try:
        contexto = contexto_cadastro if contexto_cadastro is not None else visao_produtos_cadastro_contexto_loja(client_id, lista.get("store_id") or loja)
    except RuntimeError:
        return texto
    store_id = str(contexto.get("store_id") or "").strip()
    if not store_id:
        return texto
    cadastro = ler_importador_loja_por_store_id(client_id, store_id)
    if not cadastro:
        return texto
    campos = (("Buyer", "nome_empresa"), ("TAX ID", "tax_id"), ("Phone", "telefone"),
              ("E-mail", "email"), ("Contact Personal", "contato"), ("Address", "logradouro"),
              ("Neighborhood", "bairro"), ("City", "cidade"), ("ZIP CODE", "cep"),
              ("State", "estado"), ("Country", "pais"))
    return "\n".join(f"{rotulo}: {valor}" for rotulo, chave in campos
                     if (valor := str(cadastro.get(chave) or "").strip()))


def _data_packing(lista):
    data = str(lista.get("order_date") or "").strip()
    if not data:
        data = max((str(item.get("compra_aprovada_em") or item.get("Compra aprovada em") or "")
                    for item in lista.get("itens", []) if isinstance(item, dict)), default="")
    formatada = excel._formatar_data_commercial_invoice(data)
    try:
        parsed = datetime.strptime(formatada, "%d/%m/%Y")
    except ValueError:
        return formatada
    meses = ("January", "February", "March", "April", "May", "June", "July", "August", "September", "October", "November", "December")
    return f"{meses[parsed.month - 1]} {parsed.day}, {parsed.year}"


def _preencher_cache_totais_packing(conteudo, totais):
    """Keep SUM formulas editable and their results visible without recalculation."""
    output = io.BytesIO()
    with zipfile.ZipFile(io.BytesIO(conteudo)) as source, zipfile.ZipFile(output, "w") as target:
        target.comment = source.comment
        for info in source.infolist():
            content = source.read(info)
            if info.filename == "xl/worksheets/sheet1.xml":
                for endereco, valor in totais.items():
                    padrao = re.compile(
                        rb'(<c\b[^>]*\br="' + endereco.encode("ascii") + rb'"[^>]*>)(.*?)(</c>)',
                        re.DOTALL,
                    )

                    def atualizar(match):
                        abertura = match.group(1)
                        corpo = match.group(2)
                        if not re.search(rb"<f(?:\s[^>]*)?>[^<]+</f>", corpo):
                            raise ValueError("Fórmula de total do Packing List ausente.")
                        if isinstance(valor, str):
                            abertura = re.sub(rb'\s+t="[^"]*"', b"", abertura)
                            abertura = abertura[:-1] + b' t="str">'
                            resultado = escape(valor).encode("utf-8")
                        else:
                            resultado = safe_string(valor).encode("ascii")
                        corpo, quantidade = re.subn(
                            rb"<v>[^<]*</v>", b"<v>" + resultado + b"</v>", corpo,
                        )
                        if quantidade != 1:
                            raise ValueError("Resultado de total do Packing List inválido.")
                        return abertura + corpo + match.group(3)

                    content, quantidade = padrao.subn(atualizar, content)
                    if quantidade != 1:
                        raise ValueError("Célula de total do Packing List inválida.")
            target.writestr(info, content)
    return output.getvalue()


def _gerar_packing_list_bytes(lista, dados, *, client_id=None) -> bytes:
    grupos = dados["grupos"]
    contexto_cadastro = _contexto_cadastro_documentos(
        lista, client_id, skus=[item["sku"] for grupo in grupos for item in grupo["itens"]],
    )
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Packing List"
    ws.sheet_view.showGridLines = True
    ws.page_setup.orientation = "landscape"
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_margins.header = ws.page_margins.footer = 0
    larguras = (10.25, 11, 17.88, 38.88, 8.63, 6.38, 8.38, 10.38, 4.75, 3.63, 3.63, 7.88, 10.38, 8.13, 11.38, 8.75, 9.13)
    for coluna, largura in enumerate(larguras, 1):
        ws.column_dimensions[get_column_letter(coluna)].width = largura
    for linha, altura in enumerate((132.75, 33, 18, 30, 75.75, 47.25, 30, 39), 1):
        ws.row_dimensions[linha].height = altura
    for intervalo in ("A1:Q1", "A2:Q2", "A3:D3", "E3:H6", "I3:Q3", "A4:D6", "I4:Q6",
                      "A7:A8", "B7:B8", "C7:C8", "D7:D8", "E7:F7", "I7:K8",
                      "L7:M7", "N7:O7", "P7:Q7"):
        ws.merge_cells(intervalo)
    fornecedor = excel._cadastro_fornecedor_commercial_invoice(lista, client_id)
    texto_fornecedor, _altura = excel._cabecalho_fornecedor_commercial_invoice(lista, client_id, cadastro=fornecedor)
    importer, _altura = excel._bloco_nacionalizador_commercial_invoice(lista, client_id)
    excel._atribuir_texto_excel(ws["A1"], str(texto_fornecedor))
    excel._atribuir_texto_excel(ws["A2"], "PACKING LIST")
    invoice = str(lista.get("numero_invoice") or lista.get("invoice") or "").strip()
    excel._atribuir_texto_excel(ws["A3"], f"Proforma Invoice NO.:  {invoice}")
    data = _data_packing(lista)
    excel._atribuir_texto_excel(ws["I3"], f"Date:  {data}" if data else "Date:")
    excel._atribuir_texto_excel(ws["A4"], importer or (f"LIST: {lista.get('nome_lista')}" if lista.get("nome_lista") else ""))
    excel._atribuir_texto_excel(ws["I4"], _texto_comprador_packing(lista, client_id, contexto_cadastro=contexto_cadastro))
    headers = {
        "A7": "箱号\nCARTON NO.", "B7": "款号\nITEM NO.", "C7": "图片\nPICTURE",
        "D7": "描述\nDESC.", "E7": "数量\nQTY（PCS）", "G7": "每件数量\nQTY/CTN",
        "H7": "总件数\nTOTAL CTNS", "I7": "每件尺寸\nCTN SIZE", "L7": "体积\nCBM",
        "N7": "净重\nN.W", "P7": "毛重\nG.W", "E8": "左/右\nL/R", "F8": "总数\nTOTAL",
        "G8": "（PCS）", "H8": "（CTN）", "L8": "每件体积\nCBM/CTN", "M8": "总体积\nTOTAL CBM",
        "N8": "每件净重\nN.W/CTN", "O8": "总净重\nTOTAL N.W", "P8": "每件毛重\nG.W/CTN",
        "Q8": "总毛重\nTOTAL G.W",
    }
    for endereco, texto in headers.items():
        excel._atribuir_texto_excel(ws[endereco], texto)

    por_sku = {}
    for item in lista.get("itens", []):
        if isinstance(item, dict):
            por_sku.setdefault(_chave_sku_packing(item.get("SKU") or item.get("sku")), item)
    linha = 9
    intervalos_ctns = {}
    intervalos_totais = {13: {}, 15: {}, 17: {}}
    for grupo in grupos:
        inicio = linha
        fim = inicio + len(grupo["itens"]) - 1
        for item in grupo["itens"]:
            chave_item = _chave_sku_packing(item["sku"])
            cadastro = por_sku.get(chave_item)
            if cadastro is None:
                # Supplier Item No may append an OE/model after the SKU label.
                rotulado = re.match(r"(?i)^sku[\s:]+([^\s]+)", item["sku"])
                cadastro = por_sku.get(_chave_sku_packing(rotulado.group(1)), {}) if rotulado else {}
            descricao = common._primeiro_texto_item(cadastro, [common.TITULO_PRODUTO_INGLES_KEY, "Título em inglês", "Description", "Descrição", "Descricao"])
            unidade = common._primeiro_texto_item(cadastro, ["Unit/Pair", "Unit Pair", common.COR_LADO_LISTA_PEDIDO_KEY, "Color / Side", "Cor/Lado", "Lado"])
            produto_documento = _dados_produto_documento(contexto_cadastro, item["sku"])
            if client_id:
                descricao = produto_documento["descricao"]
                unidade = produto_documento["unidade"]
            excel._atribuir_texto_excel(ws.cell(linha, 2), item["sku"])
            excel._atribuir_texto_excel(ws.cell(linha, 4), descricao)
            excel._atribuir_texto_excel(ws.cell(linha, 5), unidade)
            ws.cell(linha, 7).value = item["qty_ctn"]
            linhas_desc = sum(max(1, len(textwrap.wrap(parte, width=42))) for parte in descricao.splitlines())
            ws.row_dimensions[linha].height = max(50.25 if linha < 33 else 45, 8 + linhas_desc * 12)
            caminho_foto = None
            if client_id and produto_documento["foto"]:
                caminho_foto = excel._resolver_caminho_foto_cadastro_seguro(
                    client_id, produto_documento["foto"], produto_documento["sku"],
                    contexto_cadastro.get("store_id") or None,
                    permitir_default=False,
                )
            fonte_foto = caminho_foto if client_id else (io.BytesIO(item["foto"]) if item.get("foto") else None)
            if fonte_foto:
                try:
                    foto = XLImage(fonte_foto)
                except (OSError, ValueError):
                    foto = None
            else:
                foto = None
            if foto is not None:
                width, height = larguras[2] * 7, points_to_pixels(ws.row_dimensions[linha].height)
                escala = min((width - 8) / foto.width, (height - 8) / foto.height)
                foto.width, foto.height = foto.width * escala, foto.height * escala
                foto.anchor = OneCellAnchor(
                    _from=AnchorMarker(col=2, row=linha - 1,
                                       colOff=pixels_to_EMU((width - foto.width) / 2),
                                       rowOff=pixels_to_EMU((height - foto.height) / 2)),
                    ext=XDRPositiveSize2D(pixels_to_EMU(foto.width), pixels_to_EMU(foto.height)),
                )
                ws.add_image(foto)
            linha += 1
        for coluna in (1, 6, 9, 10, 11, 12, 14, 16):
            if fim > inicio:
                ws.merge_cells(start_row=inicio, start_column=coluna, end_row=fim, end_column=coluna)
        excel._atribuir_texto_excel(ws.cell(inicio, 1), grupo["caixas"])
        origem_ctns = grupo.get("ctns_origem") or f"grupo-{inicio}"
        if origem_ctns in intervalos_ctns:
            intervalos_ctns[origem_ctns][1] = fim
        else:
            valor_ctns = f'{grupo["ctns"]} Trays' if grupo.get("tipo_embalagem") == "trays" else grupo["ctns"]
            intervalos_ctns[origem_ctns] = [inicio, fim, valor_ctns]
        valores = {6: sum(item["quantidade"] for item in grupo["itens"]),
                   9: grupo["dimensoes"][0], 10: grupo["dimensoes"][1], 11: grupo["dimensoes"][2],
                   12: round(grupo["cbm_ctn"], 3), 14: grupo["nw_ctn"], 16: grupo["gw_ctn"]}
        for coluna, valor in valores.items():
            ws.cell(inicio, coluna).value = valor
        for coluna, campo in ((13, "cbm_total"), (15, "nw_total"), (17, "gw_total")):
            origem = grupo.get("totais_origem", {}).get(campo) or f"grupo-{inicio}"
            intervalos = intervalos_totais[coluna]
            if origem in intervalos:
                intervalos[origem][1] = fim
            else:
                valor = round(grupo[campo], 3) if coluna == 13 else grupo[campo]
                intervalos[origem] = [inicio, fim, valor]

    for inicio, fim, valor_ctns in intervalos_ctns.values():
        if fim > inicio:
            ws.merge_cells(start_row=inicio, start_column=8, end_row=fim, end_column=8)
        ws.cell(inicio, 8).value = valor_ctns
    for coluna, intervalos in intervalos_totais.items():
        for inicio, fim, valor in intervalos.values():
            if fim > inicio:
                ws.merge_cells(start_row=inicio, start_column=coluna, end_row=fim, end_column=coluna)
            ws.cell(inicio, coluna).value = valor

    total = linha
    ultima = total + 1
    for intervalo in (f"A{total}:E{ultima}", f"F{total}:G{total}", f"F{ultima}:G{ultima}",
                      f"I{total}:L{ultima}", f"N{total}:N{ultima}", f"P{total}:P{ultima}"):
        ws.merge_cells(intervalo)
    excel._atribuir_texto_excel(ws.cell(total, 1), "TOTAL:")
    totais = {}
    embalagens = _totais_embalagens_packing(grupos)
    for coluna, rotulo in ((6, "Units Total"), (8, "CTNS TOTAL"), (13, "CBM TOTAL"),
                          (15, "TOTAL NET WEIGHT"), (17, "TOTAL GROSS WEIGHT")):
        letra = get_column_letter(coluna)
        if coluna == 8 and embalagens["trays"]:
            caixas = f"SUM(H9:H{total - 1})"
            trays = f'(SUMPRODUCT(IFERROR(1*SUBSTITUTE(H9:H{total - 1}," Trays",""),0))-{caixas})'
            formula = (
                f'=IF({trays}>0,IF({caixas}>0,{caixas}&" CTNS + ","")&{trays}&" Trays",{caixas})'
            )
            ws.cell(total, coluna).value = ArrayFormula(ref=f"H{total}", text=formula)
            prefixo = f'{embalagens["ctns"]} CTNS + ' if embalagens["ctns"] else ""
            totais[f"H{total}"] = f'{prefixo}{embalagens["trays"]} Trays'
        else:
            valores = [ws.cell(r, coluna).value for r in range(9, total)
                       if ws.cell(r, coluna).value is not None]
            totais[f"{letra}{total}"] = math.fsum(valores)
            ws.cell(total, coluna).value = f"=SUM({letra}9:{letra}{total - 1})"
        excel._atribuir_texto_excel(ws.cell(ultima, coluna), rotulo)
    ws.row_dimensions[total].height = ws.row_dimensions[ultima].height = 45
    centro = Alignment(horizontal="center", vertical="center", wrap_text=True)
    borda = Border(**{edge: Side(style="thin", color="000000") for edge in ("left", "right", "top", "bottom")})
    for row in ws.iter_rows(min_row=1, max_row=ultima, min_col=1, max_col=17):
        for cell in row:
            cell.alignment = centro
            cell.border = borda
            cell.font = Font(name="Arial", size=10 if cell.column in (2, 4) else 8,
                             bold=cell.column in (1, 8, 13, 15, 17))
            if cell.column in (6, 7, 8):
                cell.number_format = "General"
            if cell.column in (12, 13):
                cell.number_format = "General"
            if cell.column in (14, 15, 16, 17):
                cell.number_format = "0.00"
    for row in ws.iter_rows(min_row=7, max_row=8, min_col=1, max_col=17):
        for cell in row:
            cell.font = Font(name="SimSun", size=10, bold=True)
    for endereco in ("A1", "A2"):
        ws[endereco].font = Font(name="Arial", size=18, bold=True)
    ws["A3"].font = Font(name="Arial", size=10)
    ws["I3"].font = Font(name="Arial", size=10, bold=True)
    ws["I3"].alignment = Alignment(horizontal="right", vertical="center", wrap_text=True)
    ws["A4"].font = Font(name="Arial", size=14, bold=True)
    ws["I4"].font = Font(name="Arial", size=11, bold=True)
    # Keep long current registered details legible in the merged header blocks.
    for endereco, tamanho, largura in (("A1", 18, 105), ("A4", 14, 47), ("I4", 11, 53)):
        texto = str(ws[endereco].value or "")
        linhas_visuais = sum(max(1, len(textwrap.wrap(parte, width=largura))) for parte in texto.splitlines())
        necessaria = 12 + linhas_visuais * tamanho * 1.3
        if endereco == "A1":
            ws.row_dimensions[1].height = min(409.5, max(132.75, necessaria))
        else:
            extra = max(0, necessaria - sum(ws.row_dimensions[r].height for r in (4, 5, 6)))
            ws.row_dimensions[6].height = min(409.5, ws.row_dimensions[6].height + extra)
    ws.print_area = f"A1:Q{ultima}"
    output = io.BytesIO()
    wb.save(output)
    wb.close()
    return _preencher_cache_totais_packing(output.getvalue(), totais)


def _arquivo_packing(client_id, documento_id):
    if not re.fullmatch(r"[0-9a-f]{32}", str(documento_id or "")):
        raise HTTPException(404, "Nenhum Packing List adicionado nesta lista.")
    tenant = os.path.realpath(os.path.dirname(common._arquivo_listas_pedidos(client_id)))
    root = os.path.realpath(os.path.join(tenant, "packing_lists"))
    if os.path.commonpath([tenant, root]) != tenant:
        raise HTTPException(404, "Nenhum Packing List adicionado nesta lista.")
    caminho = os.path.realpath(os.path.join(root, f"{documento_id}.xlsx"))
    if os.path.commonpath([root, caminho]) != root:
        raise HTTPException(404, "Nenhum Packing List adicionado nesta lista.")
    return caminho


def _documento_packing_existe(client_id, metadata):
    if not isinstance(metadata, dict) or metadata.get("schema") != PACKING_SCHEMA:
        return False
    try:
        return os.path.isfile(_arquivo_packing(client_id, metadata.get("documento_id")))
    except HTTPException:
        return False


def _importar_packing_list_sync(lista_id, conteudo, filename, client_id):
    if not str(filename or "").lower().endswith(".xlsx"):
        raise HTTPException(400, "Envie um Packing List Excel .xlsx.")
    dados = _ler_packing_list_bytes(conteudo)
    grupos = dados["grupos"]
    embalagens = _totais_embalagens_packing(grupos)
    metadata = {"schema": PACKING_SCHEMA, "documento_id": uuid.uuid4().hex,
                "sha256": hashlib.sha256(conteudo).hexdigest(),
                "nome_arquivo": str(filename).replace("\\", "/").rsplit("/", 1)[-1][:255],
                "uploaded_at": datetime.now().isoformat(timespec="seconds"),
                "total_linhas": sum(len(g["itens"]) for g in grupos),
                "total_ctns": embalagens["ctns"], "total_trays": embalagens["trays"]}
    caminho = _arquivo_packing(client_id, metadata["documento_id"])

    @common._retry_listas_pedidos
    def gravar():
        listas = common._carregar_listas_pedidos(client_id)
        lista = next((row for row in listas if str(row.get("id", "")) == str(lista_id)), None)
        if lista is None:
            raise HTTPException(404, "Lista de pedidos não encontrada.")
        if _documento_packing_existe(client_id, lista.get("packing_list")):
            raise HTTPException(409, "Esta lista já tem um Packing List. Clique novamente para baixar.")
        os.makedirs(os.path.dirname(caminho), exist_ok=True)
        descriptor, temporario = tempfile.mkstemp(prefix=".packing-", suffix=".tmp", dir=os.path.dirname(caminho))
        try:
            with os.fdopen(descriptor, "wb") as stream:
                stream.write(conteudo)
                stream.flush()
                os.fsync(stream.fileno())
            os.replace(temporario, caminho)
            lista["packing_list"] = copy.deepcopy(metadata)
            lista["updated_at"] = metadata["uploaded_at"]
            common._salvar_listas_pedidos(client_id, listas)
        except Exception:
            # A failed snapshot refresh may happen *after* JSON was committed.
            # Remove the blob only when the canonical file proves it is unreferenced.
            try:
                with common.path_lock_for(common._arquivo_listas_pedidos(client_id)):
                    atuais = common._carregar_listas_pedidos(client_id)
                    referenciado = any(
                        (row.get("packing_list") or {}).get("documento_id") == metadata["documento_id"]
                        for row in atuais if isinstance(row, dict)
                    )
                    if not referenciado and os.path.isfile(caminho):
                        os.remove(caminho)
            except (OSError, HTTPException):
                pass
            raise
        finally:
            if os.path.exists(temporario):
                os.unlink(temporario)
        return {"success": True, "packing_list_ready": True}
    return gravar()


async def api_medias_compras_lista_pedido_packing_list_importar(
    lista_id: str, file: UploadFile = File(...), client_id: str = Depends(common.get_tenant_id),
):
    from backend.services.blocking_workers import run_heavy
    conteudo = await file.read(MAX_PACKING_BYTES + 1)
    return await run_heavy(_importar_packing_list_sync, lista_id=lista_id, conteudo=conteudo,
                           filename=file.filename or "", client_id=client_id)


def _baixar_packing_list_sync(lista_id, client_id):
    listas = common._carregar_listas_pedidos(client_id)
    lista = next((row for row in listas if str(row.get("id", "")) == str(lista_id)), None)
    if lista is None:
        raise HTTPException(404, "Lista de pedidos não encontrada.")
    metadata = lista.get("packing_list") or {}
    if not _documento_packing_existe(client_id, metadata):
        raise HTTPException(404, "Nenhum Packing List adicionado nesta lista.")
    caminho = _arquivo_packing(client_id, metadata.get("documento_id"))
    try:
        with open(caminho, "rb") as stream:
            conteudo = stream.read(MAX_PACKING_BYTES + 1)
    except FileNotFoundError:
        raise HTTPException(404, "Nenhum Packing List adicionado nesta lista.") from None
    if hashlib.sha256(conteudo).hexdigest() != metadata.get("sha256"):
        raise HTTPException(409, "O documento de Packing List está inválido. Verifique o arquivo enviado.")
    dados = _ler_packing_list_bytes(conteudo)
    output = _gerar_packing_list_bytes(lista, dados, client_id=client_id)
    nome = str(lista.get("numero_invoice") or lista.get("invoice") or lista.get("nome_lista") or lista_id).strip()
    nome = re.sub(r"[^A-Za-z0-9_-]+", "_", nome).strip("_") or "lista"
    return StreamingResponse(io.BytesIO(output), media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                             headers={"Content-Disposition": f'attachment; filename="Packing_List_{nome}.xlsx"'})


async def api_medias_compras_lista_pedido_packing_list(lista_id: str, client_id: str = Depends(common.get_tenant_id)):
    from backend.services.blocking_workers import run_heavy
    return await run_heavy(_baixar_packing_list_sync, lista_id=lista_id, client_id=client_id)


def limpar_documento_packing_lista(client_id, lista):
    metadata = lista.get("packing_list") or {}
    if _documento_packing_existe(client_id, metadata):
        try:
            os.remove(_arquivo_packing(client_id, metadata.get("documento_id")))
        except OSError:
            # The list deletion is already committed; an in-flight download may hold the file.
            pass


__all__ = ["api_medias_compras_lista_pedido_packing_list", "api_medias_compras_lista_pedido_packing_list_importar"]
