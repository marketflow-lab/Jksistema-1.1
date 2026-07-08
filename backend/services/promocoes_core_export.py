"""Internal slice for promocoes_core."""

from __future__ import annotations

from __future__ import annotations
import functools
import io
import json
import logging
import math
import os
import re
import unicodedata
from typing import Any, Optional
import numpy as np
import openpyxl
import pandas as pd
import base64
import copy
import csv
import functools
import hashlib
import html as html_lib
import io
import json
import logging
import math
import os
import random
import re
import sqlite3
import tempfile
import threading
import time
import traceback
import unicodedata
import uuid
from collections import Counter, defaultdict
from concurrent.futures import ThreadPoolExecutor, as_completed
from datetime import date, datetime, timedelta, timezone
from decimal import Decimal, InvalidOperation
from pathlib import Path
from typing import Any, Optional
from urllib.parse import parse_qs, quote, unquote, urljoin, urlparse

import numpy as np
import pandas as pd
import requests
from bs4 import BeautifulSoup
from fastapi import Depends, File, Form, Header, HTTPException, Request, UploadFile
from fastapi.encoders import jsonable_encoder
from fastapi.responses import StreamingResponse
from backend.services.runtime_bridge import bind_runtime_globals
from backend.services.promocoes_common import *


def configure_promocoes_core_export_runtime(runtime_module=None, peers=None):
    _configure_common = globals().get("configure_promocoes_common_runtime")
    if callable(_configure_common):
        try:
            _configure_common(runtime_module)
        except TypeError:
            _configure_common()
    runtime = bind_runtime_globals(globals(), runtime_module)
    if peers:
        globals().update(peers)
    return runtime


configure_promocoes_core_export_runtime()


def limpar_dados_para_exportacao(df_original):
    df_limpo = df_original.copy()
    def limpar_valor(val):
        val_str = str(val)
        if "Participar" in val_str and "nÃ£o" not in val_str.lower() and "nao" not in val_str.lower(): return "Participar"
        if "NÃƒO PARTICIPAR" in val_str.upper() or "NAO PARTICIPAR" in val_str.upper(): return "NÃ£o participar"
        return val
    if "Participar ou nÃ£o" not in df_limpo.columns:
        if "AÃ§Ã£o" in df_limpo.columns:
            df_limpo["Participar ou nÃ£o"] = df_limpo["AÃ§Ã£o"]
        else:
            df_limpo["Participar ou nÃ£o"] = "Participar"
    df_limpo["Participar ou nÃ£o"] = df_limpo["Participar ou nÃ£o"].apply(limpar_valor)
    return df_limpo


def _normalizar_coluna_exportacao(valor: Any) -> str:
    texto = normalizar_texto(_limpar_nome_coluna(valor))
    return re.sub(r"[^a-z0-9]+", " ", texto).strip()


def _valor_df_por_alias(row, aliases: list[str]):
    aliases_norm = {_normalizar_coluna_exportacao(alias) for alias in aliases}
    for col in getattr(row, "index", []):
        if _normalizar_coluna_exportacao(col) in aliases_norm:
            valor = row.get(col)
            if valor is None:
                return ""
            try:
                if pd.isna(valor):
                    return ""
            except Exception:
                pass
            return valor
    return ""


def _normalizar_decisao_planilha(valor: Any) -> str:
    texto = str(valor if valor is not None else "").strip()
    if not texto or texto.lower() in {"nan", "none", "null", "nat", "<na>", "-"}:
        return ""
    norm = normalizar_texto(texto)
    if "participar" not in norm:
        return ""
    if "nao" in norm or "n o" in norm:
        return "NÃ£o participar"
    return "Participar"


def _decisao_linha_exportacao(row) -> str:
    for aliases in (
        ["AÃ§Ã£o", "Acao", "ACTION"],
        ["Participar ou nÃ£o", "Participar ou nao", "Participar"],
    ):
        decisao = _normalizar_decisao_planilha(_valor_df_por_alias(row, aliases))
        if decisao:
            return decisao
    return ""


def _localizar_coluna_em_linha(valores, aliases: list[str]) -> Optional[int]:
    aliases_norm = [_normalizar_coluna_exportacao(alias) for alias in aliases]
    valores_norm = [_normalizar_coluna_exportacao(valor) for valor in valores]
    for alias in aliases_norm:
        for idx, valor_norm in enumerate(valores_norm, start=1):
            if valor_norm == alias:
                return idx
    return None


def _localizar_cabecalho_exportacao(ws):
    key_aliases = [
        "ITEM_ID", "MLB", "ID anuncio", "ID anuncio", "ID do anuncio", "ID do anuncio",
        "NÃƒÂºmero do anuncio", "Numero do anuncio", "CÃƒÂ³digo do anuncio", "Codigo do anuncio",
        "AnÃƒÂºncio", "Anuncio", "ID",
    ]
    action_aliases = ["ACTION", "AÃ§Ã£o", "Acao"]

    max_rows = min(int(ws.max_row or 1), 40)
    for row_idx in range(1, max_rows + 1):
        valores = [cell.value for cell in ws[row_idx]]
        col_key = _localizar_coluna_em_linha(valores, key_aliases)
        if not col_key:
            continue
        col_action = _localizar_coluna_em_linha(valores, action_aliases)
        if col_action:
            return row_idx, col_key, col_action

    return None


def gerar_excel_atualizado(file_bytes, df_decisoes):
    try:
        with io.BytesIO(file_bytes) as b:
            wb = openpyxl.load_workbook(b)
            if not wb.sheetnames:
                return None
            sheet_name = find_sheet_name(wb)
            ws = wb[sheet_name]

            cabecalho = _localizar_cabecalho_exportacao(ws)
            if not cabecalho:
                raise ValueError("NÃƒÂ£o encontrei as colunas de MLB/item e ACTION/AÃ§Ã£o na planilha original.")
            header_row, col_key, col_action = cabecalho

            df_clean = limpar_dados_para_exportacao(df_decisoes)
            decisoes_por_mlb = {}
            for _, row in df_clean.iterrows():
                decisao = _decisao_linha_exportacao(row)
                if not decisao:
                    continue
                mlb = _promo_normalizar_mlb(_valor_df_por_alias(row, [
                    "MLB", "ITEM_ID", "ID anuncio", "ID anuncio", "ID do anuncio", "ID do anuncio",
                    "NÃƒÂºmero do anuncio", "Numero do anuncio", "CÃƒÂ³digo do anuncio", "Codigo do anuncio",
                    "AnÃƒÂºncio", "Anuncio", "ID",
                ]))
                if mlb:
                    decisoes_por_mlb[mlb] = decisao

            if not decisoes_por_mlb:
                raise ValueError("Nenhuma decisÃƒÂ£o valida foi enviada para exportaÃƒÂ§ÃƒÂ£o.")

            aplicadas = 0
            for row_idx in range(header_row + 1, int(ws.max_row or header_row) + 1):
                mlb = _promo_normalizar_mlb(ws.cell(row=row_idx, column=col_key).value)
                if not mlb:
                    continue
                decisao = decisoes_por_mlb.get(mlb)
                if not decisao:
                    continue
                ws.cell(row=row_idx, column=col_action).value = decisao
                aplicadas += 1

            if aplicadas == 0:
                raise ValueError("Nenhum MLB da analise foi encontrado na planilha original.")

            output = io.BytesIO()
            wb.save(output)
            output.seek(0)
            return output
    except Exception as e:
        print(f"Erro gerar excel: {e}")
        return None

PEER_EXPORTS = ['limpar_dados_para_exportacao', '_normalizar_coluna_exportacao', '_valor_df_por_alias', '_normalizar_decisao_planilha', '_decisao_linha_exportacao', '_localizar_coluna_em_linha', '_localizar_cabecalho_exportacao', 'gerar_excel_atualizado']
__all__ = PEER_EXPORTS + ["configure_promocoes_core_export_runtime"]

configure_promocoes_core_export_runtime()
