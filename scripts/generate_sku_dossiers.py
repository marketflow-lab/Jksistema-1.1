#!/usr/bin/env python3
"""Gera uma base de conhecimento rastreavel com um JSON por SKU.

O gerador combina, sem alterar as fontes de origem:

* cadastro_produtos.csv do tenant;
* vinculos locais e, opcionalmente, anuncios ativos da API do Mercado Livre;
* pesquisa publica opcional na internet.

Os campos tecnicos sao compostos somente de trechos encontrados nas fontes.
Quando uma informacao nao existe, o arquivo registra a lacuna em vez de inventar
OEM, medida, instalacao ou compatibilidade.

Para OEM, o campo ``produto_bling`` serve apenas como pista de nome. As
referencias partem do cadastro atual e do Mercado Livre e podem ser ampliadas
somente por pesquisa publica confirmatoria.
"""

from __future__ import annotations

import argparse
import csv
import hashlib
import html
import json
import os
import random
import re
import sys
import threading
import time
import unicodedata
from concurrent.futures import ThreadPoolExecutor, as_completed
from datetime import datetime, timezone
from pathlib import Path
from typing import Any, Iterable
from urllib.parse import parse_qs, unquote, urljoin, urlparse

import requests
import urllib3
from bs4 import BeautifulSoup
from openpyxl import load_workbook


SCHEMA_VERSION = 1
# O SKU permanece no cadastro comercial, mas não deve possuir dossiê publicado.
EXCLUDED_DOSSIER_SKUS = {"395"}
USER_AGENT = "Mozilla/5.0 (Windows NT 10.0; Win64; x64) JK-Sistema-SKU/1.0"
ML_API_BASE = "https://api.mercadolibre.com"
WINDOWS_INVALID_FILENAME = re.compile(r'[<>:"/\\|?*\x00-\x1f]')
SPACE_RE = re.compile(r"\s+")
YEAR_RE = re.compile(r"\b(?:19\d{2}|20\d{2})\b")
MEASURE_RE = re.compile(
    r"(?i)(?:\b(?:medidas?|dimens(?:ao|oes)|altura|largura|comprimento|diametro|espessura|peso|"
    r"distancia|rosca|encaixe|tubo|eixo|voltagem|tensao)\b[^.;\n]{0,100})?"
    r"\b\d+(?:[.,]\d+)?\s*(?:mm|cm|m|kg|g|pol(?:egadas?)?|\")\b"
)
OEM_LABEL_RE = re.compile(
    r"(?i)\b(?:oem|part\s*number|n[uú]mero\s+da\s+pe[cç]a|n[uú]mero\s+de\s+pe[cç]a|"
    r"c[oó]digo\s+original|c[oó]digo\s+oem|c[oó]digo\s+(?:de\s+)?refer[eê]ncia|"
    r"c[oó]digo\s+da\s+pe[cç]a|c[oó]digo\s+do\s+fabricante|refer[eê]ncias?)"
    r"\s*[:#-]?\s*([^\n;]{3,100})"
)
CODE_RE = re.compile(r"(?i)(?<![A-Z0-9])(?:[A-Z0-9][A-Z0-9./-]{3,28}[A-Z0-9])(?![A-Z0-9])")

# Exclusões editoriais comprovadas na revisão cruzada. Elas impedem que um
# atributo incorreto do marketplace ou um anúncio contaminado pelo reuso de SKU
# volte ao dossiê bruto antes da etapa de revisão final.
OEM_CODE_EXCLUSIONS = {
    "214": {"2573685687", "931102D000"},
    "225": {
        "06H906517B", "0280142459", "1024709650", "47372790658",
        "55559352", "55574685", "25192904",
    },
    "372": {
        "5790853526611", "55116901AA", "52028974AA", "52079880AA",
        "52079799AA", "4677493AA", "K55116901AA", "K52028974AA",
        "K52079880AA", "5058482AD", "52079778AA", "K04677493AA",
        "K52079778AA", "5058394AG", "K05058394AG", "5058446AH",
        "K05058446AH", "5058482AH",
    },
}

STOPWORDS = {
    "a", "ao", "aos", "as", "com", "cor", "da", "das", "de", "do", "dos", "e", "em",
    "lado", "modelo", "o", "os", "para", "por", "produto", "sem", "um", "uma", "unidade",
    "kit", "novo", "nova", "original", "preto", "preta", "branco", "branca",
}
MARKETING_TERMS = (
    "aproveite e compre", "atencao ao cliente", "compra efetuada", "devolucao sem custos",
    "envio imediato", "envio no mesmo dia", "estoque pode acabar", "fique a vontade",
    "garantia do vendedor", "garantia sem burocracia", "metodo de entrega", "mercado livre",
    "parcelamento", "pronta entrega", "quem somos", "sua satisfacao", "taxa de respostas",
    "transportadoras credenciadas", "campo de perguntas",
)
PURPOSE_TERMS = (
    "serve para", "utilizado para", "utilizada para", "responsavel por", "responsavel pela",
    "tem a funcao", "funcao do produto", "ideal para", "indicado para", "indicada para",
    "permite ", "protege ", "evita ", "substitui ", "reposicao", "destina-se", "repelente",
    "espanta ", "adestramento",
)
FUNCTION_TERMS = (
    "aciona", "automatic", "detecta", "envia sinal", "funciona", "funcionamento", "informa ao",
    "mede ", "movimenta", "opera ", "pressao", "regula", "responsavel por", "sensor", "transmite",
)
INSTALL_TERMS = (
    "encaix", "fixacao", "instal", "ligacao", "montagem", "plug and play", "substituicao",
    "troca", "oficina especializada", "conectar", "rosque",
)
TECH_TERMS = (
    "caracteristica", "especifica", "material", "marca", "modelo", "quantidade", "conteudo",
    "tensao", "voltagem", "potencia", "corrente", "frequencia", "pinos", "estrias", "dentes",
    "temperatura", "capacidade", "relacao de reducao", "rotacao", "lado", "cor", "acabamento",
)
COMPATIBILITY_TERMS = (
    "aplicacao", "aplica-se", "aplicavel", "compativel", "serve no", "serve para", "veiculos",
    "modelos", "anos", "ano ",
)
NON_AUTOMOTIVE_TERMS = (
    "bicicleta", "barco", "caiaque", "celular", "iphone", "maquina", "motor de popa", "piscina",
    "rocadeira", "smartphone", "tablet", "trator", "wireless",
)
VEHICLE_BRANDS = (
    "acura", "audi", "bmw", "chevrolet", "chery", "citroen", "dodge", "fiat", "ford", "honda",
    "hyundai", "jac", "jeep", "kia", "land rover", "mercedes", "mitsubishi", "nissan", "peugeot",
    "porsche", "ram", "renault", "subaru", "suzuki", "toyota", "volkswagen", "volvo", "vw",
)
PRODUCT_CLASSES: tuple[tuple[str, tuple[str, ...]], ...] = (
    ("motor_arranque", ("motor de arranque", "motor partida", "motor de partida")),
    ("arrefecimento", ("tubo de agua", "mangueira radiador", "mangueira de agua", "refrigeracao", "arrefecimento")),
    ("bomba_combustivel", ("bomba de combustivel", "bomba gasolina")),
    ("sensor", ("sensor", "cebolao", "interruptor termico")),
    ("carburador", ("carburador",)),
    ("grade_parachoque", ("grade para choque", "grade parachoque", "tela parachoque")),
    ("puxador_porta", ("puxador porta", "macaneta porta")),
    ("iluminacao", ("lampada", "luz led", "pisca", "lanterna", "farol")),
    ("transmissao", ("caixa de transferencia", "engrenagem", "cambio", "cardan")),
)

_SEARCH_LOCK = threading.Lock()
_SEARCH_NEXT_AT = 0.0


def _now_iso() -> str:
    return datetime.now(timezone.utc).isoformat(timespec="seconds")


def _fold(value: Any) -> str:
    text = str(value or "")
    text = unicodedata.normalize("NFKD", text)
    return "".join(char for char in text if not unicodedata.combining(char)).lower()


def _mojibake_score(value: str) -> int:
    return sum(value.count(marker) for marker in ("Ã", "Â", "â€", "�"))


def _repair_text(value: Any) -> str:
    text = html.unescape(str(value or "")).replace("\x00", "").strip()
    candidates = [text]
    current = text
    for _ in range(2):
        try:
            current = current.encode("latin-1").decode("utf-8")
        except (UnicodeEncodeError, UnicodeDecodeError):
            break
        candidates.append(current)
    text = min(candidates, key=lambda item: (_mojibake_score(item), len(item)))
    text = text.replace("\r\n", "\n").replace("\r", "\n")
    return SPACE_RE.sub(" ", text).strip()


def _is_missing_value(value: Any) -> bool:
    folded = re.sub(r"[_\s]+", " ", _fold(_repair_text(value))).strip(" .,:;-")
    if not folded:
        return True
    if re.search(
        r"\b(?:nao\s+(?:informad[oa]s?|localizad[oa]s?|encontrad[oa]s?|se\s+aplica)|"
        r"sem\s+informa(?:cao|coes))\b",
        folded,
    ):
        return True
    return bool(
        re.search(r"\bmedidas?.*\bembalagem\b.*\b(?:desconsiderad|descartad|removid)[oa]s?\b", folded)
    )


def _norm_sku(value: Any) -> str:
    text = _repair_text(value).strip().lstrip("'")
    if re.fullmatch(r"\d+\.0", text):
        text = text[:-2]
    return re.sub(r"\s+", "", text).upper()


def _resolve_catalog_sku(value: Any, catalog_skus: set[str]) -> str:
    normalized = _norm_sku(value)
    if normalized in catalog_skus:
        return normalized
    if normalized.isdigit():
        numeric = int(normalized)
        candidates = [item for item in catalog_skus if item.isdigit() and int(item) == numeric]
        if len(candidates) == 1:
            return candidates[0]
    compact = re.sub(r"[^A-Z0-9]", "", normalized)
    if compact:
        candidates = [item for item in catalog_skus if re.sub(r"[^A-Z0-9]", "", item) == compact]
        if len(candidates) == 1:
            return candidates[0]
    return normalized


def _safe_filename(sku: str) -> str:
    cleaned = WINDOWS_INVALID_FILENAME.sub("_", sku).strip(" .")
    return (cleaned or hashlib.sha256(sku.encode("utf-8")).hexdigest()[:16]) + ".json"


def _read_json(path: Path, default: Any) -> Any:
    try:
        return json.loads(path.read_text(encoding="utf-8-sig"))
    except (OSError, ValueError, TypeError):
        return default


def _load_web_research_overrides(path: Path, catalog_skus: set[str]) -> dict[str, dict[str, Any]]:
    raw = _read_json(path, {})
    if not isinstance(raw, dict):
        return {}
    output: dict[str, dict[str, Any]] = {}
    for raw_sku, payload in raw.items():
        sku = _resolve_catalog_sku(raw_sku, catalog_skus)
        if sku not in catalog_skus or not isinstance(payload, dict):
            continue
        results = [
            {
                "titulo": _repair_text(item.get("titulo")),
                "url": _repair_text(item.get("url")),
                "trecho": _repair_text(item.get("trecho")),
            }
            for item in payload.get("resultados") or []
            if isinstance(item, dict) and str(item.get("url") or "").startswith(("http://", "https://"))
        ]
        output[sku] = {
            "status": "pesquisado_com_resultado" if results else "pesquisado_sem_resultado",
            "provedor": _repair_text(payload.get("provedor")) or "pesquisa_curada",
            "consulta": _repair_text(payload.get("consulta")),
            "consultado_em": _repair_text(payload.get("consultado_em")) or _now_iso(),
            "resultados": results[:8],
        }
        observacao = _repair_text(payload.get("observacao"))
        if observacao:
            output[sku]["observacao"] = observacao
    return output


def _write_json_atomic(path: Path, payload: Any) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    temp = path.with_name(path.name + ".tmp")
    temp.write_text(json.dumps(payload, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    os.replace(temp, path)


def _read_csv(path: Path) -> list[dict[str, str]]:
    last_error: Exception | None = None
    for encoding in ("utf-8-sig", "cp1252"):
        try:
            with path.open("r", encoding=encoding, newline="") as handle:
                return [
                    {str(key or "").strip(): _repair_text(value) for key, value in row.items()}
                    for row in csv.DictReader(handle)
                ]
        except (OSError, UnicodeError, csv.Error) as exc:
            last_error = exc
    raise RuntimeError(f"Nao foi possivel ler {path}: {last_error}")


def _logical_source(client_id: str, filename: str) -> str:
    return f"info/{client_id}/{filename}".replace("\\", "/")


def _tokens(value: Any) -> set[str]:
    return {
        token for token in re.findall(r"[a-z0-9]{2,}", _fold(value))
        if token not in STOPWORDS and not YEAR_RE.fullmatch(token)
    }


def _relevant_to_product(product_name: str, value: str) -> bool:
    expected = _tokens(product_name)
    if not expected:
        return True
    found = _tokens(value)
    common = expected & found
    return len(common) >= min(2, max(1, len(expected) // 4))


def _split_fragments(text: str, source: str) -> list[dict[str, Any]]:
    text = _repair_text(text)
    if not text:
        return []
    expanded = re.sub(
        r"(?i)\s+(?=(?:aplicacao|aplica-se|caracteristicas?|codigo original|conteudo da embalagem|"
        r"especificacoes?|instalacao|material|medidas?|modo de uso|numero da peca|oem|produto)\s*[:?-])",
        "\n",
        text,
    )
    expanded = re.sub(r"\s*(?:[-–—]{3,}|[•●▪])\s*", "\n", expanded)
    expanded = re.sub(r"(?<=[.!?])\s+(?=[A-ZÁÉÍÓÚÃÕÇ0-9])", "\n", expanded)
    fragments: list[dict[str, Any]] = []
    for raw in expanded.split("\n"):
        value = SPACE_RE.sub(" ", raw.strip(" -*–—\t"))
        if len(value) < 4:
            continue
        if re.fullmatch(r"[^:]{1,80}:\s*", value):
            continue
        folded = _fold(value)
        if any(term in folded for term in MARKETING_TERMS):
            continue
        for start in range(0, len(value), 600):
            chunk = value[start : start + 600].strip()
            if chunk:
                fragments.append({"texto": chunk, "fontes": [source]})
    return fragments


def _merge_evidence(items: Iterable[dict[str, Any]], limit: int = 20) -> list[dict[str, Any]]:
    merged: dict[str, dict[str, Any]] = {}
    order: list[str] = []
    for raw in items:
        text = _repair_text(raw.get("texto"))
        if not text or _is_missing_value(text):
            continue
        key = _fold(text)
        if key not in merged:
            merged[key] = {"texto": text, "fontes": []}
            order.append(key)
        for source in raw.get("fontes") or []:
            source = _repair_text(source)
            if source and source not in merged[key]["fontes"]:
                merged[key]["fontes"].append(source)
        if len(order) >= limit:
            break
    return [merged[key] for key in order]


def _select_fragments(
    fragments: Iterable[dict[str, Any]], terms: Iterable[str], *, limit: int
) -> list[dict[str, Any]]:
    folded_terms = tuple(_fold(term) for term in terms)

    def matches(text: str, term: str) -> bool:
        clean = term.strip()
        if clean in {"automatic", "caracteristica", "compativel", "encaix", "especifica", "instal"}:
            return clean in text
        return bool(re.search(rf"\b{re.escape(clean)}\b", text))

    return _merge_evidence(
        (
            item for item in fragments
            if any(matches(_fold(item.get("texto")), term) for term in folded_terms)
        ),
        limit=limit,
    )


def _measure_evidence(fragments: Iterable[dict[str, Any]]) -> list[dict[str, Any]]:
    items = []
    for fragment in fragments:
        text = str(fragment.get("texto") or "")
        if MEASURE_RE.search(_fold(text)):
            items.append(fragment)
    return _merge_evidence(items, limit=14)


def _valid_oem_code(code: str, sku: str, ncm: str) -> bool:
    compact = code.strip(" .,:;()[]{}").upper()
    code_key = re.sub(r"[^A-Z0-9]", "", compact)
    if code_key in (OEM_CODE_EXCLUSIONS.get(_norm_sku(sku)) or set()):
        return False
    if len(compact) < 5 or len(compact) > 30 or not any(char.isdigit() for char in compact):
        return False
    if compact in {_norm_sku(sku), re.sub(r"\D", "", ncm or "")}:
        return False
    digits = re.sub(r"\D", "", compact)
    if len(digits) == 4 and digits.isdigit() and 1900 <= int(digits) <= 2099:
        return False
    if re.fullmatch(r"(?:19|20)\d{2}[-/](?:19|20)\d{2}", compact):
        return False
    if _looks_like_gtin(compact):
        return False
    if re.fullmatch(r"\d+(?:PCS|HP|CV|RPM)", compact) or re.fullmatch(r"ISO\d+", compact):
        return False
    if re.fullmatch(r"\d+MM(?:X\d+MM)+", compact):
        return False
    if re.fullmatch(r"\d+(?:[.,]\d+)?(?:MM|CM|M|KG|G|V|W|A)", compact):
        return False
    return any(char.isalpha() for char in compact) or len(digits) >= 7


def _looks_like_gtin(value: str) -> bool:
    digits = re.sub(r"\D", "", value)
    if value != digits or len(digits) not in {12, 13, 14}:
        return False
    # OEMs puramente numericos costumam ser bem menores; no cadastro analisado,
    # sequencias de 12 a 14 digitos vieram de EAN/GTIN ou placeholders de anuncio,
    # inclusive quando o vendedor as colocou por engano no atributo "Codigo OEM".
    return True


def _extract_oem(fragments: Iterable[dict[str, Any]], sku: str, ncm: str) -> list[dict[str, Any]]:
    found: dict[str, dict[str, Any]] = {}
    order: list[str] = []
    for fragment in fragments:
        sources = [str(source) for source in (fragment.get("fontes") or []) if source]
        # Bling, cadastro legado e planilha do fornecedor podem conter codigo
        # comercial, referencia de conjunto ou ate outro produto. Continuam
        # uteis como pistas, mas nunca iniciam um OEM: o codigo precisa vir do
        # cadastro atual, Mercado Livre ou pesquisa publica confirmatoria.
        def source_is_disallowed(source: str) -> bool:
            folded = _fold(source)
            return (
                "bling" in folded
                or folded.startswith("cadastro_legado:")
                or folded.startswith("fornecedor_planilha:")
            )

        if sources and all(source_is_disallowed(source) for source in sources):
            continue
        text = str(fragment.get("texto") or "")
        candidate_areas = [match.group(1) for match in OEM_LABEL_RE.finditer(text)]
        for area in candidate_areas:
            for match in CODE_RE.findall(area.upper()):
                code = match.strip(" .,:;()[]{}")
                if not _valid_oem_code(code, sku, ncm):
                    continue
                key = re.sub(r"[^A-Z0-9]", "", code)
                if key not in found:
                    found[key] = {"codigo": code, "fontes": []}
                    order.append(key)
                for source in fragment.get("fontes") or []:
                    if source not in found[key]["fontes"]:
                        found[key]["fontes"].append(source)
    return [found[key] for key in order[:20]]


def _is_automotive(name: str, description: str) -> bool:
    folded_name = _fold(name)
    if any(term in folded_name for term in NON_AUTOMOTIVE_TERMS):
        return False
    folded = _fold(f"{name} {description[:2500]}")
    brand_hit = any(re.search(rf"\b{re.escape(brand)}\b", folded) for brand in VEHICLE_BRANDS)
    auto_terms = (
        "automot", "carro", "veiculo", "parachoque", "radiador", "porta", "cambio", "motor",
        "retrovisor", "combustivel", "direcao", "intercooler", "freio", "suspensao", "arrefecimento",
        "refrigeracao",
    )
    auto_hit = any(
        re.search(rf"\b{re.escape(term)}\w*\b", folded) if term == "automot"
        else re.search(rf"\b{re.escape(term)}\b", folded)
        for term in auto_terms
    )
    return brand_hit or bool(auto_hit)


def _compatibility_evidence(
    fragments: Iterable[dict[str, Any]], *, automotive: bool
) -> tuple[list[dict[str, Any]], list[dict[str, Any]]]:
    vehicle: list[dict[str, Any]] = []
    equipment: list[dict[str, Any]] = []
    for fragment in fragments:
        folded = _fold(fragment.get("texto"))
        has_cue = any(term in folded for term in COMPATIBILITY_TERMS if term not in {"anos", "ano "})
        has_vehicle = any(re.search(rf"\b{re.escape(brand)}\b", folded) for brand in VEHICLE_BRANDS)
        has_year = bool(YEAR_RE.search(folded))
        has_equipment = any(term in folded for term in NON_AUTOMOTIVE_TERMS)
        if automotive and (has_vehicle or (has_cue and has_year)):
            vehicle.append(fragment)
        elif not automotive and (has_equipment or has_cue):
            equipment.append(fragment)
    return _merge_evidence(vehicle, limit=18), _merge_evidence(equipment, limit=18)


def _section(items: list[dict[str, Any]], _missing: str = "") -> dict[str, Any]:
    """Mantem o schema e deixa campos sem evidencia realmente vazios."""
    items = [
        item for item in items
        if not _is_missing_value(item.get("texto") if isinstance(item, dict) else item)
    ]
    return {
        "status": "documentado" if items else "",
        "itens": items,
        "observacao": "",
    }


def _product_class(value: str) -> str:
    folded = _fold(value)
    for class_name, terms in PRODUCT_CLASSES:
        if any(term in folded for term in terms):
            return class_name
    return ""


def _identity_check(name: str, description: str, live_titles: list[str]) -> tuple[dict[str, Any], list[str]]:
    name_tokens = _tokens(name)
    desc_tokens = _tokens(description[:2500])
    overlap = (len(name_tokens & desc_tokens) / max(1, len(name_tokens))) if name_tokens else 0.0
    name_class = _product_class(name)
    description_class = _product_class(description[:1200])
    alerts: list[str] = []
    if name_class and description_class and name_class != description_class:
        alerts.append(
            f"Possivel conflito de identidade: o nome indica {name_class} e a descricao indica {description_class}."
        )
    if live_titles and not any(_relevant_to_product(name, title) for title in live_titles):
        alerts.append("Os titulos dos anuncios ativos nao confirmaram claramente o nome do cadastro.")
    status = "conflito" if alerts else ("baixa_correspondencia" if overlap < 0.08 else "coerente")
    return {
        "status": status,
        "correspondencia_termos": round(overlap, 3),
        "classe_nome": name_class,
        "classe_descricao": description_class,
    }, alerts


def _low_quality_product_name(value: Any) -> bool:
    text = _repair_text(value)
    folded = _fold(text)
    if not text or len(text) > 180 or len(_tokens(text)) < 2:
        return True
    if any(_fold(term) in folded for term in MARKETING_TERMS):
        return True
    if re.match(r"^(?:para\b|por isso\b|nao espere\b|localizadas?\b)", folded):
        return True
    if re.match(r"^(?:19|20)\d{2}\s*[-/]", folded):
        return True
    if re.match(r"^\d+(?:[.,]\d+)?\s*(?:kg|kilos?)\b", folded):
        return True
    return False


def _select_product_name(
    row: dict[str, str],
    live_ads: list[dict[str, Any]],
    snapshot_ads: list[dict[str, Any]],
) -> tuple[str, str]:
    candidates: list[tuple[Any, str]] = [
        (row.get("nome"), "cadastro:nome"),
        (row.get("produto"), "cadastro:produto"),
        (row.get("produto_bling"), "cadastro:produto_bling"),
    ]
    candidates.extend(
        (item.get("titulo"), f"anuncio:{item.get('item_id') or index}")
        for index, item in enumerate(live_ads, start=1)
    )
    candidates.extend(
        (item.get("titulo"), f"anuncio_snapshot:{item.get('item_id') or index}")
        for index, item in enumerate(snapshot_ads, start=1)
    )
    for raw, source in candidates:
        value = _repair_text(raw)
        if value and not _low_quality_product_name(value):
            return value, source

    description = _repair_text(row.get("descricao"))
    for fragment in _split_fragments(description, "cadastro:descricao"):
        value = _repair_text(fragment.get("texto"))
        if value and not _low_quality_product_name(value):
            return value[:180], "cadastro:descricao"
    return f"Produto SKU {_repair_text(row.get('sku'))}", "cadastro:sku"


def _load_mlb_links(path: Path, catalog_skus: set[str]) -> dict[str, list[dict[str, Any]]]:
    result: dict[str, list[dict[str, Any]]] = {sku: [] for sku in catalog_skus}
    if not path.exists():
        return result
    for row in _read_csv(path):
        sku = _norm_sku(row.get("sku"))
        if sku not in result:
            continue
        ids = []
        for value in (row.get("mlb_principal"), *(row.get("mlb_ids") or "").split("|")):
            item_id = re.sub(r"[^A-Z0-9]", "", str(value or "").upper())
            if item_id.startswith("MLB") and item_id not in ids:
                ids.append(item_id)
        for item_id in ids[:20]:
            result[sku].append({
                "item_id": item_id,
                "url": f"https://produto.mercadolivre.com.br/{item_id}",
                "origem": "vinculo_local",
                "status": "nao_verificado_nesta_execucao",
            })
    return result


def _load_legacy_catalog(source_dir: Path, catalog_skus: set[str]) -> dict[str, dict[str, str]]:
    """Carrega o cadastro rico antigo apenas como pista rastreavel."""
    output: dict[str, dict[str, str]] = {}
    candidates = (
        source_dir / "legacy" / "cadastro_produtos.csv",
        source_dir / "cadastro_produtos_legacy.csv",
    )
    path = next((candidate for candidate in candidates if candidate.exists()), None)
    if not path:
        return output
    for row in _read_csv(path):
        sku = _norm_sku(row.get("sku"))
        if sku in catalog_skus and sku not in output:
            output[sku] = row
    return output


def _header_key(value: Any) -> str:
    return re.sub(r"[^a-z0-9]+", "", _fold(value))


def _load_supplier_spreadsheet(path: Path, catalog_skus: set[str]) -> dict[str, list[dict[str, str]]]:
    """Le a planilha de fornecedor/AliExpress sem executar formulas nem links."""
    output: dict[str, list[dict[str, str]]] = {sku: [] for sku in catalog_skus}
    if not path.exists():
        return output
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        worksheet = workbook["Tudo"] if "Tudo" in workbook.sheetnames else workbook[workbook.sheetnames[0]]
        rows = worksheet.iter_rows(values_only=True)
        headers = [_header_key(value) for value in next(rows, ())]
        aliases = {
            "sku": {"sku"},
            "descricao": {"description", "descricao"},
            "oem_modelo": {"oemmodel", "oemmodelo", "oem"},
            "cor_lado": {"colorside", "corlado"},
            "link": {"linkaliexpress", "aliexpress", "link"},
        }
        indexes: dict[str, int] = {}
        for logical, choices in aliases.items():
            for index, header in enumerate(headers):
                if header in choices:
                    indexes[logical] = index
                    break
        if "sku" not in indexes:
            return output
        for values in rows:
            values = tuple(values or ())
            sku_raw = values[indexes["sku"]] if indexes["sku"] < len(values) else ""
            sku = _resolve_catalog_sku(sku_raw, catalog_skus)
            if sku not in output:
                continue
            item = {}
            for logical in ("descricao", "oem_modelo", "cor_lado", "link"):
                index = indexes.get(logical)
                if index is not None and index < len(values):
                    item[logical] = _repair_text(values[index])
            if any(item.values()) and item not in output[sku]:
                output[sku].append(item)
    finally:
        workbook.close()
    return output


def _load_favorite_searches(path: Path, catalog_skus: set[str]) -> dict[str, dict[str, str]]:
    output: dict[str, dict[str, str]] = {}
    payload = _read_json(path, {})
    searches = payload.get("pesquisas") if isinstance(payload, dict) else {}
    if not isinstance(searches, dict):
        return output
    for raw in searches.values():
        if not isinstance(raw, dict):
            continue
        sku = _resolve_catalog_sku(raw.get("sku"), catalog_skus)
        if sku not in catalog_skus:
            continue
        item = {
            key: _repair_text(raw.get(key))
            for key in ("loja", "sku", "produto", "pesquisa_1", "pesquisa_2", "pesquisa_3", "updated_at")
        }
        if sku not in output or item.get("updated_at", "") > output[sku].get("updated_at", ""):
            output[sku] = item
    return output


def _load_listing_export(path: Path, catalog_skus: set[str]) -> dict[str, list[dict[str, str]]]:
    output: dict[str, list[dict[str, str]]] = {sku: [] for sku in catalog_skus}
    if not path.exists():
        return output
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        worksheet = next(
            (sheet for sheet in workbook.worksheets if _header_key(sheet.title) == "anuncios"),
            workbook.worksheets[-1],
        )
        rows = worksheet.iter_rows(values_only=True)
        headers = [_header_key(value) for value in next(rows, ())]
        by_name = {header: index for index, header in enumerate(headers) if header}

        def value(values: tuple[Any, ...], name: str) -> str:
            index = by_name.get(_header_key(name))
            return _repair_text(values[index]) if index is not None and index < len(values) else ""

        for raw_values in rows:
            values = tuple(raw_values or ())
            sku = _resolve_catalog_sku(value(values, "SKU"), catalog_skus)
            if sku not in output:
                continue
            item_id = re.sub(r"[^A-Z0-9]", "", value(values, "ITEM_ID").upper())
            item = {
                "item_id": item_id,
                "titulo": value(values, "TITLE"),
                "variacoes": value(values, "VARIATIONS"),
                "categoria": value(values, "CATEGORY"),
                "tipo_anuncio": value(values, "LISTING_TYPE"),
                "url": f"https://produto.mercadolivre.com.br/{item_id}" if item_id.startswith("MLB") else "",
                "origem": "export_anuncios_2026-04-09",
                "status": "snapshot_sem_verificacao_atual",
            }
            identity = (item["item_id"], item["titulo"], item["variacoes"])
            if any((old["item_id"], old["titulo"], old["variacoes"]) == identity for old in output[sku]):
                continue
            output[sku].append(item)
    finally:
        workbook.close()
    return output


def _ml_attribute_value(attributes: Any, wanted_id: str) -> str:
    for attribute in attributes if isinstance(attributes, list) else []:
        if str(attribute.get("id") or "").upper() == wanted_id:
            return _repair_text(attribute.get("value_name") or attribute.get("value_id"))
    return ""


def _ml_item_skus(item: dict[str, Any]) -> dict[str, dict[str, Any] | None]:
    matches: dict[str, dict[str, Any] | None] = {}
    parent = _norm_sku(
        item.get("seller_sku")
        or item.get("seller_custom_field")
        or _ml_attribute_value(item.get("attributes"), "SELLER_SKU")
    )
    if parent:
        matches[parent] = None
    for variation in item.get("variations") if isinstance(item.get("variations"), list) else []:
        sku = _norm_sku(
            variation.get("seller_sku")
            or variation.get("seller_custom_field")
            or _ml_attribute_value(variation.get("attributes"), "SELLER_SKU")
        )
        if sku:
            matches[sku] = variation
    return matches


def _ml_attributes(attributes: Any) -> list[dict[str, str]]:
    output = []
    for attribute in attributes if isinstance(attributes, list) else []:
        attr_id = _repair_text(attribute.get("id"))
        if attr_id.upper() == "SELLER_SKU":
            continue
        name = _repair_text(attribute.get("name") or attr_id)
        value = _repair_text(attribute.get("value_name") or attribute.get("value_id"))
        if name and value:
            output.append({"nome": name, "valor": value})
    return output[:80]


def _merged_ml_attributes(item: dict[str, Any], variation: dict[str, Any] | None) -> list[dict[str, str]]:
    merged: dict[str, dict[str, Any]] = {}
    anonymous = 0
    for attribute in list(item.get("attributes") or []) + list((variation or {}).get("attributes") or []):
        if not isinstance(attribute, dict):
            continue
        key = _repair_text(attribute.get("id")).upper()
        if not key:
            anonymous += 1
            key = f"__ANON_{anonymous}"
        merged[key] = attribute
    return _ml_attributes(list(merged.values()))


def _request_json(
    session: requests.Session,
    url: str,
    *,
    headers: dict[str, str] | None = None,
    params: dict[str, Any] | None = None,
    timeout: float = 25,
    verify_ssl: bool = True,
) -> tuple[int, Any]:
    response = session.get(url, headers=headers, params=params, timeout=timeout, verify=verify_ssl)
    try:
        payload = response.json()
    except ValueError:
        payload = None
    return response.status_code, payload


def _fetch_live_ml(
    tenant_dir: Path,
    catalog_skus: set[str],
    *,
    include_descriptions: bool,
    verify_ssl: bool,
) -> tuple[dict[str, list[dict[str, Any]]], dict[str, Any]]:
    output: dict[str, list[dict[str, Any]]] = {sku: [] for sku in catalog_skus}
    coverage: dict[str, Any] = {"consultado_em": _now_iso(), "lojas": []}
    config_path = tenant_dir / "lojas_config.json"
    stores = _read_json(config_path, [])
    if not isinstance(stores, list):
        coverage["erro"] = "lojas_config.json invalido"
        return output, coverage

    private_description_targets: dict[str, tuple[str, str, dict[str, Any]]] = {}
    for store in stores:
        store_name = _repair_text(store.get("nome") or store.get("name"))
        integration = ((store.get("integracoes") or {}).get("mercadolivre") or {})
        token = str(integration.get("access_token") or "").strip()
        user_id = str(integration.get("user_id") or "").strip()
        status_row: dict[str, Any] = {"loja": store_name, "status": "nao_conectada", "itens_ativos": 0}
        if not integration.get("connected") or not token or not user_id:
            coverage["lojas"].append(status_row)
            continue
        session = requests.Session()
        headers = {"Authorization": f"Bearer {token}", "User-Agent": USER_AGENT}
        ids: list[str] = []
        offset = 0
        total = None
        try:
            for _ in range(30):
                code, payload = _request_json(
                    session,
                    f"{ML_API_BASE}/users/{user_id}/items/search",
                    headers=headers,
                    params={"status": "active", "offset": offset, "limit": 100},
                    verify_ssl=verify_ssl,
                )
                if code not in (200, 206) or not isinstance(payload, dict):
                    raise RuntimeError(f"search HTTP {code}")
                page = payload.get("results") or []
                total = int((payload.get("paging") or {}).get("total") or len(page))
                for raw_id in page:
                    item_id = str(raw_id.get("id") if isinstance(raw_id, dict) else raw_id or "").upper().strip()
                    if item_id and item_id not in ids:
                        ids.append(item_id)
                offset += len(page)
                if not page or offset >= total:
                    break

            items: list[dict[str, Any]] = []
            failed_detail_batches = 0
            for start in range(0, len(ids), 20):
                code, payload = _request_json(
                    session,
                    f"{ML_API_BASE}/items",
                    headers=headers,
                    params={"ids": ",".join(ids[start : start + 20]), "include_attributes": "all"},
                    verify_ssl=verify_ssl,
                )
                if code != 200 or not isinstance(payload, list):
                    failed_detail_batches += 1
                    continue
                for entry in payload:
                    body = entry.get("body") if isinstance(entry, dict) and isinstance(entry.get("body"), dict) else entry
                    if isinstance(body, dict):
                        items.append(body)

            matched_items = 0
            for item in items:
                for sku, variation in _ml_item_skus(item).items():
                    if sku not in output:
                        continue
                    summary = {
                        "loja": store_name,
                        "item_id": _repair_text(item.get("id")),
                        "titulo": _repair_text(item.get("title")),
                        "url": _repair_text(item.get("permalink")) or f"https://produto.mercadolivre.com.br/{item.get('id', '')}",
                        "status": _repair_text(item.get("status")),
                        "condicao": _repair_text(item.get("condition")),
                        "ultima_atualizacao": _repair_text(item.get("last_updated")),
                        "atributos": _merged_ml_attributes(item, variation),
                        "variacao": _ml_attributes((variation or {}).get("attribute_combinations")) if variation else [],
                        "origem": "mercado_livre_api",
                    }
                    output[sku].append(summary)
                    matched_items += 1
                    if include_descriptions and sku not in private_description_targets and summary["item_id"]:
                        private_description_targets[sku] = (token, summary["item_id"], summary)
            status_row.update({
                "status": "consultada",
                "itens_ativos": len(ids),
                "itens_detalhados": len(items),
                "vinculos_sku": matched_items,
                "lotes_detalhes_com_falha": failed_detail_batches,
                "cobertura_completa": (
                    total is not None and offset >= total and failed_detail_batches == 0 and len(items) == len(ids)
                ),
            })
        except Exception as exc:  # rede/provedor: manter cobertura explicita
            status_row.update({"status": "falha", "erro": str(exc)[:180], "cobertura_completa": False})
        coverage["lojas"].append(status_row)
        print(f"[ML] {store_name or 'loja'}: {status_row['status']} ({status_row.get('itens_ativos', 0)} ativos)", flush=True)

    if include_descriptions and private_description_targets:
        def fetch_description(target: tuple[str, str, dict[str, Any]]) -> bool:
            token, item_id, summary = target
            session = requests.Session()
            code, payload = _request_json(
                session,
                f"{ML_API_BASE}/items/{item_id}/description",
                headers={"Authorization": f"Bearer {token}", "User-Agent": USER_AGENT},
                verify_ssl=verify_ssl,
            )
            if code == 200 and isinstance(payload, dict):
                summary["descricao"] = _repair_text(payload.get("plain_text") or payload.get("text"))
                return True
            return False

        with ThreadPoolExecutor(max_workers=6) as executor:
            futures = [executor.submit(fetch_description, target) for target in private_description_targets.values()]
            done = 0
            succeeded = 0
            for future in as_completed(futures):
                done += 1
                try:
                    succeeded += bool(future.result())
                except Exception:
                    pass
                if done % 50 == 0 or done == len(futures):
                    print(f"[ML] descricoes consultadas: {done}/{len(futures)}", flush=True)
        coverage["descricoes"] = {
            "solicitadas": len(private_description_targets),
            "obtidas": succeeded,
            "falhas": len(private_description_targets) - succeeded,
        }

    coverage["cobertura_completa"] = bool(coverage["lojas"]) and all(
        row.get("status") in {"consultada", "nao_conectada"} and row.get("cobertura_completa", True)
        for row in coverage["lojas"]
    )
    return output, coverage


def _pace_search() -> None:
    global _SEARCH_NEXT_AT
    with _SEARCH_LOCK:
        now = time.monotonic()
        wait = max(0.0, _SEARCH_NEXT_AT - now)
        _SEARCH_NEXT_AT = max(now, _SEARCH_NEXT_AT) + random.uniform(0.32, 0.58)
    if wait:
        time.sleep(wait)


def _decode_ddg_url(href: str) -> str:
    absolute = urljoin("https://duckduckgo.com", href)
    parsed = urlparse(absolute)
    query = parse_qs(parsed.query)
    if query.get("uddg"):
        return unquote(query["uddg"][0])
    return absolute


def _search_web(query: str, *, max_results: int, verify_ssl: bool) -> dict[str, Any]:
    query = SPACE_RE.sub(" ", query).strip()[:240]
    last_error = ""
    for attempt in range(3):
        try:
            _pace_search()
            response = requests.get(
                "https://html.duckduckgo.com/html/",
                params={"q": query},
                headers={"User-Agent": USER_AGENT},
                timeout=22,
                verify=verify_ssl,
            )
            if response.status_code in (202, 429, 503):
                raise RuntimeError(f"HTTP {response.status_code}")
            response.raise_for_status()
            soup = BeautifulSoup(response.text or "", "lxml")
            results: list[dict[str, str]] = []
            for node in soup.select(".result"):
                link = node.select_one(".result__a")
                snippet = node.select_one(".result__snippet")
                if not link:
                    continue
                url = _decode_ddg_url(str(link.get("href") or ""))
                if not url.startswith(("http://", "https://")):
                    continue
                result = {
                    "titulo": _repair_text(link.get_text(" ", strip=True)),
                    "url": url,
                    "trecho": _repair_text(snippet.get_text(" ", strip=True) if snippet else ""),
                }
                if result not in results:
                    results.append(result)
                if len(results) >= max_results:
                    break
            return {
                "status": "pesquisado_com_resultado" if results else "pesquisado_sem_resultado",
                "provedor": "duckduckgo_html",
                "consulta": query,
                "consultado_em": _now_iso(),
                "resultados": results,
            }
        except Exception as exc:
            last_error = str(exc)[:180]
            time.sleep(1.5 * (attempt + 1))
    return {
        "status": "indisponivel",
        "provedor": "duckduckgo_html",
        "consulta": query,
        "consultado_em": _now_iso(),
        "resultados": [],
        "erro": last_error,
    }


def _internet_query(row: dict[str, str], hint: dict[str, str] | None = None) -> str:
    name = row.get("nome") or row.get("produto") or row.get("produto_bling") or row.get("descricao", "")[:140]
    name = _repair_text(name)
    hint = hint or {}
    extra = _repair_text(hint.get("pesquisa_2") or hint.get("pesquisa_1"))
    if extra and _fold(extra) not in _fold(name):
        name = f"{name} {extra}"
    return f"{name} especificacoes aplicacao OEM"[:240]


def _search_all_skus(
    rows: list[dict[str, str]],
    existing: dict[str, dict[str, Any]],
    search_hints: dict[str, dict[str, str]],
    *,
    refresh: bool,
    workers: int,
    max_results: int,
    verify_ssl: bool,
) -> dict[str, dict[str, Any]]:
    output: dict[str, dict[str, Any]] = {}
    pending: dict[Any, str] = {}
    with ThreadPoolExecutor(max_workers=max(1, workers)) as executor:
        for row in rows:
            sku = _norm_sku(row.get("sku"))
            cached = (existing.get(sku) or {}).get("pesquisa_internet")
            if (
                not refresh
                and isinstance(cached, dict)
                and cached.get("consultado_em")
                and cached.get("status") in {"pesquisado_com_resultado", "pesquisado_sem_resultado"}
            ):
                output[sku] = cached
                continue
            future = executor.submit(
                _search_web,
                _internet_query(row, search_hints.get(sku)),
                max_results=max_results,
                verify_ssl=verify_ssl,
            )
            pending[future] = sku
        done = 0
        for future in as_completed(pending):
            sku = pending[future]
            try:
                output[sku] = future.result()
            except Exception as exc:
                output[sku] = {
                    "status": "indisponivel", "resultados": [], "consultado_em": _now_iso(), "erro": str(exc)[:180]
                }
            done += 1
            if done % 10 == 0 or done == len(pending):
                print(f"[WEB] pesquisas concluidas: {done}/{len(pending)}", flush=True)
    return output


def _internet_fragments(product_name: str, research: dict[str, Any]) -> list[dict[str, Any]]:
    fragments: list[dict[str, Any]] = []
    for index, result in enumerate(research.get("resultados") or [], start=1):
        value = _repair_text(f"{result.get('titulo', '')}. {result.get('trecho', '')}")
        if not value or not _relevant_to_product(product_name, value):
            continue
        fragments.extend(_split_fragments(value, f"internet:{index}"))
    return fragments


def _build_dossier(
    row: dict[str, str],
    *,
    client_id: str,
    live_ads: list[dict[str, Any]],
    snapshot_ads: list[dict[str, str]],
    cached_ads: list[dict[str, Any]],
    ml_coverage: dict[str, Any],
    internet_research: dict[str, Any],
    legacy: dict[str, str],
    supplier_records: list[dict[str, str]],
    favorite_search: dict[str, str],
    existing: dict[str, Any],
) -> dict[str, Any]:
    sku = _repair_text(row.get("sku"))
    description = _repair_text(row.get("descricao"))
    product_name, product_name_source = _select_product_name(row, live_ads, snapshot_ads)

    fragments = _split_fragments(product_name, product_name_source or "cadastro:nome")
    fragments.extend(_split_fragments(description, "cadastro:descricao"))
    for index, ad in enumerate(live_ads, start=1):
        source = f"anuncio:{ad.get('item_id') or index}"
        fragments.extend(_split_fragments(ad.get("titulo", ""), source))
        fragments.extend(_split_fragments(ad.get("descricao", ""), source))
        for attribute in ad.get("atributos") or []:
            fragments.append({
                "texto": f"{_repair_text(attribute.get('nome'))}: {_repair_text(attribute.get('valor'))}",
                "fontes": [source],
            })
        for attribute in ad.get("variacao") or []:
            fragments.append({
                "texto": f"{_repair_text(attribute.get('nome'))}: {_repair_text(attribute.get('valor'))}",
                "fontes": [source],
            })

    for index, ad in enumerate(snapshot_ads, start=1):
        source = f"anuncio_snapshot:{ad.get('item_id') or index}"
        fragments.extend(_split_fragments(ad.get("titulo", ""), source))
        fragments.extend(_split_fragments(ad.get("variacoes", ""), source))

    legacy_name = _repair_text(
        legacy.get("cg_product name") or legacy.get("nome") or legacy.get("produto") or legacy.get("produto_bling")
    )
    legacy_fields = (
        ("cg_product description", "Descricao legada"),
        ("cg_material ptbr", "Material"),
        ("cg_medidas individuas", "Medidas"),
        ("cg_individual size", "Dimensoes"),
        ("cg_peso individual", "Peso"),
        ("cg_para quer serve", "Para que serve"),
        ("cg_onde sera ultilizado", "Onde e utilizado"),
        ("cg_finalizade", "Finalidade"),
        ("cg_oem", "OEM"),
        ("titulos_anuncios_mlb", "Titulos de anuncios"),
    )
    legacy_summary = " ".join(_repair_text(legacy.get(field)) for field, _label in legacy_fields)
    legacy_relevant = bool(legacy) and (
        _relevant_to_product(product_name, f"{legacy_name} {legacy_summary}")
        or _norm_sku(legacy.get("sku")) == _norm_sku(sku) and not legacy_name
    )
    if legacy_relevant:
        for field, label in legacy_fields:
            value = _repair_text(legacy.get(field))
            if value:
                fragments.extend(_split_fragments(f"{label}: {value}", "cadastro_legado:pista"))

    for index, supplier in enumerate(supplier_records, start=1):
        source = f"fornecedor_planilha:{index}"
        for field, label in (
            ("descricao", "Descricao do fornecedor"),
            ("oem_modelo", "Aplicacao indicada pelo fornecedor"),
            ("cor_lado", "Cor/Lado"),
        ):
            value = _repair_text(supplier.get(field))
            if value:
                fragments.extend(_split_fragments(f"{label}: {value}", source))
    fragments.extend(_internet_fragments(product_name, internet_research))

    automotive = _is_automotive(product_name, f"{description} {row.get('categoria', '')}")
    vehicle_compatibility, equipment_compatibility = _compatibility_evidence(fragments, automotive=automotive)
    purpose = _select_fragments(fragments, PURPOSE_TERMS, limit=10)
    functioning = _select_fragments(fragments, FUNCTION_TERMS, limit=8)
    installation = _select_fragments(fragments, INSTALL_TERMS, limit=10)
    technical = _select_fragments(fragments, TECH_TERMS, limit=18)
    measures = _measure_evidence(fragments)
    oem_codes = _extract_oem(fragments, sku, row.get("ncm", ""))
    live_titles = [_repair_text(item.get("titulo")) for item in live_ads if item.get("titulo")]
    identity, alerts = _identity_check(product_name, description, live_titles)

    missing_sections = []
    for key, items in (
        ("para_que_serve", purpose),
        ("caracteristicas_tecnicas", technical),
        ("medidas", measures),
        ("modo_de_funcionamento", functioning),
        ("instalacao", installation),
        ("oem", oem_codes),
        ("veiculos_compativeis", vehicle_compatibility if automotive else equipment_compatibility),
    ):
        if not items:
            missing_sections.append(key)

    sources = [
        {
            "id": "cadastro",
            "tipo": "cadastro_local",
            "referencia": _logical_source(client_id, "cadastro_produtos.csv"),
            "consultado_em": _now_iso(),
        }
    ]
    if live_ads:
        sources.append({
            "id": "anuncios_ml",
            "tipo": "mercado_livre_api",
            "referencias": [item.get("url") for item in live_ads if item.get("url")],
            "consultado_em": ml_coverage.get("consultado_em"),
        })
    if cached_ads:
        sources.append({
            "id": "vinculos_mlb",
            "tipo": "vinculo_local_de_anuncio",
            "referencia": _logical_source(client_id, "mlb_sku_vinculos.csv"),
        })
    if snapshot_ads:
        sources.append({
            "id": "export_anuncios",
            "tipo": "snapshot_mercado_livre",
            "referencia": "JK/Anuncios-2026_04_09-15_24.xlsx#Anuncios",
            "data_snapshot": "2026-04-09",
            "referencias": [item.get("url") for item in snapshot_ads if item.get("url")],
        })
    if internet_research.get("consultado_em"):
        sources.append({
            "id": "internet",
            "tipo": "pesquisa_publica",
            "provedor": internet_research.get("provedor"),
            "consulta": internet_research.get("consulta"),
            "consultado_em": internet_research.get("consultado_em"),
        })
    if legacy:
        sources.append({
            "id": "cadastro_legado",
            "tipo": "pista_nao_autoritativa",
            "referencia": "info/default/cadastro_produtos.csv",
            "usado_nos_campos": legacy_relevant,
        })
    supplier_links = [item.get("link") for item in supplier_records if item.get("link")]
    if supplier_records:
        sources.append({
            "id": "planilha_fornecedor",
            "tipo": "planilha_fornecedor_aliexpress",
            "referencia": "Listas Diversos SKU.xlsx#Tudo",
            "links": supplier_links,
        })
    if favorite_search:
        sources.append({
            "id": "favoritos_pesquisas",
            "tipo": "campos_de_pesquisa_internos",
            "referencia": _logical_source(client_id, "favoritos_pesquisas_caio.json"),
            "atualizado_em": favorite_search.get("updated_at"),
        })

    return {
        "schema_version": SCHEMA_VERSION,
        "sku": sku,
        "client_id": client_id,
        "produto": {
            "nome": product_name,
            "o_que_e": {
                "status": "documentado",
                "texto": product_name,
                "fontes": [product_name_source],
            },
            "para_que_serve": _section(purpose, "Finalidade nao descrita de forma explicita nas fontes consultadas."),
            "caracteristicas_tecnicas": _section(technical, "Caracteristicas tecnicas nao localizadas nas fontes consultadas."),
            "medidas": _section(measures, "Medidas nao localizadas nas fontes consultadas."),
            "modo_de_funcionamento": _section(functioning, "Modo de funcionamento nao descrito nas fontes consultadas."),
            "instalacao": _section(
                installation,
                "Instalacao nao documentada. Confirmar codigo, encaixe e aplicacao antes da montagem; em veiculos, preferir profissional qualificado.",
            ),
            "oem": {
                "status": "documentado" if oem_codes else "",
                "codigos": oem_codes,
                "observacao": "",
            },
            "tipo_aplicacao": "automotiva" if automotive else "nao_automotiva_ou_nao_confirmada",
            "veiculos_compativeis": _section(
                vehicle_compatibility,
                "Nenhum veiculo compativel confirmado nas fontes consultadas."
                if automotive else "Campo nao aplicavel ou natureza automotiva nao confirmada.",
            ),
            "equipamentos_ou_aplicacoes_compativeis": _section(
                equipment_compatibility,
                "Nenhuma aplicacao nao veicular confirmada nas fontes consultadas.",
            ),
        },
        "cadastro": {
            "descricao_original": description,
            "marca": _repair_text(row.get("marca")),
            "categoria": _repair_text(row.get("categoria")),
            "ncm": _repair_text(row.get("ncm")),
            "cest": _repair_text(row.get("cest")),
            "m3_individual": _repair_text(row.get("m3 individual")),
            "lojas": [item.strip() for item in _repair_text(row.get("loja_sync")).split("|") if item.strip()],
            "atualizado_em": _repair_text(row.get("updated_at")),
        },
        "anuncios": {
            "ativos": live_ads,
            "snapshot_export": snapshot_ads,
            "vinculos_locais": cached_ads,
            "cobertura": ml_coverage,
        },
        "pesquisa_internet": internet_research,
        "pistas_complementares": {
            "cadastro_legado": {
                "encontrado": bool(legacy),
                "identidade_relevante": legacy_relevant,
                "nome": legacy_name,
                "oem": _repair_text(legacy.get("cg_oem")),
                "medidas": _repair_text(legacy.get("cg_medidas individuas") or legacy.get("cg_individual size")),
                "peso": _repair_text(legacy.get("cg_peso individual")),
                "para_que_serve": _repair_text(legacy.get("cg_para quer serve")),
                "finalidade": _repair_text(legacy.get("cg_finalizade")),
                "mlb_ids": _repair_text(legacy.get("mlb_ids")),
            },
            "planilha_fornecedor": supplier_records,
            "campos_pesquisa_favoritos": favorite_search,
        },
        "fontes_consultadas": sources,
        "qualidade": {
            "identidade": identity,
            "lacunas": missing_sections,
            "revisao_manual_recomendada": bool(alerts),
        },
        "alertas": alerts,
        "edicoes_manuais": existing.get("edicoes_manuais") if isinstance(existing.get("edicoes_manuais"), dict) else {},
        "gerado_em": _now_iso(),
    }


def _find_default_tenant(repo_root: Path, client_id: str) -> Path:
    appdata = os.environ.get("APPDATA", "")
    installed = Path(appdata) / "JK Sistema Cliente" / "local_app" / "info" / client_id if appdata else None
    if installed and (installed / "cadastro_produtos.csv").exists():
        return installed
    return repo_root / "info" / client_id


def _parse_args(argv: list[str]) -> argparse.Namespace:
    parser = argparse.ArgumentParser(description="Gera um dossie JSON rastreavel para cada SKU do cadastro.")
    parser.add_argument("--client-id", default="000002")
    parser.add_argument("--source-dir", type=Path, help="Pasta info/<client_id> de origem.")
    parser.add_argument("--output-dir", type=Path, help="Destino; padrao: <source-dir>/SKU.")
    parser.add_argument("--internet", action="store_true", help="Pesquisa cada produto na internet.")
    parser.add_argument("--refresh-internet", action="store_true", help="Ignora pesquisas ja salvas.")
    parser.add_argument("--internet-workers", type=int, default=3)
    parser.add_argument("--internet-results", type=int, default=4)
    parser.add_argument("--live-ml", action="store_true", help="Consulta anuncios ativos nas contas ML conectadas.")
    parser.add_argument("--live-ml-descriptions", action="store_true", help="Consulta uma descricao ativa por SKU.")
    parser.add_argument(
        "--insecure-web-ssl", "--insecure-ssl", dest="insecure_web_ssl", action="store_true",
        help="Desativa SSL apenas na pesquisa publica; nunca afeta chamadas ML autenticadas.",
    )
    return parser.parse_args(argv)


def main(argv: list[str] | None = None) -> int:
    args = _parse_args(list(argv if argv is not None else sys.argv[1:]))
    if args.insecure_web_ssl:
        urllib3.disable_warnings(urllib3.exceptions.InsecureRequestWarning)
    repo_root = Path(__file__).resolve().parents[1]
    source_dir = (args.source_dir or _find_default_tenant(repo_root, args.client_id)).resolve()
    client_id = str(args.client_id or "").strip()
    if not re.fullmatch(r"[A-Za-z0-9][A-Za-z0-9_.-]{0,63}", client_id):
        raise SystemExit("Client ID invalido.")
    if source_dir.name != client_id:
        raise SystemExit(
            f"A pasta de origem deve corresponder ao client-id {client_id}: {source_dir}"
        )
    output_dir = (args.output_dir or source_dir / "SKU").resolve()
    cadastro_path = source_dir / "cadastro_produtos.csv"
    if not cadastro_path.exists():
        raise SystemExit(f"Cadastro nao encontrado: {cadastro_path}")

    raw_rows = _read_csv(cadastro_path)
    rows: list[dict[str, str]] = []
    seen: set[str] = set()
    for row in raw_rows:
        sku = _norm_sku(row.get("sku"))
        if not sku or sku in seen or sku in EXCLUDED_DOSSIER_SKUS:
            continue
        seen.add(sku)
        row["sku"] = _repair_text(row.get("sku"))
        rows.append(row)
    rows.sort(key=lambda row: _fold(row.get("sku")))
    print(f"[SKU] cadastro: {len(rows)} SKU(s) unicos em {cadastro_path}", flush=True)

    output_dir.mkdir(parents=True, exist_ok=True)
    for excluded_sku in EXCLUDED_DOSSIER_SKUS:
        excluded_path = output_dir / _safe_filename(excluded_sku)
        if excluded_path.exists():
            excluded_path.unlink()
    existing: dict[str, dict[str, Any]] = {}
    for row in rows:
        sku = _norm_sku(row.get("sku"))
        existing[sku] = _read_json(output_dir / _safe_filename(row["sku"]), {})

    cached_links = _load_mlb_links(source_dir / "mlb_sku_vinculos.csv", seen)
    legacy_catalog = _load_legacy_catalog(source_dir, seen)
    supplier_catalog = _load_supplier_spreadsheet(repo_root / "Listas Diversos SKU.xlsx", seen)
    favorite_searches = _load_favorite_searches(source_dir / "favoritos_pesquisas_caio.json", seen)
    listing_snapshot = _load_listing_export(repo_root / "JK" / "Anuncios-2026_04_09-15_24.xlsx", seen)
    web_overrides = _load_web_research_overrides(
        repo_root / "scripts" / "sku_web_research_overrides.json", seen,
    )
    print(
        f"[SKU] fontes complementares: {len(legacy_catalog)} no legado; "
        f"{sum(bool(items) for items in supplier_catalog.values())} na planilha do fornecedor; "
        f"{len(favorite_searches)} nos campos de pesquisa; "
        f"{sum(bool(items) for items in listing_snapshot.values())} no snapshot de anuncios",
        flush=True,
    )
    if args.live_ml:
        live_ads, ml_coverage = _fetch_live_ml(
            source_dir,
            seen,
            include_descriptions=args.live_ml_descriptions,
            verify_ssl=True,
        )
    else:
        live_ads = {
            sku: list(((existing.get(sku) or {}).get("anuncios") or {}).get("ativos") or [])
            for sku in seen
        }
        old_coverages = [
            ((existing.get(sku) or {}).get("anuncios") or {}).get("cobertura") for sku in seen
        ]
        ml_coverage = next((item for item in old_coverages if isinstance(item, dict) and item), {
            "status": "nao_consultado_nesta_execucao", "cobertura_completa": False, "lojas": []
        })

    if args.internet:
        internet = _search_all_skus(
            rows,
            existing,
            favorite_searches,
            refresh=args.refresh_internet,
            workers=args.internet_workers,
            max_results=max(1, min(args.internet_results, 8)),
            verify_ssl=not args.insecure_web_ssl,
        )
    else:
        internet = {
            sku: (existing.get(sku) or {}).get("pesquisa_internet") or {
                "status": "nao_consultado_nesta_execucao", "resultados": []
            }
            for sku in seen
        }
    internet.update(web_overrides)

    generated = 0
    stats = {
        "com_anuncio_ativo": 0,
        "com_vinculo_mlb": 0,
        "com_snapshot_anuncio": 0,
        "com_resultado_internet": 0,
        "com_oem": 0,
        "com_medidas": 0,
        "com_compatibilidade": 0,
        "com_cadastro_legado": 0,
        "com_planilha_fornecedor": 0,
        "com_campos_pesquisa": 0,
        "com_fonte_externa": 0,
        "pesquisa_internet_indisponivel": 0,
        "pesquisa_internet_sem_resultado": 0,
        "revisao_manual": 0,
    }
    review_skus: list[str] = []
    for row in rows:
        sku_norm = _norm_sku(row.get("sku"))
        dossier = _build_dossier(
            row,
            client_id=args.client_id,
            live_ads=live_ads.get(sku_norm) or [],
            snapshot_ads=listing_snapshot.get(sku_norm) or [],
            cached_ads=cached_links.get(sku_norm) or [],
            ml_coverage=ml_coverage,
            internet_research=internet.get(sku_norm) or {"status": "indisponivel", "resultados": []},
            legacy=legacy_catalog.get(sku_norm) or {},
            supplier_records=supplier_catalog.get(sku_norm) or [],
            favorite_search=favorite_searches.get(sku_norm) or {},
            existing=existing.get(sku_norm) or {},
        )
        _write_json_atomic(output_dir / _safe_filename(row["sku"]), dossier)
        generated += 1
        product = dossier["produto"]
        stats["com_anuncio_ativo"] += bool(dossier["anuncios"]["ativos"])
        stats["com_vinculo_mlb"] += bool(dossier["anuncios"]["vinculos_locais"])
        stats["com_snapshot_anuncio"] += bool(dossier["anuncios"]["snapshot_export"])
        stats["com_resultado_internet"] += bool(dossier["pesquisa_internet"].get("resultados"))
        stats["com_oem"] += bool(product["oem"]["codigos"])
        stats["com_medidas"] += bool(product["medidas"]["itens"])
        stats["com_compatibilidade"] += bool(
            product["veiculos_compativeis"]["itens"] or product["equipamentos_ou_aplicacoes_compativeis"]["itens"]
        )
        stats["com_cadastro_legado"] += bool(legacy_catalog.get(sku_norm))
        stats["com_planilha_fornecedor"] += bool(supplier_catalog.get(sku_norm))
        stats["com_campos_pesquisa"] += bool(favorite_searches.get(sku_norm))
        stats["com_fonte_externa"] += bool(
            any(item.get("url") for item in dossier["anuncios"]["ativos"])
            or any(item.get("url") for item in dossier["anuncios"]["snapshot_export"])
            or any(item.get("link") for item in dossier["pistas_complementares"]["planilha_fornecedor"])
            or any(item.get("url") for item in dossier["pesquisa_internet"].get("resultados") or [])
        )
        stats["pesquisa_internet_indisponivel"] += dossier["pesquisa_internet"].get("status") == "indisponivel"
        stats["pesquisa_internet_sem_resultado"] += (
            dossier["pesquisa_internet"].get("status") == "pesquisado_sem_resultado"
        )
        if dossier["qualidade"]["revisao_manual_recomendada"]:
            stats["revisao_manual"] += 1
            review_skus.append(row["sku"])

    stats["sem_fonte_externa"] = generated - stats["com_fonte_externa"]

    expected_files = {_safe_filename(row["sku"]) for row in rows}
    stale_files = []
    for path in output_dir.glob("*.json"):
        if path.name.startswith("_") or path.name in expected_files:
            continue
        stale_files.append(path.name)

    index = {
        "schema_version": SCHEMA_VERSION,
        "client_id": args.client_id,
        "total_skus_cadastro": len(rows),
        "arquivos_sku_gerados": generated,
        "fonte_cadastro": _logical_source(args.client_id, "cadastro_produtos.csv"),
        "gerado_em": _now_iso(),
        "cobertura": stats,
        "skus_revisao_manual": review_skus,
        "skus_excluidos": sorted(EXCLUDED_DOSSIER_SKUS),
        "arquivos_obsoletos_nao_removidos": stale_files,
        "metodologia": (
            "Trechos tecnicos sao extraidos do cadastro, dos anuncios ativos e de resultados publicos relevantes. "
            "Campos sem evidencia permanecem vazios; compatibilidade e OEM nao sao inferidos por suposicao."
        ),
    }
    _write_json_atomic(output_dir / "_INDICE.json", index)
    print(f"[SKU] concluido: {generated} arquivos em {output_dir}", flush=True)
    print(f"[SKU] cobertura: {json.dumps(stats, ensure_ascii=False)}", flush=True)
    if stale_files:
        print(f"[SKU] aviso: {len(stale_files)} arquivo(s) obsoleto(s) preservado(s)", flush=True)
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
