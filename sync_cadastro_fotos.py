from pathlib import Path
import io
import re

import openpyxl
import requests
from google.auth.transport.requests import Request
from google.oauth2.service_account import Credentials


SPREADSHEET_ID = "1fehfm3TPRfM8PIVqRwGAZQHcSsf3p6vxKTKhM4B1lvM"
GID = "630209700"
INFO_DIR = Path("info")
CREDENTIALS_FILE = INFO_DIR / "credentials.json"
FOTOS_CONFIG_ARQUIVO = "cadastro_fotos_config.json"
FOTOS_CONFIG_SCHEMA = "jk.cadastro.fotos.v1"


def normalizar_sku(valor):
    texto = str(valor or "").strip()
    match = re.match(r"^\s*(\d+)\s*[-/]\s*([A-Za-z]{3})\s*$", texto)
    if match:
        meses = {
            "JAN": "1", "FEV": "2", "MAR": "3", "ABR": "4", "MAI": "5", "JUN": "6",
            "JUL": "7", "AGO": "8", "SET": "9", "OUT": "10", "NOV": "11", "DEZ": "12",
        }
        texto = f"{int(match.group(1))}-{meses.get(match.group(2).upper(), match.group(2))}"
    if re.fullmatch(r"\d+", texto):
        numero = int(texto)
        return str(numero).zfill(3) if 1 <= numero <= 9 else str(numero)
    return texto


def baixar_planilha_xlsx():
    creds = Credentials.from_service_account_file(
        str(CREDENTIALS_FILE),
        scopes=["https://www.googleapis.com/auth/spreadsheets.readonly"],
    )
    creds.refresh(Request())
    response = requests.get(
        f"https://docs.google.com/spreadsheets/d/{SPREADSHEET_ID}/export",
        headers={"Authorization": f"Bearer {creds.token}"},
        params={"format": "xlsx", "gid": GID},
        timeout=180,
    )
    response.raise_for_status()
    return response.content


def extrair_imagens_por_sku(xlsx_bytes):
    workbook = openpyxl.load_workbook(io.BytesIO(xlsx_bytes))
    worksheet = workbook["Dados"] if "Dados" in workbook.sheetnames else workbook[workbook.sheetnames[0]]
    imagens_por_sku = {}
    for image in getattr(worksheet, "_images", []):
        anchor = getattr(getattr(image, "anchor", None), "_from", None)
        if anchor is None:
            continue
        row_1 = int(anchor.row) + 1
        col_1 = int(anchor.col) + 1
        if col_1 != 2 or row_1 < 2:
            continue
        sku = normalizar_sku(worksheet.cell(row=row_1, column=1).value)
        if not sku:
            continue
        image_bytes = image._data()
        image_ext = str(getattr(image, "format", "png") or "png").lower().replace("jpeg", "jpg")
        atual = imagens_por_sku.get(sku)
        if atual is None or len(image_bytes) > len(atual[0]):
            imagens_por_sku[sku] = (image_bytes, image_ext)
    return imagens_por_sku


def nome_arquivo_foto_sku(sku, ext):
    sku_limpo = re.sub(r'[\\/:*?"<>|]+', '_', str(sku or '').strip())
    ext_limpa = re.sub(r'[^a-zA-Z0-9]+', '', str(ext or 'png').lower()) or 'png'
    return f"{sku_limpo}.{ext_limpa}"


def sincronizar_csv(csv_path: Path, imagens_por_sku):
    del csv_path, imagens_por_sku
    raise RuntimeError(
        "Utilitario global descontinuado: use o Cadastro por loja com store_id."
    )


def main():
    raise RuntimeError(
        "Utilitario global descontinuado: use o Cadastro por loja com store_id."
    )


if __name__ == "__main__":
    main()
